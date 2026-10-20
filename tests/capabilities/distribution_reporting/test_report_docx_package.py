"""The rendered report opens in Word without repair or field prompts."""

import io
import re
import zipfile
from importlib import import_module

from docx import Document

renderer = import_module(
    "manyselves.capabilities.distribution_reporting.runtime.rendering.v2_docx_renderer"
)

REPORT = "\n\n".join([
    "# 配电安全专家咨询报告",
    "## 1. 配电评估概述",
    "### 1.1. 评估背景",
    "正文一。",
    "## 2. 评估内容描述",
    "### 2.1. 配电系统架构问题",
    "#### 2.1.1. 配电系统负荷分配与过载风险",
    "正文二。",
    "## 3. 结论与建议",
    "### 3.2. 改善行动速查表",
    "正文三。",
    "**重要，紧急 · 芜湖区域**",
    "| 评估内容 | 发现问题 | 风险与危害 | 解决方案 | 备注 |",
    "| --- | --- | --- | --- | --- |",
    "| 低压柜 | 1AA3 温升 65K | 绝缘老化 | 一个月内紧固复测 |  |",
    "**一般 · 芜湖区域**",
    "| 评估内容 | 发现问题 | 风险与危害 | 解决方案 | 备注 |",
    "| --- | --- | --- | --- | --- |",
    "| 标识 | 缺少回路标识 | 误操作 | 三个月内补齐 |  |",
])


TEMPLATE = Document(renderer.TEMPLATE_PATH)


def _render() -> bytes:
    _, data = renderer.render_report_docx(report_text=REPORT, filename="package.docx")
    return data


def test_report_fields_are_balanced_and_not_refreshed_on_open() -> None:
    with zipfile.ZipFile(io.BytesIO(_render())) as package:
        body = package.read("word/document.xml").decode("utf-8")
        settings = package.read("word/settings.xml").decode("utf-8")

    counts = {
        kind: body.count(f'w:fldCharType="{kind}"') for kind in ("begin", "separate", "end")
    }
    assert counts["begin"] == counts["end"]
    assert counts["separate"] <= counts["begin"]
    assert "updateFields" not in settings


def test_toc_lists_only_the_headings_in_the_body() -> None:
    document = Document(io.BytesIO(_render()))
    texts = [paragraph.text for paragraph in document.paragraphs]
    title = texts.index("目录")
    start = renderer.find_content_start_paragraph_index(document)
    toc = [text for text in texts[title + 1:start] if text.strip()]
    body = [f"{number}. {name}" for number, name, _ in renderer.body_toc_entries(document)]

    assert toc == body
    assert "1. 配电评估概述" in toc and "2.1.1. 配电系统负荷分配与过载风险" in toc
    # Chapter 4 exists only when the inputs carry a special-topic file.
    assert not any(line.startswith("4.") for line in toc)
    # No page numbers copied from the template.
    assert not any(re.search(r"\t\d+$", line) for line in toc)


def test_toc_is_one_word_field_in_the_template_toc_styles() -> None:
    document = Document(io.BytesIO(_render()))
    body = document.element.body.xml
    assert body.count('TOC \\o "1-3"') == 1
    title = [paragraph.text for paragraph in document.paragraphs].index("目录")
    start = renderer.find_content_start_paragraph_index(document)
    styles = {
        paragraph.text: paragraph.style.name
        for paragraph in document.paragraphs[title + 1:start]
        if paragraph.text.strip()
    }
    assert styles["1. 配电评估概述"] == "toc 1"
    assert styles["1.1. 评估背景"] == "toc 2"
    assert styles["2.1.1. 配电系统负荷分配与过载风险"] == "toc 3"


def test_text_keeps_the_template_styles() -> None:
    document = Document(io.BytesIO(_render()))
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "表格2"):
        rendered, template = document.styles[name], TEMPLATE.styles[name]
        assert rendered.font.size == template.font.size
        assert rendered.paragraph_format.line_spacing == template.paragraph_format.line_spacing
        assert rendered.paragraph_format.space_after == template.paragraph_format.space_after
    for paragraph in document.paragraphs:
        if paragraph.style.name in {"Normal", "Heading 1", "Heading 2", "Heading 3"}:
            assert paragraph.paragraph_format.line_spacing is None
    labels = [p for p in document.paragraphs if p.text == "重要，紧急 · 芜湖区域"]
    assert labels and labels[0].style.name == "小标" and labels[0].runs[0].bold


def test_tables_follow_the_expert_table_format() -> None:
    from docx.oxml.ns import qn

    document = Document(io.BytesIO(_render()))
    assert len(document.tables) == 2
    for table in document.tables:
        properties = table._tbl.tblPr
        assert table.style.name == "Table Grid"
        layout = properties.find(qn("w:tblLayout"))
        assert layout is None or layout.get(qn("w:type")) == "autofit"
        assert properties.find(qn("w:tblLook")).get(qn("w:noHBand")) == "1"
        border = properties.find(qn("w:tblBorders")).find(qn("w:insideH"))
        assert (border.get(qn("w:val")), border.get(qn("w:sz"))) == ("single", "4")
        header = table.rows[0].cells
        assert all(
            cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "D4F4F1" for cell in header
        )
        assert all(cell.paragraphs[0].runs[0].bold for cell in header)
        row = table.rows[1].cells
        assert {cell.paragraphs[0].style.name for cell in row} == {"表格2"}
        # Short field centered, description left, empty note left empty.
        assert row[0].paragraphs[0].alignment is not None
        assert row[1].paragraphs[0].alignment is None
        assert row[4].text == ""
        following = table._tbl.getnext()
        assert following is None or following.tag != qn("w:p") or "".join(following.itertext()).strip()


def test_package_keeps_only_images_the_body_shows() -> None:
    document = Document(io.BytesIO(_render()))
    part = document.part
    referenced = set(part.element.xpath("//@r:embed | //@r:link | //@r:id"))
    images = [r_id for r_id, rel in part.rels.items() if rel.reltype.endswith("/image")]
    assert set(images) <= referenced


def test_extensionless_photo_gets_a_real_part_extension(tmp_path, monkeypatch) -> None:
    from PIL import Image

    photo = tmp_path / "photo-without-extension"
    Image.new("RGB", (8, 8), "red").save(photo, format="JPEG")
    # Compression unavailable: the original file is embedded as-is.
    monkeypatch.setattr(renderer, "compressed_photo_path", lambda path: path)

    document = Document()
    document.add_paragraph().add_run().add_picture(renderer.photo_source(photo))
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as package:
        media = [name for name in package.namelist() if name.startswith("word/media/")]

    assert media and all(re.search(r"\.[a-z]+$", name) for name in media)
