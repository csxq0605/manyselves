"""Knowledge packets must not disguise clipped standards as complete rules."""

from pathlib import Path

import pytest

from manyselves.capabilities.distribution_reporting.runtime.models.reporting import (
    SpecialTopicPlan,
    SpecialTopicSectionRequirement,
)
from manyselves.capabilities.distribution_reporting.runtime.research.knowledge_context import (
    KnowledgeContextBuilder,
)
from manyselves.capabilities.distribution_reporting.runtime.research.reference_library import (
    ReferenceDocument,
)


def _builder(tmp_path, texts):
    builder = KnowledgeContextBuilder(tmp_path, "knowledge-fidelity")
    builder._documents = tuple(
        ReferenceDocument(f"规程{index}", f"Knowledge/rule-{index}.md", text, str(index))
        for index, text in enumerate(texts)
    )
    return builder


def test_module_delivers_complete_selected_document_once_including_late_exception(tmp_path):
    text = "# 适用范围\n仅适用于工况甲。\n# 2.4.4\n" + "测量方法。" * 250 + "\n例外：工况乙不适用。"
    builder = _builder(tmp_path, [text])
    packet = builder.build_module("2.4")
    assert text in packet.text
    assert packet.text.count("例外：工况乙不适用。") == 1
    assert len(packet.source_ids) == 1
    assert "完整原文" in packet.text


def test_equal_excerpt_does_not_merge_different_authorities(tmp_path):
    prefix = "# 2.4.4\n" + "相同前文。" * 250
    packet = _builder(tmp_path, [prefix + "\n甲仅适用于A。", prefix + "\n乙仅适用于B。"]).build_module("2.4")
    assert len(packet.source_ids) == 2
    assert "甲仅适用于A。" in packet.text and "乙仅适用于B。" in packet.text


def test_oversize_document_has_explicit_unexpanded_status_and_complete_source(tmp_path):
    text = "# 2.4.4\n" + "规定。" * 8000 + "\n尾部例外。"
    builder = _builder(tmp_path, [text])
    packet = builder.build_module("2.4")
    assert len(packet.source_ids) == 1
    assert "原文未展开" in packet.text and "open_reference" in packet.text
    assert "Knowledge/rule-0.md" in packet.text
    assert "规定。" not in packet.text  # no plausible-looking partial rule
    ref = builder.ledger.content_ref(packet.source_ids[0])
    assert ref is not None and ref.as_posix() in packet.text
    assert (tmp_path / ref).read_text(encoding="utf-8") == text
    assert len(packet.text) <= builder.MAX_MODULE_CHARS


def test_quality_packet_preserves_complete_selected_rule(tmp_path):
    text = "# 报告质量\n" + "检查方法。" * 250 + "\n例外情况必须列明。"
    packet = _builder(tmp_path, [text]).build_quality()
    assert text in packet.text


def test_special_topic_never_truncates_user_requirement_to_fit_knowledge_budget(tmp_path):
    requirement = "需要核实工况。" * 2100 + "尾部用户要求。"
    plan = SpecialTopicPlan(
        source_ref=Path("Inputs/topics.md"), source_sha256="a" * 64,
        sections=[SpecialTopicSectionRequirement(section_id="4.1", title="检测", requirement=requirement)],
    )
    packet = _builder(tmp_path, ["检测\n规则原文。" * 1000]).build_special_topics(plan)
    assert requirement in packet.text
    assert "原文未展开" in packet.text
    assert len(packet.source_ids) == 1


@pytest.mark.asyncio
async def test_unexpanded_source_opens_through_real_reference_tool(tmp_path):
    from manyselves.capabilities.distribution_reporting.runtime.research_tools import (
        OpenReferenceTool,
    )

    source = tmp_path / "Knowledge" / "thermal.md"
    source.parent.mkdir()
    text = "# 2.4.4\n" + "测量条件。" * 4000 + "\n尾部适用例外。"
    source.write_text(text, encoding="utf-8")
    builder = KnowledgeContextBuilder(tmp_path, "reference-tool")
    packet = builder.build_module("2.4")
    assert "原文未展开" in packet.text
    result = await OpenReferenceTool(builder.library, builder.ledger)(locator="Knowledge/thermal.md")
    assert result["source_id"] in packet.source_ids
    assert result["text"] == text


def test_one_document_serving_multiple_sections_is_inlined_once(tmp_path):
    text = "# 2.4.1.1\n工况说明。\n# 2.4.4\n另一个主题的适用条件。"
    packet = _builder(tmp_path, [text]).build_module("2.4")
    assert packet.text.count(text) == 1
    assert packet.text.count("候选来源：R-001") >= 2


def test_long_standard_delivers_whole_relevant_clause_with_conditions_and_dependencies(tmp_path):
    text = (
        "# 合成规程\n版本：虚构测试版。\n"
        "## 1 适用范围\n仅工况甲有效。\n"
        "## 2 无关章节\n" + "无关资料。" * 4000 + "\n"
        "## 3 测量通则\n所有本章测量须稳定运行。\n"
        "### 3.1 2.4.4 检测\n判断须按第 5.1 条。\n"
        "#### 3.1.1 例外\n工况乙不得使用本条。\n"
        "### 3.2 其他方法\n不得误认为相关方法。\n"
        "## 5 配套定义\n本章给出辅助定义。\n"
        "### 5.1 修正方法\n采用现场实测环境值，不得用默认值。\n"
    )
    packet = _builder(tmp_path, [text]).build_module("2.4")
    assert "自动展开的完整章节" in packet.text
    for expected in ("版本：虚构测试版", "仅工况甲有效", "所有本章测量须稳定运行",
                     "工况乙不得使用本条", "本章给出辅助定义", "不得用默认值"):
        assert expected in packet.text
    assert "无关资料。" not in packet.text and "不得误认为相关方法" not in packet.text
    assert "不是全文" in packet.text and "原文行" in packet.text
    assert len(packet.text) <= KnowledgeContextBuilder.MAX_MODULE_CHARS


def test_structured_retrieval_never_slices_an_oversize_clause_or_hides_unresolved_reference(tmp_path):
    text = "# 其他\n" + "无关资料。" * 4000 + "\n# 2.4.4 检测\n依据第 9.9 条判断。\n"
    packet = _builder(tmp_path, [text]).build_module("2.4")
    assert "依据第 9.9 条判断。" in packet.text
    assert "未解析的条款引用：9.9" in packet.text
    assert "open_reference" in packet.text
    huge = "# 其他\n无关。\n# 2.4.4\n" + "长条款。" * 4000 + "\n末尾例外。"
    packet = _builder(tmp_path, [huge]).build_module("2.4")
    assert "原文未展开" in packet.text
    assert "长条款。" not in packet.text


def test_knowledge_docx_preserves_tabular_rule_in_original_context(tmp_path):
    from docx import Document

    source = tmp_path / "Knowledge" / "rule.docx"
    source.parent.mkdir()
    doc = Document()
    doc.add_paragraph("2.4.4：仅适用于工况甲；规则如下。")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "测试阈值=17"
    doc.add_paragraph("例外：工况乙不适用上述表格。")
    doc.save(source)
    packet = KnowledgeContextBuilder(tmp_path, "docx-rule").build_module("2.4")
    assert "测试阈值=17" in packet.text
    assert packet.text.index("仅适用于工况甲") < packet.text.index("测试阈值=17") < packet.text.index("例外：工况乙")


def test_structured_retrieval_handles_reference_cycles_and_ignores_fenced_headings():
    from manyselves.capabilities.distribution_reporting.runtime.research.reference_sections import (
        select_reference_sections,
    )

    text = ("# 1 检测\n2.4.4 依第 3.1 条。\n```text\n# 伪造边界\n```\n"
            "# 2 无关\nOTHER\n# 3 定义\n## 3.1 条件\n按第 1.1 条。\n"
            "# 1.1 补充\n见 3.1。\n末尾例外。\n")
    selected = select_reference_sections(text, ("2.4.4",))
    assert selected is not None and "末尾例外" in selected
    assert "```text\n# 伪造边界\n```" in selected
    assert "OTHER" not in selected


def test_ambiguous_clause_number_is_disclosed_instead_of_picking_one():
    from manyselves.capabilities.distribution_reporting.runtime.research.reference_sections import (
        select_reference_sections,
    )

    text = "# 1 检测\n2.4.4 参照第 3.1 条。\n# 3.1 甲\n甲规则\n# 3.1 乙\n乙规则\n"
    selected = select_reference_sections(text, ("2.4.4",))
    assert "未解析的条款引用：3.1" in selected
    assert "甲规则" not in selected and "乙规则" not in selected
