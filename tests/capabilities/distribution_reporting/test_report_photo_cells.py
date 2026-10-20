"""Photos sit in the problem table's 「现场图片」 column, as in the expert report."""

import io
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from PIL import Image

from manyselves.capabilities.distribution_reporting.runtime.rendering.handoff_docx import (
    PackagedV2DocxCore,
)
from manyselves.capabilities.distribution_reporting.runtime.rendering.pds_docx_renderer import (
    PdsDocxRenderer,
)

TEMPLATE = Path("manyselves/templates/reporting/report_template.docx")

NARRATIVE = "\n".join([
    "#### 2.4.3 带病运行问题汇总",
    "1AA3 出线开关温升偏高。",
    "| 位置 | 柜号及回路号 | 问题描述 | 建议 | 现场图片 |",
    "| --- | --- | --- | --- | --- |",
    "| 芜湖1号配电室 | 1AA3 | 温升 65K | 紧固复测 | P-0001、P-0002 |",
    "| 芜湖1号配电室 | 1AA5 | 标识缺失 | 补齐 |  |",
    "| 芜湖1号配电室 | 1AA7 | 外壳锈蚀 | 除锈 | P-0099 |",
])


def test_photo_ids_in_the_photo_column_become_tokens_only_for_report_photos() -> None:
    placed = PdsDocxRenderer._place_photo_tokens(NARRATIVE, {"P-0001", "P-0002"})

    assert "[[PHOTO:P-0001]] [[PHOTO:P-0002]]" in placed
    # An id that is not a report photo stays as written.
    assert "| P-0099 |" in placed
    assert "1AA3 出线开关温升偏高。" in placed


def test_rendered_problem_table_carries_the_photos_in_its_cells(tmp_path) -> None:
    paths = {}
    for photo_id in ("P-0001", "P-0002"):
        path = tmp_path / f"{photo_id}.jpg"
        Image.new("RGB", (40, 30), "red").save(path, format="JPEG")
        paths[photo_id] = path
    markdown = "\n\n".join([
        "# 报告",
        "## 1. 配电评估概述",
        "### 1.1 评估背景",
        "正文。",
        "## 2. 评估内容描述",
        "### 2.4 配电设备/元件风险",
        PdsDocxRenderer._place_photo_tokens(NARRATIVE, set(paths)),
    ])
    _, data = PackagedV2DocxCore(TEMPLATE).render_approved_prose(markdown)
    document = Document(io.BytesIO(data))
    report = SimpleNamespace(
        photos=[SimpleNamespace(id=photo_id, path=path) for photo_id, path in paths.items()]
    )

    PdsDocxRenderer._materialize_photos(document, report)

    table = document.tables[0]
    assert len(document.tables) == 1
    photo_cell = table.rows[1].cells[4]
    assert len(photo_cell._tc.xpath(".//w:drawing")) == 2
    assert "PHOTO" not in photo_cell.text
    assert table.rows[2].cells[4].text == ""
    assert table.rows[3].cells[4].text == "P-0099"
