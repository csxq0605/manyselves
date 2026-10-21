"""Template role Skills: extraction on the real template, publishing, lessons.

Agent turns are scripted; the outline, acceptance, workflows, tools and the
publish step are the production ones, run on the packaged expert template.
"""

from __future__ import annotations

import asyncio
import json
import re
import shutil
import tempfile
from pathlib import Path
from typing import Any
from uuid import UUID

import pytest
import yaml

from manyselves.capabilities.distribution_reporting.runtime.lesson_triage import (
    LessonTriageTools,
    lesson_admission,
)
from manyselves.capabilities.distribution_reporting.runtime.models.reporting import (
    EvidenceItem,
    SourceLocation,
)
from manyselves.capabilities.distribution_reporting.runtime.module_authoring_preparation import (
    load_template_skill,
)
from manyselves.capabilities.distribution_reporting.runtime.template_skill_binding import (
    build_template_skill_binding,
)
from manyselves.capabilities.distribution_reporting.runtime.template_skills import (
    EXPERT_REPORT_REF,
    publish_template_skills,
    skill_order,
    skills_with_lessons,
    template_policy,
    template_roles,
    writing_lessons,
)
from manyselves.config.schema import AgentDefaults
from manyselves.runtime.exemplar.extraction import role_scope
from manyselves.runtime.exemplar.lessons import LessonBook, load_lessons
from manyselves.runtime.exemplar.outline import document_outline
from manyselves.runtime.loops.bus import MessageBus
from manyselves.runtime.services import RuntimeServicesView
from tests.capabilities.distribution_reporting.test_ledger_report import (
    _input_of,
    _task,
    run_ledger_report,
)
from tests.capabilities.distribution_reporting.test_public_reporting_runtime import (
    _write_complete_template_skill_fixture,
)
from tests.runtime.scripted_loop import ScriptedLoop

EXPERT_TEXT = Path(
    "docs/research/session-records/2026-10-10_to_2026-10-15/audit_manyselves/"
    "配电安全专家咨询报告(专家优化版).docx.txt"
)
_TOC = re.compile(r"^(\d+(?:\.\d+)*)\.\s+(.+?)\d+$")


def _expert_report(path: Path) -> Path:
    """A DOCX with the expert report's headings, built from its table of contents.

    Each section gets one paragraph so it has something to cite; 3.2 also
    gets a table.  The expert report itself is not in the repository.
    """

    from docx import Document

    document = Document()
    for line in EXPERT_TEXT.read_text(encoding="utf-8").splitlines():
        match = _TOC.match(line.strip())
        if not match:
            continue
        number, title = match.groups()
        document.add_heading(title, level=number.count(".") + 1)
        if number == "2":
            continue  # chapter 2 opens straight into 2.1, as in the expert report
        document.add_paragraph(f"本节先说明{title}的现状，再分析风险并给出建议。")
        if number == "3.2":
            table = document.add_table(rows=2, cols=5)
            for cell, value in zip(table.rows[0].cells, ("评估内容", "发现问题", "风险与危害", "解决方案", "备注"), strict=True):
                cell.text = value
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


TEMPLATE = _expert_report(Path(tempfile.mkdtemp()) / "expert.docx")


def test_declared_roles_resolve_on_the_template_and_keep_the_loading_order() -> None:
    outline = document_outline(TEMPLATE)
    assert skill_order() == (
        "author-2.1", "author-2.2", "author-2.3", "author-2.4", "author-2.5",
        "chief-editor-chapter-1", "chief-editor-chapter-3",
    )
    covered: set[str] = set()
    for role in template_roles():
        scope = role_scope(outline, role)
        if not role.extract:
            assert not scope.sections
            continue
        assert scope.missing == (), role.id
        assert scope.sections, role.id
        numbers = scope.numbers
        assert not numbers & covered, role.id  # no section belongs to two roles
        covered |= numbers
    # Every section with content belongs to some role, except chapter 4: its
    # topics come from each run's input, so nothing in it is a general method.
    with_content = {s.number for s in outline.sections if s.blocks}
    assert {n for n in with_content if n.split(".")[0] != "4"} <= covered
    assert not {n for n in covered if n.split(".")[0] == "4"}


def _good_entries(packet: dict[str, Any]) -> list[dict[str, Any]]:
    """One entry per required use for each section with text, quoting its first text block."""

    required = [use["id"] for use in packet["uses"] if use["required"]]
    entries = []
    for section in packet["sections"]:
        block = next((b for b in section["blocks"] if len(re.sub(r"\s+", "", b["text"])) >= 8), None)
        if block is None:
            continue
        quote = re.sub(r"\s+", "", block["text"])[:8]
        for use in required:
            entries.append(
                {
                    "use": use,
                    "text": "先看现状，再推风险，最后对应建议，三者一一对应。",
                    "applies_when": "未明确",
                    "locator": block["locator"],
                    "quote": quote,
                }
            )
    return entries


def _skipped(packet: dict[str, Any]) -> list[dict[str, str]]:
    cited = {entry["locator"] for entry in _good_entries(packet)}
    return [
        {"section": section["number"], "reason": "只有图片"}
        for section in packet["sections"]
        if section["blocks"] and not cited & {b["locator"] for b in section["blocks"]}
    ]


def _extraction_answer(message):
    task = _task(message.content)
    value = _input_of(message.content)
    if task == "extract-template-role":
        entries = _good_entries(value)
        if value["role_id"] == "author-2.1":
            # The first answer quotes a number; the program sends it back once.
            bad = dict(entries[0], text="按110kV、10kV、0.4kV分段描述。")
            entries = [bad, *entries[1:]]
        return ("submit", [{"entries": entries, "skipped": _skipped(value)}])
    if task == "revise-template-role":
        assert value["rejected"] and "数值" in value["rejected"][0]["reasons"][0]
        return ("submit", [{"entries": REVISED["author-2.1"], "skipped": []}])
    raise AssertionError(task)


REVISED: dict[str, list[dict[str, Any]]] = {}


async def _run(workspace: Path, workflow_id: str, values: dict[str, Any], answer, command: int):
    bus = MessageBus()
    loops: list[ScriptedLoop] = []

    def build(**kwargs: Any) -> ScriptedLoop:
        loop = ScriptedLoop(bus, answer, kwargs)
        loops.append(loop)
        return loop

    binding = build_template_skill_binding(
        workspace,
        RuntimeServicesView(
            workspace=workspace,
            bus=bus,
            active_provider=object(),
            agent_defaults=AgentDefaults(),
            global_knowledge_root=None,
        ),
        loop_builder=build,
    )
    task = asyncio.create_task(bus.process_queue())
    try:
        accepted = await binding.start(
            UUID(f"70000000-0000-4000-8000-{command:012d}"), workflow_id, values
        )
    finally:
        await binding.close()
        bus.shutdown()
        await task
    return binding, accepted, loops


@pytest.mark.asyncio
async def test_extraction_checks_entries_renders_candidates_and_publishes(tmp_path: Path) -> None:
    _write_complete_template_skill_fixture(tmp_path)
    (tmp_path / EXPERT_REPORT_REF).parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, tmp_path / EXPERT_REPORT_REF)
    published_before = (tmp_path / "Inputs/report-template-role-skills/author-2.1/SKILL.md").read_text(
        encoding="utf-8"
    )
    outline = document_outline(TEMPLATE)
    role = next(r for r in template_roles() if r.id == "author-2.1")
    from manyselves.runtime.exemplar.extraction import role_packet

    REVISED["author-2.1"] = _good_entries(role_packet(outline, role_scope(outline, role), template_policy()))

    binding, accepted, loops = await _run(tmp_path, "extract-template-skills", {}, _extraction_answer, 1)
    run = binding.get_run(accepted["run_id"])
    assert run["run"]["status"] == "completed", run["state"].get("error")
    result = next(
        item["value"] for item in binding.get_outputs(accepted["run_id"])["outputs"] if item["id"] == "result"
    )
    extracted = [r.id for r in template_roles() if r.extract]
    assert sorted(result["candidates"]) == sorted(extracted)
    assert result["failed_roles"] == {}
    # One conversation per role; only author-2.1 needed a correction turn.
    turns = {loop.kwargs["agent_type"]: len(loop.received) for loop in loops}
    assert sum(turns.values()) == len(extracted) + 1
    assert result["roles"]["author-2.1"]["revised"] is True
    for role_id in extracted:
        counts = result["roles"][role_id]
        assert counts["accepted"] > 0 and counts["rejected"] == 0 and counts["uncovered"] == [], role_id

    candidate = (tmp_path / result["candidates"]["author-2.1"]).read_text(encoding="utf-8")
    # Headed by our own section titles, not the expert report's.
    assert "## 2.1.1 配电系统负荷分配与过载风险" in candidate
    assert "110kV" not in candidate
    # The no-model precheck re-reads the same run under the current rules.
    from scripts.maintenance.precheck_template_skills import precheck

    report = precheck(tmp_path, accepted["run_id"], None)
    assert "## author-2.1" in report and "只因场地词被拒的条目" in report
    assert "判断方法：" in candidate and "写法：" in candidate
    comparison = (tmp_path / result["compare_ref"]).read_text(encoding="utf-8")
    assert "## author-2.1" in comparison and "| 结果 | 现用版条目 |" in comparison
    assert set(result["compare_counts"]) >= {"author-2.1", "chief-editor-chapter-3"}
    diff = (tmp_path / result["diff_ref"]).read_text(encoding="utf-8")
    assert "author-2.1（新提炼）" in diff
    # Nothing reached Inputs before publishing.
    published = tmp_path / "Inputs/report-template-role-skills/author-2.1/SKILL.md"
    assert published.read_text(encoding="utf-8") == published_before

    outcome = publish_template_skills(tmp_path, accepted["run_id"])
    assert sorted(outcome["published"]) == sorted(extracted)
    assert published.read_text(encoding="utf-8") == candidate
    backup = tmp_path / outcome["backup"] / "author-2.1/SKILL.md"
    assert backup.read_text(encoding="utf-8") == published_before
    chapter4 = tmp_path / "Inputs/report-template-role-skills/chief-editor-chapter-4/SKILL.md"
    assert "chief-editor-chapter-4: preserve" in chapter4.read_text(encoding="utf-8")
    source = json.loads(
        (tmp_path / "Inputs/report-template-role-skills/source.json").read_text(encoding="utf-8")
    )
    assert source["last_extraction"]["run_id"] == accepted["run_id"]
    # The published package still loads for a report run.
    state: dict[str, Any] = {}
    assert load_template_skill(tmp_path, state)
    assert state["template_skill_text"]["author-2.1"] == candidate


@pytest.mark.asyncio
async def test_extraction_can_redo_only_some_roles_and_needs_the_expert_report(tmp_path: Path) -> None:
    _write_complete_template_skill_fixture(tmp_path)
    # Without the expert report in Templates/ nothing falls back to the style template.
    with pytest.raises(ValueError, match="expert report"):
        await _run(tmp_path, "extract-template-skills", {}, _extraction_answer, 2)

    (tmp_path / EXPERT_REPORT_REF).parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, tmp_path / EXPERT_REPORT_REF)
    roles = ["chief-editor-chapter-1", "chief-editor-chapter-3"]
    binding, accepted, loops = await _run(
        tmp_path, "extract-template-skills", {"roles": roles}, _extraction_answer, 3
    )
    result = next(
        item["value"] for item in binding.get_outputs(accepted["run_id"])["outputs"] if item["id"] == "result"
    )
    assert sorted(result["candidates"]) == roles and len(loops) == 2
    chapter1 = (tmp_path / result["candidates"]["chief-editor-chapter-1"]).read_text(encoding="utf-8")
    assert "## 1.2 健康度总览" in chapter1 and "章节综合：" in chapter1
    chapter3 = (tmp_path / result["candidates"]["chief-editor-chapter-3"]).read_text(encoding="utf-8")
    assert "表（行动表由系统按结构化字段生成" in chapter3 and "评估内容" not in chapter3


def _item(index: int, subject: str, **update: Any) -> EvidenceItem:
    return EvidenceItem(
        id=f"E-{index:04d}",
        subject=subject,
        fact=f"{subject}记录",
        source=SourceLocation(file_id="s4-4", path="Inputs/S4-4.xlsx", sheet="表", cell=f"A{index}"),
        module_id="2.4",
        submodule_id="2.4.1.1",
        **update,
    )


def test_lessons_join_a_skill_only_when_their_trigger_fires() -> None:
    book = LessonBook.model_validate(
        {"lessons": [
            {"id": "L-001", "roles": ["author-2.2"], "error_type": "照片缺失时把对象写成正常",
             "rule": "原表引用的照片取不出时，写明照片缺失，不据此判断正常。",
             "trigger": {"kind": "mapping_gap", "value": "photo_media_missing"},
             "source": {"run": "r", "finding": "F-001", "date": "2026-10-21"}},
            {"id": "L-002", "roles": ["author-2.2"], "error_type": "撤回的教训", "rule": "不再加载。",
             "status": "revoked", "revoked_reason": "测试",
             "source": {"run": "r", "finding": "F-002", "date": "2026-10-21"}},
        ]}
    )
    skills = {"author-2.2": "写法", "author-2.4": "写法"}
    without = skills_with_lessons(skills, items=[_item(1, "柜")], mapping_gaps=[], book=book)
    assert without == skills
    with_gap = skills_with_lessons(
        skills, items=[_item(1, "柜")], mapping_gaps=[{"code": "photo_media_missing"}], book=book
    )
    assert "易错：照片缺失时把对象写成正常" in with_gap["author-2.2"]
    assert "撤回的教训" not in with_gap["author-2.2"]
    assert with_gap["author-2.4"] == "写法"


def _variants() -> dict[str, Any]:
    import tempfile

    from scripts.benchmarks.fixed_structure_variants import mapped_variants, synthetic_workbook

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        source = synthetic_workbook(root / "S4-4.xlsx")
        (root / "v").mkdir()
        return {
            name: (items, [gap.model_dump(mode="json") for gap in gaps])
            for name, _info, items, gaps, *_rest in mapped_variants(source, root / "v")
        }


def test_adopted_lessons_are_admissible_on_fixed_structure_variants() -> None:
    report = lesson_admission(writing_lessons(), _variants())
    assert all(entry["admissible"] for entry in report.values()), report


def test_lesson_admission_refuses_a_trigger_no_variant_reaches() -> None:
    book = LessonBook.model_validate(
        {"lessons": [
            {"id": "L-001", "roles": ["author-2.2"], "error_type": "缺环境温度时仍写温升",
             "rule": "缺环境温度时写明温升无法计算。",
             "trigger": {"kind": "mapping_gap", "value": "missing_thermal_reference"},
             "source": {"run": "r", "finding": "F-001", "date": "2026-10-21"}},
            {"id": "L-002", "roles": ["author-2.2"], "error_type": "某种从不出现的缺口",
             "rule": "按通用写法处理。",
             "trigger": {"kind": "mapping_gap", "value": "no_such_gap"},
             "source": {"run": "r", "finding": "F-002", "date": "2026-10-21"}},
        ]}
    )
    report = lesson_admission(book, _variants())
    assert report["L-001"]["admissible"] and report["L-001"]["fires_on"] == ["clear_ambient"]
    assert not report["L-002"]["admissible"]


def _triage_answer(message):
    task = _task(message.content)
    value = _input_of(message.content)
    ids = [finding["id"] for finding in value.get("findings", [])]
    if task == "triage-lessons":
        assert ids, "the run's reviewer and editor findings reach the triage"
        return ("submit", [{"items": [
            {"finding_ids": ids[:1], "route": "lesson", "reason": "写法问题",
             "lesson": {"roles": ["author-2.4"], "error_type": "设备4的依据没写温度",
                        "rule": "依据里写明温度。", "trigger": {"kind": "always"}}},
            {"finding_ids": ids[1:], "route": "program_check", "reason": "程序可查"},
        ]}])
    if task == "revise-lessons":
        problem = value["problems"][0]
        assert any("本次运行的对象：设备4" in reason for reason in problem["reasons"])
        return ("submit", [{"items": [
            {"finding_ids": TRIAGE_IDS[:1], "route": "lesson", "reason": "写法问题",
             "lesson": {"roles": ["author-2.4"], "error_type": "判断依据只写结论不写读数",
                        "rule": "依据里写出所用的读数和对照的参考值。", "trigger": {"kind": "always"}}},
            {"finding_ids": TRIAGE_IDS[1:], "route": "program_check", "reason": "程序可查"},
        ]}])
    raise AssertionError(task)


TRIAGE_IDS: list[str] = []


@pytest.mark.asyncio
async def test_triage_turns_run_findings_into_checked_lesson_candidates(tmp_path: Path) -> None:
    _binding, report, _loops = await run_ledger_report(tmp_path)
    context = LessonTriageTools(tmp_path).prepare(
        {"run_id": "probe", "request": {"source_run_id": report["run_id"]}}
    )
    findings = context["packet"]["findings"]
    assert {f["origin"] for f in findings} == {"审查", "编辑"}
    assert "设备4" in json.loads((tmp_path / context["terms_ref"]).read_text(encoding="utf-8"))
    TRIAGE_IDS[:] = [f["id"] for f in findings]

    binding, accepted, _ = await _run(
        tmp_path, "triage-writing-lessons", {"source_run_id": report["run_id"]}, _triage_answer, 2
    )
    run = binding.get_run(accepted["run_id"])
    assert run["run"]["status"] == "completed", run["state"].get("error")
    result = next(
        item["value"] for item in binding.get_outputs(accepted["run_id"])["outputs"] if item["id"] == "result"
    )
    assert result["revised"] is True and result["lesson_candidates"] == 1
    candidates = load_lessons(tmp_path / result["candidates_ref"])
    lesson = candidates.lessons[0]
    assert lesson.id == "L-001" and lesson.source.run == report["run_id"]
    assert lesson.error_type == "判断依据只写结论不写读数"
    assert yaml.safe_load((tmp_path / result["candidates_ref"]).read_text(encoding="utf-8"))
    # The candidate passes the same admission the adopted file is held to.
    assert lesson_admission(candidates, _variants())["L-001"]["admissible"]
