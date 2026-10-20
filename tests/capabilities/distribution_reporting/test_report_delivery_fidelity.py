"""The delivered DOCX carries footnotes, a filled cover, the draft mark, the
Markdown's full text and correctly named media, and surfaces render warnings."""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from manyselves.capabilities.distribution_reporting.runtime.delivery_tools import (
    DeliveryTools,
)
from manyselves.capabilities.distribution_reporting.runtime.rendering.handoff_docx import (
    PackagedV2DocxCore,
)
from manyselves.capabilities.distribution_reporting.runtime.rendering.pds_docx_renderer import (
    PdsDocxRenderer,
)
from manyselves.capabilities.distribution_reporting.runtime.rendering.report_footnotes import (
    markdown_with_footnotes,
    split_footnote_definitions,
)
from manyselves.capabilities.distribution_reporting.runtime.report_front_matter import (
    report_cover_facts,
    report_is_draft,
)
from manyselves.capabilities.distribution_reporting.runtime.storage import ReportingStore
from tests.capabilities.distribution_reporting.test_ledger_report import (
    assert_media_parts_are_named_by_type,
)

TEMPLATE = Path("manyselves/templates/reporting/report_template.docx")


def _render(markdown: str) -> Document:
    _, data = PackagedV2DocxCore(TEMPLATE).render_approved_prose(markdown)
    return Document(io.BytesIO(data))


def _saved(document: Document, path: Path) -> Path:
    buffer = io.BytesIO()
    document.save(buffer)
    path.write_bytes(PdsDocxRenderer._canonical_docx(buffer.getvalue()))
    return path


def test_citations_become_numbered_markdown_footnotes() -> None:
    markdown = "\n".join([
        "# 报告",
        "",
        "第一段。",
        "",
        "[[CITE:2]]",
        "",
        "| 对象 | 问题 |",
        "| --- | --- |",
        "| 柜1 | 温升 |",
        "",
        "[[CITE:1]]",
    ])

    footnoted = markdown_with_footnotes(markdown, {1: "E-0002 表二", 2: "E-0001 表一"})

    # Numbered by position, joined to the paragraph, kept apart from a table.
    assert "第一段。[^1]" in footnoted
    assert "| 柜1 | 温升 |\n\n[^2]" in footnoted
    assert footnoted.rstrip().endswith("[^1]: E-0001 表一\n[^2]: E-0002 表二")
    body, definitions = split_footnote_definitions(footnoted)
    assert definitions == {1: "E-0001 表一", 2: "E-0002 表二"}
    assert "]:" not in body
    assert markdown_with_footnotes("无引用。", {}) == "无引用。"


@pytest.mark.parametrize("with_template", [True, False])
def test_footnote_references_become_word_footnotes(tmp_path: Path, with_template: bool) -> None:
    if with_template:
        document = _render(
            "\n\n".join([
                "# 配电安全专家咨询报告",
                "## 1. 配电评估概述",
                "### 1.1 评估背景",
                "**结论**：设备检查完成。[^1]",
                "| 对象 | 问题 |\n| --- | --- |\n| 柜1 | 温升[^2] |",
            ])
        )
    else:
        document = Document()
        document.add_paragraph("设备检查完成。[^1]")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "温升[^2]"

    count = PdsDocxRenderer._materialize_citations(
        document, {1: "E-0001 收资表（Inputs/a.xlsx）", 2: "E-0002 测温（Inputs/b.xlsx）"}
    )
    output = _saved(document, tmp_path / "footnotes.docx")

    assert count == 2
    with zipfile.ZipFile(output) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
        footnotes = archive.read("word/footnotes.xml").decode("utf-8")
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        types = archive.read("[Content_Types].xml").decode("utf-8")
    assert body.count("<w:footnoteReference ") == 2
    assert "[^" not in body
    assert "E-0001 收资表（Inputs/a.xlsx）" in footnotes
    assert "E-0002 测温（Inputs/b.xlsx）" in footnotes
    assert footnotes.count('w:type="separator"') == 1
    assert "relationships/footnotes" in rels
    assert 'PartName="/word/footnotes.xml"' in types
    reloaded = Document(output)
    texts = [p.text for p in reloaded.paragraphs]
    assert any(text.endswith("设备检查完成。") for text in texts)
    if with_template:
        # The label run keeps its bold formatting.
        paragraph = next(p for p in reloaded.paragraphs if "设备检查完成" in p.text)
        assert paragraph.runs[0].bold


def _cover_text(document: Document) -> str:
    return "".join(
        node.text or ""
        for node in document.element.body.iter(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
        )
    )


def test_cover_placeholders_are_filled_and_the_draft_is_marked(tmp_path: Path) -> None:
    document = _render("# 配电安全专家咨询报告\n\n## 1. 配电评估概述\n\n### 1.1 评估背景\n\n正文。")
    assert re.search(r"X{3,}", _cover_text(document))
    report = SimpleNamespace(
        title="配电安全专家咨询报告（草稿）",
        client_name="示例电力有限公司",
        site_name="示例市东区",
    )

    PdsDocxRenderer._fill_cover(document, report)
    PdsDocxRenderer._mark_draft_headers(document, "草稿：资料待补。")
    output = _saved(document, tmp_path / "cover.docx")

    text = _cover_text(Document(output))
    assert "示例电力有限公司" in text and "示例市东区" in text
    assert "配电安全专家咨询报告（草稿）" in text
    assert not re.search(r"X{3,}|〈[^〉]*〉", text)
    with zipfile.ZipFile(output) as archive:
        headers = [
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if re.fullmatch(r"word/header\d*\.xml", name)
        ]
    assert headers and all("草稿：资料待补。" in header for header in headers)


def test_unknown_cover_facts_say_missing() -> None:
    document = _render("# 配电安全专家咨询报告\n\n## 1. 配电评估概述\n\n正文。")

    PdsDocxRenderer._fill_cover(
        document, SimpleNamespace(title="配电安全专家咨询报告", client_name=None, site_name=None)
    )

    text = _cover_text(document)
    assert not re.search(r"X{3,}", text)
    assert text.count("待补") >= 2


def test_cover_facts_come_from_labelled_project_evidence() -> None:
    evidence = [
        {"subject": "设备1", "fact": "客户名称=别的表", "source": {"sheet": "低配评估详情"}},
        {
            "subject": "基本信息",
            "fact": "客户名称=示例电力有限公司；项目地点=示例市东区",
            "source": {"sheet": "收资表"},
        },
    ]

    facts = report_cover_facts(evidence)

    assert facts.client_name == "示例电力有限公司"
    assert facts.site_name == "示例市东区"
    missing = report_cover_facts(
        [{"subject": "x", "fact": "客户名称=XXXX", "source": {"sheet": "收资表"}}]
    )
    assert (missing.client_name, missing.site_name) == ("待补", "待补")


@pytest.mark.parametrize(
    ("policy", "has_gap", "draft"),
    [("draft", True, True), ("draft", False, False), ("skip", True, False)],
)
def test_draft_means_draft_policy_and_registered_gaps(
    tmp_path: Path, policy: str, has_gap: bool, draft: bool
) -> None:
    if has_gap:
        gap = tmp_path / "Work/runs/run-1/gaps/GAP-J-1.json"
        gap.parent.mkdir(parents=True)
        gap.write_text(json.dumps({"gap": "柜1：缺测温", "task_id": "module-2.4"}), "utf-8")
    state = {"run_id": "run-1", "request": {"missing_evidence_policy": policy}}

    assert report_is_draft(tmp_path, state) is draft


def test_renderer_keeps_text_after_lines_the_legacy_normalizer_cut_at() -> None:
    document = _render(
        "\n\n".join([
            "# 配电安全专家咨询报告",
            "## 1. 配电评估概述",
            "### 1.1 评估背景",
            "如有疑问，请联系评估团队。",
            "报告编制说明：本报告依据现场资料编写。",
            "之后的正文仍需保留。",
        ])
    )

    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "之后的正文仍需保留。" in texts
    assert any("如有疑问" in text for text in texts)


def test_photo_parts_are_named_after_their_image_format(tmp_path: Path) -> None:
    paths = {}
    for photo_id, name, image_format in (
        ("P-0001", "P-0001", "PNG"),  # no extension
        ("P-0002", "P-0002.", "JPEG"),  # trailing dot
        ("P-0003", "P-0003.png", "JPEG"),  # wrong extension
    ):
        path = tmp_path / name
        Image.new("RGB", (20, 10), "blue").save(path, format=image_format)
        paths[photo_id] = path
    rows = " ".join(f"[[PHOTO:{photo_id}]]" for photo_id in paths)
    document = _render(
        "\n\n".join([
            "# 报告",
            "## 2. 评估内容描述",
            f"| 问题 | 现场图片 |\n| --- | --- |\n| 温升 | {rows} |",
        ])
    )
    report = SimpleNamespace(
        photos=[SimpleNamespace(id=photo_id, path=path) for photo_id, path in paths.items()]
    )

    PdsDocxRenderer._materialize_photos(document, report)
    output = _saved(document, tmp_path / "photos.docx")

    with zipfile.ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) >= 3
    assert_media_parts_are_named_by_type(output)


def test_unverified_render_is_visible_in_the_delivery_result(tmp_path: Path) -> None:
    ref = Path("Work/runs/run-1/render-result.json")
    (tmp_path / ref).parent.mkdir(parents=True)
    (tmp_path / ref).write_text(
        json.dumps({
            "protected_prose_verified": False,
            "validation_warnings": ["rendered DOCX may have omitted approved Markdown content"],
        }),
        encoding="utf-8",
    )
    tools = DeliveryTools(tmp_path, ReportingStore(tmp_path))

    verification = tools._render_verification(SimpleNamespace(render_result_ref=ref))

    assert verification == {
        "protected_prose_verified": False,
        "validation_warnings": ["rendered DOCX may have omitted approved Markdown content"],
        "render_result_ref": "Work/runs/run-1/render-result.json",
    }


def test_word_tables_keep_every_markdown_row() -> None:
    from manyselves.capabilities.distribution_reporting.runtime.rendering.v2_docx_renderer import (
        write_markdown_table,
    )

    document = Document()
    lines = ["| 对象 | 问题 |", "| --- | --- |", *(f"| {n}#柜 | 问题{n} |" for n in range(1, 36))]
    write_markdown_table(document, document.add_paragraph(), lines)

    (table,) = document.tables
    assert len(table.rows) == 36
    assert table.rows[-1].cells[0].text == "35#柜"


def test_tables_apart_in_markdown_stay_apart_in_word() -> None:
    from manyselves.capabilities.distribution_reporting.runtime.rendering.v2_docx_renderer import (
        add_report_content,
    )

    document = Document()
    add_report_content(
        document,
        "总体评价。\n\n| 维度 | 状态 |\n| --- | --- |\n| 环境 | 良 |\n\n"
        "| 优先级 | 风险 |\n| --- | --- |\n| 1 | 过热 |\n",
    )

    assert [len(table.rows) for table in document.tables] == [2, 2]
