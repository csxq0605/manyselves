"""Photos survive missing or inconsistent cells in the fixed S4-4 structure."""

from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from PIL import Image

from tests.capabilities.distribution_reporting.test_intake_fidelity import _normalize

_IMAGES = (
    '<images xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<xdr:pic><xdr:cNvPr name="thermal-1"/><a:blip r:embed="rId1"/></xdr:pic></images>'
)
_RELS = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Target="media/thermal.png"/></Relationships>'
)


def _thermal_book(tmp_path, *, reading, photo_id="thermal-1", with_media=True):
    path = tmp_path / "S4-4.xlsx"
    sheet = Workbook().active
    sheet.title = "红外热成像检测记录表"
    for column, value in {1: "A房", 2: 20, 4: "柜A", 6: "进线端子", 7: reading}.items():
        sheet.cell(3, column, value)
    sheet["K3"] = f'=_xlfn.DISPIMG("{photo_id}",1)'
    sheet.parent.save(path)
    if with_media:
        image = BytesIO()
        Image.new("RGB", (2, 2)).save(image, format="PNG")
        with ZipFile(path, "a") as archive:
            archive.writestr("xl/cellimages.xml", _IMAGES)
            archive.writestr("xl/_rels/cellimages.xml.rels", _RELS)
            archive.writestr("xl/media/thermal.png", image.getvalue())
    return path


def test_shared_thermal_photo_sits_with_the_owning_submodule(tmp_path):
    context = _normalize(tmp_path, _thermal_book(tmp_path, reading=35))
    [photo] = context.photo_assets
    primary = next(item for item in context.evidence_items if item.id == photo.primary_evidence_id)
    # domain/observation-ownership.yaml: infrared temperature is owned by 2.2.2.1.
    assert primary.submodule_id == "2.2.2.1"
    assert {item.submodule_id for item in context.evidence_items if photo.id in item.photo_refs} == {
        "2.2.2.1", "2.4.4",
    }


def test_photo_of_a_checkpoint_without_reading_is_kept_and_flagged(tmp_path):
    context = _normalize(tmp_path, _thermal_book(tmp_path, reading=None))
    [photo] = context.photo_assets
    records = [item for item in context.evidence_items if photo.id in item.photo_refs]
    assert {item.submodule_id for item in records} == {"2.2.2.1", "2.4.4"}
    assert all("仅有图片" in item.fact and item.needs_confirmation for item in records)
    assert any(gap["code"] == "missing_thermal_reading" for gap in context.mapping_gaps)


def test_missing_photo_media_is_a_gap_not_a_failed_table(tmp_path):
    context = _normalize(tmp_path, _thermal_book(tmp_path, reading=35, with_media=False))
    assert not context.photo_assets
    assert not any(gap["code"] == "file_parse_failed" for gap in context.mapping_gaps)
    assert len(context.evidence_items) == 2
    assert all(
        not item.photo_refs and "图片无法从文件中取出" in item.fact and item.needs_confirmation
        for item in context.evidence_items
    )
    assert {gap["code"] for gap in context.mapping_gaps} >= {"photo_media_missing"}


def test_photo_without_any_checkpoint_is_reported(tmp_path):
    path = tmp_path / "S4-4.xlsx"
    sheet = Workbook().active
    sheet.title = "红外热成像检测记录表"
    sheet["K9"] = '=_xlfn.DISPIMG("stray",1)'
    sheet.parent.save(path)
    context = _normalize(tmp_path, path)
    assert not context.evidence_items
    assert [(gap["code"], gap["cell"]) for gap in context.mapping_gaps if gap["code"].startswith("photo")] == [
        ("photo_without_record", "K9")
    ]
