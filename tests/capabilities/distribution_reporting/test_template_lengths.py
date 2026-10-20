"""Expert template section lengths are read by heading position and title."""

from pathlib import Path

from docx import Document

from manyselves.capabilities.distribution_reporting.runtime.template_lengths import (
    matched_lengths,
    render_length_reference,
    template_sections,
)


def test_sections_are_numbered_by_position_and_counted_with_tables(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("概述", level=1)
    document.add_heading("评估背景", level=2)
    document.add_paragraph("本次评估覆盖五个维度。")
    document.add_heading("评估内容", level=1)
    document.add_heading("", level=2)  # an empty heading is not counted
    document.add_heading("系统架构问题", level=2)
    document.add_heading("负荷分配与过载风险", level=3)
    document.add_paragraph("负荷率偏高。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "对象", "结论"
    path = tmp_path / "t.docx"
    document.save(path)

    sections = template_sections(path)
    assert [(s.number, s.title) for s in sections] == [
        ("1", "概述"), ("1.1", "评估背景"), ("2", "评估内容"),
        ("2.1", "系统架构问题"), ("2.1.1", "负荷分配与过载风险"),
    ]
    titles = {"1.1": "评估背景", "2.1.1": "配电系统负荷分配与过载风险", "2.1.2": "无此节"}
    lengths = matched_lengths(sections, titles.items())
    assert lengths["1.1"].chars == 10
    assert (lengths["2.1.1"].chars, lengths["2.1.1"].tables) == (9, 1)
    assert "2.1.2" not in lengths
    assert "2.1.1 配电系统负荷分配与过载风险：约 9 字，1 张表" in render_length_reference(
        lengths, titles
    )
