"""Synthetic intake regressions: preserve observations, surface unreadable data."""

import re
import sys
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from manyselves.capabilities.distribution_reporting.runtime.intake.adapters import (
    IntakeAdapterRegistry,
)
from manyselves.capabilities.distribution_reporting.runtime.intake.manifest import _purpose_for
from manyselves.capabilities.distribution_reporting.runtime.mappers.s4_4 import _map_thermal
from manyselves.capabilities.distribution_reporting.runtime.models.preparation import (
    ManifestFile,
    PreparationContext,
    ProjectManifest,
)
from manyselves.capabilities.distribution_reporting.runtime.models.reporting import ReportRequest
from manyselves.capabilities.distribution_reporting.runtime.preparation import prepare_manifest_file
from manyselves.capabilities.distribution_reporting.runtime.preparation_tools import (
    PreparationTools,
)
from manyselves.capabilities.distribution_reporting.runtime.storage import ReportingStore


def _manifest(path):
    return ManifestFile(id="fixture", path=path, sha256="a" * 64, media_type="test", purpose=_purpose_for(path))


def _normalize(tmp_path, path):
    manifest = _manifest(path)
    worker = prepare_manifest_file(tmp_path, "fidelity", manifest, 0)
    context = PreparationContext(
        run_id="fidelity",
        request=ReportRequest(operation="module_report", instruction="test", target_modules=["2.4"]),
        project_manifest=ProjectManifest(files=[manifest]),
        preparation_worker_results=[worker], parsed_artifacts=worker.parsed_artifacts,
    )
    tools = PreparationTools(
        workspace=tmp_path, input_snapshot=lambda _: None,
        snapshot_content=lambda source, target: (target, "unused", target),
        runtime_photo_ids=lambda *_: [], store=ReportingStore(tmp_path),
    )
    return tools.evaluate_coverage(tools.normalize_evidence(context))


def test_docx_table_reaches_evidence_with_source_coordinate(tmp_path):
    path = tmp_path / "observations.docx"
    document = Document()
    document.add_paragraph("现场记录")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "设备", "电流"
    table.cell(1, 0).text, table.cell(1, 1).text = "柜A", "11.3 A"
    document.save(path)
    context = _normalize(tmp_path, path)
    evidence = next(item for item in context.evidence_items if "11.3 A" in item.fact)
    assert evidence.source.cell == "table:1:row:2"
    assert "柜A" in evidence.fact


def test_workbook_preserves_title_headers_coordinates_and_missing_formula(tmp_path):
    path = tmp_path / "measurements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["测量记录"])
    sheet.merge_cells("A1:B1")
    sheet.append(["设备", "计算值"])
    sheet.append(["柜A", "=1+2"])
    workbook.save(path)
    context = _normalize(tmp_path, path)
    facts = "\n".join(item.fact for item in context.evidence_items)
    assert "A1=测量记录" in facts
    assert "A2=设备" in facts and "B2=计算值" in facts
    assert "测量记录=柜A" not in facts
    gap = next(gap for gap in context.mapping_gaps if gap.get("reason") == "formula_result_unavailable")
    assert gap["cell"] == "B3" and gap["sheet"] == sheet.title
    assert gap["formula"] == "=1+2"
    assert context.coverage_matrix.entries["2.4"].gaps
    artifacts = IntakeAdapterRegistry().parse(path, _manifest(path))
    title = next(a for a in artifacts if a.source.cell == "A1:B1")
    assert title.payload["merged_ranges"] == ["A1:B1"]


def test_empty_pdf_page_is_a_located_gap_without_dropping_text_pages(tmp_path, monkeypatch):
    pages = [SimpleNamespace(extract_text=lambda: "可读内容"), SimpleNamespace(extract_text=lambda: "  ")]
    monkeypatch.setitem(sys.modules, "pypdf", SimpleNamespace(PdfReader=lambda _: SimpleNamespace(pages=pages)))
    context = _normalize(tmp_path, tmp_path / "mixed.pdf")
    assert any(item.fact == "可读内容" for item in context.evidence_items)
    gap = next(gap for gap in context.mapping_gaps if gap.get("reason") == "pdf_page_without_text")
    assert gap["page"] == 2
    assert any("mixed.pdf" in message for message in context.coverage_matrix.entries["2.4"].gaps)


def test_corrupt_file_surfaces_in_coverage(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_text("not a zip", encoding="utf-8")
    context = _normalize(tmp_path, path)
    gap = next(gap for gap in context.mapping_gaps if gap.get("code") == "file_parse_failed")
    assert gap["file_id"] == "fixture"
    assert "broken.docx" in gap["message"]
    assert gap["message"] in context.coverage_matrix.entries["2.4"].gaps


def test_mapper_filename_requires_workbook_suffix():
    assert _purpose_for(Path("S4-4-notes.docx")) is None
    assert _purpose_for(Path("S4-4-data.xlsx")) == "s4-4"


def test_s4_4_formula_without_cache_is_not_silently_skipped(tmp_path):
    path = tmp_path / "S4-4.xlsx"
    sheet = Workbook().active
    sheet.title = "红外热成像检测记录表"
    for column, value in {1: "A房", 2: 20, 4: "柜A", 6: "端子", 7: "=30+5"}.items():
        sheet.cell(3, column, value)
    # Image formula is handled by the separate image binding path, not numeric calculation.
    sheet["K8"] = '=_xlfn.DISPIMG("fixture",1)'
    sheet.parent.save(path)
    context = _normalize(tmp_path, path)
    gaps = [gap for gap in context.mapping_gaps if gap.get("code") == "formula_result_unavailable"]
    assert len(gaps) == 1 and gaps[0]["cell"] == "G3"
    assert "=30+5" in gaps[0]["message"]
    assert not context.evidence_items


def test_cached_formula_zero_is_a_value_and_formula_error_is_a_gap(tmp_path):
    from zipfile import ZipFile

    path = tmp_path / "cached.xlsx"
    sheet = Workbook().active
    sheet.append(["设备", "数值"])
    sheet.append(["柜A", "=1-1"])
    sheet.append(["柜B", "=1/0"])
    sheet.parent.save(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    key = "xl/worksheets/sheet1.xml"
    entries[key] = entries[key].replace(b"<f>1-1</f><v></v>", b"<f>1-1</f><v>0</v>")
    entries[key] = entries[key].replace(
        b'<c r="B3"><f>1/0</f><v></v></c>', b'<c r="B3" t="e"><f>1/0</f><v>#DIV/0!</v></c>'
    )
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    context = _normalize(tmp_path, path)
    assert any("B2=0" in item.fact for item in context.evidence_items)
    assert not any(gap.get("cell") == "B2" for gap in context.mapping_gaps)
    assert any(gap.get("reason") == "formula_result_error" and gap["cell"] == "B3"
               for gap in context.mapping_gaps)


@pytest.mark.parametrize("dimension", [None, "A1:A1", "A1:Z100"])
@pytest.mark.parametrize("specialized", [False, True])
def test_workbook_dimension_metadata_cannot_drop_observations_or_formula_gaps(
    tmp_path, dimension, specialized,
):
    path = tmp_path / ("S4-4.xlsx" if specialized else "measurements.xlsx")
    sheet = Workbook().active
    sheet.title = "红外热成像检测记录表"
    sheet.append(["标题"])
    sheet.merge_cells("A1:B1")
    for row, value in ((3, "=30+5"), (4, "=1+2"), (5, "=1/0")):
        for column, item in {1: "A房", 2: 20, 4: f"柜{row}", 6: "端子", 7: value}.items():
            sheet.cell(row, column, item)
    sheet.parent.save(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    key = "xl/worksheets/sheet1.xml"
    entries[key] = entries[key].replace(b"<f>30+5</f><v></v>", b"<f>30+5</f><v>35</v>")
    entries[key] = entries[key].replace(
        b'<c r="G5"><f>1/0</f><v></v></c>',
        b'<c r="G5" t="e"><f>1/0</f><v>#DIV/0!</v></c>',
    )
    entries[key] = re.sub(
        rb'<dimension\b[^>]*/>',
        b"" if dimension is None else f'<dimension ref="{dimension}"/>'.encode(),
        entries[key],
    )
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    context = _normalize(tmp_path, path)
    assert not any(gap.get("code") == "file_parse_failed" for gap in context.mapping_gaps)
    assert context.evidence_items
    if specialized:
        measured = [m.value for e in context.evidence_items for m in e.measurements
                    if m.quantity == "measured_temperature"]
        assert measured == [35.0, 35.0]
    else:
        assert any("G3=35" in item.fact for item in context.evidence_items)
        artifacts = IntakeAdapterRegistry().parse(path, _manifest(path))
        assert any(a.payload.get("merged_ranges") == ["A1:B1"] for a in artifacts)
    gaps = {(gap.get("code", gap.get("reason")), gap.get("cell"))
            for gap in context.mapping_gaps}
    assert ("formula_result_unavailable", "G4") in gaps
    assert ("formula_result_error", "G5") in gaps


def test_thermal_uses_only_real_merged_ranges_even_across_room_change():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "红外"
    for row, room, temp in [(3, "A房", 40), (4, "B房", 45), (5, "C房", 55)]:
        sheet.cell(row, 1, room)
        sheet.cell(row, 4, f"柜{row}")
        sheet.cell(row, 6, "端子")
        sheet.cell(row, 7, temp)
    sheet["B3"], sheet["C3"] = 20, 50
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:C4")
    evidence, gaps = _map_thermal(Path("fixture.xlsx"), "fixture", sheet, sheet)
    second = next(item for item in evidence if item.source.row == 4)
    assert "B房" in second.subject and "相对环境温升=25.0K" in second.fact
    assert second.source.cell == "A3:K4"
    third = [item for item in evidence if item.source.row == 5]
    assert len(third) == 2
    assert all("温升" not in item.fact and "环境湿度" not in item.fact for item in third)
    assert all(item.needs_confirmation for item in third)
    assert any(gap.code == "missing_thermal_reference" and gap.cell == "G5" for gap in gaps)


def test_thermal_merged_subject_does_not_leak_outside_range():
    sheet = Workbook().active
    sheet["A3"], sheet["D3"], sheet["B3"] = "A房", "柜A", 20
    for column in ("A", "B", "D"):
        sheet.merge_cells(f"{column}3:{column}4")
    for row in (3, 4, 5):
        sheet.cell(row, 6, "端子")
        sheet.cell(row, 7, 40)
    evidence, gaps = _map_thermal(Path("fixture.xlsx"), "fixture", sheet, sheet)
    assert all(item.subject == "A房/柜A" for item in evidence if item.source.row == 4)
    assert all("柜A" not in item.subject for item in evidence if item.source.row == 5)
    assert any(gap.code == "incomplete_thermal_context" for gap in gaps)


def test_thermal_typed_quantities_survive_normalization_and_fact_delivery(tmp_path):
    from manyselves.capabilities.distribution_reporting.runtime.research.fact_context import (
        render_fact_packet,
    )

    path = tmp_path / "S4-4.xlsx"
    sheet = Workbook().active
    sheet.title = "红外热成像检测记录表"
    for row, ambient in ((3, 20.2), (4, None), (5, 0)):
        for column, value in {1: "A房", 2: ambient, 4: "柜A", 6: "端子", 7: 55.3}.items():
            sheet.cell(row, column, value)
    sheet.parent.save(path)
    context = _normalize(tmp_path, path)
    for row, expected_rise in ((3, 35.1), (4, None), (5, 55.3)):
        routed = [item for item in context.evidence_items if item.source.row == row]
        assert len(routed) == 2
        assert routed[0].measurements == routed[1].measurements
        measurements = {value.quantity: value for value in routed[0].measurements}
        measured = measurements["measured_temperature"]
        ambient = measurements["ambient_temperature"]
        rise = measurements["temperature_rise"]
        assert measured.value == 55.3 and measured.source_cell == f"G{row}"
        assert ambient.source_cell == f"B{row}"
        assert rise.value == expected_rise and rise.unit == "K"
        assert rise.derived_from == ["measured_temperature", "ambient_temperature"]
        assert bool(rise.missing_reason) == (row == 4)
        packet = render_fact_packet(routed, "2.4")
        assert '"quantity":"measured_temperature"' in packet
        assert '"quantity":"temperature_rise"' in packet
        if row == 4:
            assert "missing_ambient_temperature" in packet
