"""Post-render validation must report content defects without blocking publication."""

from pathlib import Path

import pytest
from docx import Document

from manyselves.capabilities.distribution_reporting.runtime.rendering.rendered_docx_validator import (
    validate_rendered_markdown_docx,
)


def test_openable_docx_returns_content_title_and_token_warnings(tmp_path: Path) -> None:
    output = tmp_path / "warning.docx"
    document = Document()
    document.add_paragraph("Different title")
    document.add_paragraph("[[CITE:7]]")
    document.save(output)

    warnings = validate_rendered_markdown_docx(
        output,
        "# Approved title\n\nApproved body.\n",
        expected_title="Approved title",
        reject_unresolved_tokens=True,
    )

    assert any("omitted approved Markdown content" in warning for warning in warnings)
    assert "rendered DOCX is missing the approved title" in warnings
    assert any("unresolved content tokens" in warning for warning in warnings)


def test_unopenable_docx_remains_a_hard_failure(tmp_path: Path) -> None:
    output = tmp_path / "broken.docx"
    output.write_bytes(b"not a docx")

    with pytest.raises(ValueError, match="not Word/WPS-openable"):
        validate_rendered_markdown_docx(output, "# Approved\n")
