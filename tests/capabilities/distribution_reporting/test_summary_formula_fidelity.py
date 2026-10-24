"""Formula results must reach specialized intake without treating expressions as facts."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from PIL import Image

from tests.capabilities.distribution_reporting.test_intake_fidelity import _normalize

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _cache(path, cells, dimension="A1:Z100"):
    """Supply Excel-saved results; openpyxl itself does not calculate formulas."""
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    for name, data in list(entries.items()):
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        tree = ET.fromstring(data)
        declared = tree.find(f"{_NS}dimension")
        if declared is not None:
            if dimension is None:
                tree.remove(declared)
            else:
                declared.set("ref", dimension)
        for cell in tree.iter(f"{_NS}c"):
            key = (name, cell.attrib["r"])
            if key not in cells:
                continue
            kind, result = cells[key]
            cell.set("t", kind)
            value = cell.find(f"{_NS}v")
            if value is None:
                value = ET.SubElement(cell, f"{_NS}v")
            value.text = result
        entries[name] = ET.tostring(tree, encoding="utf-8")
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


def _sheet1(cells):
    return {("xl/worksheets/sheet1.xml", cell): value for cell, value in cells.items()}


@pytest.mark.parametrize("dimension", [None, "A1:A1", "A1:Z100"])
def test_collection_cached_route_status_zero_and_false_reach_normalized_facts(tmp_path, dimension):
    path = tmp_path / "S2-1.xlsx"
    sheet = Workbook().active
    sheet.title = "收资表"
    sheet["A2"], sheet["B2"] = "保护资料", '="保护定值表"'
    sheet["G2"], sheet["H2"], sheet["I2"] = '="OK"', '="OK"', "=1-1"
    sheet["B3"], sheet["G3"], sheet["H3"] = "培训记录", "=1-1", "=FALSE()"
    sheet.parent.save(path)
    _cache(
        path,
        _sheet1(
            {
                "B2": ("str", "保护定值表"),
                "G2": ("str", "OK"),
                "H2": ("str", "OK"),
                "I2": ("n", "0"),
                "G3": ("n", "0"),
                "H3": ("b", "0"),
            }
        ),
        dimension,
    )
    original = path.read_bytes()
    context = _normalize(tmp_path, path)
    assert path.read_bytes() == original
    assert len(context.evidence_items) == 2
    first, second = context.evidence_items
    assert first.submodule_id == "2.3.1" and first.source.cell == "G2:H2"
    assert first.fact == "具备情况=OK；有效性=OK；备注=0"
    assert not first.needs_confirmation
    assert second.fact == "具备情况=0；有效性=False"
    assert second.needs_confirmation
    assert not context.mapping_gaps


def test_collection_missing_and_error_results_stay_unknown_with_coordinates(tmp_path):
    path = tmp_path / "S2-1.xlsx"
    sheet = Workbook().active
    sheet.title = "收资表"
    sheet["B2"], sheet["G2"], sheet["H2"] = "保护定值表", "=1+2", "=1/0"
    sheet.parent.save(path)
    _cache(path, _sheet1({"H2": ("e", "#DIV/0!")}))
    context = _normalize(tmp_path, path)
    assert context.evidence_items[0].fact == "具备情况=未填写；有效性=未填写"
    gaps = {g["cell"]: g for g in context.mapping_gaps}
    assert gaps["G2"]["code"] == "formula_result_unavailable"
    assert gaps["H2"]["code"] == "formula_result_error"
    assert "=1+2" in gaps["G2"]["message"] and "#DIV/0!" in gaps["H2"]["message"]
    assert context.coverage_matrix.entries["2.4"].gaps


def test_unknown_category_does_not_inherit_previous_and_missing_note_needs_review(tmp_path):
    path = tmp_path / "S2-1.xlsx"
    sheet = Workbook().active
    sheet.title = "收资表"
    sheet["A2"], sheet["B2"], sheet["G2"], sheet["H2"] = "上一类别", "保护定值表", "OK", "OK"
    sheet["A3"], sheet["B3"], sheet["G3"], sheet["H3"] = '="新类别"', "培训记录", "OK", "OK"
    sheet["I3"] = '="待复核备注"'
    sheet.parent.save(path)
    context = _normalize(tmp_path, path)
    item = context.evidence_items[1]
    assert item.subject == "未分类/培训记录"
    assert item.needs_confirmation
    assert {g["cell"] for g in context.mapping_gaps} == {"A3", "I3"}


@pytest.mark.parametrize("dimension", [None, "A1:A1", "A1:Z100"])
def test_assessment_cached_text_routes_and_preserves_photo_formulas(tmp_path, dimension):
    path = tmp_path / "S4-6.xlsx"
    book = Workbook()
    summary = book.active
    summary.title = "评估信息汇总表"
    summary["B4"], summary["F4"], summary["G4"] = '="2.3.1"', '="需复核配合"', '="核查定值"'
    conclusions = book.create_sheet("结论建议汇总表")
    conclusions["B2"], conclusions["C2"] = '="额定/分断能力"', '="开关需校核"'
    conclusions["D2"], conclusions["E2"] = "=1-1", "=FALSE()"
    conclusions["F2"] = '=_xlfn.DISPIMG("photo-fixture",1)'
    # A picture/header row must not become a narrative or calculation gap.
    conclusions["B3"], conclusions["C3"] = '=DISPIMG("header-fixture",1)', "图片说明"
    book.save(path)
    cells = _sheet1(
        {"B4": ("str", "2.3.1"), "F4": ("str", "需复核配合"), "G4": ("str", "核查定值")}
    )
    cells.update(
        {
            ("xl/worksheets/sheet2.xml", c): v
            for c, v in {
                "B2": ("str", "额定/分断能力"),
                "C2": ("str", "开关需校核"),
                "D2": ("n", "0"),
                "E2": ("b", "0"),
            }.items()
        }
    )
    _cache(path, cells, dimension)
    image = BytesIO()
    Image.new("RGB", (1, 1)).save(image, format="PNG")
    with ZipFile(path, "a") as archive:
        archive.writestr(
            "xl/cellimages.xml",
            '<images xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><xdr:pic><xdr:cNvPr name="photo-fixture"/><a:blip r:embed="rId1"/></xdr:pic></images>',
        )
        archive.writestr(
            "xl/_rels/cellimages.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="media/fixture.png"/></Relationships>',
        )
        archive.writestr("xl/media/fixture.png", image.getvalue())
    original = path.read_bytes()
    context = _normalize(tmp_path, path)
    assert path.read_bytes() == original
    assert len(context.evidence_items) == 2
    first, second = context.evidence_items
    assert first.submodule_id == "2.3.1"
    assert first.fact == "既有评估结论=需复核配合；既有行动建议=核查定值"
    assert second.submodule_id == "2.4.1.1"
    assert second.fact == "既有问题=开关需校核；既有风险=0；既有建议=False"
    assert second.photo_refs == ["P-0001"]
    assert all(item.needs_confirmation for item in context.evidence_items)
    # The header picture is accounted for as an unbound photo, never as a
    # narrative or calculation gap.
    assert [(g["code"], g["cell"]) for g in context.mapping_gaps] == [
        ("photo_without_record", "B3")
    ]


def test_assessment_missing_formula_is_a_gap_even_when_route_or_whole_row_is_missing(tmp_path):
    path = tmp_path / "S4-6.xlsx"
    sheet = Workbook().active
    sheet.title = "评估信息汇总表"
    sheet["B4"], sheet["F4"] = '="2.3.1"', "尚需来源定位"
    sheet["B5"], sheet["F5"], sheet["G5"] = "2.3.1", '="未知结论"', "=1/0"
    sheet["B6"], sheet["F6"], sheet["G6"] = "2.3.1", "已知评估", '="未知建议"'
    sheet.parent.save(path)
    _cache(path, _sheet1({"G5": ("e", "#DIV/0!")}))
    context = _normalize(tmp_path, path)
    assert len(context.evidence_items) == 1
    assert context.evidence_items[0].fact == "既有评估结论=已知评估"
    formula_gaps = {
        g["cell"]: g["code"] for g in context.mapping_gaps if g["code"].startswith("formula_")
    }
    assert formula_gaps == {
        "B4": "formula_result_unavailable",
        "F5": "formula_result_unavailable",
        "G5": "formula_result_error",
        "G6": "formula_result_unavailable",
    }
    assert any(
        g["code"] == "empty_summary_assessment" and g["cell"] == "F5:G5"
        for g in context.mapping_gaps
    )
