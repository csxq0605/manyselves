"""Values in the prose are slots the program fills; the check is structural."""

from pathlib import Path

from manyselves.capabilities.distribution_reporting.domain.calculations import statistics
from manyselves.capabilities.distribution_reporting.domain.value_slots import (
    fill_slots,
    mispaired_slots,
    object_names,
    quote_findings,
    slot_findings,
    slot_table,
)
from manyselves.capabilities.distribution_reporting.runtime.models.reporting import (
    EvidenceItem,
    EvidenceMeasurement,
    SourceLocation,
)


def _item(n, subject, value, *, cell=None, submodule="2.1.1", label="负荷率"):
    return EvidenceItem(
        id=f"E-{n}", subject=subject, fact=f"运行电流=100A；原表负荷率={value}%",
        value=value, unit="%", checkpoint_field=label,
        measurements=[
            EvidenceMeasurement(quantity="running_current", value=100, unit="A",
                                label="运行电流", source_cell="C3"),
            EvidenceMeasurement(quantity="load_rate", value=value, unit="%",
                                label=label, source_cell="E3"),
        ],
        source=SourceLocation(file_id="f", path=Path("Inputs/S4-4.xlsx"), sheet="低配",
                              cell=cell or f"C{n + 10}:E{n + 10}", row=n, column="C"),
    )


def test_a_cell_routed_to_two_sections_is_counted_once_under_its_declared_column() -> None:
    items = [
        _item(1, "A房/1#", 80.0, cell="C3:E3"),
        _item(2, "A房/1#", 80.0, cell="C3:E3", submodule="2.4.4"),  # same cell, other section
        _item(3, "B房/2#", 40.0, cell="C4:E4"),
    ]
    calculations, lines, lists = statistics(items)
    counts = [c for c in calculations if c.quantity == "object_count"]
    assert [c.value for c in counts] == [2]
    # The group is named by the header of the column the value comes from.
    assert lines[0].startswith("负荷率（低配 E列）")
    (spec,) = lists.values()
    assert [entry["subject"] for entry in spec["entries"]] == ["A房/1#", "B房/2#"]
    # Every copy stays traceable.
    assert set(counts[0].evidence_ids) == {"E-1", "E-2", "E-3"}


def test_slots_fill_values_with_units_and_lists() -> None:
    items = [_item(1, "A房/1#", 80.0), _item(2, "B房/2#", 40.0)]
    calculations, _lines, lists = statistics(items)
    slots = slot_table(items, calculations, lists)
    value_slot = next(k for k, v in slots.items() if v.get("subject") == "A房/1#" and v["label"] == "负荷率")
    list_slot = next(iter(lists))
    text = f"A房/1#负荷率〔{value_slot}〕；各柜〔{list_slot}〕。"
    assert fill_slots(text, slots) == "A房/1#负荷率80%；各柜A房/1#80%、B房/2#40%。"
    assert slot_findings(text, slots, names=object_names(["A房/1#", "B房/2#"])) == []


def test_bare_numerals_and_unknown_slots_are_found_but_known_digit_texts_pass() -> None:
    slots = {"V-0001": {"text": "80%", "subject": "A房/1#", "label": "负荷率", "evidence_ids": ["E-1"]}}
    names = object_names(["A房/1#"])
    found = slot_findings("A房/1#负荷率 85%，见〔V-0009〕。", slots, names=names)
    assert any("V-0009" in p for p in found) and any("85%" in p for p in found)
    ok = "1. A房/1#负荷率〔V-0001〕（详见 2.1.1），建议 1 个月内复测，可参照 GB 50054-2011。"
    assert slot_findings(
        ok, slots, names=names, section_ids=["2.1.1"],
        digit_texts=["1 个月内", "GB 50054-2011"],
    ) == []


def test_a_count_of_picked_objects_is_its_own_slot() -> None:
    slots = {
        "V-0001": {"text": "0.8A", "subject": "1#", "label": "剩余电流", "evidence_ids": ["E-1"]},
        "V-0002": {"text": "1.0A", "subject": "2#", "label": "剩余电流", "evidence_ids": ["E-2"]},
    }
    text = "所列〔#:V-0001、V-0002〕个回路低于 1.1A 的筛选口径。"
    assert fill_slots(text, slots).startswith("所列2个回路")
    # The threshold 1.1A is a numeral outside a slot.
    assert any("1.1A" in p for p in slot_findings(text, slots))


def test_a_value_slot_next_to_another_object_is_mispaired() -> None:
    slots = {
        "V-0001": {"text": "11.3A", "subject": "A房/2#", "evidence_ids": ["E-1"]},
        "V-0002": {"text": "3.1A", "subject": "B房/3#", "evidence_ids": ["E-2"]},
    }
    assert mispaired_slots("3# 柜剩余电流〔V-0001〕。", slots)
    assert mispaired_slots("2# 柜〔V-0001〕，约为 3# 柜〔V-0002〕的几倍。", slots) == []


def test_quoted_words_must_be_the_recorded_words() -> None:
    sources = ["运维制度=不完善；评语：SOP 未覆盖倒闸操作"]
    assert quote_findings("原表评语为「SOP 未覆盖倒闸操作」。", sources) == []
    assert quote_findings("原表写明「未成文」。", sources) == ["「未成文」"]


def test_a_columns_unit_and_name_come_from_its_own_header() -> None:
    from openpyxl import Workbook

    from manyselves.capabilities.distribution_reporting.runtime.mappers.common import (
        header_label,
        header_unit,
    )

    sheet = Workbook().active
    sheet.cell(2, 3, "运行电流(A)")
    sheet.cell(2, 5, "负荷率")
    assert (header_label(sheet, 2, 3, "x"), header_unit(sheet, 2, 3, "mA")) == ("运行电流", "A")
    # A header that declares no unit keeps the mapper's.
    assert (header_label(sheet, 2, 5, "x"), header_unit(sheet, 2, 5, "%")) == ("负荷率", "%")
