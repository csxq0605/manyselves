import json

from docx import Document

from manyselves.capabilities.distribution_reporting.runtime.rendering.source_index_docx_renderer import (
    SourceIndexDocxRenderer,
)
from manyselves.capabilities.distribution_reporting.runtime.source_index_sections import (
    judgment_and_gap_index_markdown,
)


def _write_gap(workspace, run_id, name, payload):
    path = workspace / f"Work/runs/{run_id}/gaps/{name}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def test_index_lists_judgments_and_deduplicated_gaps_as_word_tables(tmp_path):
    gap = {
        "task_id": "module-2.1",
        "gap": "未提供转供方案",
        "impact": "无法判断N-1转供能力",
        "requested_input": "提供转供操作票",
    }
    _write_gap(tmp_path, "run-1", "GAP-a", gap)
    _write_gap(tmp_path, "run-1", "GAP-b", gap)
    submissions = {
        "2.1": {
            "assessments": {
                "2.1.1": {
                    "records": [
                        {
                            "evidence_ids": ["E-1", "E-2"],
                            "disposition": "undetermined",
                            "scope": "夏季高峰",
                            "basis": "缺少负荷曲线",
                        }
                    ]
                }
            }
        }
    }

    markdown = judgment_and_gap_index_markdown(tmp_path, "run-1", submissions)

    assert "| 2.1.1 | E-1、E-2 | 无法判断 | 夏季高峰 | 缺少负荷曲线 |" in markdown
    assert markdown.count("未提供转供方案") == 1
    assert "| G1 | 未提供转供方案 | 无法判断N-1转供能力 | 提供转供操作票 | 2.1 |" in markdown

    output = tmp_path / "index.docx"
    SourceIndexDocxRenderer.render("## 证据与来源索引\n\n" + markdown, output)
    tables = Document(output).tables
    assert [len(table.rows) for table in tables] == [2, 2]
    assert tables[1].cell(0, 1).text == "缺少的资料"
    assert tables[1].cell(1, 1).text == "未提供转供方案"


def test_index_says_so_when_nothing_was_recorded(tmp_path):
    markdown = judgment_and_gap_index_markdown(tmp_path, "run-1", {})

    assert "本报告没有登记判断记录。" in markdown
    assert "本报告没有登记资料缺口。" in markdown
