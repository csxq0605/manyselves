from pathlib import Path

from openpyxl import Workbook

from manyselves.core.reporting.mappers.s2_1 import map_s2_1


def test_s2_1_maps_availability_and_effectiveness_with_exact_cells(tmp_path: Path) -> None:
    path = tmp_path / "S2-1收资表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收资表"
    sheet.append(
        [
            "类型",
            "文件名称",
            "形式",
            "文件说明",
            "评估内容",
            "评估后输出内容",
            "具备情况(OK/NG)",
            "有效性(OK/NG)",
            "备注",
        ]
    )
    sheet.append(
        [
            "基础图纸与技术参数",
            "主要开关（断路器）技术参数",
            "Excel",
            None,
            None,
            None,
            "OK",
            "NG",
            "已过期",
        ]
    )
    workbook.save(path)

    result = map_s2_1(path, file_id="file-s21")

    item = result.evidence_items[0]
    assert item.module_id == "2.4"
    assert item.submodule_id == "2.4.1.1"
    assert "具备情况=OK" in item.fact
    assert "有效性=NG" in item.fact
    assert item.source.sheet == "收资表"
    assert item.source.cell == "G2:H2"
    assert item.needs_confirmation is True


def test_s2_1_routes_management_and_protection_documents_to_exact_modules(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S2-1收资表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收资表"
    sheet.append(["类型", "文件名称", None, None, None, None, "具备", "有效", "备注"])
    sheet.append(["运维资料", "SOP/EOP操作规程及应急预案", None, None, None, None, "OK", "OK"])
    sheet.append(["技术资料", "继电保护定值计算书", None, None, None, None, "OK", "OK"])
    sheet.append(["运维资料", "关键配电设备维护计划与记录", None, None, None, None, "OK", "OK"])
    workbook.save(path)

    result = map_s2_1(path, file_id="file-s21")

    routes = {
        item.subject.split("/", 1)[1]: (item.module_id, item.submodule_id)
        for item in result.evidence_items
    }
    assert routes["SOP/EOP操作规程及应急预案"] == ("2.5", "2.5.1")
    assert routes["继电保护定值计算书"] == ("2.3", "2.3.1")
    assert routes["关键配电设备维护计划与记录"] == ("2.5", "2.5.3.2")


def test_s2_1_routes_current_template_document_names_without_gaps(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S2-1收资表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收资表"
    sheet.append(["类型", "文件名称", None, None, None, None, "具备", "有效", "备注"])
    expected = {
        "竣工电气单线及控制原理图": ("2.5", "2.5.2"),
        "母线统计表": ("2.1", "2.1.1"),
        "电能质量分析报告": ("2.2", "2.2.1.1"),
        "近三年电气事故/未遂事件报告": ("2.5", "2.5.1"),
        "电气安全培训程序": ("2.5", "2.5.3.1"),
        "工作票/操作票": ("2.5", "2.5.1"),
        "上锁挂牌程序": ("2.5", "2.5.5"),
    }
    for document_name in expected:
        sheet.append(["现行模板", document_name, None, None, None, None, "OK", "OK", None])
    workbook.save(path)

    result = map_s2_1(path, file_id="file-s21")

    routes = {
        item.subject.split("/", 1)[1]: (item.module_id, item.submodule_id)
        for item in result.evidence_items
    }
    assert routes == expected
    assert result.gaps == []


def test_s2_1_routes_maintenance_reports_to_maintenance_coverage(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S2-1收资表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收资表"
    sheet.append(["类型", "文件名称", None, None, None, None, "具备", "有效", "备注"])
    sheet.append(
        [
            "运行记录",
            "维护试验与巡检报告",
            None,
            None,
            None,
            None,
            "OK",
            "OK",
            None,
        ]
    )
    workbook.save(path)

    result = map_s2_1(path, file_id="file-s21")

    item = result.evidence_items[0]
    assert (item.module_id, item.submodule_id) == ("2.5", "2.5.3.2")
