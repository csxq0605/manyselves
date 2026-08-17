from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from manyselves.core.reporting.mappers.s4_4 import map_s4_4


def _write_s4_4(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "低配评估详情"
    sheet.append([None] * 26)
    sheet.append(
        [
            "配电房",
            "低压柜",
            "运行电流(A)",
            "变压器容量(kVA)",
            "负荷率",
            "安全连锁",
            None,
            "电缆状态",
            None,
            "剩余电流(A)",
            None,
            "剩余电流占比",
            "IP与接地",
            None,
            "色标标识",
            None,
            "谐波",
            None,
            "电涌保护装置",
            None,
            "变压器噪声震动",
            "变压器温度",
            None,
            "其他",
            None,
            None,
        ]
    )
    sheet.append(["芜湖"] + [None] * 25)
    sheet.append(["宿舍配电房", "3G总柜", 33, 1250] + [None] * 22)
    sheet["F4"] = "NG"
    sheet.merge_cells("F4:G4")
    sheet.append(
        [
            "车间配电房",
            "1A2",
            2300,
            2500,
            "=C5*0.4*1.732/D5",
            "NG",
            '=DISPIMG("ID_LOCK",1)',
            "NG",
            '=DISPIMG("ID_CABLE",1)',
            46.8,
            '=DISPIMG("ID_RESIDUAL",1)',
            "=J5/C5",
            "OK",
            None,
            "OK",
            None,
            ">6.2%",
            None,
            "OK",
            None,
            "OK",
            "OK",
            None,
            None,
            None,
            None,
        ]
    )
    sheet.append(
        [
            "车间配电房",
            "2A2",
            3500,
            2500,
            "=C6*0.4*1.732/D6",
            "NG",
            None,
            "OK",
            None,
            59.6,
            '=DISPIMG("ID_RESIDUAL_2",1)',
            "=J6/C6",
            "OK",
            None,
            "OK",
            None,
            ">6.8%",
            None,
            "OK",
            None,
            "OK",
            "NG（三相温度差异大）",
            '=DISPIMG("ID_TEMP",1)',
            None,
            None,
            None,
        ]
    )
    workbook.save(path)


def test_s4_4_maps_bare_conductor_protection_to_same_row_photo(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)

    result = map_s4_4(path, file_id="file-s44")

    item = next(
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/1A2"
        and "裸露导体绝缘防护=NG" in item.fact
    )
    assert item.submodule_id == "2.4.2.1"
    assert item.photo_refs == ["ID_CABLE"]
    assert item.source.sheet == "低配评估详情"
    assert item.source.cell == "H5:I5"


def test_s4_4_calculates_96_99_percent_as_below_100_not_existing_overload(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)

    result = map_s4_4(path, file_id="file-s44")
    item = next(
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/2A2" and item.source.cell == "C6:E6"
    )

    assert item.value == pytest.approx(96.992)
    assert item.unit == "%"
    assert "过载" not in item.fact
    assert "96.99%" in item.fact


def test_s4_4_keeps_operating_current_when_capacity_is_blank(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)
    workbook = load_workbook(path)
    workbook["低配评估详情"]["D4"] = None
    workbook.save(path)
    workbook.close()

    result = map_s4_4(path, file_id="file-s44")
    items = [
        item
        for item in result.evidence_items
        if item.subject == "宿舍配电房/3G总柜" and item.source.cell == "C4:E4"
    ]

    assert {(item.module_id, item.submodule_id) for item in items} == {
        ("2.1", "2.1.1"),
        ("2.4", "2.4.4"),
    }
    assert all(item.value == 33 and item.unit == "A" for item in items)
    assert all("不能据此计算负荷率或判断过载" in item.fact for item in items)
    assert all(item.needs_confirmation for item in items)
    assert {
        (gap.module_id, gap.submodule_id)
        for gap in result.gaps
        if gap.code == "missing_capacity_for_load_rate"
    } == {("2.1", "2.1.1"), ("2.4", "2.4.1.1")}


def test_s4_4_maps_residual_current_with_its_own_photo(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)

    result = map_s4_4(path, file_id="file-s44")
    item = next(
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/1A2" and item.submodule_id == "2.4.3.1"
    )

    assert item.value == 46.8
    assert item.unit == "A"
    assert item.photo_refs == ["ID_RESIDUAL"]
    assert item.source.cell == "C5;J5:L5"
    assert "占运行电流=2.03%" in item.fact


def test_s4_4_routes_load_harmonics_and_surge_to_system_modules(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)

    result = map_s4_4(path, file_id="file-s44")

    load = next(
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/1A2" and item.submodule_id == "2.1.1"
    )
    harmonic = next(
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/1A2" and item.submodule_id == "2.2.1.1"
    )
    surge = next(
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/1A2" and item.submodule_id == "2.3.3"
    )
    assert load.module_id == "2.1"
    assert load.value == pytest.approx(63.7376)
    assert harmonic.fact == "谐波=>6.2%"
    assert surge.fact == "电涌保护装置=OK"


def test_s4_4_routes_environment_and_operations_observations(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)
    workbook = load_workbook(path)
    general = workbook.create_sheet("配电房合规性")
    general.append([None] * 21)
    general.append([None] * 21)
    general.append([None] * 21)
    general.append(
        ["车间配电房", "OK", None, None, None, None, None, None, None, "NG", None, "OK", None, "NG"]
    )
    general["M4"] = "无台账"
    thermal = workbook.create_sheet("红外热成像检测记录表")
    thermal.append([None] * 11)
    thermal.append([None] * 11)
    thermal.append(["车间配电房", 30, None, "1A2", None, "A相接头", 62, None, None, None, None])
    thermal.append([None, None, None, "1A3", None, None, 33, None, None, None, None])
    workbook.save(path)
    workbook.close()

    result = map_s4_4(path, file_id="file-s44")

    assert any(
        item.submodule_id == "2.5.5" and "安全用具=NG" in item.fact
        for item in result.evidence_items
    )
    assert any(
        item.submodule_id == "2.5.6" and "整机备件=NG" in item.fact
        for item in result.evidence_items
    )
    assert any(
        item.submodule_id == "2.5.5"
        and "操作工具=OK；补充记录=无台账" in item.fact
        for item in result.evidence_items
    )
    assert not any(
        item.source.sheet == "配电房合规性" and item.submodule_id == "2.4.4"
        for item in result.evidence_items
    )
    assert any(
        item.submodule_id == "2.2.2.1" and "相对环境温升=32.0K" in item.fact
        for item in result.evidence_items
    )
    inherited = [
        item
        for item in result.evidence_items
        if item.subject == "车间配电房/1A3" and item.source.row == 4
    ]
    assert {(item.module_id, item.submodule_id) for item in inherited} == {
        ("2.2", "2.2.2.1"),
        ("2.4", "2.4.4"),
    }
    assert all("测点温度=33℃" in item.fact for item in inherited)
    assert all("相对环境温升=3.0K" in item.fact for item in inherited)
    assert all("检测部位=未填写" in item.fact for item in inherited)
    assert all(item.needs_confirmation for item in inherited)
    assert {
        (gap.module_id, gap.submodule_id)
        for gap in result.gaps
        if gap.code == "missing_thermal_position"
    } == {("2.2", "2.2.2.1"), ("2.4", "2.4.4")}


def test_s4_4_keeps_embedded_thermal_defect_criteria_image(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)
    workbook = load_workbook(path)
    thermal = workbook.create_sheet("红外热成像检测记录表")
    thermal["J1"] = '=DISPIMG("ID_THERMAL_CRITERIA",1)'
    workbook.save(path)
    workbook.close()

    result = map_s4_4(path, file_id="file-s44")
    criteria = next(
        item
        for item in result.evidence_items
        if item.subject == "红外热成像缺陷判定标准"
    )

    assert criteria.source.sheet == "红外热成像检测记录表"
    assert criteria.source.cell == "J1"
    assert criteria.submodule_id == "2.2.2.1"
    assert criteria.photo_refs == ["ID_THERMAL_CRITERIA"]


def test_s4_4_maps_total_distribution_sheet(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    _write_s4_4(path)
    workbook = load_workbook(path)
    total = workbook.create_sheet("总配评估详情表")
    total.append([None] * 20)
    total.append([None] * 20)
    total.append(
        [
            "16号楼配电房",
            "17号楼",
            186,
            2000,
            "=C3*0.4*1.732/D3",
            "无",
            "OK",
            None,
            12.5,
            '=DISPIMG("ID_TOTAL_RESIDUAL",1)',
            "=I3/C3",
            "OK",
            None,
            "OK",
            None,
            2.3,
            None,
            "NG",
            '=DISPIMG("ID_TOTAL_SPD",1)',
            "NG",
        ]
    )
    total.append(
        [
            "老配电房",
            "占位符柜",
            10,
            1000,
            None,
            "无",
            "OK",
            None,
            None,
            None,
            None,
            "OK",
            None,
            "OK",
            None,
            "/",
            None,
            "OK",
            None,
            "/",
        ]
    )
    total.append(
        [
            "总配",
            None,
            None,
            6250,
            None,
            "无",
            "OK",
            None,
            None,
            None,
            None,
            None,
            None,
            "OK",
            None,
            None,
            None,
            "OK",
            None,
            "OK",
        ]
    )
    workbook.save(path)
    workbook.close()

    result = map_s4_4(path, file_id="file-s44")
    total_items = [
        item for item in result.evidence_items if item.source.sheet == "总配评估详情表"
    ]

    assert any(item.submodule_id == "2.1.1" and item.unit == "%" for item in total_items)
    assert any(
        item.submodule_id == "2.4.1.3" and item.fact == "安全连锁=无"
        for item in total_items
    )
    residual = next(item for item in total_items if item.submodule_id == "2.4.3.1")
    assert residual.value == 12.5
    assert residual.photo_refs == ["ID_TOTAL_RESIDUAL"]
    assert residual.source.cell == "C3;I3:K3"
    assert any(
        item.submodule_id == "2.2.1.1"
        and item.fact == "谐波=2.3%"
        and item.unit == "%"
        for item in total_items
    )
    surge = next(item for item in total_items if item.submodule_id == "2.3.3")
    assert surge.photo_refs == ["ID_TOTAL_SPD"]
    assert any(item.submodule_id == "2.2.2.3" for item in total_items)
    placeholder_items = [
        item for item in total_items if item.subject == "老配电房/占位符柜"
    ]
    assert not any(
        item.submodule_id in {"2.2.1.1", "2.2.2.3"}
        for item in placeholder_items
    )
    capacity_only = [
        item
        for item in total_items
        if item.subject == "总配" and item.source.cell == "C5:E5"
    ]
    assert {(item.module_id, item.submodule_id) for item in capacity_only} == {
        ("2.1", "2.1.1"),
        ("2.4", "2.4.1.1"),
    }
    assert all(item.value == 6250 and item.unit == "kVA" for item in capacity_only)
    assert any(
        item.subject == "总配"
        and item.submodule_id == "2.4.1.3"
        and item.fact == "安全连锁=无"
        for item in total_items
    )
    assert {
        (gap.module_id, gap.submodule_id)
        for gap in result.gaps
        if gap.code == "missing_current_for_load_rate"
    } == {("2.1", "2.1.1"), ("2.4", "2.4.1.1")}
