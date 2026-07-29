from pathlib import Path

from openpyxl import Workbook

from manyselves.core.reporting.mappers.s4_6 import map_s4_6


def test_s4_6_maps_summary_taxonomy_without_using_filename_as_location(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S4-6评估总表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估信息汇总表"
    summary.append([None] * 20)
    summary.append([None] * 20)
    summary.append([None] * 20)
    row = [None] * 20
    row[1] = "2.4.2.2"
    row[4] = "等电位连接与接地问题"
    row[5] = "总电箱进线缺少地线"
    row[6] = "补敷接地线"
    summary.append(row)
    architecture = [None] * 20
    architecture[1] = "2.1.1"
    architecture[5] = "负荷分配待优化"
    architecture[6] = "复核运行方式"
    summary.append(architecture)
    conclusion = workbook.create_sheet("结论建议汇总表")
    conclusion.append([None] * 8)
    conclusion.append(
        [
            None,
            None,
            "接地扁铁断开",
            "存在触电风险",
            "修复接地",
            None,
            None,
            "配电设备＞等电位连接与接地问题",
        ]
    )
    conclusion.append(
        [
            None,
            None,
            "操作规程未更新",
            "误操作风险",
            "修订规程",
            None,
            None,
            "运维管理＞SOP/EOP",
        ]
    )
    workbook.save(path)

    result = map_s4_6(path, file_id="file-s46")

    summary_item = next(item for item in result.evidence_items if item.source.cell == "F4:G4")
    assert summary_item.submodule_id == "2.4.2.2"
    assert summary_item.subject == "评估总表 2.4.2.2"
    assert path.name not in summary_item.subject
    assert summary_item.needs_confirmation is True
    conclusion_item = next(item for item in result.evidence_items if item.source.cell == "C2:E2")
    assert conclusion_item.submodule_id == "2.4.2.2"
    architecture_item = next(item for item in result.evidence_items if item.source.cell == "F5:G5")
    assert (architecture_item.module_id, architecture_item.submodule_id) == ("2.1", "2.1.1")
    operations_item = next(item for item in result.evidence_items if item.source.cell == "C3:E3")
    assert (operations_item.module_id, operations_item.submodule_id) == ("2.5", "2.5.1")


def test_s4_6_maps_actual_b_to_f_conclusion_layout_and_photo(tmp_path: Path) -> None:
    path = tmp_path / "S4-6评估总表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估信息汇总表"
    conclusion = workbook.create_sheet("结论建议汇总表")
    conclusion.append([None] * 6)
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞配置与选型问题＞电压事件（过压）的防范",
            "浪涌保护器未投用",
            "过电压风险",
            "恢复投用",
            '=DISPIMG("ID_SPD",1)',
        ]
    )
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞安装规范性问题＞设备分合/储能/工作位置显示",
            "状态指示灯损坏",
            "误操作风险",
            "修复指示灯",
            None,
        ]
    )
    workbook.save(path)

    result = map_s4_6(path, file_id="file-s46")

    surge = next(item for item in result.evidence_items if item.source.cell == "B2:F2")
    assert (surge.module_id, surge.submodule_id) == ("2.3", "2.3.3")
    assert surge.photo_refs == ["ID_SPD"]
    display = next(item for item in result.evidence_items if item.source.cell == "B3:F3")
    assert (display.module_id, display.submodule_id) == ("2.4", "2.4.1.4")
    assert result.gaps == []


def test_s4_6_uses_observation_specific_routes_for_broad_source_categories(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S4-6评估总表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估信息汇总表"
    conclusion = workbook.create_sheet("结论建议汇总表")
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞带病运行问题",
            "断路器开关损坏，无法实现分合闸操作",
            "保护失效",
            "更换断路器",
            None,
        ]
    )
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞带病运行问题",
            "变压器异响、震动",
            "设备损坏",
            "检测维修",
            None,
        ]
    )
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞安装规范性问题＞电缆、桥架、母线安装问题",
            "母线接头裸露",
            "触电风险",
            "加装防护罩",
            None,
        ]
    )
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞带病运行问题",
            "低压回路流过接地线的电流过高",
            "漏电风险",
            "查明原因",
            None,
        ]
    )
    conclusion.append(
        [
            None,
            "配电设备/元件内在风险＞安装规范性问题＞电缆、桥架、母线安装问题",
            "控制器接线端子存在铜线裸露现象",
            "短路风险",
            "绝缘处理",
            None,
        ]
    )
    workbook.save(path)

    result = map_s4_6(path, file_id="file-s46")

    routes = [
        (item.module_id, item.submodule_id) for item in result.evidence_items
    ]
    assert routes == [
        ("2.4", "2.4.1.4"),
        ("2.2", "2.2.2.3"),
        ("2.4", "2.4.2.1"),
        ("2.4", "2.4.3.1"),
        ("2.4", "2.4.2.1"),
    ]


def test_s4_6_preserves_parent_only_paths_with_explicit_inference_gap(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S4-6评估总表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估信息汇总表"
    conclusion = workbook.create_sheet("结论建议汇总表")
    conclusion.append(
        [
            None,
            None,
            "车间配电室2#变压器负载率达96.99%",
            "过负荷风险",
            "调整负荷",
            None,
            None,
            "配电系统架构问题＞配电系统负荷分配与过载风险",
        ]
    )
    conclusion.append(
        [
            None,
            None,
            "联络柜开关为3极开关",
            "选型不匹配",
            "改为4极开关",
            None,
            None,
            "配电设备/元件内在风险＞配置与选型问题",
        ]
    )
    workbook.save(path)

    result = map_s4_6(path, file_id="file-s46")

    assert [
        (item.module_id, item.submodule_id) for item in result.evidence_items
    ] == [("2.1", "2.1.1"), ("2.4", "2.4.1.1")]
    assert {(gap.code, gap.module_id, gap.submodule_id) for gap in result.gaps} == {
        ("inferred_conclusion_submodule", "2.4", "2.4.1.1")
    }


def test_s4_6_rejects_placeholder_summary_and_flags_source_taxonomy_conflict(
    tmp_path: Path,
) -> None:
    path = tmp_path / "S4-6评估总表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估信息汇总表"
    summary.append([None] * 7)
    summary.append([None] * 7)
    summary.append([None] * 7)
    summary.append([None, "2.3.1", None, None, None, "--", "--"])
    summary.append([None, "2.4.1.1", None, None, None, "断路器无法分合", "更换"])
    workbook.save(path)

    result = map_s4_6(path, file_id="file-s46")

    assert len(result.evidence_items) == 1
    assert result.evidence_items[0].submodule_id == "2.4.1.1"
    assert {(gap.code, gap.submodule_id) for gap in result.gaps} == {
        ("empty_summary_assessment", "2.3.1"),
        ("summary_taxonomy_conflict", "2.4.1.1"),
    }
