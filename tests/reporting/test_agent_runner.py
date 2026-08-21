import asyncio
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

import manyselves.core.reporting.agent_runner as agent_runner_module
from manyselves.capabilities.distribution_reporting import (
    load_distribution_reporting_capability,
)
from manyselves.config.schema import AgentDefaults
from manyselves.core.artifacts import ToolContractError
from manyselves.core.artifacts.content_store import ContentAddressedStore
from manyselves.core.loops.bus import MessageBus
from manyselves.core.providers.base import (
    LLMProvider,
    LLMResponse,
    LLMToolCall,
    build_provider_request_metrics,
)
from manyselves.core.providers.base import (
    Message as LLMMessage,
)
from manyselves.core.reporting.agent_runner import (
    InspectDocumentTool,
    InspectImageTool,
    ReportingAgentRunner,
    load_conversation_trace,
)
from manyselves.core.reporting.agentic_models import (
    SUBMISSION_INPUT_TYPES,
    TEMPLATE_ROLE_SKILL_IDS,
    AgentRunStatus,
    CrossReviewFinding,
    ModuleSubmission,
    TaskEnvelope,
)
from manyselves.core.reporting.config import load_packaged_agents
from manyselves.core.reporting.input_contracts import (
    INPUT_CONTRACT_EXAMPLES,
    INPUT_CONTRACT_TYPES,
    ChiefEditorInput,
    CrossOwnerInput,
    ModuleAuthoringInput,
    ModuleContentView,
    ModuleReviewInput,
    TemplateDistillationInput,
    ValidationReport,
)
from manyselves.core.reporting.taxonomy import REPORT_TAXONOMY, compose_module_markdown
from manyselves.core.reporting.workflow import ReportingRunBudget
from manyselves.core.usage_ledger import UsageLedger
from manyselves.interfaces.types import (
    AgentResultMessage,
    UserMessage,
)
from manyselves.kernel.definitions import DefinitionKind, ToolDefinition
from manyselves.kernel.recovery import RecoveryActionKind, RecoveryEventKind
from manyselves.kernel.workflow import ResolvedPlan
from manyselves.runtime.state_store import FileWorkflowStateStore


def _write_template_contract(
    workspace: Path,
    *,
    run_id: str,
    template_ref: str,
) -> str:
    contract_ref = (
        f"Work/runs/{run_id}/context/template-distillation-input.json"
    )
    target = workspace / contract_ref
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        TemplateDistillationInput(
            run_id=run_id,
            template_ref=template_ref,
            inspect_max_chars=100_000,
            required_part_ids=list(TEMPLATE_ROLE_SKILL_IDS),
        ).model_dump_json(),
        encoding="utf-8",
    )
    return contract_ref


@pytest.mark.parametrize("kind", sorted(INPUT_CONTRACT_TYPES))
def test_runner_loads_every_registered_model_input_contract(
    tmp_path: Path,
    kind: str,
) -> None:
    contract_ref = f"Work/runs/run-input-matrix/context/{kind}.json"
    target = tmp_path / contract_ref
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(INPUT_CONTRACT_EXAMPLES[kind], ensure_ascii=False),
        encoding="utf-8",
    )
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id=f"input-matrix-{kind}",
        run_id="run-input-matrix",
        agent_id="main",
        objective="验证所有模型输入身份共用的契约加载路径",
        input_refs=[contract_ref],
        input_contract_kind=kind,
        input_contract_ref=contract_ref,
    )

    loaded = runner._input_contract(envelope)

    assert isinstance(loaded, INPUT_CONTRACT_TYPES[kind])


@pytest.mark.asyncio
async def test_inspect_image_requires_current_run_photo_scope(tmp_path: Path) -> None:
    image_path = tmp_path / "Work/runs/run-photo/photos/photo.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (2, 2), color="white").save(image_path)
    tool = InspectImageTool(
        tmp_path,
        allowed_refs=("Work/runs/run-photo/photos/photo.png",),
        photo_refs={"P-001": "Work/runs/run-photo/photos/photo.png"},
    )

    inspected = await tool(path="P-001")
    assert inspected["path"] == "Work/runs/run-photo/photos/photo.png"
    with pytest.raises(ToolContractError, match="current run PhotoAsset map"):
        await tool(path="P-999")
    with pytest.raises(PermissionError):
        await tool(path="../outside.png")


class DirectSubmissionProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="scripted")
        self.system_prompts: list[str] = []
        self.task_messages: list[str] = []
        self.message_snapshots: list[list[LLMMessage]] = []
        self.max_tokens_seen: list[int] = []

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.message_snapshots.append(list(messages))
        self.max_tokens_seen.append(max_tokens)
        self.system_prompts.append(messages[0].content)
        self.task_messages.append(messages[1].content)
        if messages[-1].is_tool_result and '"status": "completed"' in messages[-1].content:
            return LLMResponse(content="已提交。")
        if not any(message.is_tool_result for message in messages):
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{submodule_id}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 完整分析正文。",
                            "evidence_ids": [],
                        },
                    )
                    for submodule_id in REPORT_TAXONOMY["2.1"].submodules
                ],
            )
        return LLMResponse(
            content="",
            tool_calls=[
                LLMToolCall(
                    id="submit-1",
                    name="submit_result",
                    arguments={
                        "kind": "module_submission",
                        "module_id": "2.1",
                        "unresolved_questions": ["待补充一次系统图"],
                        "revision": 0,
                        "revision_responses": [],
                    },
                )
            ],
        )


class MeasuredDirectSubmissionProvider(DirectSubmissionProvider):
    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        response = await super().chat(
            messages,
            tools=tools,
            temperature=temperature,
            max_tokens=max_tokens,
        )
        response.request_metrics = build_provider_request_metrics(
            {
                "model": self.model,
                "messages": [
                    {
                        "role": message.role,
                        "content": message.content,
                        "tool_calls": [
                            {
                                "id": call.id,
                                "name": call.name,
                                "arguments": call.arguments,
                            }
                            for call in (message.tool_calls or [])
                        ],
                        "tool_call_id": message.tool_call_id,
                    }
                    for message in messages
                ],
                "tools": tools or [],
                "temperature": temperature,
                "max_tokens": max_tokens,
            },
            representation="scripted_provider_payload_v1",
        )
        return response


class MaxTokensThenDirectSubmissionProvider(DirectSubmissionProvider):
    def __init__(self):
        super().__init__()
        self.truncated_once = False

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        if not self.truncated_once:
            self.truncated_once = True
            self.max_tokens_seen.append(max_tokens)
            return LLMResponse(
                content="",
                thinking="unfinished reasoning",
                usage={
                    "input_tokens": 100,
                    "output_tokens": max_tokens,
                    "total_tokens": 100 + max_tokens,
                },
                stop_reason="max_tokens",
            )
        return await super().chat(
            messages,
            tools=tools,
            temperature=temperature,
            max_tokens=max_tokens,
        )


class ModuleReviewSubmissionProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="scripted-auditor")

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        if any(message.is_tool_result for message in messages):
            return LLMResponse(content="审计结果已提交。")
        content = "\n".join(message.content for message in messages)
        module_id = "2.1" if "audit-2.1" in content else "2.2"
        return LLMResponse(
            content="",
            tool_calls=[
                LLMToolCall(
                    id=f"submit-audit-{module_id}",
                    name="submit_result",
                    arguments={
                        "kind": "module_review_finding_submission",
                        "findings": [],
                    },
                )
            ],
        )


def test_submit_result_exposes_only_the_role_output_schema(tmp_path: Path) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-schema",
        agent_id="module-2.1-specialist",
        objective="分析 2.1",
        allowed_outputs=["module_submission"],
    )

    registry = runner._tools(
        load_packaged_agents()["module-2.1-specialist"],
        envelope,
        "session-schema",
        "workflow-schema",
    )
    submission = registry._schema_cache["submit_result"]

    assert submission["properties"]["kind"]["const"] == "module_submission"
    assert "tool arguments are this submission object itself" in submission["description"]
    assert "payload" not in submission["properties"]
    assert "module_id" in submission["properties"]
    assert "$defs" in submission
    assert "ModuleDispatchPlan" not in submission["$defs"]
    assert submission["additionalProperties"] is False


@pytest.mark.parametrize("kind", sorted(SUBMISSION_INPUT_TYPES))
def test_every_submission_identity_uses_one_flat_provider_schema(
    tmp_path: Path,
    kind: str,
) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["module-2.1-specialist"]
    envelope = TaskEnvelope(
        task_id=f"flat-schema-{kind}",
        run_id="run-flat-schema",
        agent_id=definition.id,
        objective="验证统一扁平提交接口",
        allowed_outputs=[kind],
    )

    schema = runner._tools(
        definition,
        envelope,
        "session-flat-schema",
        "workflow-flat-schema",
    )._schema_cache["submit_result"]

    assert schema["type"] == "object"
    assert schema["properties"]["kind"]["const"] == kind
    assert "kind" in schema["required"]
    assert "payload" not in schema["properties"]
    assert schema["additionalProperties"] is False


def test_provider_manifest_distinguishes_pre_adapter_and_provider_payload(
    tmp_path: Path,
) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["module-2.1-specialist"]
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-provider-manifest",
        agent_id=definition.id,
        objective="分析 2.1",
        allowed_outputs=["module_submission"],
    )
    manifest_path = runner._write_provider_call_manifest(
        definition=definition,
        envelope=envelope,
        identity_key=definition.id,
        session_id="session-manifest",
        messages=[LLMMessage(role="user", content="task")],
        tool_definitions=[
            {
                "name": "submit_result",
                "description": "submit",
                "input_schema": {"type": "object"},
            }
        ],
        phase="initial",
        attempt=1,
        call_index=1,
    )
    before = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert before["provider_context_manifest_version"] == 3
    assert before["task_attempt_id"] == envelope.task_attempt_id
    assert before["request_sha256_scope"] == "agent_pre_adapter"
    assert before["provider_payload_status"] == "pending"
    assert before["provider_payload"] is None

    relative = manifest_path.relative_to(tmp_path).as_posix()
    runner._finalize_provider_call_manifest(
        {
            "context_manifest_ref": relative,
            "provider_call_id": manifest_path.stem,
            "request_metric_source": "provider_adapter_payload",
            "provider_request_representation": "test_provider_payload_v1",
            "request_fingerprint": "a" * 64,
            "message_fingerprint": "b" * 64,
            "tool_schema_fingerprint": "c" * 64,
            "request_chars": 700,
            "message_chars": 500,
            "tool_schema_chars": 100,
            "status": "success",
            "usage_source": "provider",
            "attempt_disposition": "completed",
            "retry_decision": "completed",
        }
    )
    after = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert after["provider_payload_status"] == "observed"
    assert after["provider_payload"] == {
        "representation": "test_provider_payload_v1",
        "request_sha256": "a" * 64,
        "message_sha256": "b" * 64,
        "tool_schema_sha256": "c" * 64,
        "request_chars": 700,
        "message_chars": 500,
        "tool_schema_chars": 100,
    }
    assert after["pre_adapter_request"]["request_sha256"] == before[
        "request_sha256"
    ]
    assert after["attempt_disposition"] == "completed"


@pytest.mark.asyncio
async def test_new_dispatch_preserves_old_journal_and_starts_fresh_attempt(
    tmp_path: Path,
) -> None:
    provider = DirectSubmissionProvider()
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    definition = load_packaged_agents()["module-2.1-specialist"]
    workflow_id = "workflow-fresh-attempt-recovery"
    envelope = TaskEnvelope(
        task_id="module-2.1-fresh-attempt",
        run_id="run-fresh-attempt-recovery",
        agent_id=definition.id,
        objective="显式恢复后创建新的物理请求尝试",
        allowed_outputs=["module_submission"],
    )
    session_id = "session-" + hashlib.sha256(
        f"{workflow_id}:{definition.id}".encode("utf-8")
    ).hexdigest()[:12]
    manifest_path = runner._write_provider_call_manifest(
        definition=definition,
        envelope=envelope,
        identity_key=definition.id,
        session_id=session_id,
        messages=[LLMMessage(role="user", content="task")],
        tool_definitions=[],
        phase="initial",
        attempt=1,
        call_index=1,
    )

    try:
        result = await runner.run(
            definition,
            envelope,
            [],
            workflow_id=workflow_id,
        )
    finally:
        await runner.close_workflow(workflow_id)
        bus.shutdown()
        await bus_task

    assert result.status == AgentRunStatus.COMPLETED
    assert provider.max_tokens_seen
    assert manifest_path.is_file()


@pytest.mark.asyncio
async def test_agent_reference_tools_receive_the_optional_global_root(tmp_path: Path) -> None:
    global_root = tmp_path / "global"
    global_root.mkdir()
    (global_root / "standard.md").write_text(
        "shared protection baseline", encoding="utf-8"
    )
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
        global_root=global_root,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-global-tools",
        agent_id="module-2.1-specialist",
        objective="analyse 2.1",
        allowed_outputs=["module_submission"],
    )

    registry = runner._tools(
        load_packaged_agents()["module-2.1-specialist"],
        envelope,
        "session-global-tools",
        "workflow-global-tools",
    )
    search = registry.get("search_reference_library")
    assert search is not None
    result = await search(query="protection baseline")

    assert result["hits"][0]["namespace"] == "global"
    assert result["hits"][0]["locator"] == "GlobalKnowledge/standard.md"


def test_module_authoring_schema_and_example_use_current_identity(
    tmp_path: Path,
) -> None:
    run_id = "run-module-schema"
    contract = ModuleAuthoringInput(
        run_id=run_id,
        module_id="2.4",
        revision=3,
        required_submodule_ids=list(REPORT_TAXONOMY["2.4"].submodules),
        coverage_ref=f"Work/runs/{run_id}/preparation/coverage.json",
        evidence_ref=f"Work/runs/{run_id}/preparation/evidence.jsonl",
        manifest_ref=f"Work/runs/{run_id}/preparation/manifest.json",
        knowledge_ref=f"Work/runs/{run_id}/context/module-2.4-knowledge.md",
    )
    contract_ref = f"Work/runs/{run_id}/context/authoring.json"
    target = tmp_path / contract_ref
    target.parent.mkdir(parents=True)
    target.write_text(contract.model_dump_json(), encoding="utf-8")
    contract_payload = json.loads(target.read_text(encoding="utf-8"))
    assert "cross_synthesis_inputs" not in contract_payload
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="module-2.4",
        run_id=run_id,
        agent_id="module-2.4-specialist",
        objective="完成模块 2.4",
        input_refs=[contract_ref],
        allowed_outputs=["module_submission"],
        input_contract_kind="module_authoring_input",
        input_contract_ref=contract_ref,
    )

    registry = runner._tools(
        load_packaged_agents()["module-2.4-specialist"],
        envelope,
        "session-module-schema",
        "workflow-module-schema",
    )
    payload = registry._schema_cache["submit_result"]
    writer_schema = registry._schema_cache["write_result_part"]

    assert payload["properties"]["module_id"]["const"] == "2.4"
    assert payload["properties"]["revision"]["const"] == 3
    assert payload["examples"][0]["module_id"] == "2.4"
    assert payload["examples"][0]["revision"] == 3
    expected_part_ids = list(REPORT_TAXONOMY["2.4"].submodules)
    assert writer_schema["properties"]["part_id"]["enum"] == expected_part_ids
    assert writer_schema["required"] == ["part_id", "content", "evidence_ids"]
    assert writer_schema["additionalProperties"] is False
    content_description = writer_schema["properties"]["content"]["description"]
    assert "Always supply the full intended prose" in content_description
    assert "Every call must contain all required arguments" in content_description
    assert "may omit content" not in content_description
    assert "persisted_result_part" not in content_description
    assert "list_result_parts" in content_description
    assert registry.get("write_result_parts") is None
    assert "write_result_parts" not in registry._schema_cache


def test_cross_owner_reviewers_use_one_identity_per_owner_and_reuse_on_recheck() -> None:
    definition = load_packaged_agents()["cross-module-reviewer"]
    initial = TaskEnvelope(
        task_id="cross-owner-2.1-r0-initial",
        run_id="run-cross-owner-identity",
        agent_id=definition.id,
        objective="检查模块 2.1 的跨模块关系",
        allowed_outputs=["cross_owner_finding_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )
    recheck = initial.model_copy(
        update={
            "task_id": "cross-owner-2.1-r1-recheck",
            "allowed_outputs": ["cross_owner_verdict_submission"],
        }
    )
    other_owner = initial.model_copy(
        update={"task_id": "cross-owner-2.3-r0-initial"}
    )

    assert ReportingAgentRunner._identity_key(
        definition, initial, "cross-owner-2.1"
    ) == "cross-owner-2.1"
    assert ReportingAgentRunner._identity_key(
        definition, recheck, "cross-owner-2.1"
    ) == "cross-owner-2.1"
    assert ReportingAgentRunner._identity_key(
        definition, other_owner, "cross-owner-2.3"
    ) == "cross-owner-2.3"
    assert ReportingAgentRunner._identity_key(
        definition, initial, "cross-module-reviewer"
    ) == definition.id


def test_reporting_lane_runtime_ids_expose_the_durable_identity() -> None:
    agents = load_packaged_agents()

    assert ReportingAgentRunner._runtime_id(
        agents["evidence-auditor"],
        "module-auditor-2.1",
        "session-a1b2c3d4",
    ) == "module-auditor-2.1--session-a1b2c3d4"
    assert ReportingAgentRunner._runtime_id(
        agents["cross-module-reviewer"],
        "cross-owner-2.4",
        "session-a1b2c3d4",
    ) == "cross-owner-2.4--session-a1b2c3d4"
    assert ReportingAgentRunner._runtime_id(
        agents["chief-editor"],
        "chief-chapter-3",
        "session-a1b2c3d4",
    ) == "chief-chapter-3--session-a1b2c3d4"
    assert ReportingAgentRunner._runtime_id(
        agents["chief-editor-auditor"],
        "final-chapter-4",
        "session-a1b2c3d4",
    ) == "final-chapter-4--session-a1b2c3d4"
    assert ReportingAgentRunner._runtime_id(
        agents["module-2.1-specialist"],
        "module-2.1-specialist",
        "session-a1b2c3d4",
    ) == "module-2.1-specialist--session-a1b2c3d4"


def test_cross_owner_system_prompt_does_not_get_packaged_module_skills(tmp_path: Path) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["cross-module-reviewer"]
    envelope = TaskEnvelope(
        task_id="cross-owner-2.4-r0-initial",
        run_id="run-cross-owner-skill",
        agent_id=definition.id,
        objective="发现模块 2.4 的跨模块关系问题",
        allowed_outputs=["cross_owner_finding_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.4"].submodules),
    )

    selected = runner._selected_module_skills(definition, envelope)
    prompt = runner._system_prompt(definition, envelope)

    assert selected == []
    assert "配置与选型核查" not in prompt
    assert "电力系统架构核查" not in prompt
    assert "固定模块的 Cross owner" in prompt


def test_cross_owner_submission_schema_binds_owner_and_required_verdict_ids() -> None:
    contract = CrossOwnerInput.model_validate(
        INPUT_CONTRACT_EXAMPLES["cross_owner_input"]
    )
    initial = ReportingAgentRunner._task_submission_schema(
        "cross_owner_finding_submission",
        contract,
    )
    assert initial["properties"]["owner_module_id"]["const"] == "2.1"
    assert (
        initial["$defs"]["CrossReviewCoverageEntry"]["properties"]["module_id"]["const"]
        == "2.1"
    )
    assert "id" not in initial["$defs"]["CrossReviewFinding"]["required"]
    assert "id" not in initial["$defs"]["CrossReviewFinding"]["properties"]

    required = CrossReviewFinding.model_construct(id="X-2.1-001")
    recheck_contract = contract.model_copy(
        update={
            "phase": "recheck",
            "review_round": 1,
            "required_findings": [required],
        }
    )
    recheck = ReportingAgentRunner._task_submission_schema(
        "cross_owner_verdict_submission",
        recheck_contract,
    )
    assert recheck["properties"]["verdicts"]["minItems"] == 1
    assert recheck["properties"]["verdicts"]["maxItems"] == 1
    assert "new_findings" in recheck["properties"]
    assert "new_findings" in recheck["examples"][0]
    assert "synthesis_inputs" not in recheck["properties"]
    assert "synthesis_inputs" not in recheck["examples"][0]
    assert "interface_closures" not in recheck["properties"]
    assert "interface_closures" not in recheck["examples"][0]
    assert recheck["$defs"]["ResolutionVerdict"]["properties"]["finding_id"]["enum"] == [
        "X-2.1-001"
    ]
    assert recheck["examples"][0]["verdicts"][0]["finding_id"] == "X-2.1-001"


def test_module_review_tool_schema_omits_runtime_owned_fields(
    tmp_path: Path,
) -> None:
    run_id = "run-review-schema"
    contract = ModuleReviewInput(
        phase="initial",
        run_id=run_id,
        module_id="2.1",
        lifecycle_id="initial",
        review_round=0,
        subject_ref=f"Work/runs/{run_id}/modules/2.1-r0.json",
        subject_revision=0,
        subject=ModuleContentView(
            module_id="2.1",
            revision=0,
            submodule_narratives={"2.1.1": "完整审查正文"},
            evidence_ids_by_submodule={"2.1.1": []},
        ),
        evidence=[],
        required_submodule_ids=["2.1.1"],
        validation_report_ref=f"Work/runs/{run_id}/validations/2.1.json",
        validation_report=ValidationReport(
            validation_protocol_version=2,
            run_id=run_id,
            subject_ref=f"Work/runs/{run_id}/modules/2.1-r0.json",
            subject_revision=0,
            validator="test/v2",
            check_ids=["structure"],
            passed=True,
        ),
    )
    contract_ref = f"Work/runs/{run_id}/reviews/input.json"
    target = tmp_path / contract_ref
    target.parent.mkdir(parents=True)
    target.write_text(contract.model_dump_json(), encoding="utf-8")
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="module-2.1-review-r0",
        run_id=run_id,
        agent_id="evidence-auditor",
        objective="审查模块 2.1",
        allowed_outputs=["module_review_finding_submission"],
        allowed_tools=["submit_result"],
        input_refs=[contract_ref],
        input_contract_kind="module_review_input",
        input_contract_ref=contract_ref,
        target_submodule_ids=["2.1.1"],
    )

    registry = runner._tools(
        load_packaged_agents()["evidence-auditor"],
        envelope,
        "session-review-schema",
        "workflow-review-schema",
    )
    payload = registry._schema_cache["submit_result"]

    assert "coverage" not in payload["properties"]
    finding = payload["$defs"]["ModuleReviewFinding"]
    assert "id" not in finding["properties"]
    assert finding["properties"]["target_submodule_id"]["enum"] == ["2.1.1"]
    assert registry.get("write_result_part") is None
    assert registry.get("write_result_parts") is None


def test_all_audit_roles_override_the_90_second_provider_idle_default(
    tmp_path: Path,
) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    agents = load_packaged_agents()

    assert {
        agent_id: runner._provider_stream_idle_timeout(agents[agent_id])
        for agent_id in (
            "evidence-auditor",
            "cross-module-reviewer",
            "chief-editor-auditor",
        )
    } == {
        "evidence-auditor": 600.0,
        "cross-module-reviewer": 600.0,
        "chief-editor-auditor": 600.0,
    }


def test_template_skill_submission_schema_and_tools_are_exposed_to_distiller(
    tmp_path: Path,
) -> None:
    template_ref = (
        "Work/runs/run-skill-schema/templates/template-for-skill.docx"
    )
    template = tmp_path / template_ref
    template.parent.mkdir(parents=True)
    Document().save(template)
    contract_ref = _write_template_contract(
        tmp_path,
        run_id="run-skill-schema",
        template_ref=template_ref,
    )
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="template-skill-distillation",
        run_id="run-skill-schema",
        agent_id="template-distiller",
        objective="分析模板并生成写作 Skill",
        allowed_outputs=["template_skill_submission"],
        allowed_tools=[
            "inspect_document",
            "write_result_part",
            "list_result_parts",
            "submit_result",
            "report_blocked",
        ],
        input_refs=[contract_ref, template_ref],
        input_contract_kind="template_distillation_input",
        input_contract_ref=contract_ref,
    )

    registry = runner._tools(
        load_packaged_agents()["template-distiller"],
        envelope,
        "session-schema",
        "workflow-schema",
    )
    payload = registry._schema_cache["submit_result"]

    assert set(registry.get_all()) == {
        "inspect_document",
        "write_result_part",
        "list_result_parts",
        "submit_result",
        "report_blocked",
        "open_tool_result",
    }
    assert payload["properties"]["kind"]["const"] == "template_skill_submission"
    assert "skills" in payload["properties"]
    assert "boundary_manifest" not in payload["properties"]
    assert "TextArtifactRef" in payload["$defs"]
    inspection = registry.get("inspect_document")
    assert isinstance(inspection, InspectDocumentTool)
    assert inspection.required_path == template_ref
    assert inspection.required_max_chars == 100_000
    assert inspection.cache_ref == (
        "Work/runs/run-skill-schema/context/template-inspection.json"
    )
    expected_parts = TEMPLATE_ROLE_SKILL_IDS
    assert registry.get("write_result_part").expected_part_ids == expected_parts
    assert registry.get("write_result_parts") is None
    assert registry.get("list_result_parts").expected_part_ids == expected_parts
    writer_schema = registry._schema_cache["write_result_part"]
    assert writer_schema["properties"]["part_id"]["enum"] == list(expected_parts)
    assert "evidence_ids" not in writer_schema["properties"]
    assert writer_schema["required"] == ["part_id", "content"]
    assert (
        "Every call must contain all required arguments"
        in writer_schema["properties"]["content"]["description"]
    )
    assert "may omit content" not in writer_schema["properties"]["content"]["description"]
    assert "persisted_result_part" not in writer_schema["properties"]["content"]["description"]
    assert "write_result_parts" not in registry._schema_cache


@pytest.mark.asyncio
async def test_template_distiller_accepts_one_verified_cas_template_view(
    tmp_path: Path,
) -> None:
    source = tmp_path / "Templates/report_template.docx"
    source.parent.mkdir(parents=True)
    Document().save(source)
    template_ref = (
        "Work/runs/run-cas-skill/templates/template-for-skill.docx"
    )
    template = tmp_path / template_ref
    content_store = ContentAddressedStore(tmp_path)
    content_store.link_view(content_store.ingest_file(source), template)
    contract_ref = _write_template_contract(
        tmp_path,
        run_id="run-cas-skill",
        template_ref=template_ref,
    )
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="template-skill-distillation",
        run_id="run-cas-skill",
        agent_id="template-distiller",
        objective="分析模板并生成写作 Skill",
        allowed_outputs=["template_skill_submission"],
        allowed_tools=["inspect_document", "submit_result", "report_blocked"],
        input_refs=[contract_ref, template_ref],
        input_contract_kind="template_distillation_input",
        input_contract_ref=contract_ref,
    )

    registry = runner._tools(
        load_packaged_agents()["template-distiller"],
        envelope,
        "session-cas",
        "workflow-cas",
    )

    inspection = registry.get("inspect_document")
    assert isinstance(inspection, InspectDocumentTool)
    assert inspection.required_path == template_ref
    assert template.is_symlink()
    assert template.resolve().is_relative_to(
        (tmp_path / "Work/content/sha256").resolve()
    )
    result = await inspection(template_ref, max_chars=100_000)
    assert result["path"] == template_ref
    assert result["kind"] == "docx"
    assert result["error"] is None


def test_template_distiller_rejects_non_cas_template_symlink(
    tmp_path: Path,
) -> None:
    source = tmp_path / "Templates/report_template.docx"
    source.parent.mkdir(parents=True)
    Document().save(source)
    template_ref = "Work/runs/run-linked-skill/templates/template.docx"
    template = tmp_path / template_ref
    template.parent.mkdir(parents=True)
    template.symlink_to(source)
    contract_ref = _write_template_contract(
        tmp_path,
        run_id="run-linked-skill",
        template_ref=template_ref,
    )
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="template-skill-distillation",
        run_id="run-linked-skill",
        agent_id="template-distiller",
        objective="分析模板并生成写作 Skill",
        allowed_outputs=["template_skill_submission"],
        input_refs=[contract_ref, template_ref],
        input_contract_kind="template_distillation_input",
        input_contract_ref=contract_ref,
    )

    with pytest.raises(ValueError, match="verified CAS view"):
        runner._tools(
            load_packaged_agents()["template-distiller"],
            envelope,
            "session-linked",
            "workflow-linked",
        )


def test_chief_tools_expose_only_current_report_parts(
    tmp_path: Path,
) -> None:
    run_id = "run-chief-tools"
    modules = {}
    for module_id, definition in REPORT_TAXONOMY.items():
        submodule_id = next(iter(definition.submodules))
        modules[module_id] = ModuleContentView(
            module_id=module_id,
            revision=0,
            submodule_narratives={submodule_id: f"{module_id} 已批准正文"},
            evidence_ids_by_submodule={submodule_id: []},
        )
    contract = ChiefEditorInput(
        run_id=run_id,
        approved_module_markers={
            module_id: f"[[APPROVED_MODULE:{module_id}]]"
            for module_id in REPORT_TAXONOMY
        },
        modules=modules,
        cross_review_completion_ref=(
            f"Work/runs/{run_id}/reviews/cross-completion.json"
        ),
        special_topic_plan={
            "source_ref": "Inputs/专项问题分析.md",
            "source_sha256": "0" * 64,
            "sections": [
                {
                    "section_id": "4.1",
                    "title": "动态专项问题",
                    "requirement": "分析项目边界、方案条件和验证方法。",
                }
            ],
        },
    )
    contract_ref = f"Work/runs/{run_id}/context/chief-editor-input.json"
    target = tmp_path / contract_ref
    target.parent.mkdir(parents=True)
    target.write_text(contract.model_dump_json(), encoding="utf-8")
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="chief-edit",
        run_id=run_id,
        agent_id="chief-editor",
        objective="整合报告",
        input_refs=[contract_ref],
        allowed_outputs=["edited_report_submission"],
        allowed_tools=[
            "write_result_part",
            "list_result_parts",
            "submit_result",
        ],
        input_contract_kind="chief_editor_input",
        input_contract_ref=contract_ref,
    )

    registry = runner._tools(
        load_packaged_agents()["chief-editor"],
        envelope,
        "session-chief-tools",
        "workflow-chief-tools",
    )
    expected_parts = (
        "assessment_background",
        "findings_overview",
        "regional_executive_summary",
        "risk_panorama",
        "dimension_risk_analysis",
        "data_gap_analysis",
        "improvement_action_plan",
        "special_topic_analysis",
    )
    writer = registry.get("write_result_part")
    listing = registry.get("list_result_parts")
    assert writer.expected_part_ids == expected_parts
    assert registry.get("write_result_parts") is None
    assert listing.expected_part_ids == expected_parts
    assert listing.required_synthesis_input_ids == ()
    writer_schema = registry._schema_cache["write_result_part"]
    assert writer_schema["properties"]["part_id"]["enum"] == list(expected_parts)
    assert "evidence_ids" not in writer_schema["properties"]
    assert "write_result_parts" not in registry._schema_cache


@pytest.mark.asyncio
async def test_inspect_document_exposes_raw_docx_text_and_visual_structure(
    tmp_path: Path,
) -> None:
    image_path = tmp_path / "evidence.png"
    Image.new("RGB", (20, 20), color="red").save(image_path)
    document = Document()
    document.add_heading("风险分析", level=1)
    document.add_paragraph("观察之后说明原因、影响与建议。")
    document.add_picture(str(image_path), width=Inches(1))
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "问题"
    table.rows[0].cells[1].text = "行动"
    table.rows[1].cells[0].text = "风险"
    table.rows[1].cells[1].text = "复核"
    template = tmp_path / "template.docx"
    document.save(template)

    result = await InspectDocumentTool(tmp_path)("template.docx")

    assert "观察之后说明原因" in result["text"]
    assert result["structure"]["table_count"] == 1
    assert result["structure"]["image_count"] == 1
    assert result["structure"]["headings"][0]["text"] == "风险分析"
    assert result["structure"]["tables"][0]["headers"] == ["问题", "行动"]

    one_shot = InspectDocumentTool(tmp_path, one_shot=True)
    await one_shot("template.docx", max_chars=100_000)
    with pytest.raises(RuntimeError, match="TEMPLATE_ALREADY_INSPECTED"):
        await one_shot("template.docx", max_chars=100_000)


@pytest.mark.asyncio
async def test_inspect_document_never_discards_text_beyond_inline_hint(
    tmp_path: Path,
) -> None:
    source = tmp_path / "long.txt"
    exact = "A" * 128 + "TAIL-MUST-REMAIN"
    source.write_text(exact, encoding="utf-8")

    result = await InspectDocumentTool(tmp_path)("long.txt", max_chars=16)

    assert result["truncated"] is False
    assert result["requested_max_chars"] == 16
    assert result["text_chars"] >= len(exact)
    assert "TAIL-MUST-REMAIN" in result["text"]
    assert result["text_sha256"] == hashlib.sha256(
        result["text"].encode("utf-8")
    ).hexdigest()


@pytest.mark.asyncio
async def test_template_inspection_rejects_wrong_path_without_consuming_one_shot(
    tmp_path: Path,
) -> None:
    contract = tmp_path / "contract.json"
    contract.write_text('{"template_ref":"template.docx"}', encoding="utf-8")
    document = Document()
    document.add_paragraph("唯一模板正文。")
    document.save(tmp_path / "template.docx")
    tool = InspectDocumentTool(
        tmp_path,
        one_shot=True,
        required_path="template.docx",
        required_max_chars=100_000,
    )

    with pytest.raises(RuntimeError, match="TEMPLATE_INSPECTION_PATH_MISMATCH"):
        await tool("contract.json", max_chars=100_000)
    with pytest.raises(RuntimeError, match="TEMPLATE_INSPECTION_LIMIT_MISMATCH"):
        await tool("template.docx", max_chars=50_000)

    result = await tool("template.docx", max_chars=100_000)
    assert "唯一模板正文" in result["text"]
    with pytest.raises(RuntimeError, match="TEMPLATE_ALREADY_INSPECTED"):
        await tool("template.docx", max_chars=100_000)


@pytest.mark.asyncio
async def test_template_inspection_cache_survives_a_new_tool_instance(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("需要跨进程恢复的模板正文。")
    document.save(tmp_path / "template.docx")
    kwargs = {
        "one_shot": True,
        "required_path": "template.docx",
        "required_max_chars": 100_000,
        "cache_ref": "Work/runs/run-template/context/template-inspection.json",
    }
    first = await InspectDocumentTool(tmp_path, **kwargs)(
        "template.docx", max_chars=100_000
    )
    cache = (
        tmp_path
        / "Work/runs/run-template/context/template-inspection.json"
    )
    assert cache.is_file()

    def must_not_parse(_path):
        raise AssertionError("a resumed process must reuse the durable inspection")

    monkeypatch.setattr(
        "manyselves.core.tools.document_tool.parse_artifact",
        must_not_parse,
    )
    resumed = await InspectDocumentTool(tmp_path, **kwargs)(
        "template.docx", max_chars=100_000
    )

    assert resumed == first


@pytest.mark.asyncio
async def test_template_parse_failure_does_not_consume_one_shot(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("解析重试正文。")
    document.save(tmp_path / "template.docx")
    tool = InspectDocumentTool(
        tmp_path,
        one_shot=True,
        required_path="template.docx",
        required_max_chars=100_000,
    )
    from manyselves.core.tools import document_tool

    original = document_tool.parse_artifact

    def fail_once(_path):
        raise ValueError("synthetic parse failure")

    monkeypatch.setattr(document_tool, "parse_artifact", fail_once)
    with pytest.raises(ValueError, match="synthetic parse failure"):
        await tool("template.docx", max_chars=100_000)
    monkeypatch.setattr(document_tool, "parse_artifact", original)

    result = await tool("template.docx", max_chars=100_000)
    assert "解析重试正文" in result["text"]


@pytest.mark.asyncio
async def test_template_inspection_cache_rejects_changed_source(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("第一版。")
    document.save(template)
    kwargs = {
        "one_shot": True,
        "required_path": "template.docx",
        "required_max_chars": 100_000,
        "cache_ref": "Work/runs/run-template/context/template-inspection.json",
    }
    await InspectDocumentTool(tmp_path, **kwargs)(
        "template.docx", max_chars=100_000
    )
    changed = Document()
    changed.add_paragraph("第二版。")
    changed.save(template)

    with pytest.raises(RuntimeError, match="TEMPLATE_INSPECTION_CACHE_MISMATCH"):
        await InspectDocumentTool(tmp_path, **kwargs)(
            "template.docx", max_chars=100_000
        )


@pytest.mark.asyncio
async def test_only_template_distiller_can_inspect_expert_source(tmp_path: Path) -> None:
    source = tmp_path / "Templates/配电安全专家咨询报告(专家优化版).docx"
    source.parent.mkdir(parents=True)
    document = Document()
    document.add_paragraph("仅用于测试隔离边界。")
    document.save(source)

    with pytest.raises(PermissionError, match="EXPERT_TEMPLATE_AGENT_ACCESS_FORBIDDEN"):
        await InspectDocumentTool(tmp_path)(source.relative_to(tmp_path).as_posix())

    result = await InspectDocumentTool(
        tmp_path,
        one_shot=True,
        allow_template_distiller_source=True,
    )(source.relative_to(tmp_path).as_posix())
    assert "仅用于测试隔离边界" in result["text"]


class NonRetryableFailureProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="failing")

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        raise ValueError("invalid provider request")


class NaturalCompletionProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="natural-completion")
        self.calls = 0

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        return LLMResponse(content="已完成当前分析，但没有提交结构化结果。")


class ToolSliceContinuationProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="tool-slice-continuation")
        self.calls = 0

    @staticmethod
    def submission_call(call_id: str) -> LLMToolCall:
        return LLMToolCall(
            id=call_id,
            name="submit_result",
            arguments={
                "kind": "module_submission",
                "module_id": "2.1",
                "unresolved_questions": [],
                "revision": 0,
                "revision_responses": [],
            },
        )

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        if self.calls == 1:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{submodule_id}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 已保存正文",
                            "evidence_ids": [],
                        },
                    )
                    for submodule_id in REPORT_TAXONOMY["2.1"].submodules
                ],
            )
        return LLMResponse(
            content="",
            tool_calls=[self.submission_call(f"submit-{self.calls}")],
        )


class ProductiveManyToolSlicesProvider(ToolSliceContinuationProvider):
    """Make durable progress beyond the former profile slice ceiling."""

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        submodule_ids = list(REPORT_TAXONOMY["2.1"].submodules)
        # The response immediately following a tool result is observed before
        # the inner loop reports its boundary, so each outer slice consumes a
        # pair of provider calls while only the first tool call is executed.
        slice_index = (self.calls - 1) // 2
        if slice_index < len(submodule_ids):
            submodule_id = submodule_ids[slice_index]
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{submodule_id}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 已保存正文",
                            "evidence_ids": [],
                        },
                    )
                ],
            )
        return LLMResponse(
            content="",
            tool_calls=[self.submission_call(f"submit-{self.calls}")],
        )


class CorrectionToolSliceContinuationProvider(LLMProvider):
    """Reach the tool boundary only after the runner asks for typed correction."""

    def __init__(self):
        super().__init__("test", model="correction-tool-slice-continuation")
        self.calls = 0
        self.correction = ""

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        if self.calls == 1:
            return LLMResponse(content="总编分析已完成，但尚未提交结构化结果。")
        if self.calls == 2:
            self.correction = messages[-1].content
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-after-correction-{submodule_id}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 已保存正文",
                            "evidence_ids": [],
                        },
                    )
                    for submodule_id in REPORT_TAXONOMY["2.1"].submodules
                ],
            )
        return LLMResponse(
            content="",
            tool_calls=[ToolSliceContinuationProvider.submission_call(f"submit-{self.calls}")],
        )


class RepeatingNoProgressToolProvider(LLMProvider):
    """Always request the same pure read so the outer harness must stop it."""

    def __init__(self):
        super().__init__("test", model="repeating-no-progress")
        self.calls = 0
        self.max_tokens_seen: list[int] = []

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        self.max_tokens_seen.append(max_tokens)
        return LLMResponse(
            content="",
            tool_calls=[
                LLMToolCall(
                    id=f"repeat-{self.calls}",
                    name="calculate",
                    arguments={"expression": "1 + 1"},
                )
            ],
        )


class TemplateSkillCorrectionProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="template-skill-correction")
        self.calls = 0
        self.inspection_result = ""

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        if self.calls == 1:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id="inspect-template",
                        name="inspect_document",
                        arguments={"path": "template.docx", "max_chars": 100_000},
                    )
                ],
            )
        if self.calls == 2:
            self.inspection_result = messages[-1].content
            return LLMResponse(content="已读取模板，现在开始提炼写作 Skill。")
        if self.calls == 3:
            assert "submission_correction" in messages[-1].content
            repeated_inspection = any(
                call.get("function", {}).get("name") == "inspect_document"
                for call in (tools or [])
                if isinstance(call, dict)
            )
            assert not repeated_inspection or "不要重新读取文件" in messages[-1].content
            description = "指导专业评估报告从证据限定推进到风险判断、跨章节综合、行动建议与图证叙事，并在写作和审查报告时使用。"
            parts = {
                skill_id: (
                    f"---\nname: report-template-{skill_id}\ndescription: {description}\n---\n\n"
                    f"# {skill_id} 模板方法\n\n"
                    + "先界定证据，再形成判断，说明原因影响并提出可验收行动；检查失败表现并执行职责内修正。" * 8
                )
                for skill_id in TEMPLATE_ROLE_SKILL_IDS
            }
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{part_id}",
                        name="write_result_part",
                        arguments={"part_id": part_id, "content": content},
                    )
                    for part_id, content in parts.items()
                ],
            )
        if self.calls == 4:
            root = "Work/runs/run-template-skill/drafts/template-skill-distillation/r0"

            def ref(part_id):
                return {"artifact_refs": [f"{root}/{part_id}.md"]}

            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id="submit-template-skill",
                        name="submit_result",
                        arguments={
                            "kind": "template_skill_submission",
                            "name": "report-template-role-skills",
                            "skills": {
                                skill_id: ref(skill_id)
                                for skill_id in TEMPLATE_ROLE_SKILL_IDS
                            },
                        },
                    )
                ],
            )
        return LLMResponse(content="已提交。")


def test_reporting_agent_runner_has_no_absolute_task_deadline(tmp_path: Path) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        NaturalCompletionProvider(),
        AgentDefaults(),
    )

    assert runner.timeout is None


class PartSubmissionProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="part-submission")
        self.calls = 0
        self.submodule_ids = list(REPORT_TAXONOMY["2.1"].submodules)

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        if self.calls == 1:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{index}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 分析正文",
                            "evidence_ids": [],
                        },
                    )
                    for index, submodule_id in enumerate(self.submodule_ids[:4])
                ],
            )
        if self.calls == 2:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{index + 4}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 分析正文",
                            "evidence_ids": [],
                        },
                    )
                    for index, submodule_id in enumerate(self.submodule_ids[4:])
                ],
            )
        if self.calls == 3:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id="list-parts",
                        name="list_result_parts",
                        arguments={},
                    )
                ],
            )
        if self.calls == 4:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id="submit-parts",
                        name="submit_result",
                        arguments={
                            "kind": "module_submission",
                            "module_id": "2.1",
                            "unresolved_questions": [],
                            "revision": 0,
                            "revision_responses": [],
                        },
                    )
                ],
            )
        return LLMResponse(content="分段正文已经提交。")


class RepeatedInvalidSubmissionProvider(LLMProvider):
    """Repeats one invalid current-contract submission."""

    def __init__(self):
        super().__init__("test", model="repeated-invalid")
        self.calls = 0

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        return LLMResponse(
            content="",
            tool_calls=[
                LLMToolCall(
                    id=f"invalid-{self.calls}",
                    name="submit_result",
                    arguments={
                        "kind": "module_submission",
                        "module_id": "2.1",
                        "submodule_narratives": {"wrong": []},
                    },
                )
            ],
        )


class WrappedThenFlatSubmissionProvider(LLMProvider):
    """Retry one forbidden wrapper using the correction's flat example."""

    def __init__(self):
        super().__init__("test", model="wrapped-then-flat")
        self.calls = 0
        self.correction = ""

    @staticmethod
    def candidate() -> dict:
        return {
            "kind": "module_submission",
            "module_id": "2.1",
            "unresolved_questions": [],
            "revision": 0,
            "revision_responses": [],
        }

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        self.calls += 1
        submit_definition = next(
            tool for tool in (tools or []) if tool["name"] == "submit_result"
        )
        submission_schema = submit_definition["input_schema"]
        assert "tool arguments are this submission object itself" in submission_schema[
            "description"
        ]
        assert "payload" not in submission_schema["properties"]
        if self.calls == 1:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id=f"write-{submodule_id}",
                        name="write_result_part",
                        arguments={
                            "part_id": submodule_id,
                            "content": f"{submodule_id} 分析正文",
                            "evidence_ids": [],
                        },
                    )
                    for submodule_id in REPORT_TAXONOMY["2.1"].submodules
                ],
            )
        if self.calls == 2:
            return LLMResponse(
                content="",
                tool_calls=[
                    LLMToolCall(
                        id="wrapped-submit",
                        name="submit_result",
                        arguments={"payload": json.dumps(self.candidate())},
                    )
                ],
            )
        self.correction = messages[-1].content
        assert '"field": "payload"' in self.correction
        assert '"received_type": "string"' in self.correction
        assert '"kind": "module_submission"' in self.correction
        assert "Remove the outer payload property" in self.correction
        return LLMResponse(
            content="",
            tool_calls=[
                LLMToolCall(
                    id="flat-submit",
                    name="submit_result",
                    arguments=self.candidate(),
                )
            ],
        )


@pytest.mark.asyncio
async def test_reporting_agent_runner_uses_real_isolated_loop_and_can_finish_without_research(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    workflow_messages: list[UserMessage] = []

    def capture_workflow_message(message: UserMessage) -> None:
        if message.source == "workflow":
            workflow_messages.append(message)

    bus.subscribe(UserMessage, capture_workflow_message)
    provider = MeasuredDirectSubmissionProvider()
    agents = load_packaged_agents()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=5), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-test",
        agent_id="module-2.1-specialist",
        objective="完成 2.1 分析",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            agents["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-test",
            session_key="specialist-2.1",
        )
    finally:
        await runner.close_workflow("wf-test")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert workflow_messages
    assert workflow_messages[0].provider_stream_idle_timeout_seconds == 600.0
    assert len(provider.system_prompts) == 2
    memory_ref = "Work/runs/run-test/context/module-2.1-evidence-memory.json"
    assert memory_ref in provider.task_messages[0]
    assert (tmp_path / memory_ref).is_file()
    assert isinstance(result.payload, ModuleSubmission)
    assert "负荷率必须保留计算口径" in provider.system_prompts[0]
    assert "剩余电流大于 10A" not in provider.system_prompts[0]
    assert provider.max_tokens_seen == [32768, 32768]
    assert (tmp_path / "Work/runs/run-test/results/module-2.1.json").is_file()
    conversation_blobs = [
        path
        for path in (tmp_path / "Work/content/sha256").glob("*/*/*")
        if path.is_file()
    ]
    assert conversation_blobs == []
    conversation_manifests = list(
        (tmp_path / "Work/runs/run-test/agent-conversations").glob("*.json")
    )
    assert conversation_manifests
    restart_state = json.loads(conversation_manifests[0].read_text(encoding="utf-8"))
    assert restart_state["manifest_version"] == 4
    assert restart_state["encoding"] == "identity+refs"
    assert "compaction_summary" not in restart_state
    assert "restart_state_sha256" not in restart_state
    assert "task_boundaries" not in restart_state
    assert restart_state["business_refs"]["completed_result_ref"] == (
        "Work/runs/run-test/results/module-2.1.json"
    )
    summaries = list((tmp_path / "Work/runs/run-test/session-summaries").glob("*.json"))
    assert len(summaries) == 1
    summary = json.loads(summaries[0].read_text(encoding="utf-8"))
    assert summary["agent_id"] == "module-2.1-specialist"
    assert summary["status"] == "completed"
    assert summary["context_only"] is True
    usage_rows = [
        json.loads(line)
        for line in (tmp_path / ".manyselves/usage/run-test.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert len(usage_rows) == 2
    assert all(row["status"] == "success" for row in usage_rows)
    assert all(
        row["context_manifest_ref"]
        and (tmp_path / row["context_manifest_ref"]).is_file()
        for row in usage_rows
    )
    assert all(
        row["provider_call_id"] == Path(row["context_manifest_ref"]).stem
        for row in usage_rows
    )
    assert len({row["provider_call_id"] for row in usage_rows}) == len(usage_rows)
    assert all(
        json.loads(
            (tmp_path / row["context_manifest_ref"]).read_text(encoding="utf-8")
        )["provider_call_id"]
        == row["provider_call_id"]
        for row in usage_rows
    )
    assert all(
        row["uncached_input_tokens"]
        == max(
            0,
            row["input_tokens"]
            - row["cached_input_tokens"]
            - row["cache_write_input_tokens"],
        )
        for row in usage_rows
    )
    assert all(
        row["request_metric_source"] == "provider_adapter_payload"
        for row in usage_rows
    )
    assert [row["turn_kind"] for row in usage_rows] == [
        "task_initial",
        "tool_followup",
    ]
    assert all(len(row["execution_profile_sha256"]) == 64 for row in usage_rows)
    assert all(row["resolved_provider_route"] == "inherit" for row in usage_rows)
    assert all(row["resolved_model"] == "scripted" for row in usage_rows)
    for row in usage_rows:
        for field in (
            "queue_wait_ms",
            "context_build_ms",
            "serialization_ms",
            "provider_active_ms",
            "tool_time_ms",
        ):
            assert row[field] >= 0
    for row in usage_rows:
        provider_manifest = json.loads(
            (tmp_path / row["context_manifest_ref"]).read_text(
                encoding="utf-8"
            )
        )
        assert provider_manifest["provider_payload_status"] == "observed"
        assert (
            provider_manifest["provider_payload"]["request_sha256"]
            == row["request_fingerprint"]
        )


@pytest.mark.asyncio
async def test_max_tokens_continues_same_identity_without_submission_correction(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    internal_messages: list[UserMessage] = []

    def capture_internal(message: UserMessage) -> None:
        if message.internal:
            internal_messages.append(message)

    bus.subscribe(UserMessage, capture_internal)
    provider = MaxTokensThenDirectSubmissionProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=5), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-max-tokens-continuation",
        agent_id="module-2.1-specialist",
        objective="完成 2.1 分析",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-max-tokens-continuation",
            session_key="specialist-2.1",
        )
    finally:
        await runner.close_workflow("wf-max-tokens-continuation")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert provider.truncated_once is True
    assert all(value == 32768 for value in provider.max_tokens_seen)
    continuation_messages = [
        message.content
        for message in internal_messages
        if "same_identity_continuation" in message.content
    ]
    assert len(continuation_messages) == 1
    assert "max_tokens" in continuation_messages[0]
    assert "submission_correction" not in continuation_messages[0]
    continuation_state = json.loads(
        (
            tmp_path
            / "Work/runs/run-max-tokens-continuation/continuations/"
            f"module-2.1/{envelope.task_attempt_id}.json"
        ).read_text(encoding="utf-8")
    )
    assert continuation_state["status"] == "typed_result"
    assert continuation_state["events"][0]["turn_kind"] == (
        "max_tokens_continuation"
    )


def test_evidence_auditor_uses_long_reasoning_output_limit(tmp_path: Path) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        ModuleReviewSubmissionProvider(),
        AgentDefaults(max_tokens=8192),
    )
    definition = load_packaged_agents()["evidence-auditor"]
    envelope = TaskEnvelope(
        task_id="module-2.4-review-r0",
        run_id="run-auditor-limit",
        agent_id=definition.id,
        objective="审计模块 2.4",
        allowed_outputs=["module_review_finding_submission"],
    )

    assert runner._loop_config(definition, envelope).max_tokens == 32768


def test_cross_reviewer_uses_long_reasoning_output_limit(tmp_path: Path) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        ModuleReviewSubmissionProvider(),
        AgentDefaults(max_tokens=8192),
    )
    definition = load_packaged_agents()["cross-module-reviewer"]
    envelope = TaskEnvelope(
        task_id="cross-review-r0",
        run_id="run-cross-limit",
        agent_id=definition.id,
        objective="执行跨模块审查",
        allowed_outputs=["cross_review_finding_submission"],
    )

    assert runner._loop_config(definition, envelope).max_tokens == 32768
    assert runner._loop_config(definition, envelope).max_tool_iterations == 28


def test_chief_editor_uses_long_synthesis_output_limit(tmp_path: Path) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        ModuleReviewSubmissionProvider(),
        AgentDefaults(max_tokens=8192),
    )
    definition = load_packaged_agents()["chief-editor"]
    envelope = TaskEnvelope(
        task_id="chief-edit",
        run_id="run-chief-limit",
        agent_id=definition.id,
        objective="执行总编整合",
        allowed_outputs=["edited_report_submission"],
    )

    config = runner._loop_config(definition, envelope)
    assert config.max_tokens == 32768
    assert config.max_tool_iterations == 28


@pytest.mark.asyncio
async def test_final_auditor_open_artifact_honors_requested_page_size(
    tmp_path: Path,
) -> None:
    run_id = "run-final-auditor-page"
    artifact_ref = f"Work/runs/{run_id}/reviews/final-review-input-r0.json"
    artifact = tmp_path / artifact_ref
    artifact.parent.mkdir(parents=True)
    artifact.write_text("x" * 1000, encoding="utf-8")
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        ModuleReviewSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["chief-editor-auditor"]
    envelope = TaskEnvelope(
        task_id="final-review-r0",
        run_id=run_id,
        agent_id=definition.id,
        objective="执行最终审计",
        input_refs=[artifact_ref],
        allowed_outputs=["final_review_finding_submission"],
    )
    registry = runner._tools(
        definition,
        envelope,
        "session-final-auditor-page",
        "workflow-final-auditor-page",
    )
    tool = registry.get("open_artifact")

    assert tool is not None
    page = await tool(ref=artifact_ref, offset=0, limit=500)
    assert page["returned"] == 500
    assert len(page["content"]) == 500


@pytest.mark.asyncio
async def test_artifact_delivery_modes_enforce_tool_read_boundary(
    tmp_path: Path,
) -> None:
    run_id = "run-delivery-modes"
    contract_ref = f"Work/runs/{run_id}/context/authoring.json"
    prior_ref = f"Work/runs/{run_id}/modules/2.4-r0.json"
    summary_ref = f"Work/runs/{run_id}/context/session-summary.json"
    for ref, content in (
        (
            contract_ref,
            ModuleAuthoringInput(
                run_id=run_id,
                module_id="2.4",
                revision=0,
                required_submodule_ids=list(REPORT_TAXONOMY["2.4"].submodules),
                coverage_ref="coverage.json",
                evidence_ref="evidence.jsonl",
                manifest_ref="manifest.json",
                knowledge_ref="knowledge.md",
            ).model_dump_json(),
        ),
        (prior_ref, '{"revision":0}'),
        (summary_ref, '{"summary":"retained context"}'),
    ):
        target = tmp_path / ref
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["module-2.4-specialist"]
    envelope = TaskEnvelope(
        task_id="module-2.4-revision-r1",
        run_id=run_id,
        agent_id=definition.id,
        objective="执行定向修订",
        input_refs=[contract_ref],
        allowed_outputs=["module_submission"],
        prior_result_ref=prior_ref,
        context_summary_refs=[summary_ref],
        artifact_delivery_modes={
            contract_ref: "inline",
            prior_ref: "hash_retained",
            summary_ref: "reference",
        },
        input_contract_kind="module_authoring_input",
        input_contract_ref=contract_ref,
    )
    registry = runner._tools(
        definition,
        envelope,
        "session-delivery-modes",
        "workflow-delivery-modes",
    )
    opener = registry.get("open_artifact")

    assert opener is not None
    assert "retained context" in (await opener(ref=summary_ref))["content"]
    with pytest.raises(PermissionError, match="not delivered by reference"):
        await opener(ref=contract_ref)
    with pytest.raises(PermissionError, match="not delivered by reference"):
        await opener(ref=prior_ref)


@pytest.mark.asyncio
async def test_reporting_identity_keeps_one_stable_session_across_workflow_turns(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = DirectSubmissionProvider()
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_iterations=5),
        timeout=None,
    )
    definition = load_packaged_agents()["module-2.1-specialist"]
    first = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-stable-identity",
        agent_id=definition.id,
        objective="完成初稿",
        allowed_outputs=["module_submission"],
    )
    second = first.model_copy(update={"objective": "根据审查继续修订", "revision": 1})

    try:
        first_result = await runner.run(
            definition,
            first,
            [],
            workflow_id="workflow-stable-identity",
            session_key="initial-draft",
        )
        second_result = await runner.run(
            definition,
            second,
            [],
            workflow_id="workflow-stable-identity",
            session_key="revision-r1",
        )

        assert first_result.session_id == second_result.session_id
        assert len(runner._sessions) == 1
        assert next(iter(runner._sessions))[1] == definition.id
        assert runner.timeout is None
        identities = json.loads(
            (tmp_path / "Work/runs/run-stable-identity/agent-identities.json").read_text(
                encoding="utf-8"
            )
        )
        assert identities["created_by"] == "main"
        assert set(identities["identities"]) == {definition.id}
        identity = identities["identities"][definition.id]
        assert identity["session_id"] == first_result.session_id
        assert identity["first_task_id"] == "module-2.1"
        assert identity["last_revision"] == 1
        assert identity["status"] == "waiting"
        manifests = [
            json.loads(path.read_text(encoding="utf-8"))
            for path in sorted(
                (
                    tmp_path
                    / "Work/runs/run-stable-identity/context-manifests"
                ).glob("module-2.1-r*-session-*.json")
            )
        ]
        assert [item["revision"] for item in manifests] == [0, 1]
        assert manifests[0]["session_id"] == manifests[1]["session_id"]
        first_components = {
            item["kind"]: item for item in manifests[0]["prompt_components"]
        }
        second_components = {
            item["kind"]: item for item in manifests[1]["prompt_components"]
        }
        assert first_components["system_prompt"]["repeated_content"] is False
        assert second_components["system_prompt"]["repeated_content"] is True
        assert second_components["task_message"]["repeated_content"] is False
        assert {
            submodule_id
            for skill in manifests[1]["module_skills"]
            for submodule_id in skill["submodules"]
        }.issubset(set(REPORT_TAXONOMY["2.1"].submodules))
        assert all("content" not in item for item in manifests[1]["prompt_components"])
        provider_manifests = [
            json.loads(path.read_text(encoding="utf-8"))
            for path in sorted(
                (
                    tmp_path
                    / "Work/runs/run-stable-identity/context-manifests/provider-calls"
                ).glob("module-2.1-r*-session-*.json")
            )
        ]
        assert len(provider_manifests) == 4
        assert [item["provider_call_index"] for item in provider_manifests] == [1, 2, 1, 2]
        assert [item["phase"] for item in provider_manifests] == [
            "initial",
            "tool_followup",
            "initial",
            "tool_followup",
        ]
        assert all(item["attempt"] == 1 for item in provider_manifests)
        assert all(len(item["request_sha256"]) == 64 for item in provider_manifests)
        initial_provider_manifests = [
            item for item in provider_manifests if item["phase"] == "initial"
        ]
        assert [item["message_count"] for item in initial_provider_manifests][0] == 2
        assert [item["message_count"] for item in initial_provider_manifests][1] > 2
        assert [
            item["provider_message_count"] for item in initial_provider_manifests
        ][0] == 2
        assert [
            item["provider_message_count"] for item in initial_provider_manifests
        ][1] > 2
        assert [
            item["logical_task_message_count"] for item in initial_provider_manifests
        ][0] == 2
        assert [
            item["logical_task_message_count"] for item in initial_provider_manifests
        ][1] > 2
        assert all(
            "content" not in message
            for manifest in provider_manifests
            for message in manifest["messages"]
        )
        second_initial = provider.message_snapshots[2]
        boundary = next(
            message
            for message in second_initial
            if message.role == "user" and "<task_boundary>" in message.content
        )
        assert "根据审查继续修订" in boundary.content
        assert "<allowed_output>module_submission</allowed_output>" not in boundary.content
        assert "<input_contract" not in boundary.content
        assert second_initial[0].content == provider.message_snapshots[0][0].content
    finally:
        await runner.close_workflow("workflow-stable-identity")
        bus.shutdown()
        await bus_task
    assert runner._routers == {}


@pytest.mark.asyncio
async def test_new_process_style_runner_recovers_persisted_attempt_without_provider_call(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    first_provider = DirectSubmissionProvider()
    definition = load_packaged_agents()["module-2.1-specialist"]
    envelope = TaskEnvelope(
        task_id="module-2.1-recoverable",
        run_id="run-attempt-recovery",
        agent_id=definition.id,
        objective="完成可恢复的模块任务",
        allowed_outputs=["module_submission"],
    )
    first_runner = ReportingAgentRunner(
        tmp_path,
        bus,
        first_provider,
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    def fail_manifest_finalization(_record: dict) -> None:
        raise OSError("injected provider manifest finalization failure")

    monkeypatch.setattr(
        first_runner,
        "_finalize_provider_call_manifest",
        fail_manifest_finalization,
    )
    try:
        first = await first_runner.run(
            definition,
            envelope,
            [],
            workflow_id="workflow-attempt-recovery",
        )
        await first_runner.close_workflow("workflow-attempt-recovery")
        provider_manifests = [
            json.loads(path.read_text(encoding="utf-8"))
            for path in (
                tmp_path
                / "Work/runs/run-attempt-recovery/context-manifests/provider-calls"
            ).glob("*.json")
        ]
        assert provider_manifests
        assert all(
            item["provider_payload_status"] == "pending"
            for item in provider_manifests
        )

        recovery_provider = DirectSubmissionProvider()
        recovered_runner = ReportingAgentRunner(
            tmp_path,
            bus,
            recovery_provider,
            AgentDefaults(max_tool_iterations=5),
            timeout=5,
        )
        recovered = await recovered_runner.run(
            definition,
            envelope.model_copy(),
            [],
            workflow_id="workflow-attempt-recovery",
        )

        assert recovered == first
        assert recovery_provider.max_tokens_seen == []
        assert recovered_runner._sessions == {}
    finally:
        bus.shutdown()
        await bus_task


@pytest.mark.asyncio
async def test_module_auditors_use_isolated_per_module_sessions_and_reuse_on_review(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        ModuleReviewSubmissionProvider(),
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    definition = load_packaged_agents()["evidence-auditor"]
    audit_21 = TaskEnvelope(
        task_id="audit-2.1-r0-a0",
        run_id="run-auditor-identities",
        agent_id=definition.id,
        objective="审计模块 2.1",
        allowed_outputs=["module_review_finding_submission"],
    )
    audit_22 = audit_21.model_copy(
        update={"task_id": "audit-2.2-r0-a0", "objective": "审计模块 2.2"}
    )
    audit_21_review = audit_21.model_copy(
        update={"task_id": "audit-2.1-r1-a1", "objective": "复审模块 2.1", "revision": 1}
    )

    try:
        first_21 = await runner.run(
            definition,
            audit_21,
            [],
            workflow_id="workflow-auditor-identities",
            session_key="auditor-2.1",
        )
        first_22 = await runner.run(
            definition,
            audit_22,
            [],
            workflow_id="workflow-auditor-identities",
            session_key="auditor-2.2",
        )
        second_21 = await runner.run(
            definition,
            audit_21_review,
            [],
            workflow_id="workflow-auditor-identities",
            session_key="auditor-2.1",
        )

        assert first_21.session_id != first_22.session_id
        assert second_21.session_id == first_21.session_id
        assert set(runner._sessions) == {
            ("workflow-auditor-identities", "auditor-2.1"),
            ("workflow-auditor-identities", "auditor-2.2"),
        }
        identities = json.loads(
            (
                tmp_path
                / "Work/runs/run-auditor-identities/agent-identities.json"
            ).read_text(encoding="utf-8")
        )
        assert set(identities["identities"]) == {
            "auditor-2.1",
            "auditor-2.2",
        }
    finally:
        await runner.close_workflow("workflow-auditor-identities")
        bus.shutdown()
        await bus_task


@pytest.mark.asyncio
async def test_reporting_agent_runner_fails_immediately_when_isolated_loop_errors(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        NonRetryableFailureProvider(),
        AgentDefaults(max_tool_iterations=5),
        timeout=600,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-fail-fast",
        agent_id="module-2.1-specialist",
        objective="分析模块 2.1",
        allowed_outputs=["module_submission"],
    )

    try:
        with pytest.raises(RuntimeError, match="数据或参数校验失败"):
            await asyncio.wait_for(
                runner.run(
                    load_packaged_agents()["module-2.1-specialist"],
                    envelope,
                    [],
                    workflow_id="wf-fail-fast",
                ),
                timeout=1,
            )
    finally:
        await runner.close_workflow("wf-fail-fast")
        bus.shutdown()
        await bus_task


@pytest.mark.asyncio
async def test_reporting_agent_runner_reprompts_untyped_completion_once_then_stops(
    tmp_path: Path,
) -> None:
    template_ref = "template.docx"
    Document().save(tmp_path / template_ref)
    contract_ref = _write_template_contract(
        tmp_path,
        run_id="run-natural-completion",
        template_ref=template_ref,
    )
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = NaturalCompletionProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=5), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="template-skill-distillation",
        run_id="run-natural-completion",
        agent_id="template-distiller",
        objective="提炼模板写作 Skill",
        input_refs=[contract_ref, template_ref],
        allowed_outputs=["template_skill_submission"],
        input_contract_kind="template_distillation_input",
        input_contract_ref=contract_ref,
    )
    published: list[AgentResultMessage] = []
    published_event = asyncio.Event()

    def capture(message: AgentResultMessage) -> None:
        if message.run_id == envelope.run_id and message.task_id == envelope.task_id:
            published.append(message)
            published_event.set()

    bus.subscribe(AgentResultMessage, capture)
    requeue_runner: ReportingAgentRunner | None = None
    requeue_provider: NaturalCompletionProvider | None = None
    requeued = None

    try:
        result = await runner.run(
            load_packaged_agents()["template-distiller"],
            envelope,
            [],
            workflow_id="wf-natural-completion",
        )
        await asyncio.wait_for(published_event.wait(), timeout=1)
        await runner.close_workflow("wf-natural-completion")
        requeue_provider = NaturalCompletionProvider()
        requeue_runner = ReportingAgentRunner(
            tmp_path,
            bus,
            requeue_provider,
            AgentDefaults(max_tool_iterations=5),
            timeout=5,
        )
        requeued = await requeue_runner.run(
            load_packaged_agents()["template-distiller"],
            envelope.model_copy(),
            [],
            workflow_id="wf-natural-completion",
        )
    finally:
        if requeue_runner is not None:
            await requeue_runner.close_workflow("wf-natural-completion")
        await runner.close_workflow("wf-natural-completion")
        bus.shutdown()
        await bus_task

    assert provider.calls == 2
    assert requeue_provider is not None
    assert requeue_provider.calls == 2
    assert result.status is AgentRunStatus.INCOMPLETE
    assert requeued is not None
    assert requeued.status is AgentRunStatus.INCOMPLETE
    assert result.raw_output == "已完成当前分析，但没有提交结构化结果。"
    assert result.reason == "agent ended without a typed submission"
    assert len(published) == 2
    assert published[0].status == "incomplete"
    assert published[0].content == result.raw_output
    persisted = json.loads(
        (
            tmp_path
            / "Work/runs/run-natural-completion/results/template-skill-distillation.json"
        ).read_text(encoding="utf-8")
    )
    assert persisted["raw_output"] == result.raw_output


@pytest.mark.asyncio
async def test_tool_iteration_boundary_continues_same_identity_until_typed_submission(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = ToolSliceContinuationProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=1), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-continuation",
        agent_id="module-2.1-specialist",
        objective="分段完成模块 2.1",
        allowed_outputs=["module_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )

    try:
        definition = load_packaged_agents()["module-2.1-specialist"].model_copy(
            update={"max_turns": 1}
        )
        result = await runner.run(
            definition,
            envelope,
            [],
            workflow_id="wf-continuation",
        )
        session_id = runner._sessions[("wf-continuation", "module-2.1-specialist")][1]
    finally:
        await runner.close_workflow("wf-continuation")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert result.session_id == session_id
    assert provider.calls == 3
    turn_kinds = [
        row["turn_kind"]
        for row in UsageLedger(tmp_path, "run-continuation").rows()
    ]
    assert "tool_slice_continuation" in turn_kinds
    continuation_state = json.loads(
        (
            tmp_path
            / "Work/runs/run-continuation/continuations/"
            f"module-2.1/{envelope.task_attempt_id}.json"
        ).read_text(encoding="utf-8")
    )
    assert continuation_state["status"] == "typed_result"
    assert continuation_state["events"][0]["turn_kind"] == (
        "tool_slice_continuation"
    )
    assert (
        tmp_path
        / "Work/runs/run-continuation/drafts/module-2.1/r0/2.1.1.md"
    ).is_file()


@pytest.mark.asyncio
async def test_productive_tool_slices_have_no_profile_count_limit(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = ProductiveManyToolSlicesProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=1), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        task_attempt_id="attempt-many-productive-slices",
        run_id="run-many-productive-slices",
        agent_id="module-2.1-specialist",
        objective="逐段完成模块 2.1",
        allowed_outputs=["module_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )

    try:
        definition = load_packaged_agents()["module-2.1-specialist"].model_copy(
            update={"max_turns": 1}
        )
        result = await runner.run(
            definition,
            envelope,
            [],
            workflow_id="wf-many-productive-slices",
        )
    finally:
        await runner.close_workflow("wf-many-productive-slices")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    continuation_paths = list(
        (
            tmp_path
            / "Work/runs/run-many-productive-slices/continuations/module-2.1"
        ).glob("*.json")
    )
    assert len(continuation_paths) == 1
    state = json.loads(continuation_paths[0].read_text(encoding="utf-8"))
    assert "tool_slice_continuation" not in state["limits"]
    assert state["continuation_counts"]["tool_slice_continuation"] > 3
    assert state["status"] == "typed_result"


@pytest.mark.asyncio
async def test_repeated_no_progress_continuation_stops_at_profile_harness_boundary(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = RepeatingNoProgressToolProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=1), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        task_attempt_id="attempt-no-progress",
        run_id="run-no-progress-continuation",
        agent_id="module-2.1-specialist",
        objective="验证无进展 continuation 的有限终止",
        allowed_outputs=["module_submission"],
    )

    try:
        definition = load_packaged_agents()["module-2.1-specialist"].model_copy(
            update={"max_turns": 1}
        )
        result = await runner.run(
            definition,
            envelope,
            [],
            workflow_id="wf-no-progress-continuation",
        )
    finally:
        await runner.close_workflow("wf-no-progress-continuation")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.INCOMPLETE
    assert result.reason.startswith("continuation harness stopped")
    assert provider.calls == 4
    assert all(value == 32768 for value in provider.max_tokens_seen)
    state = json.loads(
        (
            tmp_path
            / "Work/runs/run-no-progress-continuation/continuations/"
            f"module-2.1/{envelope.task_attempt_id}.json"
        ).read_text(encoding="utf-8")
    )
    assert state["status"] == "stopped"
    assert state["stop_reason"] == "repeated_no_progress"
    assert state["stop_turn_kind"] == "tool_slice_continuation"
    assert state["continuation_counts"]["tool_slice_continuation"] == 1
    assert state["events"][-1]["decision"] == "stop"
    turn_kinds = [
        row["turn_kind"]
        for row in UsageLedger(tmp_path, envelope.run_id).rows()
    ]
    assert turn_kinds == [
        "task_initial",
        "tool_followup",
        "tool_slice_continuation",
        "tool_followup",
    ]


@pytest.mark.asyncio
async def test_submission_correction_boundary_continues_until_typed_submission(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = CorrectionToolSliceContinuationProvider()
    internal_messages: list[UserMessage] = []

    def capture_internal(message: UserMessage) -> None:
        if message.internal:
            internal_messages.append(message)

    bus.subscribe(UserMessage, capture_internal)
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=1), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-correction-continuation",
        agent_id="module-2.1-specialist",
        objective="完成模块并提交 typed result",
        allowed_outputs=["module_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )

    try:
        definition = load_packaged_agents()["module-2.1-specialist"].model_copy(
            update={"max_turns": 1}
        )
        result = await runner.run(
            definition,
            envelope,
            [],
            workflow_id="wf-correction-continuation",
        )
        session_id = runner._sessions[
            ("wf-correction-continuation", "module-2.1-specialist")
        ][1]
    finally:
        await runner.close_workflow("wf-correction-continuation")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert result.session_id == session_id
    assert provider.calls == 4
    assert "evidence_ids" in provider.correction
    assert "submit_result 只发送 schema 声明的字段" in provider.correction
    assert "先在内部压缩措辞" not in provider.correction
    assert result.raw_output != "AGENT_TURN_CONTINUATION_REQUIRED"
    turn_kinds = [
        row["turn_kind"]
        for row in UsageLedger(tmp_path, "run-correction-continuation").rows()
    ]
    assert "submission_correction" in turn_kinds
    assert turn_kinds.count("submission_correction") == 2
    assert "tool_slice_continuation" not in turn_kinds
    continuation_state = json.loads(
        (
            tmp_path
            / "Work/runs/run-correction-continuation/continuations/"
            f"module-2.1/{envelope.task_attempt_id}.json"
        ).read_text(encoding="utf-8")
    )
    assert continuation_state["status"] == "typed_result"
    assert continuation_state["events"][0]["turn_kind"] == (
        "submission_correction"
    )
    submission_messages = [
        message
        for message in internal_messages
        if "submission_correction" in message.content
        or "same_identity_continuation" in message.content
    ]
    assert submission_messages
    assert all(
        message.provider_stream_idle_timeout_seconds == 600.0
        for message in submission_messages
    )


@pytest.mark.asyncio
async def test_template_distillation_gets_full_inspection_and_one_submission_correction(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_heading("模板分析", level=1)
    document.add_paragraph("前段" + ("证据—判断—原因—影响—建议。" * 700))
    document.add_paragraph("模板尾部标记：TAIL-CONTENT-MUST-BE-VISIBLE")
    document.save(tmp_path / "template.docx")
    contract_ref = _write_template_contract(
        tmp_path,
        run_id="run-template-skill",
        template_ref="template.docx",
    )
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = TemplateSkillCorrectionProvider()
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_result_chars=6000, max_tool_iterations=6),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="template-skill-distillation",
        run_id="run-template-skill",
        agent_id="template-distiller",
        objective="蒸馏模板写作 Skill",
        input_refs=[contract_ref, "template.docx"],
        allowed_outputs=["template_skill_submission"],
        allowed_tools=[
            "inspect_document",
            "write_result_part",
            "list_result_parts",
            "submit_result",
            "report_blocked",
        ],
        input_contract_kind="template_distillation_input",
        input_contract_ref=contract_ref,
    )

    try:
        result = await runner.run(
            load_packaged_agents()["template-distiller"],
            envelope,
            ["template.docx"],
            workflow_id="wf-template-skill",
        )
    finally:
        await runner.close_workflow("wf-template-skill")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert result.payload.boundary_manifest.policy_version == 1
    assert set(result.payload.boundary_manifest.transferred_categories) == {
        "analysis_method",
        "synthesis_method",
        "visual_method",
        "quality_check",
    }
    assert provider.calls == 4
    assert "TAIL-CONTENT-MUST-BE-VISIBLE" in provider.inspection_result
    assert '"full_result_ref"' not in provider.inspection_result
    for part_id in TEMPLATE_ROLE_SKILL_IDS:
        assert (
            tmp_path
            / f"Work/runs/run-template-skill/drafts/template-skill-distillation/r0/{part_id}.md"
        ).is_file()


@pytest.mark.asyncio
async def test_reporting_agent_runner_writes_parts_then_returns_materialized_result(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = PartSubmissionProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=8), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-parts-e2e",
        agent_id="module-2.1-specialist",
        objective="分段完成 2.1 分析",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-parts-e2e",
        )
    finally:
        await runner.close_workflow("wf-parts-e2e")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert provider.calls == 4
    assert isinstance(result.payload, ModuleSubmission)
    assert result.payload.markdown == compose_module_markdown(
        "2.1",
        {
            submodule_id: f"{submodule_id} 分析正文"
            for submodule_id in provider.submodule_ids
        },
    )
    for submodule_id in provider.submodule_ids:
        assert (
            tmp_path
            / f"Work/runs/run-parts-e2e/drafts/module-2.1/r0/{submodule_id}.md"
        ).is_file()


@pytest.mark.asyncio
async def test_reporting_usage_thresholds_do_not_stop_provider_attempts(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = PartSubmissionProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=8), timeout=5
    )
    budget = ReportingRunBudget(tmp_path, "run-parts-e2e", 2, 100_000)
    runner.set_provider_attempt_guard(budget.acquire_provider_attempt)
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-parts-e2e",
        agent_id="module-2.1-specialist",
        objective="在预算内分段完成 2.1 分析",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-provider-budget",
        )
    finally:
        await runner.close_workflow("wf-provider-budget")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert provider.calls == 4
    assert len(UsageLedger(tmp_path, envelope.run_id).rows()) == 4
    assert budget.snapshot()["limits_enforced"] is False
    assert list(
        (tmp_path / "Work/runs/run-parts-e2e/drafts/module-2.1/r0").glob("*.md")
    )


def test_auditor_research_usage_is_telemetry_across_runner_instances(tmp_path: Path) -> None:
    envelope = TaskEnvelope(
        task_id="audit-2.4-r0",
        run_id="run-audit-guard",
        agent_id="evidence-auditor",
        objective="审计 2.4",
        allowed_outputs=["module_review_finding_submission"],
    )
    definition = load_packaged_agents()["evidence-auditor"]
    first = ReportingAgentRunner(
        tmp_path, MessageBus(), DirectSubmissionProvider(), AgentDefaults()
    )
    guard = first._reporting_research_guard(definition, envelope, "wf-audit-guard")
    assert guard is not None
    for _ in range(8):
        guard()

    restarted = ReportingAgentRunner(
        tmp_path, MessageBus(), DirectSubmissionProvider(), AgentDefaults()
    )
    resumed_guard = restarted._reporting_research_guard(
        definition, envelope, "wf-audit-guard"
    )
    assert resumed_guard is not None
    resumed_guard()
    usage = json.loads(
        (tmp_path / "Work/runs/run-audit-guard/research-tool-usage.json").read_text(
            encoding="utf-8"
        )
    )
    assert usage["evidence-auditor:audit-2.4-r0:r0"] == 9


@pytest.mark.asyncio
async def test_reporting_agent_runner_stops_after_repeated_identical_submission_validation_error(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = RepeatedInvalidSubmissionProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=10), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-invalid-submission",
        agent_id="module-2.1-specialist",
        objective="完成模块 2.1",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-invalid-submission",
        )
    finally:
        await runner.close_workflow("wf-invalid-submission")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.FAILED
    assert "the same validation defect was repeated" in (result.reason or "")
    assert provider.calls == 2


@pytest.mark.asyncio
async def test_reporting_agent_runner_rejects_wrapper_then_accepts_flat_submission(
    tmp_path: Path,
) -> None:
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = WrappedThenFlatSubmissionProvider()
    runner = ReportingAgentRunner(
        tmp_path, bus, provider, AgentDefaults(max_tool_iterations=10), timeout=5
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-string-correction",
        agent_id="module-2.1-specialist",
        objective="完成模块 2.1",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-string-correction",
        )
    finally:
        await runner.close_workflow("wf-string-correction")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert result.payload is not None
    assert result.payload.module_id == "2.1"
    assert provider.calls == 3
    assert provider.correction


def test_conversation_trace_persists_only_bounded_restart_state(
    tmp_path: Path,
) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    envelope = TaskEnvelope(
        task_id="audit-2.1",
        run_id="run-compressed-trace",
        agent_id="evidence-auditor",
        objective="测试对话压缩",
        allowed_outputs=["module_submission"],
        input_contract_kind="module_review_input",
        input_contract_ref="Work/runs/run-compressed-trace/reviews/audit-2.1.json",
        input_refs=["Work/runs/run-compressed-trace/reviews/audit-2.1.json"],
        target_submodule_ids=["2.1.1"],
    )
    long_content = "重复的长对话正文。" * 2000
    loop = SimpleNamespace(
        _conversation_history=[
            LLMMessage(role="user", content=long_content),
            LLMMessage(role="assistant", content="已处理。"),
        ]
    )

    manifest_path = runner._save_conversation_trace(
        loop,
        envelope,
        "module-2.1-specialist--session-test",
        "session-test",
        status="waiting",
    )

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["manifest_version"] == 4
    assert manifest["encoding"] == "identity+refs"
    assert (
        manifest["transcript_semantics"]
        == "durable_identity_reference_state_v1"
    )
    assert "messages" not in manifest
    assert "transcript_ref" not in manifest
    assert "compaction_summary" not in manifest
    assert "restart_state_sha256" not in manifest
    assert "task_boundaries" not in manifest
    assert "active_task_identity" not in manifest
    assert "input_contract_payload" not in manifest_path.read_text(encoding="utf-8")
    assert manifest["business_refs"]["attention_scope_ref"] == (
        "Work/runs/run-compressed-trace/reviews/audit-2.1.json"
    )
    assert manifest["identity_state"]["attention_scope_ref"] == (
        "Work/runs/run-compressed-trace/reviews/audit-2.1.json"
    )
    assert "target_submodule_ids" not in manifest["identity_state"]
    decoded = load_conversation_trace(tmp_path, manifest_path)
    assert decoded["messages"] == []
    assert decoded["status"] == "waiting"
    assert long_content not in manifest_path.read_text(encoding="utf-8")

    unsupported = tmp_path / "Work/runs/run-compressed-trace/unsupported.json"
    unsupported.write_text(
        json.dumps({"status": "completed", "messages": [{"role": "user"}]}),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="unsupported conversation identity state"):
        load_conversation_trace(tmp_path, unsupported)


@pytest.mark.parametrize(
    ("agent_id", "task_id", "contract_kind", "expects_attention_scope"),
    [
        ("chief-editor", "chief-chapter-1", "chief_chapter_input", False),
        ("chief-editor", "aggregate-existing", "aggregate_editor_input", False),
        (
            "chief-editor-auditor",
            "final-chapter-1-r0",
            "final_chapter_lane_input",
            True,
        ),
        (
            "cross-module-reviewer",
            "cross-owner-2.4-r0-initial",
            "cross_owner_input",
            True,
        ),
    ],
)
def test_editor_chief_final_and_cross_use_identity_reference_persistence(
    tmp_path: Path,
    agent_id: str,
    task_id: str,
    contract_kind: str,
    expects_attention_scope: bool,
) -> None:
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    contract_ref = f"Work/runs/run-role-context/context/{task_id}.json"
    envelope = TaskEnvelope(
        task_id=task_id,
        run_id="run-role-context",
        agent_id=agent_id,
        objective="验证角色持久状态",
        input_refs=[contract_ref],
        input_contract_kind=contract_kind,
        input_contract_ref=contract_ref,
    )
    manifest_path = runner._save_conversation_trace(
        SimpleNamespace(_conversation_history=[]),
        envelope,
        f"{agent_id}--session-test",
        "session-test",
        status="waiting",
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["manifest_version"] == 4
    assert manifest["encoding"] == "identity+refs"
    assert manifest["business_refs"]["input_contract_ref"] == contract_ref
    assert manifest["business_refs"]["attention_scope_ref"] == (
        contract_ref if expects_attention_scope else None
    )
    assert "compaction_summary" not in manifest
    assert "restart_state_sha256" not in manifest


def _current_reporting_recovery_policy():
    _capability, registry = load_distribution_reporting_capability()
    return registry.require(
        DefinitionKind.RECOVERY,
        "current-reporting-recovery",
    )


def _save_reporting_tool_plan(
    workspace: Path,
    run_id: str,
    *definitions: ToolDefinition,
) -> None:
    plan = ResolvedPlan(
        workflow_id="declarative-reporting-tool-test",
        workflow_version="1.0.0",
        actions=[],
        agent_tool_ids=[definition.id for definition in definitions],
        definition_snapshots={
            f"tool:{definition.id}": definition.model_dump(mode="json")
            for definition in definitions
        },
    )
    FileWorkflowStateStore(workspace).save_plan(run_id, plan)


@pytest.mark.asyncio
async def test_saved_declarative_tool_implementation_selects_current_python_tool(
    tmp_path: Path,
) -> None:
    run_id = "run-declarative-tool-implementation"
    declared = ToolDefinition(
        id="saved-calculator",
        version="1.0.0",
        description="Saved calculator description",
        instructions="Use the saved calculator instructions.",
        implementation="capability:distribution-reporting:calculate",
        input_contract="reporting_tool_input",
        output_contract="reporting_tool_output",
        side_effect="pure_read",
        parallel_safe=True,
        reuse_result=True,
    )
    _save_reporting_tool_plan(tmp_path, run_id, declared)
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["module-2.1-specialist"].model_copy(
        update={"tools": [declared.id]}
    )
    envelope = TaskEnvelope(
        task_id="saved-tool-selection",
        run_id=run_id,
        agent_id=definition.id,
        objective="验证保存计划的工具实现绑定",
        allowed_tools=[declared.id],
    )

    registry = runner._tools(
        definition,
        envelope,
        "session-saved-tool",
        "workflow-saved-tool",
        recovery_policy=_current_reporting_recovery_policy(),
    )

    tool = registry.get(declared.id)
    assert tool is not None
    assert await tool(expression="1 + 2") == {"expression": "1 + 2", "result": 3}
    assert tool.name == declared.id
    assert tool.description == declared.instructions
    assert tool.side_effect == declared.side_effect
    assert tool.parallel_safe is True
    assert tool.reuse_result is True


def test_saved_declarative_hidden_tool_is_not_model_registered(
    tmp_path: Path,
) -> None:
    run_id = "run-declarative-hidden-tool"
    hidden = ToolDefinition(
        id="calculate",
        version="1.0.0",
        description="Hidden saved calculator",
        implementation="capability:distribution-reporting:calculate",
        input_contract="reporting_tool_input",
        output_contract="reporting_tool_output",
        model_visible=False,
    )
    _save_reporting_tool_plan(tmp_path, run_id, hidden)
    runner = ReportingAgentRunner(
        tmp_path,
        MessageBus(),
        DirectSubmissionProvider(),
        AgentDefaults(),
    )
    definition = load_packaged_agents()["module-2.1-specialist"].model_copy(
        update={"tools": [hidden.id]}
    )
    envelope = TaskEnvelope(
        task_id="saved-hidden-tool",
        run_id=run_id,
        agent_id=definition.id,
        objective="验证隐藏工具不进入模型 Registry",
        allowed_tools=[hidden.id],
    )

    registry = runner._tools(
        definition,
        envelope,
        "session-hidden-tool",
        "workflow-hidden-tool",
        recovery_policy=_current_reporting_recovery_policy(),
    )

    assert registry.get(hidden.id) is None
    assert hidden.id not in {item["name"] for item in registry.get_definitions()}

    legacy_registry = runner._tools(
        definition,
        envelope,
        "session-legacy-tool",
        "workflow-legacy-tool",
    )
    assert legacy_registry.get(hidden.id) is not None


def _install_recovery_controller_spy(
    monkeypatch: pytest.MonkeyPatch,
    calls: list[tuple[str, str]],
    observations: list[tuple[bool, str | None]],
) -> None:
    real_controller = agent_runner_module.RecoveryController

    class SpyRecoveryController:
        def __init__(self) -> None:
            self.delegate = real_controller()

        def decide(self, event, policy, state):
            decision = self.delegate.decide(event, policy, state)
            calls.append((event.kind.value, decision.action.value))
            return decision

        def observe_progress(self, observation, policy, state):
            decision = self.delegate.observe_progress(observation, policy, state)
            observations.append(
                (observation.progressed, decision.action.value if decision else None)
            )
            if decision is not None:
                calls.append((decision.event.kind.value, decision.action.value))
            return decision

    monkeypatch.setattr(
        agent_runner_module,
        "RecoveryController",
        SpyRecoveryController,
    )


@pytest.mark.asyncio
async def test_reporting_provider_error_uses_declared_recovery_controller(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple[str, str]] = []
    observations: list[tuple[bool, str | None]] = []
    _install_recovery_controller_spy(monkeypatch, calls, observations)
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        NonRetryableFailureProvider(),
        AgentDefaults(max_tool_iterations=1),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-provider-recovery",
        agent_id="module-2.1-specialist",
        objective="验证 Provider 失败进入声明式恢复策略",
        allowed_outputs=["module_submission"],
    )

    try:
        with pytest.raises(RuntimeError, match="数据或参数校验失败"):
            await runner.run(
                load_packaged_agents()["module-2.1-specialist"],
                envelope,
                [],
                workflow_id="wf-declarative-provider-recovery",
                recovery_policy=_current_reporting_recovery_policy(),
            )
    finally:
        await runner.close_workflow("wf-declarative-provider-recovery")
        bus.shutdown()
        await bus_task

    assert (
        RecoveryEventKind.PROVIDER_ERROR.value,
        RecoveryActionKind.CONTINUE.value,
    ) in calls


@pytest.mark.asyncio
async def test_declarative_recovery_policy_actions_drive_reporting_paths(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Characterize continuation and typed-result actions at Reporting boundary."""

    calls: list[tuple[str, str]] = []
    observations: list[tuple[bool, str | None]] = []
    _install_recovery_controller_spy(monkeypatch, calls, observations)
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = ToolSliceContinuationProvider()
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_iterations=1),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-recovery-actions",
        agent_id="module-2.1-specialist",
        objective="验证声明式 recovery action 驱动续接与复用",
        allowed_outputs=["module_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )

    recovery_runner = None
    try:
        first = await runner.run(
            load_packaged_agents()["module-2.1-specialist"].model_copy(
                update={"max_turns": 1}
            ),
            envelope,
            [],
            workflow_id="wf-declarative-recovery-actions",
            recovery_policy=_current_reporting_recovery_policy(),
        )
        await runner.close_workflow("wf-declarative-recovery-actions")

        class NoCallProvider(LLMProvider):
            def __init__(self) -> None:
                super().__init__("test", model="tool-slice-continuation")

            async def chat(self, *_args, **_kwargs):
                raise AssertionError("persisted typed result must bypass Provider")

        recovery_runner = ReportingAgentRunner(
            tmp_path,
            bus,
            NoCallProvider(),
            AgentDefaults(max_tool_iterations=1),
            timeout=5,
        )
        recovered = await recovery_runner.run(
            load_packaged_agents()["module-2.1-specialist"].model_copy(
                update={"max_turns": 1}
            ),
            envelope.model_copy(),
            [],
            workflow_id="wf-declarative-recovery-actions",
            recovery_policy=_current_reporting_recovery_policy(),
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-actions")
        if recovery_runner is not None:
            await recovery_runner.close_workflow("wf-declarative-recovery-actions")
        bus.shutdown()
        await bus_task

    assert first.status is AgentRunStatus.COMPLETED
    assert recovered == first
    assert provider.calls > 0
    assert (
        RecoveryEventKind.TOOL_SLICE_BOUNDARY.value,
        RecoveryActionKind.CONTINUE.value,
    ) in calls
    assert (
        RecoveryEventKind.COMPLETED_TOOL_RESULT.value,
        RecoveryActionKind.REUSE_RESULT.value,
    ) in calls


@pytest.mark.asyncio
async def test_declarative_recovery_policy_handles_max_tokens_and_natural_text(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Characterize max-token continuation and ordinary-text correction actions."""

    calls: list[tuple[str, str]] = []
    observations: list[tuple[bool, str | None]] = []
    _install_recovery_controller_spy(monkeypatch, calls, observations)
    policy = _current_reporting_recovery_policy()
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        MaxTokensThenDirectSubmissionProvider(),
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-recovery-max-tokens",
        agent_id="module-2.1-specialist",
        objective="验证声明式 max_tokens continuation",
        allowed_outputs=["module_submission"],
    )

    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-declarative-recovery-max-tokens",
            recovery_policy=policy,
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-max-tokens")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert (
        RecoveryEventKind.MAX_TOKENS.value,
        RecoveryActionKind.CONTINUE.value,
    ) in calls

    template_ref = "template-natural.docx"
    Document().save(tmp_path / template_ref)
    contract_ref = _write_template_contract(
        tmp_path,
        run_id="run-declarative-recovery-natural",
        template_ref=template_ref,
    )
    natural_bus = MessageBus()
    natural_bus_task = asyncio.create_task(natural_bus.process_queue())
    natural_runner = ReportingAgentRunner(
        tmp_path,
        natural_bus,
        NaturalCompletionProvider(),
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    natural_envelope = TaskEnvelope(
        task_id="template-skill-distillation",
        run_id="run-declarative-recovery-natural",
        agent_id="template-distiller",
        objective="验证普通文字 completion 的 correction",
        input_refs=[contract_ref, template_ref],
        allowed_outputs=["template_skill_submission"],
        input_contract_kind="template_distillation_input",
        input_contract_ref=contract_ref,
    )
    try:
        natural = await natural_runner.run(
            load_packaged_agents()["template-distiller"],
            natural_envelope,
            [],
            workflow_id="wf-declarative-recovery-natural",
            recovery_policy=policy,
        )
    finally:
        await natural_runner.close_workflow("wf-declarative-recovery-natural")
        natural_bus.shutdown()
        await natural_bus_task

    assert natural.status is AgentRunStatus.INCOMPLETE
    assert (
        RecoveryEventKind.NATURAL_LANGUAGE_WITHOUT_SUBMISSION.value,
        RecoveryActionKind.CORRECT.value,
    ) in calls


@pytest.mark.asyncio
async def test_declarative_recovery_policy_uses_no_progress_observer(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Characterize no-progress STOP as a RecoveryController observation."""

    calls: list[tuple[str, str]] = []
    observations: list[tuple[bool, str | None]] = []
    _install_recovery_controller_spy(monkeypatch, calls, observations)
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        RepeatingNoProgressToolProvider(),
        AgentDefaults(max_tool_iterations=1),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        task_attempt_id="attempt-declarative-no-progress",
        run_id="run-declarative-recovery-no-progress",
        agent_id="module-2.1-specialist",
        objective="验证声明式 no-progress stop",
        allowed_outputs=["module_submission"],
    )
    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"].model_copy(
                update={"max_turns": 1}
            ),
            envelope,
            [],
            workflow_id="wf-declarative-recovery-no-progress",
            recovery_policy=_current_reporting_recovery_policy(),
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-no-progress")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.INCOMPLETE
    assert (False, RecoveryActionKind.STOP.value) in observations


@pytest.mark.asyncio
async def test_declarative_recovery_policy_observes_invalid_structured_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Characterize SubmitResultTool schema correction through the Controller."""

    calls: list[tuple[str, str]] = []
    observations: list[tuple[bool, str | None]] = []
    _install_recovery_controller_spy(monkeypatch, calls, observations)
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        WrappedThenFlatSubmissionProvider(),
        AgentDefaults(max_tool_iterations=10),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-recovery-invalid-structured",
        agent_id="module-2.1-specialist",
        objective="验证 invalid structured output correction",
        allowed_outputs=["module_submission"],
    )
    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-declarative-recovery-invalid-structured",
            recovery_policy=_current_reporting_recovery_policy(),
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-invalid-structured")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert (
        RecoveryEventKind.INVALID_STRUCTURED_OUTPUT.value,
        RecoveryActionKind.CORRECT.value,
    ) in calls


@pytest.mark.asyncio
async def test_declarative_recovery_policy_drives_tool_contract_correction(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Characterize Reporting-private ToolContractError policy dispatch."""

    class ToolContractThenTextProvider(LLMProvider):
        def __init__(self) -> None:
            super().__init__("test", model="scripted")
            self.calls = 0
            self.tool_definitions: list[dict] = []

        async def chat(self, *_args, **_kwargs):
            self.calls += 1
            self.tool_definitions = list(_kwargs.get("tools") or [])
            if self.calls == 1:
                return LLMResponse(
                    content="",
                    tool_calls=[
                        LLMToolCall(
                            id="inspect-invalid-photo",
                            name="inspect_image",
                            arguments={"path": "P-999"},
                        )
                    ],
                )
            return LLMResponse(content="finished without a typed submission")

    calls: list[tuple[str, str]] = []
    observations: list[tuple[bool, str | None]] = []
    _install_recovery_controller_spy(monkeypatch, calls, observations)
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = ToolContractThenTextProvider()
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="tool-contract-correction",
        run_id="run-declarative-recovery-tool-contract",
        agent_id="module-2.1-specialist",
        objective="验证 tool contract correction",
    )
    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"].model_copy(
                update={"tools": ["inspect_image"]}
            ),
            envelope,
            [],
            workflow_id="wf-declarative-recovery-tool-contract",
            recovery_policy=_current_reporting_recovery_policy(),
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-tool-contract")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.INCOMPLETE
    assert (
        RecoveryEventKind.TOOL_CONTRACT_ERROR.value,
        RecoveryActionKind.CORRECT.value,
    ) in calls
    inspect_schema = next(
        item["input_schema"]
        for item in provider.tool_definitions
        if item["name"] == "inspect_image"
    )
    assert inspect_schema["anyOf"] == [{"required": ["path"]}, {"required": ["ref"]}]


@pytest.mark.asyncio
async def test_recovery_policy_stop_ends_tool_slice_without_another_provider_call(
    tmp_path: Path,
) -> None:
    """A declared STOP ends the current recovery instead of continuing implicitly."""

    policy = _current_reporting_recovery_policy()
    rules = dict(policy.rules)
    rules[RecoveryEventKind.TOOL_SLICE_BOUNDARY.value] = rules[
        RecoveryEventKind.TOOL_SLICE_BOUNDARY.value
    ].model_copy(update={"action": RecoveryActionKind.STOP.value})
    mismatched_policy = policy.model_copy(update={"rules": rules})
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    provider = ToolSliceContinuationProvider()
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_iterations=1),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-recovery-mismatch",
        agent_id="module-2.1-specialist",
        objective="验证不匹配 recovery action 不会隐式续接",
        allowed_outputs=["module_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )
    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"].model_copy(
                update={"max_turns": 1}
            ),
            envelope,
            [],
            workflow_id="wf-declarative-recovery-mismatch",
            recovery_policy=mismatched_policy,
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-mismatch")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.INCOMPLETE
    # The AgentLoop's bounded slice uses two physical turns before publishing
    # its continuation sentinel; STOP prevents the Reporting adapter's third.
    assert provider.calls == 2


@pytest.mark.asyncio
async def test_recovery_policy_fail_fails_tool_slice_explicitly(
    tmp_path: Path,
) -> None:
    policy = _current_reporting_recovery_policy()
    rules = dict(policy.rules)
    rules[RecoveryEventKind.TOOL_SLICE_BOUNDARY.value] = rules[
        RecoveryEventKind.TOOL_SLICE_BOUNDARY.value
    ].model_copy(update={"action": RecoveryActionKind.FAIL.value})
    failing_policy = policy.model_copy(update={"rules": rules})
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        ToolSliceContinuationProvider(),
        AgentDefaults(max_tool_iterations=1),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-recovery-fail",
        agent_id="module-2.1-specialist",
        objective="验证 recovery fail",
        allowed_outputs=["module_submission"],
        target_submodule_ids=list(REPORT_TAXONOMY["2.1"].submodules),
    )
    try:
        with pytest.raises(RuntimeError, match="failed tool_slice_boundary"):
            await runner.run(
                load_packaged_agents()["module-2.1-specialist"].model_copy(
                    update={"max_turns": 1}
                ),
                envelope,
                [],
                workflow_id="wf-declarative-recovery-fail",
                recovery_policy=failing_policy,
            )
    finally:
        await runner.close_workflow("wf-declarative-recovery-fail")
        bus.shutdown()
        await bus_task


@pytest.mark.asyncio
async def test_recovery_policy_prompt_is_sent_on_same_conversation_continuation(
    tmp_path: Path,
) -> None:
    policy = _current_reporting_recovery_policy()
    rules = dict(policy.rules)
    prompt = "POLICY_PROMPT_CONTINUE_EXACT_TASK"
    rules[RecoveryEventKind.MAX_TOKENS.value] = rules[
        RecoveryEventKind.MAX_TOKENS.value
    ].model_copy(update={"prompt": prompt})
    prompted_policy = policy.model_copy(update={"rules": rules})
    provider = MaxTokensThenDirectSubmissionProvider()
    bus = MessageBus()
    bus_task = asyncio.create_task(bus.process_queue())
    runner = ReportingAgentRunner(
        tmp_path,
        bus,
        provider,
        AgentDefaults(max_tool_iterations=5),
        timeout=5,
    )
    envelope = TaskEnvelope(
        task_id="module-2.1",
        run_id="run-declarative-recovery-prompt",
        agent_id="module-2.1-specialist",
        objective="验证 recovery prompt",
        allowed_outputs=["module_submission"],
    )
    try:
        result = await runner.run(
            load_packaged_agents()["module-2.1-specialist"],
            envelope,
            [],
            workflow_id="wf-declarative-recovery-prompt",
            recovery_policy=prompted_policy,
        )
    finally:
        await runner.close_workflow("wf-declarative-recovery-prompt")
        bus.shutdown()
        await bus_task

    assert result.status is AgentRunStatus.COMPLETED
    assert provider.message_snapshots
    assert any(
        prompt in message.content
        for message in provider.message_snapshots[0]
    )
