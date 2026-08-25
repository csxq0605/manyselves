import hashlib
import importlib.util
import io
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from manyselves.core.reporting.agentic_models import ClaimRecord, SourceKind, SourceRecord
from manyselves.core.reporting.claim_ledger import ClaimLedger
from manyselves.core.reporting.rendering.handoff_docx import PackagedV2DocxCore
from manyselves.core.reporting.rendering.packaged_docx import PackagedDocxCore
from manyselves.core.reporting.rendering.pds_docx_renderer import (
    ApprovedReport,
    PdsDocxRenderer,
    ReportPhoto,
    ReportTable,
)
from manyselves.core.reporting.rendering.rendered_docx_validator import (
    validate_rendered_markdown_docx,
)
from manyselves.core.reporting.taxonomy import REPORT_TAXONOMY


def _set_style_east_asia_font(style, font_name: str) -> None:
    run_properties = style.element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), font_name)


def _style_east_asia_font(style) -> str | None:
    run_properties = style.element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    return fonts.get(qn("w:eastAsia")) if fonts is not None else None


def test_packaged_v2_core_matches_normalized_handoff_source() -> None:
    core = PackagedV2DocxCore(Path("unused-template.docx"))
    assert hashlib.sha256(core.source_path.read_bytes()).hexdigest() == (
        "41a4e09573d347ff3f795feccd831a364e7d447f9a0986a56d423019937438c4"
    )


def test_packaged_v2_core_supports_legacy_top_level_dynamic_loading() -> None:
    core = PackagedV2DocxCore(Path("unused-template.docx"))
    spec = importlib.util.spec_from_file_location(
        "_legacy_v2_docx_renderer",
        core.source_path,
    )
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)

    spec.loader.exec_module(module)

    assert callable(module.render_report_docx)


def test_report_embedding_removes_only_duplicate_module_heading() -> None:
    narrative = (
        "## 2.5 运维管理与风险管控机制\n\n"
        "### 2.5.1 SOP/EOP\n\n完整分析正文。"
    )

    embedded = PdsDocxRenderer._strip_leading_module_heading(narrative, "2.5")

    assert "## 2.5 运维管理与风险管控机制" not in embedded
    assert embedded.startswith("#### 2.5.1 SOP/EOP")


def test_report_embedding_removes_module_heading_after_editor_transition() -> None:
    narrative = (
        "以下为运维模块的批准正文。\n\n"
        "## 2.5 运维管理与风险管控机制\n\n"
        "### 2.5.1 SOP/EOP\n\n完整分析正文。"
    )

    embedded = PdsDocxRenderer._strip_leading_module_heading(narrative, "2.5")

    assert "## 2.5 运维管理与风险管控机制" not in embedded
    assert "以下为运维模块的批准正文。" in embedded
    assert "#### 2.5.1 SOP/EOP" in embedded


def test_report_embedding_does_not_mistake_first_submodule_for_module_heading() -> None:
    narrative = (
        "#### 2.5.1 SOP/EOP\n\n完整分析正文。\n\n"
        "#### 2.5.2 图纸资料\n\n图纸分析正文。"
    )

    embedded = PdsDocxRenderer._strip_leading_module_heading(narrative, "2.5")

    assert embedded == narrative
    assert embedded.count("#### 2.5.1 SOP/EOP") == 1
    assert embedded.count("#### 2.5.2 图纸资料") == 1


def test_packaged_v2_core_removes_markdown_markers_and_uses_one_label_style(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    markdown = """# 配电安全评估报告

## 1. 配电评估概述

## 2. 评估内容描述

### 2.1 配电系统架构问题

### 现状描述

现场记录显示一项异常。

**判断：** 该项需要整改。

### 风险与影响

- 可能导致供电中断。

### 建议

1. 完成复核并验收。

> **判定：** NG，优先级高。
"""

    _, data = PackagedV2DocxCore(template).render_approved_prose(markdown)

    rendered = Document(io.BytesIO(data))
    visible = [paragraph.text for paragraph in rendered.paragraphs if paragraph.text.strip()]
    assert "配电安全评估报告" not in visible
    assert "现状描述：" in visible
    assert "判断：该项需要整改。" in visible
    assert "风险与影响：" in visible
    assert "建议：" in visible
    assert "判定：NG，优先级高。" in visible
    assert not any(text.lstrip().startswith(("#", ">", "**")) for text in visible)
    assert not any("【" in text or "】" in text for text in visible)


def test_packaged_v2_core_preserves_template_fonts_and_uses_consistent_type_scale(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    template_document = Document()
    _set_style_east_asia_font(template_document.styles["Normal"], "Template Body CJK")
    _set_style_east_asia_font(template_document.styles["Title"], "Template Title CJK")
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        _set_style_east_asia_font(
            template_document.styles[style_name],
            "Template Heading CJK",
        )
    template_document.save(template)

    markdown = """# 配电安全评估报告

## 2. 评估内容描述

### 2.1 配电系统架构问题

#### 2.1.1 配电系统负荷分配与过载风险

正文段落用于验证统一排版节奏。
"""

    _, data = PackagedV2DocxCore(template).render_approved_prose(markdown)

    rendered = Document(io.BytesIO(data))
    assert _style_east_asia_font(rendered.styles["Normal"]) == "Template Body CJK"
    assert _style_east_asia_font(rendered.styles["Title"]) == "Template Title CJK"
    assert all(
        _style_east_asia_font(rendered.styles[style_name]) == "Template Heading CJK"
        for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4")
    )
    assert rendered.styles["Normal"].font.size.pt == 11
    assert rendered.styles["Heading 1"].font.size.pt == 18
    assert rendered.styles["Heading 2"].font.size.pt == 16
    assert rendered.styles["Heading 3"].font.size.pt == 14
    assert rendered.styles["Heading 4"].font.size.pt == 12
    body = next(
        paragraph
        for paragraph in rendered.paragraphs
        if paragraph.text == "正文段落用于验证统一排版节奏。"
    )
    assert body.paragraph_format.line_spacing == 1.3
    assert body.paragraph_format.space_after.pt == 4


def test_packaged_v2_core_uses_real_word_numbering_for_lists(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    markdown = """# 配电安全评估报告

## 2. 评估内容描述

- 项目符号内容换行后应保持悬挂对齐。
1. 第一项整改措施。
3. 第三项整改措施。
"""

    _, data = PackagedV2DocxCore(template).render_approved_prose(markdown)

    rendered = Document(io.BytesIO(data))
    items = {
        paragraph.text: paragraph
        for paragraph in rendered.paragraphs
        if paragraph.text.strip()
    }
    assert "项目符号内容换行后应保持悬挂对齐。" in items
    assert "第一项整改措施。" in items
    assert "第三项整改措施。" in items
    assert all(
        items[text]._p.pPr.numPr is not None
        for text in (
            "项目符号内容换行后应保持悬挂对齐。",
            "第一项整改措施。",
            "第三项整改措施。",
        )
    )
    assert not any(
        paragraph.text.startswith(("•", "1.", "3."))
        for paragraph in items.values()
    )
    numbering_xml = rendered.part.numbering_part.element.xml
    assert 'w:numFmt w:val="bullet"' in numbering_xml
    assert 'w:numFmt w:val="decimal"' in numbering_xml
    assert 'w:startOverride w:val="3"' in numbering_xml


def test_shared_validator_accepts_arrow_bullet_rendered_as_native_word_list(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    markdown = """# 配电安全评估报告

## 3. 结论与建议

➢ 总体评价：整改措施应按风险等级闭环验证。
"""
    _, data = PackagedV2DocxCore(template).render_approved_prose(markdown)
    output = tmp_path / "arrow-bullet.docx"
    output.write_bytes(data)

    validate_rendered_markdown_docx(
        output,
        "➢ 总体评价：整改措施应按风险等级闭环验证。",
    )

    paragraph = next(
        item
        for item in Document(output).paragraphs
        if item.text == "总体评价：整改措施应按风险等级闭环验证。"
    )
    assert paragraph._p.pPr.numPr is not None


def test_shared_validator_warns_for_reordered_approved_prose(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    output = tmp_path / "reordered.docx"
    document = Document()
    document.add_paragraph("第二项必须随后出现。")
    document.add_paragraph("第一项必须先出现。")
    document.save(output)

    warnings = validate_rendered_markdown_docx(
        output,
        "第一项必须先出现。\n\n第二项必须随后出现。",
    )

    assert any("omitted approved Markdown content" in item for item in warnings)
    assert "DOCX post-render validation warning" in caplog.text


def test_shared_validator_warns_for_collapsed_duplicate_prose(tmp_path: Path) -> None:
    output = tmp_path / "collapsed-duplicate.docx"
    document = Document()
    document.add_paragraph("同一句批准正文必须保留两次。")
    document.save(output)

    warnings = validate_rendered_markdown_docx(
        output,
        "同一句批准正文必须保留两次。\n\n同一句批准正文必须保留两次。",
    )

    assert any("omitted approved Markdown content" in item for item in warnings)


def test_shared_validator_warns_for_unresolved_delivery_tokens(tmp_path: Path) -> None:
    output = tmp_path / "unresolved-token.docx"
    document = Document()
    document.add_paragraph("批准正文。[[CITE:1]]")
    document.save(output)

    warnings = validate_rendered_markdown_docx(
        output,
        "批准正文。[[CITE:1]]",
        reject_unresolved_tokens=True,
    )

    assert any("unresolved content tokens" in item for item in warnings)


def test_shared_validator_still_blocks_an_unopenable_docx(tmp_path: Path) -> None:
    output = tmp_path / "not-a-docx.docx"
    output.write_bytes(b"not a zip archive")

    with pytest.raises(ValueError, match="not Word/WPS-openable"):
        validate_rendered_markdown_docx(output, "批准正文。")


def test_packaged_v2_core_preserves_dotted_section_references_in_bullets(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    markdown = """# 配电安全评估报告

## 3. 改善建议

- 2.2.2.1节红外热缺陷：整改后复测；
- 2.3.2节/2.4.3.1节剩余电流异常：排查后复测。
"""

    _, data = PackagedV2DocxCore(template).render_approved_prose(markdown)

    rendered = Document(io.BytesIO(data))
    visible = {paragraph.text for paragraph in rendered.paragraphs}
    assert "2.2.2.1节红外热缺陷：整改后复测；" in visible
    assert "2.3.2节/2.4.3.1节剩余电流异常：排查后复测。" in visible


def test_packaged_v2_core_constrains_table_geometry_to_page_body(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    template_document = Document()
    section = template_document.sections[0]
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    template_document.save(template)
    headers = [f"字段{i}" for i in range(1, 9)]
    markdown = "\n".join(
        [
            "# 配电安全评估报告",
            "",
            "## 2. 评估内容描述",
            "",
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
            "| " + " | ".join(f"内容{i}" for i in range(1, 9)) + " |",
        ]
    )

    _, data = PackagedV2DocxCore(template).render_approved_prose(markdown)

    rendered = Document(io.BytesIO(data))
    table = rendered.tables[0]
    usable_twips = int(
        (
            rendered.sections[0].page_width.cm
            - rendered.sections[0].left_margin.cm
            - rendered.sections[0].right_margin.cm
        )
        * 567
    )
    table_width = int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")))
    grid_widths = [
        int(column.get(qn("w:w")))
        for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    assert table_width <= usable_twips
    assert sum(grid_widths) == table_width
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    assert table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")) == "0"
    assert all(
        [
            int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w")))
            for cell in row.cells
        ]
        == grid_widths
        for row in table.rows
    )
    table_style = rendered.styles["表格2"]
    assert table_style.font.size.pt == 10


def _approved_report(photo_path: Path) -> ApprovedReport:
    narratives = {
        module_id: (
            f"模块 {module_id} 的专家自然分析保留原样。"
            f"[[CLAIM:C-00{index}]]"
        )
        for index, module_id in enumerate(
            ("2.1", "2.2", "2.3", "2.4", "2.5"), 1
        )
    }
    claims = [
        ClaimRecord(
            id=f"C-00{index}",
            module_id=module_id,
            submodule_id=next(iter(REPORT_TAXONOMY[module_id].submodules)),
            text=narratives[module_id],
            claim_type="risk_judgment",
            source_ids=["E-001"],
        )
        for index, module_id in enumerate(("2.1", "2.2", "2.3", "2.4", "2.5"), 1)
    ]
    source = SourceRecord(
        id="E-001",
        kind=SourceKind.PROJECT_EVIDENCE,
        title="现场照片与检测表",
        locator="Inputs/设备.xlsx；工作表=问题；单元格=A2:F2；图片=IMG-1",
    )
    return ApprovedReport(
        title="配电安全专家咨询报告",
        assessment_background="总编形成的评估背景保持不变。",
        findings_overview="总编形成的健康度总览保持不变。",
        regional_executive_summary="总编形成的区域摘要保持不变。",
        module_narratives=narratives,
        risk_panorama="总编形成的风险全景保持不变。",
        dimension_risk_analysis="总编形成的维度风险分析保持不变。",
        data_gap_analysis="总编形成的数据缺口分析保持不变。",
        improvement_action_plan="总编形成的改善行动计划保持不变。",
        special_topic_plan={
            "source_ref": "Inputs/专项问题分析.md",
            "source_sha256": "0" * 64,
            "sections": [
                {
                    "section_id": "4.1",
                    "title": "动态专项问题",
                    "requirement": "分析项目边界、方案条件和验证方法。",
                }
            ],
        },
        special_topic_analysis=(
            "### 4.1 动态专项问题\n\n"
            "总编形成的动态专项分析保持不变，并说明项目边界、方案条件和验证方法。"
        ),
        ledger=ClaimLedger(claims=claims, sources=[source]),
        tables=[
            ReportTable(
                title="整改行动表",
                headers=["对象", "动作"],
                rows=[["1A2", "复核连接"]],
                source_ids=["E-001"],
                claim_ids=["C-004"],
            )
        ],
        photos=[
            ReportPhoto(
                id="IMG-1",
                path=photo_path,
                caption="1A2 柜现场状态",
                source_id="E-001",
                claim_ids=["C-004"],
                submodule_id=next(iter(REPORT_TAXONOMY["2.4"].submodules)),
            )
        ],
    )


def test_renderer_preserves_prose_and_keeps_source_index_outside_docx(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (30, 20), color="red").save(photo)
    renderer = PdsDocxRenderer(
        PackagedDocxCore(template),
    )
    output = tmp_path / "final.docx"

    result = renderer.render(_approved_report(photo), output)

    assert result.output_path == output
    assert result.handoff_core_used is True
    assert result.protected_prose_verified is True
    assert result.validation_warnings == []
    rendered = Document(output)
    all_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "模块 2.4 的专家自然分析保留原样。" in all_text
    assert "配电安全专家咨询报告" in all_text
    assert "总编形成的风险全景保持不变。" in all_text
    assert "总编形成的动态专项分析保持不变" in all_text
    assert "项目证据 E-*" not in all_text
    assert "Inputs/设备.xlsx；工作表=问题；单元格=A2:F2；图片=IMG-1" not in all_text
    assert any(
        run.font.superscript and run.text == "1" for p in rendered.paragraphs for run in p.runs
    )
    assert any(table.cell(0, 0).text == "对象" for table in rendered.tables)
    assert len(rendered.inline_shapes) == 1

    second_output = tmp_path / "final-second.docx"
    second = renderer.render(_approved_report(photo), second_output)
    assert second.output_sha256 == result.output_sha256
    assert second_output.read_bytes() == output.read_bytes()


def test_renderer_omits_chapter_four_when_special_topic_plan_is_absent(
    tmp_path: Path,
) -> None:
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (30, 20), color="red").save(photo)
    payload = _approved_report(photo).model_dump(mode="python")
    payload["special_topic_plan"] = None
    payload["special_topic_analysis"] = None
    report = ApprovedReport.model_validate(payload)

    markdown = PdsDocxRenderer._compose_markdown(report)

    assert "## 4. 专项问题分析" not in markdown
    assert "### 4." not in markdown


def test_renderer_preserves_body_labels_that_collide_with_report_heading_numbers(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (30, 20), color="red").save(photo)
    report = _approved_report(photo).model_copy(
        update={
            "improvement_action_plan": (
                "**1.1 负荷均衡调整**\n\n"
                "负荷均衡措施必须完整保留。\n\n"
                "**2.1 设备防护等级恢复**\n\n"
                "设备防护措施必须完整保留。\n\n"
                "**3.2 智能化平台功能完善**\n\n"
                "平台完善措施必须完整保留。"
            )
        }
    )
    output = tmp_path / "final.docx"

    result = PdsDocxRenderer(PackagedV2DocxCore(template)).render(report, output)

    assert result.protected_prose_verified is True
    visible = "\n".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )
    assert "1.1 负荷均衡调整" in visible
    assert "2.1 设备防护等级恢复" in visible
    assert "3.2 智能化平台功能完善" in visible


def test_renderer_accepts_the_exact_citation_bound_delivery_markdown(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (30, 20), color="red").save(photo)
    report = _approved_report(photo)
    delivery_markdown = report.ledger.bind_citations(
        PdsDocxRenderer._compose_markdown(report)
    )
    output = tmp_path / "delivery.docx"

    PdsDocxRenderer(PackagedDocxCore(template)).render(
        report,
        output,
        approved_markdown=delivery_markdown,
    )

    rendered = Document(output)
    assert not any(
        "[[CLAIM:" in paragraph.text or "[[CITE:" in paragraph.text
        for paragraph in rendered.paragraphs
    )
    assert any(
        run.font.superscript and run.text == "1"
        for paragraph in rendered.paragraphs
        for run in paragraph.runs
    )


def test_renderer_warns_but_publishes_missing_regional_executive_summary(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (30, 20), color="red").save(photo)
    report = _approved_report(photo)
    packaged_core = PackagedDocxCore(template)

    class DropRegionalSummaryCore:
        def render_approved_prose(self, report_text: str, **kwargs):
            return packaged_core.render_approved_prose(
                report_text.replace(report.regional_executive_summary, ""),
                **kwargs,
            )

    output = tmp_path / "published-with-warning.docx"
    result = PdsDocxRenderer(DropRegionalSummaryCore()).render(report, output)

    assert output.is_file()
    assert result.protected_prose_verified is False
    assert any(
        "omitted approved Markdown content" in item
        for item in result.validation_warnings
    )


def test_delivery_markdown_places_structured_tables_once(
    tmp_path: Path,
) -> None:
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (30, 20), color="red").save(photo)
    fixture = json.loads(
        (
            Path(__file__).parents[1]
            / "fixtures/report-730c5d83f6-table-placement.json"
        ).read_text(encoding="utf-8")
    )
    report = _approved_report(photo).model_copy(
        update={
            "risk_panorama": fixture["risk_panorama"],
            "improvement_action_plan": fixture["improvement_action_plan"],
            "tables": [
                ReportTable(
                    **table,
                    source_ids=["E-001"],
                    claim_ids=["C-004"],
                )
                for table in fixture["tables"]
            ],
        }
    )

    markdown = PdsDocxRenderer._compose_markdown(report)

    assert markdown.count("\n风险簇矩阵（来源：E-001）\n") == 1
    assert markdown.count("\n行动依赖矩阵（来源：E-001）\n") == 1
    assert markdown.count("2A2变压器单一故障点") == 1
    assert markdown.count("2A2连续负荷监测") == 1
    assert markdown.index("## 4. 专项问题分析") < markdown.index(
        "2A2变压器单一故障点"
    )
    assert markdown.index("2A2变压器单一故障点") < markdown.index(
        "2A2连续负荷监测"
    )


def test_photo_token_is_placed_inside_linked_submodule() -> None:
    first, second = list(REPORT_TAXONOMY["2.4"].submodules)[:2]
    narrative = f"### {first} 第一项\n\n第一项分析。\n\n### {second} 第二项\n\n第二项分析。"
    photo = ReportPhoto(
        id="IMG-1",
        path=Path("photo.png"),
        caption="第一项图证",
        source_id="E-001",
        claim_ids=["C-004"],
        submodule_id=first,
    )

    placed = PdsDocxRenderer._place_photo_tokens(narrative, [photo])

    assert placed.index("第一项分析。") < placed.index("[[PHOTO:IMG-1]]")
    assert placed.index("[[PHOTO:IMG-1]]") < placed.index(f"### {second}")


def test_multiple_adjacent_photos_become_a_submodule_evidence_summary_table(
    tmp_path: Path,
) -> None:
    from PIL import Image

    first_path = tmp_path / "first.png"
    second_path = tmp_path / "second.png"
    Image.new("RGB", (30, 20), color="red").save(first_path)
    Image.new("RGB", (30, 20), color="blue").save(second_path)
    report = _approved_report(first_path)
    second = report.photos[0].model_copy(
        update={"id": "IMG-2", "path": second_path, "caption": "第二幅图证"}
    )
    report = report.model_copy(update={"photos": [report.photos[0], second]})
    document = Document()
    document.add_paragraph("[[PHOTO:IMG-1]]")
    document.add_paragraph("[[PHOTO:IMG-2]]")
    document.add_paragraph("后续分析。")

    PdsDocxRenderer._materialize_photos(document, report)

    assert len(document.tables) == 1
    assert len(document.inline_shapes) == 2
    assert document.tables[0].cell(0, 0).text == "原表对应内容"
    assert document.tables[0].cell(0, 1).text == "图证"
    assert "1A2 柜现场状态" in document.tables[0].cell(1, 0).text
    assert "原表图证" in document.tables[0].cell(1, 1).text
    assert "IMG-1" not in document.tables[0].cell(1, 1).text
    assert "第二幅图证" in document.tables[0].cell(2, 0).text
    assert "原表图证" in document.tables[0].cell(2, 1).text
    assert "IMG-2" not in document.tables[0].cell(2, 1).text
    assert all("[[PHOTO:" not in paragraph.text for paragraph in document.paragraphs)


def test_large_photo_set_is_split_into_page_safe_summary_tables(
    tmp_path: Path,
) -> None:
    from PIL import Image

    photos = []
    for index in range(5):
        path = tmp_path / f"portrait-{index}.png"
        Image.new("RGB", (300, 600), color=(index * 20, 0, 0)).save(path)
        photos.append(
            _approved_report(path).photos[0].model_copy(
                update={
                    "id": f"IMG-{index + 1}",
                    "path": path,
                    "caption": f"第 {index + 1} 幅图证",
                }
            )
        )
    report = _approved_report(photos[0].path).model_copy(update={"photos": photos})
    document = Document()
    for photo in photos:
        document.add_paragraph(f"[[PHOTO:{photo.id}]]")
    document.add_paragraph("后续分析。")

    PdsDocxRenderer._materialize_photos(document, report)

    assert len(document.tables) == 3
    assert [len(table.rows) - 1 for table in document.tables] == [2, 2, 1]
    assert len(document.inline_shapes) == 5
    assert sum(
        paragraph.text == "原表图证汇总（续）"
        for paragraph in document.paragraphs
    ) == 2
    assert all(
        shape.height <= Cm(7.0)
        for shape in document.inline_shapes
    )
    assert all("[[PHOTO:" not in paragraph.text for paragraph in document.paragraphs)


def test_photo_without_claim_still_routes_by_smallest_submodule(
    tmp_path: Path,
) -> None:
    from PIL import Image

    photo_path = tmp_path / "criteria.png"
    Image.new("RGB", (120, 30), color="white").save(photo_path)
    report = _approved_report(photo_path)
    photo = report.photos[0].model_copy(
        update={
            "claim_ids": [],
            "submodule_id": "2.2.2.1",
            "caption": "原表红外缺陷判定标准",
        }
    )
    narratives = dict(report.module_narratives)
    narratives["2.2"] = (
        "### 2.2.2.1 低压配电设备发热情况\n\n发热分析。\n\n"
        "### 2.2.2.2 高压配电设备局放情况\n\n局放分析。"
    )
    report = report.model_copy(
        update={"photos": [photo], "module_narratives": narratives}
    )

    markdown = PdsDocxRenderer._compose_markdown(report)

    assert markdown.index("[[PHOTO:IMG-1]]") < markdown.index(
        "### 2.2.2.2 高压配电设备局放情况"
    )


def test_renderer_rejects_untraceable_table() -> None:
    with pytest.raises(ValueError, match="table.*source"):
        ReportTable(title="无来源表", headers=["值"], rows=[["1"]])


def test_approved_report_requires_hidden_claim_ledger_across_all_five_modules(
    tmp_path: Path,
) -> None:
    report = _approved_report(tmp_path / "photo.png")
    report.ledger.claims.pop()

    with pytest.raises(ValueError, match="claim ledger.*2.1-2.5"):
        ApprovedReport.model_validate(report.model_dump())


def test_title_insertion_falls_back_when_handoff_template_has_no_title_style() -> None:
    document = Document()
    document.styles["Title"].delete()

    PdsDocxRenderer._ensure_title(document, "配电安全专家咨询报告")

    assert document.paragraphs[0].text == "配电安全专家咨询报告"


def test_post_render_content_warning_does_not_block_output(tmp_path: Path) -> None:
    class CorruptCore:
        def render_approved_prose(self, *args, **kwargs):
            buffer = io.BytesIO()
            Document().save(buffer)
            return "broken.docx", buffer.getvalue()

    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (20, 20), color="blue").save(photo)
    output = tmp_path / "published-with-warning.docx"

    result = PdsDocxRenderer(CorruptCore()).render(_approved_report(photo), output)

    assert output.is_file()
    assert result.protected_prose_verified is False
    assert result.validation_warnings


def test_stale_fence_before_publish_never_exposes_rendered_output(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    photo = tmp_path / "photo.png"
    from PIL import Image

    Image.new("RGB", (20, 20), color="green").save(photo)
    output = tmp_path / "stale-worker-must-not-publish.docx"

    def reject_stale_worker() -> None:
        raise RuntimeError("stale project write lease")

    with pytest.raises(RuntimeError, match="stale project write lease"):
        PdsDocxRenderer(PackagedDocxCore(template)).render(
            _approved_report(photo),
            output,
            before_publish=reject_stale_worker,
        )

    assert not output.exists()


def test_packaged_core_explicitly_forbids_structured_model_prose_generation(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.docx"
    Document().save(template)
    core = PackagedDocxCore(template)

    with pytest.raises(ValueError, match="structured-model prose generation is disabled"):
        core.render_approved_prose("approved", report_model={"issues": []})


def test_render_verifier_accepts_bold_numbered_markdown_labels(tmp_path: Path) -> None:
    output = tmp_path / "numbered-labels.docx"
    document = Document()
    document.add_paragraph("1. 多处剩余电流严重超标：")
    document.add_paragraph(
        '2. 剩余电流监测体系缺失： 两厂均"目前没有监测"，未建立持续监测机制。'
    )
    document.save(output)

    validate_rendered_markdown_docx(
        output,
        "\n\n".join(
            [
                "**1. 多处剩余电流严重超标：**",
                '**2. 剩余电流监测体系缺失：** 两厂均"目前没有监测"，未建立持续监测机制。',
            ]
        ),
    )


def test_packaged_core_renders_markdown_without_external_source(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    template_document = Document()
    template_document.add_paragraph("TEMPLATE SAMPLE BODY MUST NOT SURVIVE")
    template_document.sections[0].header.paragraphs[0].text = "保留的模板页眉"
    template_document.save(template)

    name, data = PackagedDocxCore(template).render_approved_prose(
        "# 标题\n\n正文\n\n| 对象 | 动作 |\n| --- | --- |\n| 1A2 | 复核连接 |"
    )

    rendered = Document(io.BytesIO(data))
    assert name == "配电安全专家咨询报告.docx"
    rendered_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "正文" in rendered_text
    assert "TEMPLATE SAMPLE BODY MUST NOT SURVIVE" not in rendered_text
    assert rendered.sections[0].header.paragraphs[0].text == "保留的模板页眉"
    assert rendered.tables[0].cell(1, 1).text == "复核连接"
