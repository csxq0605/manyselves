from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from autoreport.core.reporting.coverage import evaluate_coverage
from autoreport.core.reporting.models import (
    Claim,
    ClaimKind,
    EvidenceItem,
    ModuleDraft,
    ProjectManifest,
    ReportRequest,
    ReviewIssue,
    SourceLocation,
)
from autoreport.core.reporting.rendering.docx import DocxRenderer
from autoreport.core.reporting.report_state import build_report_state


def _state():
    request = ReportRequest(
        instruction="生成 2.4",
        target_modules=["2.4"],
        missing_evidence_policy="skip",
    )
    evidence = EvidenceItem(
        id="ev-cable",
        subject="车间配电房/1A2",
        fact="电缆状态=NG",
        source=SourceLocation(
            file_id="file-s44",
            path=Path("Inputs/S4-4诊断工作用表.xlsx"),
            sheet="低配评估详情",
            cell="H5:I5",
        ),
        module_id="2.4",
        submodule_id="2.4.2.5",
        photo_refs=["ID_MISSING"],
    )
    claim = Claim(
        id="claim-cable",
        module_id="2.4",
        submodule_id="2.4.2.5",
        kind=ClaimKind.CONCLUSION,
        text="客户诊断表将该检查项记录为异常。",
        evidence_ids=["ev-cable"],
        skill_ids=["pds.module24.installation@1.0.0"],
    )
    draft = ModuleDraft(
        module_id="2.4",
        markdown="draft",
        evidence_ids=["ev-cable"],
        claims=[claim],
        approved=True,
    )
    issue = ReviewIssue(
        module_id="2.4",
        submodule_id="2.4.2.5",
        claim_id="claim-cable",
        kind="missing_photo",
        message="NG 但缺少同一行照片",
        severity="warning",
    )
    return build_report_state(
        title="配电安全评估报告",
        request=request,
        manifest=ProjectManifest(),
        coverage=evaluate_coverage(request, [evidence]),
        evidence_items=[evidence],
        photo_assets=[],
        module_drafts=[draft],
        review_issues=[issue],
    )


def test_docx_renderer_writes_headings_issue_table_and_image_fallback(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    Document().save(template_path)
    output_path = tmp_path / "Outputs" / "配电安全评估报告.docx"

    result = DocxRenderer(template_path, asset_root=tmp_path).render(_state(), output_path)

    assert result.output_path == output_path
    assert output_path.exists()
    assert result.issue_count == 1
    assert result.missing_photo_ids == ["ID_MISSING"]
    rendered = Document(output_path)
    texts = [paragraph.text for paragraph in rendered.paragraphs]
    assert "配电安全评估报告" in texts
    assert "1 配电评估概述" in texts
    assert "1.1 评估背景" in texts
    assert "1.2 待提升问题与建议概览" in texts
    assert "2 评估内容描述" in texts
    assert any("2.4 配电设备/元件风险" in text for text in texts)
    assert any("2.4.2.5 电缆、桥架、母线安装问题" in text for text in texts)
    assert "3 结论与建议" in texts
    assert "3.1 风险/问题汇总与概览" in texts
    assert any("图片缺失：ID_MISSING" in text for text in texts)
    assert rendered.tables
    assert any(table.cell(0, 0).text == "3.2 改善行动列表与优先级" for table in rendered.tables)
    assert rendered.tables[-1].cell(0, 0).text == "审校与补证事项"
    assert rendered.tables[-1].cell(1, 0).text == "子模块"
    for style_name in ("Normal", "Heading 1", "Heading 2"):
        style = rendered.styles[style_name]
        assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == "Hiragino Sans GB"
