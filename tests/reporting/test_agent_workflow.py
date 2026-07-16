import asyncio
import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from manyselves.core.loops.bus import MessageBus
from manyselves.core.providers.base import LLMProvider
from manyselves.core.reporting.agentic_models import (
    AgentResult,
    AgentRunStatus,
    AuditSubmission,
    ClaimRecord,
    CrossReviewSubmission,
    EditedReportSubmission,
    ModuleSubmission,
    PlanSubmission,
    TaskEnvelope,
    WorkflowDecisionSubmission,
)
from manyselves.core.reporting.models import (
    REPORT_MODULE_IDS,
    EvidenceItem,
    PhotoAsset,
    ReportRequest,
    ReviewIssue,
    RevisionRequest,
    SourceLocation,
)
from manyselves.core.reporting.revisions import RevisionCoordinator
from manyselves.core.reporting.service import ReportingService
from manyselves.core.reporting.session_summary import AgentSessionSummary, SessionSummaryStore
from manyselves.core.reporting.taxonomy import REPORT_TAXONOMY
from manyselves.core.reporting.workflow import ReportWorkflowRunner
from manyselves.core.tools.task_board import TaskBoard
from manyselves.interfaces.types import TaskStatus


def blocking_issue(
    *,
    kind: str,
    message: str,
    claim_id: str = "C-2.4-001",
    round: int = 0,
) -> ReviewIssue:
    return ReviewIssue(
        module_id="2.4",
        submodule_id="2.4.1.1",
        claim_id=claim_id,
        kind=kind,
        message=message,
        severity="blocking",
        round=round,
        affected_claim_ids=[claim_id],
        evidence_refs=["E-0001"],
        blocking_reason="继续交付会形成未经解决的专业错误。",
        resolution_criteria=["责任专家修订后由 Evidence Auditor 复审确认问题已解决"],
        owner_agent_id="module-2.4-specialist",
    )


class ScriptedWorkflowAgents:
    def __init__(
        self, target_modules=REPORT_MODULE_IDS, photo_ids=(), workspace: Path | None = None
    ):
        self.target_modules = tuple(target_modules)
        self.photo_ids = tuple(photo_ids)
        self.active_specialists = 0
        self.max_active_specialists = 0
        self.closed = False
        self.specialist_ids: list[str] = []
        self.envelopes: dict[str, TaskEnvelope] = {}
        self.decision_calls = 0
        self.workspace = workspace

    async def run(self, definition, envelope, shared_artifacts, *, workflow_id, session_key=None):
        agent_id = definition.id
        if agent_id == "report-planner":
            payload = PlanSubmission(
                module_tasks=[
                    TaskEnvelope(
                        task_id=f"planned-{module_id}",
                        run_id=envelope.run_id,
                        agent_id=f"module-{module_id}-specialist",
                        objective=f"分析模块 {module_id}",
                    )
                    for module_id in self.target_modules
                ],
                rationale="五个专业并行分析后汇合。",
            )
        elif agent_id.startswith("module-"):
            self.specialist_ids.append(agent_id)
            self.envelopes[agent_id] = envelope
            module_id = agent_id.removeprefix("module-").removesuffix("-specialist")
            self.active_specialists += 1
            self.max_active_specialists = max(self.max_active_specialists, self.active_specialists)
            await asyncio.sleep(0.02)
            self.active_specialists -= 1
            narrative = f"模块 {module_id} 从本专业机理出发形成主动分析。"
            photo_claim = module_id == "2.4" and bool(self.photo_ids)
            claim_text = "1A2 柜连接状态需要复核" if photo_claim else narrative
            if photo_claim:
                narrative += claim_text + "。"
            payload = ModuleSubmission(
                module_id=module_id,
                markdown=narrative,
                submodule_narratives={
                    submodule_id: f"{submodule_id}：{narrative}"
                    for submodule_id in REPORT_TAXONOMY[module_id].submodules
                },
                claims=[
                    ClaimRecord(
                        id=f"C-{module_id}-001",
                        module_id=module_id,
                        submodule_id=next(iter(REPORT_TAXONOMY[module_id].submodules)),
                        text=claim_text,
                        claim_type="technical_interpretation",
                        source_ids=["E-0001"] if photo_claim else [],
                        confidence=1.0 if photo_claim else 0.5,
                        footnote_required=photo_claim,
                        unresolved=not photo_claim,
                    )
                ],
                source_ids=[],
                unresolved_questions=["尚无客户事实，保留判断边界"],
                revision=envelope.revision,
            )
        elif agent_id == "evidence-auditor":
            module_id = envelope.task_id.split("-")[1]
            payload = AuditSubmission(
                module_id=module_id,
                approved=True,
                issues=[],
                checked_claim_ids=[f"C-{module_id}-001"],
            )
        elif agent_id == "main-agent":
            self.decision_calls += 1
            module_id = envelope.task_id.split("module-")[-1].split("-r")[0]
            payload = WorkflowDecisionSubmission(
                decision="revise",
                rationale="审计提出了可处理的新问题。",
                target_module_ids=[module_id],
            )
        elif agent_id == "cross-module-reviewer":
            payload = CrossReviewSubmission(
                approved=True,
                global_constraints=["所有未知均保持为未知"],
            )
        elif agent_id == "chief-editor":
            module_narratives = {
                module_id: f"模块 {module_id} 从本专业机理出发形成主动分析。"
                for module_id in REPORT_MODULE_IDS
            }
            if self.photo_ids:
                module_narratives["2.4"] += "1A2 柜连接状态需要复核。"
            payload = EditedReportSubmission(
                title="配电安全专家咨询报告",
                overview="本报告按五个专业视角综合审视配电安全。",
                module_narratives=module_narratives,
                conclusion="现有资料不足以形成客户现场事实结论，后续应补充核验。",
                protected_claim_ids=[f"C-{module_id}-001" for module_id in REPORT_MODULE_IDS],
                citation_anchors=(
                    {"C-2.4-001": "1A2 柜连接状态需要复核"} if self.photo_ids else {}
                ),
                photo_ids=list(self.photo_ids),
            )
        else:
            raise AssertionError(agent_id)
        result = AgentResult(
            task_id=envelope.task_id,
            run_id=envelope.run_id,
            agent_id=agent_id,
            session_id=session_key or envelope.task_id,
            status=AgentRunStatus.COMPLETED,
            payload=payload,
        )
        if self.workspace is not None:
            SessionSummaryStore(self.workspace).save(
                AgentSessionSummary(
                    summary_id=f"summary-{envelope.task_id}-{envelope.revision}",
                    run_id=envelope.run_id,
                    task_id=envelope.task_id,
                    agent_id=agent_id,
                    session_id=session_key or envelope.task_id,
                    status="completed",
                    objective=envelope.objective,
                    input_refs=[],
                    output_refs=[],
                    issue_refs=[],
                    message_count=1,
                    tool_names=["submit_result"],
                    agent_rationale="测试摘要",
                )
            )
        return result

    async def close_workflow(self, workflow_id):
        self.closed = True


class NeverCalledProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="never-called")

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        raise AssertionError("ReportWorkflowRunner must use the injected scripted agents")


class PostDeliveryRevisionAgents(ScriptedWorkflowAgents):
    def __init__(self, *, drift: bool = False):
        super().__init__(target_modules=("2.4",))
        self.drift = drift

    async def run(self, definition, envelope, shared_artifacts, *, workflow_id, session_key=None):
        result = await super().run(
            definition,
            envelope,
            shared_artifacts,
            workflow_id=workflow_id,
            session_key=session_key,
        )
        if definition.id == "module-2.4-specialist":
            payload = result.payload
            assert isinstance(payload, ModuleSubmission)
            narratives = dict(payload.submodule_narratives)
            narratives["2.4.1.1"] = "2.4.1.1：已根据反馈修订设备状态边界。"
            if self.drift:
                narratives["2.4.2.2"] = "2.4.2.2：额外改变了未授权范围。"
            result = result.model_copy(
                update={
                    "payload": payload.model_copy(
                        update={
                            "markdown": "模块 2.4 已根据交付后反馈完成局部修订。",
                            "submodule_narratives": narratives,
                        }
                    )
                }
            )
        elif definition.id == "chief-editor":
            payload = result.payload
            assert isinstance(payload, EditedReportSubmission)
            narratives = dict(payload.module_narratives)
            narratives["2.4"] = "模块 2.4 已根据交付后反馈完成局部修订。"
            result = result.model_copy(
                update={"payload": payload.model_copy(update={"module_narratives": narratives})}
            )
        return result


class DriftingRevisionAgents(ScriptedWorkflowAgents):
    async def run(self, definition, envelope, shared_artifacts, *, workflow_id, session_key=None):
        result = await super().run(
            definition,
            envelope,
            shared_artifacts,
            workflow_id=workflow_id,
            session_key=session_key,
        )
        if definition.id == "evidence-auditor" and envelope.revision == 0:
            result.payload = AuditSubmission(
                module_id="2.4",
                approved=False,
                issues=[blocking_issue(kind="unsupported", message="只修订容量判断")],
                checked_claim_ids=[],
            )
        if definition.id == "module-2.4-specialist" and envelope.revision == 1:
            assert isinstance(result.payload, ModuleSubmission)
            payload = result.payload.model_dump(mode="python")
            payload["submodule_narratives"]["2.4.2.2"] = "越界修改接地子模块"
            result.payload = ModuleSubmission.model_validate(payload)
        return result


class RepeatedAuditAgents(ScriptedWorkflowAgents):
    async def run(self, definition, envelope, shared_artifacts, *, workflow_id, session_key=None):
        result = await super().run(
            definition,
            envelope,
            shared_artifacts,
            workflow_id=workflow_id,
            session_key=session_key,
        )
        if definition.id == "evidence-auditor":
            result.payload = AuditSubmission(
                module_id="2.4",
                approved=False,
                issues=[
                    blocking_issue(
                        kind="unsupported",
                        message="仍需定向修订",
                        round=envelope.revision,
                    )
                ],
                checked_claim_ids=[],
            )
        if definition.id == "main-agent" and self.decision_calls >= 2:
            result.payload = WorkflowDecisionSubmission(
                decision="stop_incomplete",
                rationale="相同审计问题重复出现且没有新证据，继续返工不会收敛。",
            )
        return result


class AcceptingBlockingAgents(ScriptedWorkflowAgents):
    async def run(self, definition, envelope, shared_artifacts, *, workflow_id, session_key=None):
        result = await super().run(
            definition,
            envelope,
            shared_artifacts,
            workflow_id=workflow_id,
            session_key=session_key,
        )
        if definition.id == "evidence-auditor":
            result.payload = AuditSubmission(
                module_id="2.4",
                approved=False,
                issues=[blocking_issue(kind="unsupported", message="现场事实缺少证据。")],
                checked_claim_ids=[],
            )
        if definition.id == "main-agent":
            result.payload = WorkflowDecisionSubmission(
                decision="accept",
                rationale="尝试直接绕过阻断问题。",
                target_module_ids=["2.4"],
            )
        return result


class CrossRevisionAuditAgents(ScriptedWorkflowAgents):
    def __init__(self):
        super().__init__()
        self.cross_round = 0

    async def run(self, definition, envelope, shared_artifacts, *, workflow_id, session_key=None):
        if definition.id == "main-agent":
            self.decision_calls += 1
            payload = (
                WorkflowDecisionSubmission(
                    decision="revise",
                    rationale="先定向修订 2.4。",
                    target_module_ids=["2.4"],
                    target_submodule_ids=["2.4.1.1"],
                )
                if self.decision_calls == 1
                else WorkflowDecisionSubmission(
                    decision="stop_incomplete",
                    rationale="责任审计确认冲突仍未解决。",
                )
            )
            return AgentResult(
                task_id=envelope.task_id,
                run_id=envelope.run_id,
                agent_id=definition.id,
                session_id=session_key or envelope.task_id,
                status=AgentRunStatus.COMPLETED,
                payload=payload,
            )
        result = await super().run(
            definition,
            envelope,
            shared_artifacts,
            workflow_id=workflow_id,
            session_key=session_key,
        )
        if definition.id == "cross-module-reviewer":
            self.cross_round += 1
            result.payload = CrossReviewSubmission(
                approved=self.cross_round > 1,
                issues=(
                    []
                    if self.cross_round > 1
                    else [
                        blocking_issue(
                            kind="cross_conflict",
                            message="2.4 判断与其他模块冲突。",
                        )
                    ]
                ),
            )
        if definition.id == "evidence-auditor" and "-cross-" in envelope.task_id:
            result.payload = AuditSubmission(
                module_id="2.4",
                approved=False,
                issues=[
                    blocking_issue(
                        kind="cross_conflict_unresolved",
                        message="定向修订仍未解决冲突。",
                    )
                ],
                checked_claim_ids=[],
            )
        return result


@pytest.mark.asyncio
async def test_full_five_module_workflow_runs_parallel_barrier_editor_and_handoff_docx(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv(
        "MANYSELVES_HANDOFF_DOCX_CORE",
        str(tmp_path / "external-core-must-not-be-used.py"),
    )
    project_template = tmp_path / "Templates/report_template.docx"
    project_template.parent.mkdir(parents=True)
    template_document = Document()
    template_document.add_paragraph("PROJECT TEMPLATE MARKER")
    template_document.save(project_template)
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    agents = ScriptedWorkflowAgents()
    state = {
        "run_id": "run-full",
        "request": ReportRequest(
            instruction="生成完整配电安全专家报告",
            missing_evidence_policy="draft",
        ),
    }

    await ReportWorkflowRunner(service, agents).run(state)

    assert agents.max_active_specialists == 5
    assert agents.closed is True
    assert set(state["module_submissions"]) == set(REPORT_MODULE_IDS)
    report_path = tmp_path / "Outputs/Reports/配电安全专家咨询报告.docx"
    assert report_path.is_file()
    rendered = Document(report_path)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "PROJECT TEMPLATE MARKER" in text
    assert "模块 2.1 从本专业机理出发形成主动分析。" in text
    assert "模块 2.5 从本专业机理出发形成主动分析。" in text
    version_path = tmp_path / "Work/report-versions/run-full/version.json"
    assert version_path.is_file()
    version = json.loads(version_path.read_text(encoding="utf-8"))
    assert version["parent_version_id"] is None
    assert version["artifact_refs"]["final_docx"].startswith(
        "Work/report-versions/run-full/artifacts/"
    )
    assert version["artifact_refs"]["report_template"].startswith(
        "Work/report-versions/run-full/artifacts/"
    )
    provenance = json.loads(
        (tmp_path / "Work/runs/run-full/template-provenance.json").read_text(
            encoding="utf-8"
        )
    )
    assert provenance["source"] == "project"
    assert provenance["selected_path"] == "Templates/report_template.docx"
    assert len(provenance["sha256"]) == 64
    render_log = json.loads(
        (tmp_path / "Outputs/Reports/render-log.json").read_text(encoding="utf-8")
    )
    assert render_log["template"]["source"] == "project"
    assert render_log["template"]["sha256"] == provenance["sha256"]
    assert version["skill_provenance"]


@pytest.mark.asyncio
async def test_partial_request_runs_only_target_module_and_propagates_requirements(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    agents = ScriptedWorkflowAgents(target_modules=("2.4",))
    state = {
        "run_id": "run-partial",
        "request": ReportRequest(
            instruction="只重写设备模块",
            target_modules=["2.4"],
            execution_requirements=["深度核对现场图片"],
            missing_evidence_policy="draft",
        ),
    }

    await ReportWorkflowRunner(service, agents).run(state)

    assert agents.specialist_ids == ["module-2.4-specialist"]
    assert "深度核对现场图片" in agents.envelopes["module-2.4-specialist"].constraints
    assert (tmp_path / "Outputs/Modules/2.4.md").is_file()
    assert not (tmp_path / "Outputs/Reports/配电安全专家咨询报告.docx").exists()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "policy,required_constraint",
    [
        ("draft", "待核实或不确定性"),
        ("skip", "标注“未评估”"),
    ],
)
async def test_missing_evidence_choice_reaches_specialist_as_contract(
    tmp_path: Path, policy: str, required_constraint: str
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    agents = ScriptedWorkflowAgents(target_modules=("2.4",))
    state = {
        "run_id": f"run-{policy}",
        "request": ReportRequest(
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy=policy,
        ),
    }

    await ReportWorkflowRunner(service, agents).run(state)

    constraints = agents.envelopes["module-2.4-specialist"].constraints
    assert any(required_constraint in value for value in constraints)


@pytest.mark.asyncio
async def test_workflow_records_revision_diff_instead_of_rejecting_business_drift(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    agents = DriftingRevisionAgents(target_modules=("2.4",))
    state = {
        "run_id": "run-drift",
        "request": ReportRequest(
            instruction="修订设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="draft",
        ),
    }

    await ReportWorkflowRunner(service, agents).run(state)

    diff = tmp_path / "Work/runs/run-drift/reviews/diff-2.4-r1.json"
    assert diff.is_file()
    assert "2.4.2.2" in diff.read_text(encoding="utf-8")


@pytest.mark.asyncio
async def test_active_workflow_passes_traceable_photo_to_renderer(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    photo_path = tmp_path / "Work/assets/IMG-1.png"
    photo_path.parent.mkdir(parents=True)
    Image.new("RGB", (40, 30), color="red").save(photo_path)
    template_photo_count = len(Document(service.report_template_path).inline_shapes)

    async def normalize(state: dict) -> None:
        state["evidence_items"] = [
            EvidenceItem(
                id="E-0001",
                subject="1A2 柜",
                fact="连接点存在异常",
                source=SourceLocation(file_id="F-1", path=Path("Inputs/check.xlsx"), cell="A2"),
                module_id="2.4",
                submodule_id="2.4.1.1",
                photo_refs=["IMG-1"],
            )
        ]
        state["photo_assets"] = [
            PhotoAsset(
                id="IMG-1",
                path=photo_path.relative_to(tmp_path),
                sha256="test",
                media_type="image/png",
                source_member="media/image1.png",
            )
        ]

    monkeypatch.setattr(service, "_normalize_evidence", normalize)
    state = {
        "run_id": "run-photo",
        "request": ReportRequest(
            instruction="生成带现场图片的完整报告",
            missing_evidence_policy="draft",
        ),
    }

    await ReportWorkflowRunner(service, ScriptedWorkflowAgents(photo_ids=("IMG-1",))).run(state)

    rendered = Document(tmp_path / "Outputs/Reports/配电安全专家咨询报告.docx")
    assert len(rendered.inline_shapes) == template_photo_count + 1


@pytest.mark.asyncio
async def test_lead_agent_stops_repeated_revision_loop_and_task_board_closes(
    tmp_path: Path,
) -> None:
    board = TaskBoard()
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=board,
        llm_provider=NeverCalledProvider(),
    )
    state = {
        "run_id": "run-budget",
        "request": ReportRequest(
            instruction="验证主决策 Agent 识别重复返工",
            target_modules=["2.4"],
            missing_evidence_policy="draft",
        ),
    }

    agents = RepeatedAuditAgents(target_modules=("2.4",))
    with pytest.raises(Exception, match="相同审计问题重复出现"):
        await ReportWorkflowRunner(service, agents).run(state)

    assert agents.decision_calls == 2

    tasks = board.get_all()
    assert tasks
    assert all(
        task.status in {TaskStatus.COMPLETED, TaskStatus.FAILED, TaskStatus.BLOCKED}
        for task in tasks
    )
    assert any("main-agent:decision-module-2.4-r1" in task.brief for task in tasks)


@pytest.mark.asyncio
async def test_main_agent_cannot_accept_unresolved_blocking_module_issue(tmp_path: Path) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    state = {
        "run_id": "run-blocking-accept",
        "request": ReportRequest(
            instruction="验证阻断问题不可绕过",
            target_modules=["2.4"],
            missing_evidence_policy="draft",
        ),
    }

    with pytest.raises(Exception, match="blocking"):
        await ReportWorkflowRunner(
            service,
            AcceptingBlockingAgents(target_modules=("2.4",)),
        ).run(state)


@pytest.mark.asyncio
async def test_cross_revision_must_pass_responsibility_audit_before_state_update(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    state = {
        "run_id": "run-cross-audit",
        "request": ReportRequest(
            instruction="验证跨模块返修重新审计",
            missing_evidence_policy="draft",
        ),
    }
    agents = CrossRevisionAuditAgents()

    with pytest.raises(Exception, match="责任审计确认冲突仍未解决"):
        await ReportWorkflowRunner(service, agents).run(state)

    assert agents.decision_calls == 2


@pytest.mark.asyncio
async def test_post_delivery_revision_restores_parent_reuses_untouched_modules_and_delivers_full_docx(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    baseline_state = {
        "run_id": "run-baseline",
        "request": ReportRequest(
            instruction="生成完整基线报告",
            missing_evidence_policy="draft",
        ),
    }
    await ReportWorkflowRunner(service, ScriptedWorkflowAgents(workspace=tmp_path)).run(
        baseline_state
    )
    baseline_version = baseline_state["report_version"]
    untouched_hash = baseline_version.artifact_sha256["module_submission:2.1"]

    agents = PostDeliveryRevisionAgents()
    result = await RevisionCoordinator(service, agents).run(
        RevisionRequest(
            baseline_version_id="run-baseline",
            feedback="只修订设备状态判断边界。",
            target_module_ids=["2.4"],
            target_submodule_ids=["2.4.1.1"],
            promote_to_skill=True,
            promote_skill_id="pds.module24.device-risk",
        )
    )

    assert result.status == "completed"
    assert result.feedback_record_id
    assert not (tmp_path / "Capabilities/skills/manifest.json").exists()
    child = baseline_version.__class__.model_validate_json(
        (tmp_path / f"Work/report-versions/{result.run_id}/version.json").read_text(
            encoding="utf-8"
        )
    )
    assert child.parent_version_id == "run-baseline"
    assert child.artifact_sha256["module_submission:2.1"] == untouched_hash
    assert (
        child.artifact_sha256["module_submission:2.4"]
        != baseline_version.artifact_sha256["module_submission:2.4"]
    )
    assert agents.envelopes["module-2.4-specialist"].context_summary_refs
    assert (tmp_path / "Outputs/Reports/配电安全专家咨询报告.docx").is_file()


@pytest.mark.asyncio
async def test_post_delivery_scope_drift_creates_reviewable_expansion_request(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    baseline_state = {
        "run_id": "run-baseline",
        "request": ReportRequest(
            instruction="生成完整基线报告",
            missing_evidence_policy="draft",
        ),
    }
    await ReportWorkflowRunner(service, ScriptedWorkflowAgents()).run(baseline_state)

    result = await RevisionCoordinator(service, PostDeliveryRevisionAgents(drift=True)).run(
        RevisionRequest(
            baseline_version_id="run-baseline",
            feedback="只修订设备状态判断边界。",
            target_module_ids=["2.4"],
            target_submodule_ids=["2.4.1.1"],
        )
    )

    assert result.status == "needs_scope_expansion"
    assert result.scope_expansion_request_id
    request_path = (
        tmp_path
        / f"Work/runs/{result.run_id}/scope-expansions/{result.scope_expansion_request_id}.json"
    )
    expansion = json.loads(request_path.read_text(encoding="utf-8"))
    assert expansion["unexpected_submodule_ids"] == ["2.4.2.2"]
    assert not (tmp_path / f"Work/report-versions/{result.run_id}/version.json").exists()
