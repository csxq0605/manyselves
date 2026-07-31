import asyncio
import fcntl
import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from manyselves.core.loops.bus import MessageBus
from manyselves.core.providers.base import LLMProvider
from manyselves.core.reporting.decisions import EvidenceDecisionStore
from manyselves.core.reporting.mappers.common import MappingResult
from manyselves.core.reporting.models import (
    EvidenceDecisionRequest,
    EvidenceItem,
    ManifestFile,
    OutputArtifact,
    PhotoAsset,
    ProjectManifest,
    ReportRequest,
    RevisionRequest,
    SourceLocation,
    UserSupplement,
)
from manyselves.core.reporting.revisions import RevisionCoordinator
from manyselves.core.reporting.service import ReportingRunResult, ReportingService
from manyselves.core.reporting.store import ReportingStore
from manyselves.core.reporting.workflow import (
    AgentWorkflowError,
    AgentWorkflowBlocked,
    ReportingNeedsDecisionError,
    ReportWorkflowRunner,
)
from manyselves.core.tools.reporting_tool import (
    ResumeReportingWorkflowTool,
    RunReportingWorkflowTool,
)
from manyselves.core.tools.task_board import TaskBoard
from manyselves.core.usage_ledger import UsageLedger


class TemplateResolutionProvider(LLMProvider):
    def __init__(self):
        super().__init__("test", model="never-called")

    async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
        raise AssertionError("template resolution must not call the provider")


@pytest.mark.asyncio
async def test_evidence_normalization_scopes_canonical_photos_to_run_and_remaps_primary(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    input_path = tmp_path / "Inputs/S4-4诊断工作用表.xlsx"
    input_path.parent.mkdir(parents=True)
    input_path.write_bytes(b"fixture")
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    manifest_file = ManifestFile(
        id="file-s44",
        path=input_path.relative_to(tmp_path),
        sha256="abc",
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        purpose="s4-4",
    )
    state = {
        "run_id": "report-photo-bindings",
        "project_manifest": ProjectManifest(files=[manifest_file]),
        "parsed_artifacts": [],
    }
    evidence = [
        EvidenceItem(
            id=f"ev-source-{index}",
            subject=f"检查项 {index}",
            fact=f"结果 {index}=NG",
            source=SourceLocation(
                file_id=manifest_file.id,
                path=manifest_file.path,
                sheet="配电房合规性",
                cell=f"{column}6:{photo_column}6",
            ),
            module_id="2.2",
            submodule_id="2.2.2.3",
            photo_refs=["ID_SHARED"],
        )
        for index, column, photo_column in (
            (1, "B", "C"),
            (2, "D", "E"),
        )
    ]
    observed_output_dirs: list[Path] = []

    def fake_extract(_path: Path, *, output_dir: Path) -> dict[str, PhotoAsset]:
        observed_output_dirs.append(output_dir)
        output_dir.mkdir(parents=True)
        source_path = output_dir / "source-0001.jpeg"
        source_path.write_bytes(b"image")
        return {
            "ID_SHARED": PhotoAsset(
                id="ID_SHARED",
                path=source_path,
                sha256="def",
                media_type="image/jpeg",
                source_member="xl/media/image1.jpeg",
                source_image_id="ID_SHARED",
            )
        }

    monkeypatch.setattr(
        "manyselves.core.reporting.service.extract_wps_images",
        fake_extract,
    )
    monkeypatch.setattr(
        "manyselves.core.reporting.service.map_s4_4",
        lambda *_args, **_kwargs: MappingResult(evidence_items=evidence, gaps=[]),
    )

    await service._normalize_evidence(state)

    expected_root = (
        tmp_path
        / "Work/runs/report-photo-bindings/assets/file-s44"
    )
    assert observed_output_dirs == [expected_root]
    assert [item.id for item in state["evidence_items"]] == ["E-0001", "E-0002"]
    assert [item.photo_refs for item in state["evidence_items"]] == [
        ["P-0001"],
        ["P-0001"],
    ]
    assert len(state["photo_assets"]) == 1
    photo = state["photo_assets"][0]
    assert photo.id == "P-0001"
    assert photo.path == Path(
        "Work/runs/report-photo-bindings/assets/file-s44/P-0001.jpeg"
    )
    assert photo.source_image_id == "ID_SHARED"
    assert photo.primary_evidence_id == "E-0001"
    photo_view = tmp_path / photo.path
    assert photo_view.is_symlink()
    blobs = list((tmp_path / "Work/content/sha256").glob("*/*/*"))
    assert len(blobs) == 1
    assert photo_view.resolve() == blobs[0]


def test_service_template_snapshots_share_one_content_blob(tmp_path: Path) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    source = tmp_path / "Templates/report_template.docx"
    source.parent.mkdir(parents=True)
    Document().save(source)
    first = tmp_path / "Work/runs/run-a/templates/report_template.docx"
    second = tmp_path / "Work/runs/run-b/templates/report_template.docx"

    first_path, first_sha, first_blob = service.snapshot_content(source, first)
    second_path, second_sha, second_blob = service.snapshot_content(source, second)

    assert first_path.is_symlink()
    assert second_path.is_symlink()
    assert first_path.resolve() == second_path.resolve()
    assert first_sha == second_sha
    assert first_blob == second_blob
    assert len(list((tmp_path / "Work/content/sha256").glob("*/*/*"))) == 1


def test_snapshot_content_atomically_replaces_matching_file_with_cas_view(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    source = tmp_path / "Inputs/photo.jpeg"
    target = tmp_path / "Work/runs/run-photo/assets/P-0001.jpeg"
    source.parent.mkdir(parents=True)
    target.parent.mkdir(parents=True)
    source.write_bytes(b"same-photo-bytes")
    target.write_bytes(source.read_bytes())

    path, digest, blob_ref = service.snapshot_content(
        source,
        target,
        replace_existing_with_view=True,
    )

    assert path == target
    assert path.is_symlink()
    assert path.read_bytes() == b"same-photo-bytes"
    assert path.resolve() == (tmp_path / blob_ref).resolve()
    assert path.resolve().name == digest
    assert list(target.parent.glob(".*.cas-view")) == []


def test_snapshot_content_keeps_existing_file_if_view_staging_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    source = tmp_path / "Inputs/photo.jpeg"
    target = tmp_path / "Work/runs/run-photo/assets/P-0001.jpeg"
    source.parent.mkdir(parents=True)
    target.parent.mkdir(parents=True)
    source.write_bytes(b"same-photo-bytes")
    target.write_bytes(source.read_bytes())

    def fail_staging(*_args, **_kwargs):
        raise OSError("staging failed")

    monkeypatch.setattr(service.content_store, "link_view", fail_staging)
    with pytest.raises(OSError, match="staging failed"):
        service.snapshot_content(
            source,
            target,
            replace_existing_with_view=True,
        )

    assert not target.is_symlink()
    assert target.read_bytes() == b"same-photo-bytes"
    assert list(target.parent.glob(".*.cas-view")) == []


@pytest.mark.asyncio
async def test_service_retains_same_identity_registry_while_waiting_for_user(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    request = ReportRequest(instruction="生成完整报告", missing_evidence_policy="draft")
    run_id = "report-retained-identities"
    runner_ids: list[int] = []

    async def scripted_run(workflow, state: dict) -> None:
        runner_ids.append(id(workflow.agent_runner))
        if len(runner_ids) == 1:
            raise ReportingNeedsDecisionError("等待用户确认命名。")

    delivered = tmp_path / "Outputs/Reports/result.docx"
    delivered.parent.mkdir(parents=True)
    delivered.write_bytes(b"verified")
    monkeypatch.setattr(ReportWorkflowRunner, "run", scripted_run)
    monkeypatch.setattr(
        "manyselves.core.reporting.service.verify_current_run_outputs",
        lambda *_args, **_kwargs: [delivered],
    )

    first = await service._execute(request, run_id)
    assert first.status == "needs_decision"
    assert list(service._active_agent_runners) == [
        f"full-power-distribution-report:{run_id}"
    ]

    second = await service._execute(request, run_id, resume=True)
    assert second.status == "completed"
    assert runner_ids[0] == runner_ids[1]
    assert service._active_agent_runners == {}


@pytest.mark.asyncio
async def test_agent_blocked_state_is_preserved(tmp_path: Path, monkeypatch) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def blocked(_self, _state: dict) -> None:
        raise AgentWorkflowBlocked("chief-editor", "edit", "missing reviewed modules")

    monkeypatch.setattr(ReportWorkflowRunner, "run", blocked)
    result = await service.run(ReportRequest(instruction="生成报告"))
    assert result.status == "blocked"
    assert result.error == "missing reviewed modules"


@pytest.mark.asyncio
async def test_render_existing_bypasses_all_analysis_agents(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "Work/drafts/approved.md"
    source.parent.mkdir(parents=True)
    source.write_text(
        "# 配电安全专家咨询报告\n\n## 1. 配电评估概述\n\n这是无需重新分析的正文。\n",
        encoding="utf-8",
    )
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def must_not_run(_self, _state: dict) -> None:
        raise AssertionError("render_existing must not start the analysis workflow")

    monkeypatch.setattr(ReportWorkflowRunner, "run", must_not_run)
    result = await service.run(
        ReportRequest(
            operation="render_existing",
            instruction="将已有 Markdown 转为 Word",
            source_markdown_ref=Path("Work/drafts/approved.md"),
            output_filename="approved.docx",
        )
    )

    assert result.status == "completed", result.error
    assert result.output_paths == [tmp_path / "Outputs/Reports/approved.docx"]
    rendered = Document(result.output_paths[0])
    rendered_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "配电安全专家咨询报告" in rendered_text
    assert "这是无需重新分析的正文。" in rendered_text
    assert "2025年8月" not in rendered_text
    render_request = json.loads(
        (tmp_path / f"Work/runs/{result.run_id}/render-request.json").read_text(
            encoding="utf-8"
        )
    )
    assert render_request["source_markdown_ref"] == "Work/drafts/approved.md"
    assert render_request["output_ref"] == "Outputs/Reports/approved.docx"


@pytest.mark.asyncio
async def test_render_existing_preserves_approved_sources_and_positive_advice(
    tmp_path: Path,
) -> None:
    source = tmp_path / "Work/drafts/approved-with-sources.md"
    source.parent.mkdir(parents=True)
    source.write_text(
        """# 配电安全专家咨询报告

## 1. 配电评估概述

# 1 评估概述

## 1.1 评估范围与方法

评估方法包括现场测量（S4-4诊断工作用表）和文件审阅（S2-1收资表）。

### 2.1.5 系统无功补偿与电容柜问题

【结论】

OK，评估过程中未发现异常（S4-6评估信息汇总表）。

【建议】

保持年度检查并记录电容柜投切状态。
""",
        encoding="utf-8",
    )
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    result = await service.run(
        ReportRequest(
            operation="render_existing",
            instruction="原样渲染批准正文",
            source_markdown_ref=Path("Work/drafts/approved-with-sources.md"),
            output_filename="approved-with-sources.docx",
        )
    )

    assert result.status == "completed", result.error
    rendered = Document(result.output_paths[0])
    rendered_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "S4-4诊断工作用表" in rendered_text
    assert "S2-1收资表" in rendered_text
    assert "S4-6评估信息汇总表" in rendered_text
    assert "保持年度检查并记录电容柜投切状态" in rendered_text
    assert "1.1. 评估背景" in rendered_text
    assert "1.1 评估范围与方法" not in rendered_text


@pytest.mark.asyncio
async def test_distill_template_skill_dispatches_only_the_standalone_action(
    tmp_path: Path, monkeypatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def distill(_self, state: dict) -> None:
        path = Path("Work/report-template-writing/SKILL.md")
        service.store.write_text(path.as_posix(), "---\nname: report-template-writing\n---\n")
        state["output_artifacts"] = [OutputArtifact(kind="skill", path=path)]

    async def must_not_run(_self, _state: dict) -> None:
        raise AssertionError("distill_template_skill must not start report writing")

    monkeypatch.setattr(ReportWorkflowRunner, "distill_template_skill", distill)
    monkeypatch.setattr(ReportWorkflowRunner, "run", must_not_run)
    monkeypatch.setattr(ReportWorkflowRunner, "aggregate_existing", must_not_run)
    result = await service.run(
        ReportRequest(
            operation="distill_template_skill",
            instruction="只更新模板写作 Skill",
            target_modules=[],
        )
    )

    assert result.status == "completed"
    assert result.output_paths == [tmp_path / "Work/report-template-writing/SKILL.md"]


def test_legacy_template_skill_reports_the_missing_boundary_manifest(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    root = tmp_path / "Work/report-template-writing"
    for relative in (
        "SKILL.md",
        "references/analysis-language.md",
        "references/synthesis.md",
        "references/visual-organization.md",
        "references/quality-rubric.md",
    ):
        target = root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(f"# {relative}\n", encoding="utf-8")
    (root / "source.json").write_text(
        json.dumps(
            {
                "source": "legacy-template-skill",
                "template_ref": "Work/runs/old/templates/template.docx",
            }
        ),
        encoding="utf-8",
    )
    runner = object.__new__(ReportWorkflowRunner)
    runner.service = service

    with pytest.raises(
        AgentWorkflowError,
        match=r"缺少文件：Work/report-template-writing/boundary\.json",
    ):
        runner._require_template_skill({})


def test_template_skill_loader_rejects_persisted_result_part_marker(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    root = tmp_path / "Work/report-template-writing"
    files = {
        "SKILL.md": "---\nname: report-template-writing\n---\n# Skill\n",
        "references/analysis-language.md": (
            "<persisted_result_part sha256=" + "a" * 64 + " characters=100>\n"
        ),
        "references/synthesis.md": "# Synthesis\n",
        "references/visual-organization.md": "# Visual\n",
        "references/quality-rubric.md": "# Rubric\n",
    }
    for relative, content in files.items():
        target = root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
    boundary = {
        "policy_version": 1,
        "transferred_categories": [
            "analysis_method",
            "synthesis_method",
            "visual_method",
            "quality_check",
        ],
        "excluded_categories": [
            "domain_knowledge",
            "domain_standard_or_threshold",
            "project_fact_or_number",
            "customer_identity",
            "project_finding_or_risk",
            "project_conclusion_or_recommendation",
            "evidence_or_claim_identifier",
        ],
        "boundary_statement": (
            "本固定 Skill 仅迁移可跨项目复用的分析、综合、视觉组织与质量检查方法，"
            "不迁移任何客户身份、项目事实、具体数值、风险结论、行动建议、专业机理、"
            "标准阈值或证据标识；这些内容必须在当前运行中分别由 Evidence 和 Knowledge 提供。"
        ),
    }
    (root / "boundary.json").write_text(
        json.dumps(boundary, ensure_ascii=False),
        encoding="utf-8",
    )
    artifact_paths = [root / relative for relative in files]
    artifact_paths.append(root / "boundary.json")
    (root / "source.json").write_text(
        json.dumps(
            {
                "boundary_policy_version": 1,
                "boundary_ref": "Work/report-template-writing/boundary.json",
                "artifact_sha256": {
                    path.relative_to(root).as_posix(): hashlib.sha256(
                        path.read_bytes()
                    ).hexdigest()
                    for path in artifact_paths
                },
            }
        ),
        encoding="utf-8",
    )
    runner = object.__new__(ReportWorkflowRunner)
    runner.service = service

    with pytest.raises(
        AgentWorkflowError,
        match="含有内部 persisted_result_part 历史标记",
    ):
        runner._require_template_skill({})


@pytest.mark.asyncio
async def test_distill_template_skill_preserves_the_original_failure_in_checkpoint(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def fail_distillation(_self, _state: dict, _workflow_id: str) -> None:
        raise ValueError("template snapshot validation failed")

    monkeypatch.setattr(
        ReportWorkflowRunner,
        "_distill_template_skill",
        fail_distillation,
    )

    result = await service.run(
        ReportRequest(
            operation="distill_template_skill",
            instruction="只更新模板写作 Skill",
            target_modules=[],
        )
    )
    checkpoint = json.loads(
        (
            tmp_path
            / f"Work/runs/{result.run_id}/workflow-state.json"
        ).read_text(encoding="utf-8")
    )

    assert result.status == "failed"
    assert result.error == "template snapshot validation failed"
    assert result.usage["provider_attempts"] == 0
    assert checkpoint["activity"] == "template-skill-distillation"
    assert checkpoint["status"] == "failed"
    assert checkpoint["error"] == "template snapshot validation failed"


@pytest.mark.asyncio
async def test_aggregate_existing_uses_aggregation_route_then_render(
    tmp_path: Path, monkeypatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def aggregate(_self, state: dict) -> None:
        markdown_ref = Path("Outputs/Reports/combined.md")
        service.store.write_text(markdown_ref.as_posix(), "# 汇总报告\n\n已汇总正文。\n")
        state["aggregate_markdown_ref"] = markdown_ref
        state["output_artifacts"] = [OutputArtifact(kind="report", path=markdown_ref)]

    async def must_not_run(_self, _state: dict) -> None:
        raise AssertionError("aggregate_existing must not start the full analysis workflow")

    monkeypatch.setattr(ReportWorkflowRunner, "aggregate_existing", aggregate)
    monkeypatch.setattr(ReportWorkflowRunner, "run", must_not_run)
    result = await service.run(
        ReportRequest(
            operation="aggregate_existing",
            instruction="汇总已有模块并生成 Word",
            output_filename="combined.docx",
        )
    )

    assert result.status == "completed"
    assert result.output_paths == [
        tmp_path / "Outputs/Reports/combined.md",
        tmp_path / "Outputs/Reports/combined.docx",
    ]
    Document(result.output_paths[1])


def test_project_report_template_takes_priority_without_restarting_service(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    assert service.report_template_source == "packaged"
    assert service.report_template_path == service.packaged_report_template_path

    project_template = tmp_path / "Templates/report_template.docx"
    project_template.parent.mkdir(parents=True)
    Document().save(project_template)

    assert service.report_template_source == "project"
    assert service.report_template_path == project_template


def test_expert_document_is_selected_only_for_skill_distillation(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    project_template = tmp_path / "Templates/report_template.docx"
    expert_source = (
        tmp_path / "Templates/配电安全专家咨询报告(专家优化版).docx"
    )
    project_template.parent.mkdir(parents=True)
    Document().save(project_template)
    Document().save(expert_source)

    assert service.resolve_report_template() == (project_template, "project")
    assert service.resolve_skill_distillation_template() == (
        expert_source,
        "expert-skill-source",
    )


def test_project_report_template_path_must_be_a_file(tmp_path: Path) -> None:
    (tmp_path / "Templates/report_template.docx").mkdir(parents=True)
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    with pytest.raises(ValueError, match="must be a DOCX file"):
        service.resolve_report_template()


def test_reporting_service_rejects_missing_provider(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="requires an LLM provider"):
        ReportingService(
            tmp_path,
            bus=MessageBus(),
            task_board=TaskBoard(),
            llm_provider=None,  # type: ignore[arg-type]
        )


def test_reporting_tool_rejects_missing_provider(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="requires an LLM provider"):
        RunReportingWorkflowTool(
            workspace=tmp_path,
            bus=MessageBus(),
            task_board=TaskBoard(),
            llm_provider=None,  # type: ignore[arg-type]
        )


@pytest.mark.asyncio
async def test_reporting_tool_defaults_new_report_to_draft_policy(
    tmp_path: Path,
) -> None:
    captured: dict[str, ReportRequest] = {}

    class CapturingController:
        def start(self, request: ReportRequest) -> dict:
            captured["request"] = request
            return {"status": "running", "run_id": "report-default-draft"}

    tool = RunReportingWorkflowTool(
        workspace=tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
        controller=CapturingController(),  # type: ignore[arg-type]
    )

    payload = await tool(
        instruction="从 Inputs 生成完整报告",
        operation="full_report",
    )

    assert payload["status"] == "running"
    assert captured["request"].missing_evidence_policy == "draft"


@pytest.mark.asyncio
async def test_service_returns_resumable_decision_before_calling_provider_when_evidence_requires_confirmation(
    tmp_path: Path,
) -> None:
    class NeverCalledProvider(LLMProvider):
        def __init__(self):
            super().__init__("test", model="never-called")

        async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
            raise AssertionError("blocked workflow must not call the provider")

    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )

    result = await service.run(
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="ask",
        )
    )

    assert result.status == "needs_user_decision"
    assert result.decision_id
    assert result.missing_evidence
    assert result.error is None
    workflow_state = json.loads(
        (tmp_path / f"Work/runs/{result.run_id}/workflow-state.json").read_text(encoding="utf-8")
    )
    assert workflow_state["status"] == "blocked"
    decision = json.loads(
        (tmp_path / f"Work/runs/{result.run_id}/decisions/{result.decision_id}.json").read_text(
            encoding="utf-8"
        )
    )
    assert decision["allowed_actions"] == ["supplement", "draft", "skip", "stop"]


@pytest.mark.asyncio
async def test_supplement_rescans_the_same_run_after_process_restart(tmp_path: Path) -> None:
    class NeverCalledProvider(LLMProvider):
        def __init__(self):
            super().__init__("test", model="never-called")

        async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
            raise AssertionError("coverage should block before provider use")

    first = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    pending = await first.run(
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="ask",
        )
    )

    restarted = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    rescanned = await restarted.resume(
        pending.decision_id,
        "supplement",
        [UserSupplement(id="US-rescan", content="已补充资料")],
    )

    assert rescanned.run_id == pending.run_id
    assert rescanned.status == "needs_user_decision"
    assert rescanned.decision_id != pending.decision_id


@pytest.mark.asyncio
async def test_user_decision_resume_syncs_new_facts_into_same_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    run_id = "report-user-fact-resume"
    request = ReportRequest(
        operation="module_report",
        instruction="恢复 2.4",
        target_modules=["2.4"],
        missing_evidence_policy="draft",
        max_provider_attempts=1,
        max_total_tokens=10_000,
    )
    service.store.write_json(
        f"Work/runs/{run_id}/request.json", request.model_dump(mode="json")
    )
    service._save_run(
        ReportingRunResult(
            run_id=run_id,
            status="needs_decision",
            error="XMR-001：需要用户确认两个工厂名称是否指同一实体。",
        )
    )
    service.store.write_json(
        f"Work/runs/{run_id}/workflow-state.json",
        {"run_id": run_id, "activity": "cross-module-review", "status": "failed"},
    )
    seen = {}

    async def resumed_execute(request, resumed_run_id, *, resume=False):
        seen["request"] = request
        seen["run_id"] = resumed_run_id
        seen["resume"] = resume
        return ReportingRunResult(run_id=resumed_run_id, status="needs_decision")

    monkeypatch.setattr(service, "_execute", resumed_execute)

    result = await service.resume_run(
        run_id,
        supplements=[
            UserSupplement(
                id="US-factory-name",
                content="嘉仕工厂与芜湖工厂是同一实体，全报告统一使用芜湖工厂。",
            )
        ],
    )

    assert result.run_id == run_id
    assert seen["run_id"] == run_id
    assert seen["resume"] is True
    assert seen["request"].user_supplements[0].id == "US-factory-name"
    assert "嘉仕工厂" in seen["request"].user_supplements[0].content
    persisted = json.loads(
        (tmp_path / f"Work/runs/{run_id}/user-supplements.json").read_text(
            encoding="utf-8"
        )
    )
    assert persisted["supplements"] == [
        item.model_dump(mode="json")
        for item in seen["request"].user_supplements
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("previous_status", "previous_error"),
    [("failed", "submission failed"), ("cancelled", "interrupted by user")],
)
async def test_checkpointed_terminal_run_can_resume_same_run(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    previous_status: str,
    previous_error: str,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    run_id = "report-failed-resume"
    request = ReportRequest(
        operation="module_report",
        instruction="恢复失败的 2.4",
        target_modules=["2.4"],
        missing_evidence_policy="draft",
        max_provider_attempts=2,
        max_total_tokens=10_000,
    )
    service.store.write_json(
        f"Work/runs/{run_id}/request.json", request.model_dump(mode="json")
    )
    service.store.write_json(
        f"Work/runs/{run_id}/workflow-state.json",
        {"run_id": run_id, "activity": "module-pipelines", "status": "failed"},
    )
    service._save_run(
        ReportingRunResult(
            run_id=run_id, status=previous_status, error=previous_error
        )
    )
    seen = {}

    async def resumed_execute(request, resumed_run_id, *, resume=False):
        seen["request"] = request
        seen["run_id"] = resumed_run_id
        seen["resume"] = resume
        return ReportingRunResult(run_id=resumed_run_id, status="needs_decision")

    monkeypatch.setattr(service, "_execute", resumed_execute)

    result = await service.resume_run(
        run_id, max_provider_attempts=6, max_total_tokens=20_000
    )

    assert result.run_id == run_id
    assert seen["run_id"] == run_id
    assert seen["resume"] is True


@pytest.mark.asyncio
async def test_crashed_in_progress_checkpoint_without_terminal_result_can_resume(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    run_id = "report-crashed-resume"
    request = ReportRequest(
        operation="module_report",
        instruction="恢复崩溃中的 2.4",
        target_modules=["2.4"],
        missing_evidence_policy="draft",
        max_provider_attempts=2,
        max_total_tokens=10_000,
    )
    service.store.write_json(
        f"Work/runs/{run_id}/request.json", request.model_dump(mode="json")
    )
    service.store.write_json(
        f"Work/runs/{run_id}/workflow-state.json",
        {"run_id": run_id, "activity": "module-work", "status": "in_progress"},
    )
    seen = {}

    async def resumed_execute(request, resumed_run_id, *, resume=False):
        seen["run_id"] = resumed_run_id
        seen["resume"] = resume
        return ReportingRunResult(run_id=resumed_run_id, status="needs_decision")

    monkeypatch.setattr(service, "_execute", resumed_execute)

    result = await service.resume_run(
        run_id, max_provider_attempts=6, max_total_tokens=20_000
    )

    assert result.run_id == run_id
    assert seen == {"run_id": run_id, "resume": True}


def test_run_lock_rejects_a_second_live_process_and_releases_on_close(
    tmp_path: Path,
) -> None:
    first = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    second = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    handle = first._acquire_run_lock("report-live")
    try:
        with pytest.raises(RuntimeError, match="already active"):
            second._acquire_run_lock("report-live")
    finally:
        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        handle.close()

    released = second._acquire_run_lock("report-live")
    fcntl.flock(released.fileno(), fcntl.LOCK_UN)
    released.close()


@pytest.mark.asyncio
async def test_budget_resume_reuses_same_revision_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    run_id = "report-revision-budget-resume"
    request = RevisionRequest(
        baseline_version_id="baseline-version",
        feedback="恢复局部修订",
        target_module_ids=["2.4"],
        max_provider_attempts=1,
        max_total_tokens=10_000,
    )
    service.store.write_json(
        f"Work/runs/{run_id}/revision-request.json", request.model_dump(mode="json")
    )
    service._save_run(
        ReportingRunResult(
            run_id=run_id,
            status="needs_decision",
            error="报告运行预算已耗尽，已保存完成模块和写作分段。",
        )
    )
    UsageLedger(tmp_path, run_id).record_attempt(
        run_id=run_id,
        task_id="module-2.4-post-delivery-r1",
        agent_id="module-2.4-specialist",
        input_tokens=800,
        output_tokens=200,
        total_tokens=1000,
        status="success",
    )
    seen = {}

    async def resumed_revision(_coordinator, revision, *, run_id=None, resume=False):
        seen["request"] = revision
        seen["run_id"] = run_id
        seen["resume"] = resume
        return ReportingRunResult(run_id=run_id, status="needs_decision")

    monkeypatch.setattr(RevisionCoordinator, "run", resumed_revision)
    result = await service.resume_run(
        run_id,
        cost_control_mode="warn",
        max_provider_attempts=4,
        max_total_tokens=20_000,
    )

    assert result.run_id == run_id
    assert seen["run_id"] == run_id
    assert seen["resume"] is True
    assert seen["request"].cost_control_mode == "warn"
    assert seen["request"].max_provider_attempts == 4
    assert seen["request"].max_total_tokens == 20_000


@pytest.mark.asyncio
async def test_revision_resume_keeps_new_supplements_structured(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    run_id = "report-revision-structured-supplement"
    request = RevisionRequest(
        baseline_version_id="baseline-version",
        feedback="恢复局部修订",
        target_module_ids=["2.4"],
    )
    service.store.write_json(
        f"Work/runs/{run_id}/revision-request.json", request.model_dump(mode="json")
    )
    service._save_run(
        ReportingRunResult(
            run_id=run_id,
            status="needs_decision",
            error="等待用户确认设备名称。",
        )
    )
    service.store.write_json(
        f"Work/runs/{run_id}/workflow-state.json",
        {"run_id": run_id, "activity": "module-work", "status": "failed"},
    )
    seen = {}

    async def resumed_revision(_coordinator, revision, *, run_id=None, resume=False):
        seen["request"] = revision
        return ReportingRunResult(run_id=run_id, status="needs_decision")

    monkeypatch.setattr(RevisionCoordinator, "run", resumed_revision)
    supplement = UserSupplement(
        id="US-revision-device-name",
        content="设备名称确认为 1A2 进线柜。",
        scope="module",
        target_ids=["2.4"],
        stages=["module_authoring", "module_review"],
    )

    await service.resume_run(run_id, supplements=[supplement])

    assert seen["request"].feedback == "恢复局部修订"
    assert seen["request"].user_supplements == [supplement]
    persisted = json.loads(
        (tmp_path / f"Work/runs/{run_id}/user-supplements.json").read_text(
            encoding="utf-8"
        )
    )
    assert persisted["supplements"] == [supplement.model_dump(mode="json")]


@pytest.mark.asyncio
@pytest.mark.parametrize("action", ["draft", "skip"])
async def test_draft_and_skip_resume_same_run_with_explicit_policy(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, action: str
) -> None:
    class NeverCalledProvider(LLMProvider):
        def __init__(self):
            super().__init__("test", model="never-called")

        async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
            raise AssertionError("provider is replaced by workflow spy")

    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    pending = await service.run(
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="ask",
        )
    )
    seen: dict[str, str] = {}

    async def completed(self, state: dict) -> None:
        seen["policy"] = state["request"].missing_evidence_policy
        output = self.service.workspace / "Outputs/Reports/current-run.txt"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text("verified report output", encoding="utf-8")
        state["output_artifacts"] = [
            OutputArtifact(kind="report", path=output.relative_to(self.service.workspace))
        ]

    monkeypatch.setattr(ReportWorkflowRunner, "run", completed)
    restarted = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    result = await restarted.resume(pending.decision_id, action)

    assert result.run_id == pending.run_id
    assert result.status == "completed"
    assert seen["policy"] == action
    choice = json.loads(
        (tmp_path / f"Work/runs/{pending.run_id}/evidence-choice.json").read_text(encoding="utf-8")
    )
    assert choice["selected_action"] == action


@pytest.mark.asyncio
async def test_invalid_draft_supplement_does_not_consume_evidence_decision(
    tmp_path: Path,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    pending = await service.run(
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="ask",
        )
    )

    with pytest.raises(
        ValueError,
        match="supplements are only valid with action=supplement",
    ):
        await service.resume(
            pending.decision_id,
            "draft",
            ["按 draft 继续"],  # type: ignore[list-item]
        )

    decision = service.decisions.load(pending.decision_id)
    assert decision.status == "pending"
    assert decision.selected_action is None
    request = ReportRequest.model_validate_json(
        (
            tmp_path / f"Work/runs/{pending.run_id}/request.json"
        ).read_text(encoding="utf-8")
    )
    assert request.missing_evidence_policy == "ask"
    assert not (
        tmp_path / f"Work/runs/{pending.run_id}/evidence-choice.json"
    ).exists()


@pytest.mark.asyncio
async def test_resolved_draft_decision_reconciles_stale_ask_request(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )
    pending = await service.run(
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="ask",
        )
    )
    service.decisions.resolve(pending.decision_id, "draft")
    seen: dict[str, str] = {}

    async def completed(self, state: dict) -> None:
        seen["policy"] = state["request"].missing_evidence_policy
        output = self.service.workspace / "Outputs/Reports/reconciled.txt"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text("verified report output", encoding="utf-8")
        state["output_artifacts"] = [
            OutputArtifact(kind="report", path=output.relative_to(self.service.workspace))
        ]

    monkeypatch.setattr(ReportWorkflowRunner, "run", completed)

    result = await service.resume(pending.decision_id, "draft")

    assert result.run_id == pending.run_id
    assert result.status == "completed"
    assert seen["policy"] == "draft"
    request = ReportRequest.model_validate_json(
        (
            tmp_path / f"Work/runs/{pending.run_id}/request.json"
        ).read_text(encoding="utf-8")
    )
    assert request.missing_evidence_policy == "draft"


@pytest.mark.asyncio
async def test_service_never_marks_a_run_completed_without_verified_current_outputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def no_outputs(_self, _state: dict) -> None:
        return None

    monkeypatch.setattr(ReportWorkflowRunner, "run", no_outputs)

    result = await service.run(ReportRequest(instruction="生成报告"))

    assert result.status == "failed"
    assert result.output_paths == []
    assert result.error == "workflow finished without verified output artifacts"


@pytest.mark.asyncio
async def test_stop_marks_run_incomplete_without_success_artifact(tmp_path: Path) -> None:
    class NeverCalledProvider(LLMProvider):
        def __init__(self):
            super().__init__("test", model="never-called")

        async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
            raise AssertionError("stop must not call provider")

    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )
    pending = await service.run(
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
            missing_evidence_policy="ask",
        )
    )

    result = await service.resume(pending.decision_id, "stop")

    assert result.run_id == pending.run_id
    assert result.status == "stopped_incomplete"
    assert result.output_paths == []


@pytest.mark.asyncio
async def test_user_cancellation_is_saved_as_cancelled_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=TemplateResolutionProvider(),
    )

    async def cancelled(_self, _state: dict) -> None:
        raise asyncio.CancelledError

    monkeypatch.setattr(ReportWorkflowRunner, "run", cancelled)

    result = await service.run(ReportRequest(instruction="生成报告"))

    assert result.status == "cancelled"
    saved = json.loads((tmp_path / f"Work/runs/{result.run_id}.json").read_text(encoding="utf-8"))
    assert saved["status"] == "cancelled"
    assert saved["error"] == "interrupted by user"


@pytest.mark.asyncio
async def test_resume_tool_uses_persisted_decision_id(tmp_path: Path) -> None:
    class NeverCalledProvider(LLMProvider):
        def __init__(self):
            super().__init__("test", model="never-called")

        async def chat(self, messages, tools=None, temperature=0.1, max_tokens=8192):
            raise AssertionError("stop must not call provider")

    run_id = "report-tool-resume"
    ReportingStore(tmp_path).write_json(
        f"Work/runs/{run_id}/request.json",
        ReportRequest(
            operation="module_report",
            instruction="生成设备模块",
            target_modules=["2.4"],
        ).model_dump(mode="json"),
    )
    decision = EvidenceDecisionStore(tmp_path).create(
        EvidenceDecisionRequest(
            decision_id="evidence-tool-resume",
            run_id=run_id,
            missing_items=["2.4.1.1 缺少证据"],
            affected_modules=["2.4"],
        )
    )
    tool = ResumeReportingWorkflowTool(
        tmp_path,
        MessageBus(),
        TaskBoard(),
        llm_provider=NeverCalledProvider(),
    )

    payload = await tool(decision.decision_id, "stop")

    assert payload["run_id"] == run_id
    assert payload["status"] == "running"
    task = tool.controller._tasks[run_id]
    await task
    persisted = ReportingRunResult.model_validate_json(
        (tmp_path / f"Work/runs/{run_id}.json").read_text(encoding="utf-8")
    )
    assert persisted.status == "stopped_incomplete"
