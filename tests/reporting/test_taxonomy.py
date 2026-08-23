import json
import subprocess
import sys
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from manyselves.capabilities.distribution_reporting.domain.taxonomy import (
    REPORT_TAXONOMY,
    reset_report_taxonomy,
    resolve_submodule,
)
from manyselves.core.reporting.models import ManifestFile, ProjectManifest
from manyselves.core.reporting.workflow import ReportWorkflowRunner


def test_capability_taxonomy_import_is_independent_of_core_reporting() -> None:
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            "\n".join(
                (
                    "import json",
                    "import sys",
                    "from manyselves.capabilities.distribution_reporting.domain "
                    "import taxonomy",
                    "print(json.dumps({",
                    "    'module': taxonomy.__name__,",
                    "    'reporting_modules': sorted(",
                    "        name for name in sys.modules",
                    "        if name.startswith('manyselves.core.reporting')",
                    "    ),",
                    "}))",
                )
            ),
        ],
        check=True,
        capture_output=True,
        text=True,
    )

    assert json.loads(completed.stdout) == {
        "module": (
            "manyselves.capabilities.distribution_reporting.domain.taxonomy"
        ),
        "reporting_modules": [],
    }


def test_taxonomy_contains_fixed_module_24_submodules() -> None:
    module = REPORT_TAXONOMY["2.4"]

    assert "2.4.2.2" in module.submodules
    assert module.submodules["2.4.2.2"].title == "等电位连接与接地问题"
    assert module.submodules["2.4.4"].title == "末端配电抽查情况"


def test_taxonomy_matches_v2_handoff_for_all_five_modules() -> None:
    assert {module_id: module.title for module_id, module in REPORT_TAXONOMY.items()} == {
        "2.1": "配电系统架构问题",
        "2.2": "环境工况风险",
        "2.3": "针对故障的保护",
        "2.4": "配电设备/元件内在风险",
        "2.5": "运维管理与风险管控机制",
    }
    assert REPORT_TAXONOMY["2.1"].submodules["2.1.1"].title == "配电系统负荷分配与过载风险"
    assert REPORT_TAXONOMY["2.2"].submodules["2.2.1.1"].title == "谐波风险情况"
    assert REPORT_TAXONOMY["2.3"].submodules["2.3.2"].title == "零序/漏电的防范"
    assert REPORT_TAXONOMY["2.5"].submodules["2.5.1"].title == "SOP/EOP"
    assert REPORT_TAXONOMY["2.5"].submodules["2.5.3.3"].title == "配电设备维保覆盖"


def test_resolve_submodule_rejects_unknown_identifier() -> None:
    with pytest.raises(ValueError, match="unknown report submodule"):
        resolve_submodule("2.4.99")


def test_xlsx_taxonomy_drives_all_five_complete_modules(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评估信息汇总表"
    row = 1
    expected_titles = {}
    for module_number in range(1, 6):
        module_id = f"2.{module_number}"
        sheet.cell(row, 2, float(module_id))
        sheet.cell(row, 3, f"动态模块 {module_id}")
        expected_titles[module_id] = f"动态模块 {module_id}"
        row += 1
        leaf_count = 14 if module_id == "2.4" else 2
        for leaf in range(1, leaf_count + 1):
            sheet.cell(row, 2, f"{module_id}.{leaf}")
            sheet.cell(row, 4, f"动态标题 {module_id}.{leaf}")
            row += 1
    path = tmp_path / "S4-6动态目录.xlsx"
    workbook.save(path)
    workbook.close()

    runner = object.__new__(ReportWorkflowRunner)
    runner.service = SimpleNamespace(workspace=tmp_path)
    state = {
        "project_manifest": ProjectManifest(
            files=[
                ManifestFile(
                    id="file-s4-6",
                    path=Path("Inputs/S4-6动态目录.xlsx"),
                    snapshot_ref=path.relative_to(tmp_path),
                    sha256="0" * 64,
                    media_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    purpose="s4-6",
                )
            ]
        )
    }
    token = runner._prepare_report_taxonomy(state)
    try:
        assert {key: value.title for key, value in REPORT_TAXONOMY.items()} == expected_titles
        assert REPORT_TAXONOMY["2.4"].submodules["2.4.14"].title == "动态标题 2.4.14"
        assert "_report_taxonomy_token" not in state
        assert deepcopy(state)["report_taxonomy"] == state["report_taxonomy"]
    finally:
        reset_report_taxonomy(token)
