import errno
import fcntl
import hashlib
import json
import os
import stat
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from manyselves.core.reporting.delivery import DeliveryPackage, ProjectDelivery
from manyselves.core.reporting.locks import exclusive_reporting_writer_lock
from manyselves.core.reporting.models import REPORT_MODULE_IDS
from manyselves.core.reporting.output_ownership import (
    OutputOwnerStore,
    build_output_owner,
)
from manyselves.core.reporting.storage_compaction import (
    CompletedRunOutputCompactor,
    StorageCompactionError,
)
from manyselves.core.reporting.store import ReportingStore
from manyselves.core.reporting.versions import ReportVersionStore


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _snapshot_tree(workspace: Path) -> dict[str, tuple[str, str]]:
    snapshot: dict[str, tuple[str, str]] = {}
    for path in sorted(workspace.rglob("*")):
        relative = path.relative_to(workspace).as_posix()
        if path.is_symlink():
            snapshot[relative] = ("symlink", os.readlink(path))
        elif path.is_file():
            snapshot[relative] = ("file", _sha256(path))
    return snapshot


def _build_completed_workspace(
    tmp_path: Path,
    *,
    run_id: str = "report-complete",
    publish_owner: bool = True,
    revision: bool = False,
) -> tuple[Path, str, Path, Path]:
    workspace = tmp_path / "project"
    run_root = workspace / "Work/runs" / run_id
    target = workspace / CompletedRunOutputCompactor.FINAL_DOCX_REF
    target.parent.mkdir(parents=True, exist_ok=True)
    report = Document()
    report.add_heading("Verified final report", level=1)
    report.add_paragraph("Stable completed-run content.")
    report.save(target)

    modules: dict[str, Path] = {}
    for module_id in REPORT_MODULE_IDS:
        module_path = run_root / "approved-modules" / f"{module_id}.md"
        module_path.parent.mkdir(parents=True, exist_ok=True)
        module_path.write_text(
            f"# Module {module_id}\n\nApproved.",
            encoding="utf-8",
        )
        modules[module_id] = module_path
    report_state = run_root / "report-state.json"
    report_state.write_text('{"status":"complete"}\n', encoding="utf-8")
    source_index = run_root / "evidence-source-index.md"
    source_index.write_text("# Sources\n", encoding="utf-8")
    source_index_docx = run_root / "evidence-source-index.docx"
    Document().save(source_index_docx)

    receipt = ProjectDelivery(run_root / "delivery").deliver(
        DeliveryPackage(
            report_id=run_id,
            version=run_id,
            module_files=modules,
            final_docx=target,
            report_state=report_state,
            source_index=source_index,
            source_index_docx=source_index_docx,
        )
    )
    store = ReportingStore(workspace)
    receipt_ref = Path(f"Work/runs/{run_id}/delivery-receipt.json")
    receipt_path = store.write_json(
        receipt_ref.as_posix(),
        receipt.model_dump(mode="json"),
    )
    ReportVersionStore(workspace).publish_from_delivery(
        receipt,
        receipt_path,
        {},
        [],
        [],
    )
    completion_ref = Path(f"Work/runs/{run_id}/delivery-completion.json")
    store.write_json(
        completion_ref.as_posix(),
        {
            "status": "completed",
            "run_id": run_id,
            "report_version_id": run_id,
            "delivery_receipt_ref": receipt_ref.as_posix(),
        },
    )
    workflow = {
        "run_id": run_id,
        "status": "completed",
        "activity": "revision-delivery" if revision else "delivery",
        "cross_review_completed": True,
        "final_review_completed": True,
        "delivery_completion_ref": completion_ref.as_posix(),
    }
    if revision:
        workflow["completed_revision_modules"] = ["2.4"]
        store.write_json(
            f"Work/runs/{run_id}/revision-request.json",
            {
                "baseline_version_id": "report-baseline",
                "feedback": "Revise module 2.4",
                "target_module_ids": ["2.4"],
            },
        )
    else:
        workflow["completed_modules"] = list(REPORT_MODULE_IDS)
    store.write_json(
        f"Work/runs/{run_id}/workflow-state.json",
        workflow,
    )
    store.write_json(
        f"Work/runs/{run_id}.json",
        {
            "run_id": run_id,
            "status": "completed",
            "output_paths": [
                str(target.resolve()),
                str(Path(receipt.manifest_path).resolve()),
            ],
            "error": None,
        },
    )
    if publish_owner:
        OutputOwnerStore(workspace).publish(
            build_output_owner(
                run_id=run_id,
                report_version_id=run_id,
                final_docx_sha256=_sha256(target),
                delivery_receipt_ref=receipt_ref,
            )
        )
    blob = workspace / receipt.artifact_refs["final_docx"]
    return workspace, run_id, target, blob


def test_storage_compaction_dry_run_is_strictly_read_only(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    before = _snapshot_tree(workspace)

    record = CompletedRunOutputCompactor(workspace).execute(run_id)

    assert record.action == "dry_run"
    assert record.status == "planned"
    assert record.dry_run is True
    assert record.pre_state == "materialized"
    assert record.post_state == "materialized"
    assert record.owner_verified is True
    assert record.mutation_allowed is True
    assert record.logical_size_bytes == target.stat().st_size
    assert record.estimated_reclaimed_logical_bytes == target.stat().st_size
    assert record.estimated_reclaimed_allocated_bytes == (
        target.stat().st_blocks * 512
    )
    assert _snapshot_tree(workspace) == before
    assert not (
        workspace
        / f"Work/runs/{run_id}/{CompletedRunOutputCompactor.MANIFEST_NAME}"
    ).exists()


def test_storage_compaction_apply_and_rollback_are_verified_and_idempotent(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, blob = _build_completed_workspace(tmp_path)
    receipt_path = workspace / f"Work/runs/{run_id}/delivery-receipt.json"
    version_path = workspace / f"Work/report-versions/{run_id}/version.json"
    immutable_hashes = (_sha256(receipt_path), _sha256(version_path))
    original_size = target.stat().st_size
    compactor = CompletedRunOutputCompactor(workspace)

    applied = compactor.execute(run_id, action="apply")

    assert applied.pre_state == "materialized"
    assert applied.post_state == "cas_view"
    assert applied.status == "completed"
    assert applied.estimated_reclaimed_logical_bytes == original_size
    assert target.is_symlink()
    assert target.resolve() == blob.resolve()
    Document(target)
    assert (_sha256(receipt_path), _sha256(version_path)) == immutable_hashes

    applied_again = compactor.execute(run_id, action="apply")
    assert applied_again.pre_state == "cas_view"
    assert applied_again.estimated_reclaimed_logical_bytes == 0

    rolled_back = compactor.execute(run_id, action="rollback")

    assert rolled_back.pre_state == "cas_view"
    assert rolled_back.post_state == "materialized"
    assert rolled_back.materialized_logical_bytes == original_size
    assert not target.is_symlink()
    assert _sha256(target) == applied.sha256
    Document(target)
    assert (_sha256(receipt_path), _sha256(version_path)) == immutable_hashes

    rolled_back_again = compactor.execute(run_id, action="rollback")
    assert rolled_back_again.pre_state == "materialized"
    assert rolled_back_again.materialized_logical_bytes == 0


def test_storage_compaction_refuses_an_active_run(tmp_path: Path) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    lock_path = workspace / f"Work/runs/{run_id}/.active.lock"
    lock_handle = lock_path.open("a+", encoding="utf-8")
    fcntl.flock(lock_handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    try:
        with pytest.raises(StorageCompactionError, match="active"):
            CompletedRunOutputCompactor(workspace).execute(
                run_id,
                action="apply",
            )
    finally:
        fcntl.flock(lock_handle.fileno(), fcntl.LOCK_UN)
        lock_handle.close()

    assert not target.is_symlink()


def test_storage_compaction_rejects_symlink_run_lock_without_nofollow(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    lock_path = workspace / f"Work/runs/{run_id}/.active.lock"
    outside = tmp_path / "outside-run.lock"
    outside.write_text("untouched", encoding="utf-8")
    lock_path.symlink_to(outside)
    monkeypatch.delattr(os, "O_NOFOLLOW", raising=False)

    with pytest.raises(StorageCompactionError, match="symbolic link"):
        CompletedRunOutputCompactor(workspace).execute(run_id, action="apply")

    assert outside.read_text(encoding="utf-8") == "untouched"
    assert not target.is_symlink()


def test_storage_compaction_lock_fallback_rejects_opened_inode_mismatch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, _target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    lock_path = workspace / f"Work/runs/{run_id}/.active.lock"
    lock_path.write_text("expected", encoding="utf-8")
    outside = tmp_path / "substituted-run.lock"
    outside.write_text("outside", encoding="utf-8")
    outside.chmod(0o644)
    original_mode = stat.S_IMODE(outside.stat().st_mode)
    real_open = os.open

    def substitute_open(path, flags, mode=0o777):
        if Path(path) == lock_path:
            return real_open(outside, flags, mode)
        return real_open(path, flags, mode)

    monkeypatch.delattr(os, "O_NOFOLLOW", raising=False)
    monkeypatch.setattr(os, "open", substitute_open)

    with pytest.raises(StorageCompactionError, match="verified regular file"):
        with compactor._exclusive_run_lock(run_id):
            pass

    assert lock_path.read_text(encoding="utf-8") == "expected"
    assert outside.read_text(encoding="utf-8") == "outside"
    assert stat.S_IMODE(outside.stat().st_mode) == original_mode


def test_storage_compaction_rejects_a_hardlinked_run_lock_before_chmod(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    lock_path = workspace / f"Work/runs/{run_id}/.active.lock"
    outside = tmp_path / "outside-hardlink.lock"
    outside.write_text("outside", encoding="utf-8")
    outside.chmod(0o644)
    original_mode = stat.S_IMODE(outside.stat().st_mode)
    lock_path.hardlink_to(outside)

    with pytest.raises(StorageCompactionError, match="verified regular file"):
        CompletedRunOutputCompactor(workspace).execute(run_id, action="apply")

    assert not target.is_symlink()
    assert outside.read_text(encoding="utf-8") == "outside"
    assert stat.S_IMODE(outside.stat().st_mode) == original_mode


def test_storage_compaction_refuses_another_workspace_writer(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)

    with exclusive_reporting_writer_lock(workspace):
        with pytest.raises(StorageCompactionError, match="shared workspace"):
            CompletedRunOutputCompactor(workspace).execute(
                run_id,
                action="apply",
            )

    assert not target.is_symlink()


@pytest.mark.parametrize("run_id", [".", ".."])
def test_storage_compaction_rejects_dot_run_ids_before_creating_locks(
    tmp_path: Path,
    run_id: str,
) -> None:
    workspace = tmp_path / "project"
    (workspace / "Work/runs").mkdir(parents=True)

    with pytest.raises(ValueError, match="safe, real path segment"):
        CompletedRunOutputCompactor(workspace).execute(
            run_id,
            action="apply",
        )

    assert not (workspace / "Work/.active.lock").exists()
    assert not (workspace / "Work/runs/.active.lock").exists()
    assert not (workspace / "Work/.reporting-writer.lock").exists()


def test_storage_compaction_legacy_run_is_dry_run_only(tmp_path: Path) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(
        tmp_path,
        publish_owner=False,
    )
    compactor = CompletedRunOutputCompactor(workspace)

    planned = compactor.execute(run_id)

    assert planned.owner_verified is False
    assert planned.mutation_allowed is False
    assert planned.estimated_reclaimed_logical_bytes == target.stat().st_size
    with pytest.raises(StorageCompactionError, match="output-owner.json"):
        compactor.execute(run_id, action="apply")
    assert not target.is_symlink()
    assert not (workspace / f"Work/runs/{run_id}/storage-compaction.json").exists()


def test_storage_compaction_accepts_current_completed_revision(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, blob = _build_completed_workspace(
        tmp_path,
        run_id="report-revision-current",
        revision=True,
    )
    compactor = CompletedRunOutputCompactor(workspace)

    planned = compactor.execute(run_id)
    applied = compactor.execute(run_id, action="apply")

    assert planned.completion_kind == "revision"
    assert planned.owner_verified is True
    assert applied.status == "completed"
    assert target.is_symlink()
    assert target.resolve() == blob.resolve()


def test_storage_compaction_refuses_old_run_with_identical_visible_bytes(
    tmp_path: Path,
) -> None:
    workspace, old_run_id, target, _blob = _build_completed_workspace(
        tmp_path,
        run_id="report-old",
        publish_owner=False,
    )
    _workspace, current_run_id, _target, _current_blob = (
        _build_completed_workspace(
            tmp_path,
            run_id="report-current",
            publish_owner=True,
        )
    )
    assert current_run_id != old_run_id

    with pytest.raises(StorageCompactionError, match="does not bind"):
        CompletedRunOutputCompactor(workspace).execute(old_run_id)

    assert not target.is_symlink()


def test_storage_compaction_reports_hardlink_allocation_as_not_reclaimable(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, blob = _build_completed_workspace(tmp_path)
    target.unlink()
    os.link(blob, target)
    assert target.stat().st_nlink > 1

    planned = CompletedRunOutputCompactor(workspace).execute(run_id)

    assert planned.estimated_reclaimed_logical_bytes == target.stat().st_size
    assert planned.allocated_size_bytes == target.stat().st_blocks * 512
    assert planned.target_link_count > 1
    assert planned.estimated_reclaimed_allocated_bytes == 0


def test_storage_compaction_refuses_incomplete_workflow(tmp_path: Path) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    state_path = workspace / f"Work/runs/{run_id}/workflow-state.json"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["final_review_completed"] = False
    state_path.write_text(json.dumps(state), encoding="utf-8")

    with pytest.raises(StorageCompactionError, match="Cross and Final review"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_refuses_failed_or_unbound_top_level_result(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    result_path = workspace / f"Work/runs/{run_id}.json"
    result = json.loads(result_path.read_text(encoding="utf-8"))
    result["status"] = "failed"
    result["error"] = "injected final verification failure"
    result_path.write_text(json.dumps(result), encoding="utf-8")

    with pytest.raises(StorageCompactionError, match="top-level report run"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    result["status"] = "completed"
    result["error"] = None
    result["output_paths"] = [str(target.resolve())]
    result_path.write_text(json.dumps(result), encoding="utf-8")
    with pytest.raises(StorageCompactionError, match="current receipt"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_refuses_receipt_version_mismatch(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    receipt_path = workspace / f"Work/runs/{run_id}/delivery-receipt.json"
    receipt_path.write_text(
        receipt_path.read_text(encoding="utf-8") + "\n",
        encoding="utf-8",
    )

    with pytest.raises(StorageCompactionError, match="report version"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_refuses_tampered_cas_blob(tmp_path: Path) -> None:
    workspace, run_id, target, blob = _build_completed_workspace(tmp_path)
    blob.chmod(0o644)
    blob.write_bytes(b"tampered")

    with pytest.raises(StorageCompactionError, match="hashes|CAS blob"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_refuses_outputs_from_another_run(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    replacement = Document()
    replacement.add_paragraph("A different run replaced the shared output.")
    replacement.save(target)

    with pytest.raises(StorageCompactionError, match="visible final DOCX"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_refuses_delivery_symlink_outside_workspace(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    receipt_path = workspace / f"Work/runs/{run_id}/delivery-receipt.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    delivery_docx = Path(receipt["final_docx"])
    outside = tmp_path / "outside.docx"
    outside.write_bytes(target.read_bytes())
    delivery_docx.unlink()
    delivery_docx.symlink_to(outside)

    with pytest.raises(StorageCompactionError, match="escapes"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_refuses_internal_symlink_delivery_ancestor(
    tmp_path: Path,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    delivery_root = workspace / f"Work/runs/{run_id}/delivery"
    relocated = workspace / "Work/relocated-delivery"
    delivery_root.rename(relocated)
    delivery_root.symlink_to(relocated, target_is_directory=True)

    with pytest.raises(StorageCompactionError, match="symbolic-link ancestor"):
        CompletedRunOutputCompactor(workspace).execute(run_id)

    assert not target.is_symlink()


def test_storage_compaction_post_validation_failure_restores_and_records(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_validate = compactor._validate
    calls = 0

    def fail_second_validation(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("injected post-validation failure")
        return original_validate(*args, **kwargs)

    monkeypatch.setattr(compactor, "_validate", fail_second_validation)

    with pytest.raises(StorageCompactionError, match="original output state"):
        compactor.execute(run_id, action="apply")

    assert not target.is_symlink()
    record = json.loads(
        (workspace / f"Work/runs/{run_id}/storage-compaction.json").read_text(
            encoding="utf-8"
        )
    )
    assert record["status"] == "rolled_back"
    assert record["pre_state"] == "materialized"
    assert record["post_state"] == "materialized"
    assert "injected post-validation failure" in record["error"]


def test_storage_compaction_intent_fsync_failure_prevents_mutation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_inode = target.lstat().st_ino
    original_hash = _sha256(target)
    mutation_called = False

    def fail_directory_fsync(_path: Path) -> None:
        raise OSError(errno.EIO, "injected intent directory fsync failure")

    def mutation_spy(*_args, **_kwargs) -> None:
        nonlocal mutation_called
        mutation_called = True

    monkeypatch.setattr(compactor.store, "fsync_directory", fail_directory_fsync)
    monkeypatch.setattr(compactor, "_mutate", mutation_spy)

    with pytest.raises(OSError) as captured:
        compactor.execute(run_id, action="apply")

    assert captured.value.errno == errno.EIO
    assert mutation_called is False
    assert not target.is_symlink()
    assert target.lstat().st_ino == original_inode
    assert _sha256(target) == original_hash


def test_storage_compaction_terminal_fsync_failure_is_compensated(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_fsync = compactor.store.fsync_directory
    original_hash = _sha256(target)
    fsync_calls = 0

    def fail_second_directory_fsync(path: Path) -> None:
        nonlocal fsync_calls
        fsync_calls += 1
        if fsync_calls == 2:
            raise OSError(errno.EIO, "injected terminal directory fsync failure")
        original_fsync(path)

    monkeypatch.setattr(compactor.store, "fsync_directory", fail_second_directory_fsync)

    with pytest.raises(StorageCompactionError, match="original output state"):
        compactor.execute(run_id, action="apply")

    assert fsync_calls == 3
    assert not target.is_symlink()
    assert _sha256(target) == original_hash
    record = json.loads(
        (workspace / f"Work/runs/{run_id}/storage-compaction.json").read_text(
            encoding="utf-8"
        )
    )
    assert record["status"] == "rolled_back"
    assert "terminal directory fsync failure" in record["error"]


def test_storage_compaction_cas_fsync_failure_is_compensated(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_fsync = compactor.content_store._fsync_directory
    original_hash = _sha256(target)
    fsync_calls = 0

    def fail_first_cas_fsync(path: Path) -> None:
        nonlocal fsync_calls
        fsync_calls += 1
        if fsync_calls == 1:
            raise OSError(errno.EIO, "injected CAS replace fsync failure")
        original_fsync(path)

    monkeypatch.setattr(
        compactor.content_store,
        "_fsync_directory",
        fail_first_cas_fsync,
    )

    with pytest.raises(StorageCompactionError, match="original output state"):
        compactor.execute(run_id, action="apply")

    assert fsync_calls == 2
    assert not target.is_symlink()
    assert _sha256(target) == original_hash
    record = json.loads(
        (workspace / f"Work/runs/{run_id}/storage-compaction.json").read_text(
            encoding="utf-8"
        )
    )
    assert record["status"] == "rolled_back"
    assert "CAS replace fsync failure" in record["error"]


def test_storage_compaction_persistent_cas_fsync_failure_requires_recovery(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_hash = _sha256(target)

    def fail_cas_fsync(_path: Path) -> None:
        raise OSError(errno.EIO, "persistent CAS replace fsync failure")

    monkeypatch.setattr(
        compactor.content_store,
        "_fsync_directory",
        fail_cas_fsync,
    )

    with pytest.raises(StorageCompactionError, match="recovery is required"):
        compactor.execute(run_id, action="apply")

    assert not target.is_symlink()
    assert _sha256(target) == original_hash
    record = json.loads(
        (workspace / f"Work/runs/{run_id}/storage-compaction.json").read_text(
            encoding="utf-8"
        )
    )
    assert record["status"] == "recovery_required"
    assert record["post_state"] == "materialized"
    assert "persistent CAS replace fsync failure" in record["error"]


def test_storage_compaction_completion_record_failure_is_compensated(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_write = compactor._write_record
    writes = 0

    def fail_second_record(record):
        nonlocal writes
        writes += 1
        if writes == 2:
            raise OSError("injected completion journal failure")
        return original_write(record)

    monkeypatch.setattr(compactor, "_write_record", fail_second_record)

    with pytest.raises(StorageCompactionError, match="original output state"):
        compactor.execute(run_id, action="apply")

    assert not target.is_symlink()
    record = json.loads(
        (workspace / f"Work/runs/{run_id}/storage-compaction.json").read_text(
            encoding="utf-8"
        )
    )
    assert writes == 3
    assert record["status"] == "rolled_back"
    assert "injected completion journal failure" in record["error"]


def test_storage_compaction_leaves_intent_when_terminal_journal_is_unavailable(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, run_id, target, _blob = _build_completed_workspace(tmp_path)
    compactor = CompletedRunOutputCompactor(workspace)
    original_write = compactor._write_record
    writes = 0

    def fail_after_intent(record):
        nonlocal writes
        writes += 1
        if writes > 1:
            raise OSError("injected terminal journal outage")
        return original_write(record)

    monkeypatch.setattr(compactor, "_write_record", fail_after_intent)

    with pytest.raises(StorageCompactionError, match="recovery is required"):
        compactor.execute(run_id, action="apply")

    # Compensation restored the bytes, while the durable non-terminal intent
    # remains as an explicit recovery signal because terminal writes failed.
    assert not target.is_symlink()
    record = json.loads(
        (workspace / f"Work/runs/{run_id}/storage-compaction.json").read_text(
            encoding="utf-8"
        )
    )
    assert record["status"] == "intent"
    assert writes == 4


def test_storage_compaction_cli_defaults_to_dry_run(tmp_path: Path) -> None:
    workspace, run_id, _target, _blob = _build_completed_workspace(tmp_path)
    script = (
        Path(__file__).resolve().parents[2]
        / "scripts"
        / "compact_reporting_storage.py"
    )

    result = subprocess.run(
        [
            sys.executable,
            str(script),
            "--workspace",
            str(workspace),
            "--run-id",
            run_id,
        ],
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["action"] == "dry_run"
    assert payload["dry_run"] is True
    assert not (
        workspace
        / f"Work/runs/{run_id}/{CompletedRunOutputCompactor.MANIFEST_NAME}"
    ).exists()
