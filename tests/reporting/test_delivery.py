import hashlib
import json
import shutil
from pathlib import Path

import pytest
from docx import Document

from manyselves.core.reporting.delivery import DeliveryPackage, ProjectDelivery


def _package(tmp_path: Path) -> DeliveryPackage:
    modules = {}
    for module_id in ("2.1", "2.2", "2.3", "2.4", "2.5"):
        path = tmp_path / f"module-{module_id}.md"
        path.write_text(f"# {module_id}\n\n专业正文", encoding="utf-8")
        modules[module_id] = path
    final = tmp_path / "final.docx"
    Document().save(final)
    state = tmp_path / "report-state.json"
    state.write_text('{"approved": true}', encoding="utf-8")
    source_index = tmp_path / "证据与来源索引.md"
    source_index.write_text("## 证据与来源索引\n\n- E-001：测试来源\n", encoding="utf-8")
    source_index_docx = tmp_path / "证据与来源索引.docx"
    source_index_document = Document()
    source_index_document.add_heading("证据与来源索引", level=1)
    source_index_document.add_paragraph("E-001：测试来源")
    source_index_document.save(source_index_docx)
    return DeliveryPackage(
        report_id="report-001",
        version="v1",
        module_files=modules,
        final_docx=final,
        report_state=state,
        source_index=source_index,
        source_index_docx=source_index_docx,
    )


def _publish_legacy_v1(delivery_root: Path, package: DeliveryPackage) -> Path:
    destination = delivery_root / f"{package.report_id}-{package.version}"
    modules_dir = destination / "modules"
    modules_dir.mkdir(parents=True)
    targets = {
        "final_docx": destination / "配电安全专家咨询报告.docx",
        "report_state": destination / "report-state.json",
        "source_index": destination / "证据与来源索引.md",
        "source_index_docx": destination / "证据与来源索引.docx",
        **{
            f"module:{module_id}": modules_dir / f"{module_id}.md"
            for module_id in package.module_files
        },
    }
    sources = {
        "final_docx": package.final_docx,
        "report_state": package.report_state,
        "source_index": package.source_index,
        "source_index_docx": package.source_index_docx,
        **{
            f"module:{module_id}": source
            for module_id, source in package.module_files.items()
        },
    }
    for key, source in sources.items():
        shutil.copyfile(source, targets[key])
    hashes = {
        key: hashlib.sha256(path.read_bytes()).hexdigest()
        for key, path in targets.items()
    }
    (destination / "delivery-manifest.json").write_text(
        json.dumps(
            {
                "report_id": package.report_id,
                "version": package.version,
                "status": "success",
                "modules": list(package.module_files),
                "artifacts": hashes,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return destination


def test_delivery_publishes_complete_five_module_package_atomically(tmp_path: Path) -> None:
    workspace = tmp_path / "project"
    delivery = ProjectDelivery(workspace / "Outputs")

    receipt = delivery.deliver(_package(tmp_path))

    assert receipt.success is True
    assert receipt.final_docx.is_file()
    assert set(receipt.module_files) == {"2.1", "2.2", "2.3", "2.4", "2.5"}
    assert Document(receipt.final_docx)
    assert receipt.source_index.read_text(encoding="utf-8").startswith(
        "## 证据与来源索引"
    )
    assert Document(receipt.source_index_docx).paragraphs[0].text == "证据与来源索引"
    assert receipt.manifest_path.is_file()
    manifest = json.loads(receipt.manifest_path.read_text(encoding="utf-8"))
    assert receipt.storage_version == 2
    assert manifest["manifest_version"] == 2
    assert manifest["storage"] == "sha256-cas"
    assert set(receipt.artifact_refs) == set(manifest["artifacts"])
    assert set(receipt.trusted_handle_refs) == set(manifest["artifacts"])
    assert manifest["trusted_handle_refs"] == {
        key: value.as_posix()
        for key, value in receipt.trusted_handle_refs.items()
    }
    blobs = [path for path in (workspace / "Work/content/sha256").rglob("*") if path.is_file()]
    assert len(blobs) == len(manifest["artifacts"])
    for key, relative in receipt.artifact_refs.items():
        blob = workspace / relative
        assert blob.is_file()
        assert blob.name == receipt.artifact_sha256[key]


def test_delivery_reuses_only_an_identical_complete_package(tmp_path: Path) -> None:
    delivery = ProjectDelivery(tmp_path / "project" / "Outputs")
    package = _package(tmp_path)
    first = delivery.deliver(package)

    resumed = delivery.deliver(package)

    assert resumed.delivery_dir == first.delivery_dir
    assert resumed.artifact_sha256 == first.artifact_sha256

    package.module_files["2.1"].write_text("changed", encoding="utf-8")
    with pytest.raises(ValueError, match="does not match"):
        delivery.deliver(package)


def test_delivery_reuses_legacy_v1_snapshot_without_migrating_it(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "project"
    delivery_root = workspace / "Outputs"
    package = _package(tmp_path)
    destination = _publish_legacy_v1(delivery_root, package)

    receipt = ProjectDelivery(delivery_root).deliver(package)

    assert receipt.delivery_dir == destination
    assert receipt.storage_version == 1
    assert receipt.artifact_refs == {}
    assert not (workspace / "Work/content").exists()


def test_delivery_rejects_corrupted_v2_blob_on_reuse(tmp_path: Path) -> None:
    workspace = tmp_path / "project"
    delivery = ProjectDelivery(workspace / "Outputs")
    package = _package(tmp_path)
    receipt = delivery.deliver(package)
    blob = workspace / receipt.artifact_refs["report_state"]
    blob.chmod(0o644)
    blob.write_text('{"tampered": true}', encoding="utf-8")

    with pytest.raises(ValueError, match="does not match"):
        delivery.deliver(package)


def test_delivery_rejects_incomplete_or_invalid_docx_without_success_receipt(
    tmp_path: Path,
) -> None:
    package = _package(tmp_path)
    package.module_files.pop("2.5")
    delivery_root = tmp_path / "Outputs"

    with pytest.raises(ValueError, match="exactly modules 2.1-2.5"):
        ProjectDelivery(delivery_root).deliver(package)
    assert not delivery_root.exists()

    package = _package(tmp_path)
    package.final_docx.write_text("not a docx", encoding="utf-8")
    with pytest.raises(ValueError, match="Word/WPS-openable"):
        ProjectDelivery(delivery_root).deliver(package)
    assert not delivery_root.exists()
