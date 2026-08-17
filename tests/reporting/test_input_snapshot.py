from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from manyselves.core.loops.bus import MessageBus
from manyselves.core.providers.base import LLMProvider
from manyselves.core.reporting.input_snapshot import RunInputSnapshotStore
from manyselves.core.reporting.models import ReportRequest
from manyselves.core.reporting.research.reference_library import ReferenceLibrary
from manyselves.core.reporting.service import ReportingService
from manyselves.core.tools.task_board import TaskBoard


class NoCallProvider(LLMProvider):
    def __init__(self) -> None:
        super().__init__("test", model="no-call")

    async def chat(self, *args, **kwargs):
        raise AssertionError("input snapshot tests do not call the Provider")


@pytest.mark.asyncio
async def test_queued_run_uses_frozen_inputs_knowledge_and_template(
    tmp_path: Path,
) -> None:
    input_file = tmp_path / "Inputs/site.txt"
    knowledge_file = tmp_path / "Knowledge/guide.md"
    template_file = tmp_path / "Templates/report_template.docx"
    input_file.parent.mkdir(parents=True)
    knowledge_file.parent.mkdir(parents=True)
    template_file.parent.mkdir(parents=True)
    input_file.write_text("ORIGINAL INPUT", encoding="utf-8")
    knowledge_file.write_text("ORIGINAL KNOWLEDGE", encoding="utf-8")
    original_template = Document()
    original_template.add_paragraph("ORIGINAL TEMPLATE")
    original_template.save(template_file)
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NoCallProvider(),
    )
    request = ReportRequest(
        operation="module_report",
        instruction="freeze",
        target_modules=["2.1"],
    )
    run_id = service.prepare_run(request)

    input_file.write_text("CHANGED INPUT", encoding="utf-8")
    (tmp_path / "Inputs/later.txt").write_text("LATE FILE", encoding="utf-8")
    knowledge_file.write_text("CHANGED KNOWLEDGE", encoding="utf-8")
    changed_template = Document()
    changed_template.add_paragraph("CHANGED TEMPLATE")
    changed_template.save(template_file)
    state = {"run_id": run_id, "request": request}

    await service._build_manifest(state)
    snapshot = RunInputSnapshotStore(tmp_path).load(run_id)
    frozen_knowledge = snapshot.scope_root(tmp_path, "Knowledge")
    library = ReferenceLibrary(
        tmp_path,
        knowledge_root=frozen_knowledge,
        index_root=tmp_path / f"Work/runs/{run_id}/indexes/knowledge",
    )
    selected_template, source = service.resolve_report_template(run_id)

    assert [item.path.as_posix() for item in state["project_manifest"].files] == [
        "Inputs/site.txt"
    ]
    manifest_file = state["project_manifest"].files[0]
    assert manifest_file.sha256 == hashlib.sha256(b"ORIGINAL INPUT").hexdigest()
    assert manifest_file.snapshot_ref is not None
    assert (tmp_path / manifest_file.snapshot_ref).read_text(encoding="utf-8") == (
        "ORIGINAL INPUT"
    )
    assert library.open("Knowledge/guide.md").text == "ORIGINAL KNOWLEDGE"
    assert source == "project"
    assert Document(selected_template).paragraphs[0].text == "ORIGINAL TEMPLATE"


def test_aggregate_copies_referenced_legacy_assets_without_hash_gate(
    tmp_path: Path,
) -> None:
    for module_id in ("2.1", "2.2", "2.3", "2.4", "2.5"):
        module = tmp_path / f"Outputs/Modules/{module_id}.md"
        module.parent.mkdir(parents=True, exist_ok=True)
        module.write_text(
            f"## {module_id}\n\n引用 E-0001 和缺失的旧来源 W-0999。\n",
            encoding="utf-8",
        )
    evidence = {
        "id": "E-0001",
        "subject": "旧运行证据",
        "fact": "可继续用于汇总",
        "source": {"file_id": "file-a", "path": "Inputs/a.xlsx"},
        "confidence": 1.0,
        "needs_confirmation": False,
        "photo_refs": ["P-0001"],
    }
    work = tmp_path / "Work"
    work.mkdir(exist_ok=True)
    (work / "evidence.jsonl").write_text(
        json.dumps(evidence, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    source_run = work / "runs/report-source"
    (source_run / "ledgers").mkdir(parents=True)
    (source_run / "sources").mkdir(parents=True)
    (source_run / "ledgers/sources.json").write_text(
        json.dumps(
            [
                {
                    "id": "E-0001",
                    "kind": "project_evidence",
                    "title": "旧运行证据",
                    "locator": "Inputs/a.xlsx#A1",
                    "content_sha256": "deliberately-not-checked",
                }
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (source_run / "sources/E-0001.txt").write_text(
        "旧运行证据正文",
        encoding="utf-8",
    )
    old_photo = source_run / "assets/P-0001.png"
    old_photo.parent.mkdir(parents=True)
    old_photo.write_bytes(b"legacy-photo")
    (work / "photo-manifest.json").write_text(
        json.dumps(
            {
                "assets": [
                    {
                        "id": "P-0001",
                        "path": old_photo.relative_to(tmp_path).as_posix(),
                        "sha256": "also-not-checked",
                        "media_type": "image/png",
                        "source_member": "xl/media/image1.png",
                        "primary_evidence_id": "E-0001",
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    service = ReportingService(
        tmp_path,
        bus=MessageBus(),
        task_board=TaskBoard(),
        llm_provider=NoCallProvider(),
    )

    run_id = service.prepare_run(
        ReportRequest(operation="aggregate_existing", instruction="aggregate")
    )

    imported_source = tmp_path / f"Work/runs/{run_id}/sources/E-0001.txt"
    imported_photo = tmp_path / f"Work/runs/{run_id}/assets/imported/P-0001.png"
    assert imported_source.read_text(encoding="utf-8") == "旧运行证据正文"
    assert imported_photo.read_bytes() == b"legacy-photo"
    assert not imported_source.is_symlink()
    assert not imported_photo.is_symlink()
    manifest = json.loads(
        (tmp_path / f"Work/runs/{run_id}/context/photo-manifest.json").read_text(
            encoding="utf-8"
        )
    )
    assert manifest["assets"][0]["path"] == (
        f"Work/runs/{run_id}/assets/imported/P-0001.png"
    )
    provenance = json.loads(
        (tmp_path / f"Work/runs/{run_id}/context/imported-provenance.json").read_text(
            encoding="utf-8"
        )
    )
    assert provenance["gaps"] == [
        {"id": "W-0999", "kind": "source", "reason": "not_found"}
    ]
