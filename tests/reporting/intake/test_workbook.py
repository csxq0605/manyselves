from pathlib import Path

from openpyxl import Workbook

from manyselves.core.reporting.intake.workbook import inspect_workbook


def test_workbook_artifact_preserves_sheet_dimensions_and_formulas(tmp_path: Path) -> None:
    path = tmp_path / "S4-4诊断工作用表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "低配评估详情"
    sheet["A1"] = "设备"
    sheet["M5"] = '=DISPIMG("ID_ABC",1)'
    workbook.save(path)

    artifact = inspect_workbook(path, file_id="file-test")

    assert artifact.file_id == "file-test"
    assert artifact.sheets["低配评估详情"].max_row == 5
    assert artifact.sheets["低配评估详情"].max_column == 13
    assert artifact.sheets["低配评估详情"].formulas["M5"] == '=DISPIMG("ID_ABC",1)'
