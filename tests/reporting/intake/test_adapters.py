import hashlib
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

from autoreport.core.reporting.intake.adapters import IntakeAdapterRegistry
from autoreport.core.reporting.models import ManifestFile


def _manifest(path: Path) -> ManifestFile:
    return ManifestFile(
        id=f"file-{path.stem}",
        path=Path("Inputs") / path.name,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        media_type="application/octet-stream",
    )


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_text_adapter_preserves_content(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"notes{suffix}"
    path.write_text("接地连续性需要复核", encoding="utf-8")

    artifacts = IntakeAdapterRegistry().parse(path, _manifest(path))

    assert artifacts[0].kind == "text"
    assert artifacts[0].payload["text"] == "接地连续性需要复核"


def test_docx_adapter_preserves_paragraph_locations(tmp_path: Path) -> None:
    path = tmp_path / "inspection.docx"
    document = Document()
    document.add_paragraph("第一项检查")
    document.add_paragraph("第二项检查")
    document.save(path)

    artifacts = IntakeAdapterRegistry().parse(path, _manifest(path))

    assert [artifact.payload["text"] for artifact in artifacts] == [
        "第一项检查",
        "第二项检查",
    ]
    assert [artifact.source.row for artifact in artifacts] == [1, 2]


def test_image_adapter_records_metadata(tmp_path: Path) -> None:
    path = tmp_path / "site.png"
    Image.new("RGB", (32, 24), color="blue").save(path)

    artifact = IntakeAdapterRegistry().parse(path, _manifest(path))[0]

    assert artifact.kind == "image_metadata"
    assert artifact.payload == {"width": 32, "height": 24, "format": "PNG", "mode": "RGB"}


def test_workbook_adapter_returns_located_rows(tmp_path: Path) -> None:
    path = tmp_path / "records.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "检查"
    sheet.append(["对象", "状态"])
    sheet.append(["1A2", "NG"])
    workbook.save(path)

    artifacts = IntakeAdapterRegistry().parse(path, _manifest(path))

    assert artifacts[0].kind == "workbook_row"
    assert artifacts[0].source.sheet == "检查"
    assert artifacts[0].source.cell == "A2"
    assert artifacts[0].payload == {"headers": ["对象", "状态"], "values": ["1A2", "NG"]}


@pytest.mark.parametrize("suffix", [".dwg", ".mp4", ".mov", ".avi"])
def test_unsupported_binary_format_is_explicit_manual_required(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"asset{suffix}"
    path.write_bytes(b"binary")

    artifact = IntakeAdapterRegistry().parse(path, _manifest(path))[0]

    assert artifact.kind == "manual_required"
    assert artifact.payload["reason"] == "unsupported_binary_format"
