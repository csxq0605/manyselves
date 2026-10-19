from pathlib import Path

from docx import Document
from openpyxl import Workbook

from manyselves.runtime.artifacts import parse_artifact


def test_docx_tables_and_xlsx_cells_are_explicit_blocks(tmp_path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    docx = tmp_path / "sample.docx"
    doc.save(docx)
    assert any("A\tB" in block.text for block in parse_artifact(docx).blocks)

    workbook = Workbook()
    workbook.active["A1"] = "value"
    xlsx = tmp_path / "sample.xlsx"
    workbook.save(xlsx)
    assert parse_artifact(xlsx).blocks[0].text == "value"


def test_image_metadata_does_not_claim_visual_verification(tmp_path: Path) -> None:
    from PIL import Image

    image = tmp_path / "sample.png"
    Image.new("RGB", (10, 20)).save(image)
    parsed = parse_artifact(image)
    assert parsed.kind == "image_metadata"
    assert parsed.visual_verified is False


def test_document_blocks_keep_tables_between_their_governing_paragraphs(tmp_path):
    doc = Document()
    doc.add_paragraph("条件甲，仅适用于下表")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "甲阈值"
    doc.add_paragraph("条件乙，仅适用于下表")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "乙阈值"
    path = tmp_path / "ordered.docx"
    doc.save(path)
    blocks = parse_artifact(path).blocks
    assert [b.text for b in blocks] == ["条件甲，仅适用于下表", "甲阈值", "条件乙，仅适用于下表", "乙阈值"]
    assert [b.locator for b in blocks] == ["paragraph:1", "table:1:row:1", "paragraph:2", "table:2:row:1"]
