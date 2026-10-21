"""The generic exemplar mechanism: outline, acceptance, rendering and lessons.

The exemplar here is a small synthetic DOCX about nothing in particular, so
nothing in the mechanism can depend on the report capability.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from manyselves.runtime.exemplar.extraction import (
    ExemplarPolicy,
    ExemplarRole,
    ExemplarSubmission,
    ExemplarUse,
    Target,
    accept_submission,
    instance_terms,
    render_role_skill,
    role_packet,
    role_scope,
)
from manyselves.runtime.exemplar.lessons import (
    LessonBook,
    lesson_problems,
    lessons_for,
    next_lesson_ids,
    render_lessons,
    specific_terms,
)
from manyselves.runtime.exemplar.outline import document_outline

POLICY = ExemplarPolicy(mandatory_words=["必须", "一律"], organisation_suffixes=["有限公司"])


def _exemplar(path: Path) -> Path:
    document = Document()
    document.add_paragraph("受示例花园有限公司委托编写")
    document.add_heading("概况", level=1)
    document.add_heading("背景", level=2)
    document.add_paragraph("本次对园区东侧12株乔木进行了巡查，由示例花园有限公司陪同。")
    document.add_heading("巡查结果", level=1)
    document.add_heading("", level=2)  # an empty heading is skipped, as in the template
    document.add_heading("树木健康", level=2)
    document.add_paragraph("现状描述：")
    document.add_paragraph("抽查发现3株树干开裂，详见下表：")
    table = document.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, (("位置", "问题", "建议"), ("东门", "开裂", "支撑")), strict=False):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    merged = table.cell(2, 0).merge(table.cell(2, 2))
    merged.text = "合并说明单元格"
    document.add_paragraph("建议：")
    document.add_paragraph("尽快加固开裂树干", style="List Bullet")
    document.add_heading("土壤", level=2)
    document.add_paragraph("结论：土壤湿度适中。")
    document.save(str(path))
    return path


@pytest.fixture()
def outline(tmp_path: Path):
    return document_outline(_exemplar(tmp_path / "exemplar.docx"))


def test_outline_numbers_headings_and_locates_blocks(outline) -> None:
    numbers = [(s.number, s.title) for s in outline.sections]
    assert numbers == [("1", "概况"), ("1.1", "背景"), ("2", "巡查结果"), ("2.1", "树木健康"), ("2.2", "土壤")]
    health = outline.section("2.1")
    assert health.labels == ("现状描述", "建议")
    assert health.tables == 1
    table = next(outline.block(loc) for loc in health.blocks if loc.startswith("t-"))
    assert table.columns == ("位置", "问题", "建议")
    assert table.rows == 3
    # A merged cell appears once per row and is counted once.
    assert table.text.splitlines()[-1] == "合并说明单元格"
    assert outline.section("2.2").labels == ("结论",)
    # List items are not labels, and the paragraph before the first heading is front matter.
    assert any(outline.block(loc).list_item for loc in health.blocks)
    assert outline.front_matter == ("受示例花园有限公司委托编写",)


def test_outline_is_deterministic(tmp_path: Path) -> None:
    path = _exemplar(tmp_path / "exemplar.docx")
    assert document_outline(path) == document_outline(path)


def _role() -> ExemplarRole:
    return ExemplarRole.model_validate(
        {"id": "writer", "title": "巡查结果的写法", "purpose": "写巡查结果时参考。",
         "sections": [{"number": "2", "title": "巡查结果"}]}
    )


def _locator(outline, section: str, text: str) -> str:
    return next(
        block.locator for block in outline.blocks if block.section == section and text in block.text
    )


def test_role_scope_and_packet_cover_the_subtree(outline) -> None:
    scope = role_scope(outline, _role())
    assert [s.number for s in scope.sections] == ["2", "2.1", "2.2"]
    packet = role_packet(outline, scope)
    health = next(s for s in packet["sections"] if s["number"] == "2.1")
    assert "分段小标题：现状描述 → 建议" in health["structure"]
    assert "表（3 行，列：位置、问题、建议）" in health["structure"]
    assert all(block["locator"] for block in health["blocks"])
    # A declared title that matches nothing is reported, not guessed.
    missing = ExemplarRole.model_validate(
        {"id": "x", "title": "x", "purpose": "x", "sections": [{"number": "9", "title": "不存在的小节"}]}
    )
    assert role_scope(outline, missing).missing == ("9",)


def test_acceptance_keeps_methods_and_refuses_facts_and_hard_rules(outline) -> None:
    scope = role_scope(outline, _role())
    lead = _locator(outline, "2.1", "抽查发现")
    advice = _locator(outline, "2.1", "尽快加固")
    soil = _locator(outline, "2.2", "土壤湿度")
    submission = ExemplarSubmission.model_validate(
        {
            "entries": [
                {"use": "method", "text": "先用一句话交代抽查范围和发现，再用表格列出位置、问题和建议。",
                 "applies_when": "发现多处同类问题时", "locator": lead, "quote": "抽查发现3株树干开裂"},
                {"use": "example", "text": "抽查发现〈数量〉处〈问题〉，详见下表。",
                 "applies_when": "未明确", "locator": lead, "quote": "树干开裂，详见下表"},
                {"use": "method", "text": "列出3株问题树。", "applies_when": "未明确",
                 "locator": lead, "quote": "抽查发现3株树干开裂"},
                {"use": "method", "text": "建议必须写成加固措施。", "applies_when": "未明确",
                 "locator": advice, "quote": "尽快加固开裂树干"},
                {"use": "example", "text": "抽查发现3株树干开裂。", "applies_when": "未明确",
                 "locator": lead, "quote": "抽查发现3株树干开裂"},
                {"use": "method", "text": "结论一句话说明状态。", "applies_when": "未明确",
                 "locator": soil, "quote": "土壤湿度偏高"},
                {"use": "method", "text": "交代陪同单位示例花园有限公司。", "applies_when": "未明确",
                 "locator": _locator(outline, "1.1", "陪同"), "quote": "示例花园有限公司陪同"},
            ],
            "skipped": [{"section": "2", "reason": "只有标题"}],
        }
    )
    result = accept_submission(outline, scope, submission, POLICY)
    assert [e.text for e in result.accepted] == [
        "先用一句话交代抽查范围和发现，再用表格列出位置、问题和建议。",
        "抽查发现〈数量〉处〈问题〉，详见下表。",
    ]
    reasons = {r.index: " ".join(r.reasons) for r in result.rejected}
    assert "数值或编号：3" in reasons[2]
    assert "强制词：必须" in reasons[3]
    assert "占位：3" in reasons[4]
    assert "找不到" in reasons[5]
    assert "不在本角色的小节里" in reasons[6]
    # 2.2 had only a refused entry, so it is uncovered; 2 is skipped with a reason.
    assert result.uncovered == ["2.2"]
    assert not result.passed


def test_uses_unavailable_features_and_generated_tables(outline) -> None:
    policy = ExemplarPolicy(
        mandatory_words=["必须"],
        uses=[
            ExemplarUse(id="judge", title="判断方法", description="怎样下判断"),
            ExemplarUse(id="write", title="写法", description="怎样组织"),
            ExemplarUse(id="example", title="范例", description="句式", placeholders=True),
        ],
        unavailable={"附件": "成品不带附件"},
    )
    role = ExemplarRole.model_validate(
        {"id": "writer", "title": "巡查结果", "purpose": "写巡查结果时参考。",
         "sections": [{"number": "2", "title": "巡查结果", "generated": "表由系统生成"}],
         "uses": ["judge", "write"], "required_uses": ["judge"]}
    )
    scope = role_scope(outline, role)
    packet = role_packet(outline, scope, policy)
    assert [use["id"] for use in packet["uses"]] == ["judge", "write"]
    assert packet["uses"][0]["required"] is True
    health = next(s for s in packet["sections"] if s["number"] == "2.1")
    assert "表（表由系统生成）" in health["structure"] and "位置" not in health["structure"]
    lead = _locator(outline, "2.1", "抽查发现")
    soil = _locator(outline, "2.2", "土壤湿度")
    submission = ExemplarSubmission.model_validate(
        {"entries": [
            {"use": "write", "text": "先交代抽查范围，再列出问题。", "applies_when": "未明确",
             "locator": lead, "quote": "抽查发现3株树干开裂"},
            {"use": "example", "text": "抽查发现〈数量〉处〈问题〉。", "applies_when": "未明确",
             "locator": lead, "quote": "抽查发现3株树干开裂"},
            {"use": "write", "text": "细节放进附件。", "applies_when": "未明确",
             "locator": soil, "quote": "土壤湿度适中"},
        ]}
    )
    result = accept_submission(outline, scope, submission, policy)
    reasons = {r.index: " ".join(r.reasons) for r in result.rejected}
    assert "不是本角色可用的用途" in reasons[1]
    assert "附件" in reasons[2]
    assert result.missing_uses == ["judge"] and not result.passed
    text = render_role_skill(outline, scope, result.accepted, source="示例", policy=policy)
    assert "写法：\n- 先交代抽查范围" in text and "判断方法：" not in text


def test_organisation_names_are_refused_even_in_scope(tmp_path: Path) -> None:
    outline = document_outline(_exemplar(tmp_path / "e.docx"))
    role = ExemplarRole.model_validate(
        {"id": "bg", "title": "背景", "purpose": "背景", "sections": [{"number": "1"}]}
    )
    scope = role_scope(outline, role)
    locator = _locator(outline, "1.1", "陪同")
    submission = ExemplarSubmission.model_validate(
        {"entries": [{"use": "example", "text": "由示例花园有限公司陪同完成巡查。",
                      "applies_when": "未明确", "locator": locator, "quote": "示例花园有限公司陪同"}]}
    )
    result = accept_submission(outline, scope, submission, POLICY)
    assert "单位名称" in result.rejected[0].reasons[0]


def test_rendered_skill_keeps_locators_and_drops_quotes(outline) -> None:
    scope = role_scope(outline, _role())
    lead = _locator(outline, "2.1", "抽查发现")
    entries = ExemplarSubmission.model_validate(
        {"entries": [{"use": "method", "text": "先交代范围，再列表。", "applies_when": "未明确",
                      "locator": lead, "quote": "抽查发现3株树干开裂"}]}
    ).entries
    text = render_role_skill(outline, scope, entries, source="示例", policy=POLICY)
    assert "表（列：位置、问题、建议）" in text and "3 行" not in text
    assert text.startswith("---\nname: writer\n")
    assert "## 2.1 树木健康" in text
    assert f"方法：\n- 先交代范围，再列表。 〔模板 {lead}〕" in text
    assert "适用：未明确" not in text
    assert "3株" not in text and "抽查发现" not in text


def test_lessons_refuse_run_objects_numbers_and_hard_words() -> None:
    terms = specific_terms(["东门老槐树", "乔木", "园区"], vocabulary=["乔木养护", "园区巡查"])
    assert terms == {"东门老槐树"}
    problems = lesson_problems(
        roles=["writer", "ghost"],
        error_type="东门老槐树写成了正常",
        rule="必须把3处开裂写进表格",
        trigger={"kind": "sometimes"},
        known_roles=["writer"],
        run_terms=terms,
        mandatory_words=["必须"],
        trigger_problem=lambda trigger: None if trigger.get("kind") == "always" else "触发条件不对",
    )
    joined = " ".join(problems)
    for expected in ("角色不存在：ghost", "数值或编号：3", "本次运行的对象：东门老槐树", "强制词：必须", "触发条件不对"):
        assert expected in joined
    assert lesson_problems(
        roles=["writer"], error_type="把未见记录写成正常", rule="缺资料时写明缺什么、影响什么判断。",
        trigger={"kind": "always"}, known_roles=["writer"], run_terms=terms,
        mandatory_words=["必须"], trigger_problem=lambda trigger: None,
    ) == []


def test_lessons_merge_only_active_and_firing() -> None:
    book = LessonBook.model_validate(
        {"lessons": [
            {"id": "L-001", "roles": ["writer"], "error_type": "甲", "rule": "乙",
             "source": {"run": "r", "finding": "F-001", "date": "2026-10-21"}},
            {"id": "L-002", "roles": ["writer"], "error_type": "丙", "rule": "丁",
             "trigger": {"kind": "never"},
             "source": {"run": "r", "finding": "F-002", "date": "2026-10-21"}},
            {"id": "L-003", "roles": ["writer"], "error_type": "戊", "rule": "己", "status": "revoked",
             "revoked_reason": "不再适用", "source": {"run": "r", "finding": "F-003", "date": "2026-10-21"}},
        ]}
    )
    chosen = lessons_for(book, "writer", lambda trigger: trigger.get("kind") == "always")
    assert [lesson.id for lesson in chosen] == ["L-001"]
    assert "易错：甲。做法：乙" in render_lessons(chosen)
    assert next_lesson_ids(book, 2) == ["L-004", "L-005"]


def test_comparison_sorts_old_items_by_their_closest_match() -> None:
    from manyselves.runtime.exemplar.compare import compare_items, render_comparison, skill_items

    old = skill_items(
        "---\nname: x\n---\n# 写法\n- 先交代抽查范围，再列出发现的问题。\n"
        "- 缺资料时只写一句资料未提供。\n- 树木的养护记录按季度归档备查。\n| 表 | 行 |\n"
    )
    assert [item for _heading, item in old][0] == "先交代抽查范围，再列出发现的问题。"
    new = [("writer", "先交代抽查范围，再列出发现的问题和建议。")]
    rules = [("合同", "缺资料时只写一句“资料未提供”，不展开说明。")]
    matches = compare_items(old, new, rules)
    text, counts = render_comparison({"writer": matches}, threshold=0.4)
    assert counts["writer"] == {"新版有对应": 1, "合同已有": 1, "未对应": 1}
    assert text.index("未对应") < text.index("新版有对应 |")


def _sites_exemplar(path: Path) -> Path:
    """Two parks visited; each section opens with the same labels, one with its own finding."""

    document = Document()
    document.add_heading("巡查结果", level=1)
    document.add_heading("灌溉情况", level=2)  # our product calls it 浇水
    document.add_paragraph("现状描述：东湖园区与西山园区的喷头均已检查。")
    document.add_paragraph("喷头堵塞较多：西山园区北坡最明显。")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("位置", "问题"), ("北坡花坛", "堵塞"))):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("建议：分区清洗喷头。")
    document.add_heading("东湖园区", level=3)  # one place of the exemplar, no counterpart
    document.add_paragraph("现状描述：东湖园区喷头完好。")
    document.add_heading("修剪", level=2)
    document.add_paragraph("现状描述：枝条过密。")
    document.add_paragraph("建议：分批修剪。")
    document.save(str(path))
    return path


SITE_POLICY = ExemplarPolicy(site_suffixes=["园区"], instance_columns=["位置"])
TARGETS = [Target("1", "巡查结果"), Target("1.1", "修剪情况"), Target("1.2", "浇水与灌溉情况")]


def test_sections_are_aligned_to_our_titles_and_unknown_ones_fold(tmp_path: Path) -> None:
    outline = document_outline(_sites_exemplar(tmp_path / "e.docx"))
    role = ExemplarRole.model_validate(
        {"id": "w", "title": "巡查", "purpose": "写巡查时参考。", "sections": [{"number": "1", "title": "巡查结果"}]}
    )
    scope = role_scope(outline, role, TARGETS)
    # Matched by title, not by the exemplar's position.
    assert scope.target(outline.section("1.1")).heading == "1.2 浇水与灌溉情况"
    assert scope.target(outline.section("1.2")).heading == "1.1 修剪情况"
    # The exemplar's own place has no counterpart: it folds into its parent.
    assert scope.target(outline.section("1.1.1")).heading == "1.2 浇水与灌溉情况"
    assert scope.folded == ("1.1.1",)
    packet = role_packet(outline, scope, SITE_POLICY)
    assert packet["sections"][1]["product_section"] == "1.2 浇水与灌溉情况"

    entries = ExemplarSubmission.model_validate(
        {"entries": [
            {"use": "method", "text": "先写现状，再给建议。", "applies_when": "未明确",
             "locator": _locator(outline, "1.1.1", "喷头完好"), "quote": "东湖园区喷头完好"},
            {"use": "method", "text": "分批写出修剪建议。", "applies_when": "未明确",
             "locator": _locator(outline, "1.2", "分批修剪"), "quote": "建议：分批修剪"},
        ]}
    ).entries
    text = render_role_skill(outline, scope, entries, source="示例", policy=SITE_POLICY)
    # One heading per product section; the exemplar's numbering is gone.
    assert text.count("## 1.2 浇水与灌溉情况") == 1 and "## 1.1.1" not in text
    assert text.index("## 1.2 浇水") < text.index("## 1.1 修剪情况")  # exemplar order kept
    # Structure shows the labels reused across sections, never one section's own subheading.
    assert "分段标签：现状描述 → 建议" in text
    assert "喷头堵塞较多" not in text and "另有 1 个分项小标题" in text


def test_the_exemplars_own_places_and_objects_are_refused(tmp_path: Path) -> None:
    outline = document_outline(_sites_exemplar(tmp_path / "e.docx"))
    role = ExemplarRole.model_validate(
        {"id": "w", "title": "巡查", "purpose": "写巡查时参考。", "sections": [{"number": "1", "title": "巡查结果"}]}
    )
    scope = role_scope(outline, role, TARGETS)
    terms = instance_terms(outline, scope, SITE_POLICY)
    assert {"东湖园区", "西山园区", "北坡花坛"} <= set(terms)
    lead = _locator(outline, "1.1", "喷头均已检查")
    submission = ExemplarSubmission.model_validate(
        {"entries": [
            {"use": "method", "text": "先写西山园区，再写其他地方。", "applies_when": "未明确",
             "locator": lead, "quote": "喷头均已检查"},
            {"use": "example", "text": "〈位置〉北坡花坛喷头〈问题〉。", "applies_when": "未明确",
             "locator": lead, "quote": "喷头均已检查"},
            {"use": "method", "text": "逐个园区说明喷头状态。", "applies_when": "未明确",
             "locator": lead, "quote": "喷头均已检查"},
            {"use": "method", "text": "按区域说明喷头状态，再给建议。", "applies_when": "未明确",
             "locator": lead, "quote": "喷头均已检查"},
            {"use": "example", "text": "〈园区〉喷头〈问题〉。", "applies_when": "未明确",
             "locator": lead, "quote": "喷头均已检查"},
        ]}
    )
    result = accept_submission(outline, scope, submission, SITE_POLICY)
    reasons = {r.index: " ".join(r.reasons) for r in result.rejected}
    assert "西山园区（专家版项目的地点）" in reasons[0]
    assert "北坡花坛（专家版表格里的位置或对象）" in reasons[1]
    assert "场地划分" in reasons[2]
    # Our own word and a placeholder are fine.
    assert [e.text for e in result.accepted] == ["按区域说明喷头状态，再给建议。", "〈园区〉喷头〈问题〉。"]
