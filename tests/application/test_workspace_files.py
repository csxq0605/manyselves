"""Security and concurrency contracts for project workspace file access."""

import asyncio
import base64
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from manyselves.application.preview_service import PreviewEncodingError, PreviewService
from manyselves.application.workspace_files import (
    DestinationExists,
    FileRevisionConflict,
    UnsafeWorkspacePath,
    UploadTooLarge,
    WorkspaceFiles,
)


@pytest.fixture
def workspace_files(tmp_path: Path) -> WorkspaceFiles:
    project = tmp_path / "p1"
    (project / "Inputs").mkdir(parents=True)
    (project / "Inputs" / "a.txt").write_text("initial", encoding="utf-8")
    return WorkspaceFiles(project, max_text_bytes=1024, max_upload_bytes=8)


@pytest.mark.parametrize(
    "unsafe",
    [
        "../outside.txt",
        "/tmp/outside.txt",
        r"C:\outside.txt",
        "%2e%2e%2foutside.txt",
        "Inputs%2fa.txt",
        r"Inputs%5ca.txt",
    ],
)
def test_resolve_rejects_traversal_and_absolute_paths(
    workspace_files: WorkspaceFiles, unsafe: str
) -> None:
    """Normalizing attacker-controlled paths must never escape or reinterpret separators."""
    with pytest.raises(UnsafeWorkspacePath):
        workspace_files.resolve(unsafe)


def test_resolve_rejects_existing_parent_symlink_escape(
    workspace_files: WorkspaceFiles, tmp_path: Path
) -> None:
    """A symlink parent must not redirect a validated path outside the workspace."""
    outside = tmp_path / "outside"
    outside.mkdir()
    link = workspace_files.project_root / "Inputs" / "linked"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"platform denied symlink creation: {error}")

    with pytest.raises(UnsafeWorkspacePath):
        workspace_files.resolve("Inputs/linked/secret.txt")


def test_write_rejects_stale_revision(workspace_files: WorkspaceFiles) -> None:
    """A client must not overwrite content that changed after its read."""
    current = workspace_files.read_text("Inputs/a.txt")
    workspace_files.write_text("Inputs/a.txt", "server", current.revision)

    with pytest.raises(FileRevisionConflict):
        workspace_files.write_text("Inputs/a.txt", "client", current.revision)

    assert workspace_files.read_text("Inputs/a.txt").content == "server"


def test_write_is_atomic_and_cleans_failed_sibling_temp(
    workspace_files: WorkspaceFiles, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A failed replacement must preserve the prior file and leave no upload artifact."""
    current = workspace_files.read_text("Inputs/a.txt")

    def fail_replace(self: Path, target: Path) -> Path:
        raise OSError("injected replace interruption")

    monkeypatch.setattr(Path, "replace", fail_replace)
    with pytest.raises(OSError, match="replace interruption"):
        workspace_files.write_text("Inputs/a.txt", "replacement", current.revision)

    assert (workspace_files.project_root / "Inputs" / "a.txt").read_text() == "initial"
    assert list((workspace_files.project_root / "Inputs").glob(".manyselves-tmp-*")) == []


def test_create_and_rename_reject_duplicate_destinations(workspace_files: WorkspaceFiles) -> None:
    """Create or rename must never silently replace an existing entry."""
    workspace_files.create_file("Inputs/b.txt", "b")

    with pytest.raises(DestinationExists):
        workspace_files.create_file("Inputs/b.txt", "other")
    with pytest.raises(DestinationExists):
        workspace_files.rename("Inputs/a.txt", "Inputs/b.txt")

    assert workspace_files.read_text("Inputs/b.txt").content == "b"


@pytest.mark.asyncio
async def test_interrupted_upload_cleans_temp_and_destination(
    workspace_files: WorkspaceFiles,
) -> None:
    """Cancellation must not publish a partial upload or strand a sibling temp file."""

    async def interrupted():
        yield b"abc"
        raise asyncio.CancelledError()

    with pytest.raises(asyncio.CancelledError):
        await workspace_files.upload("Inputs/interrupted.bin", interrupted())

    inputs = workspace_files.project_root / "Inputs"
    assert not (inputs / "interrupted.bin").exists()
    assert list(inputs.glob(".manyselves-tmp-*")) == []


@pytest.mark.asyncio
async def test_upload_limit_is_checked_while_streaming(workspace_files: WorkspaceFiles) -> None:
    """An upload crossing the configured limit must stop without publishing the file."""

    async def chunks():
        yield b"1234"
        yield b"56789"
        raise AssertionError("stream should stop after crossing the limit")

    with pytest.raises(UploadTooLarge):
        await workspace_files.upload("Inputs/large.bin", chunks())

    assert not (workspace_files.project_root / "Inputs" / "large.bin").exists()


def test_preview_dtos_are_bounded_passive_and_path_sanitized(
    workspace_files: WorkspaceFiles,
) -> None:
    """Preview responses must not expose server paths or active SVG/Markdown content."""
    svg = '<svg xmlns="http://www.w3.org/2000/svg" onload="evil()"><script>x</script><a href="javascript:evil()">x</a></svg>'
    workspace_files.create_file("Inputs/unsafe.svg", svg)
    workspace_files.create_file("Inputs/unsafe.md", "# ok\n<script>evil()</script>")
    service = PreviewService(workspace_files, max_preview_bytes=1024, row_limit=2, column_limit=2)

    svg_dto = service.preview(
        "Inputs/unsafe.svg", content_url="/api/v1/projects/p1/files/download?path=Inputs%2Funsafe.svg"
    )
    markdown_dto = service.preview(
        "Inputs/unsafe.md", content_url="/api/v1/projects/p1/files/download?path=Inputs%2Funsafe.md"
    )
    serialized = f"{svg_dto!r}{markdown_dto!r}"

    assert svg_dto["type"] == "image"
    assert "script" not in svg_dto["content"].casefold()
    assert "onload" not in svg_dto["content"].casefold()
    assert "javascript:" not in svg_dto["content"].casefold()
    assert markdown_dto["type"] == "markdown"
    assert "<script" not in markdown_dto["content"].casefold()
    assert str(workspace_files.project_root) not in serialized


def test_preview_reports_invalid_text_encoding_explicitly(workspace_files: WorkspaceFiles) -> None:
    """Binary bytes must not be silently replaced and presented as valid text."""
    (workspace_files.project_root / "Inputs" / "bad.md").write_bytes(b"\xff\xfe")

    with pytest.raises(PreviewEncodingError):
        PreviewService(workspace_files).preview(
            "Inputs/bad.md", content_url="/api/v1/projects/p1/files/download?path=Inputs%2Fbad.md"
        )


def test_image_preview_reports_dimensions_without_embedding_server_path(
    workspace_files: WorkspaceFiles,
) -> None:
    """Image preview metadata must be passive and client-addressable."""
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    (workspace_files.project_root / "Inputs" / "pixel.png").write_bytes(png)

    dto = PreviewService(workspace_files).preview(
        "Inputs/pixel.png", content_url="/content/pixel.png"
    )

    assert dto == {
        "type": "image",
        "path": "Inputs/pixel.png",
        "mimeType": "image/png",
        "width": 1,
        "height": 1,
        "contentUrl": "/content/pixel.png",
    }


def test_spreadsheet_preview_limits_rows_columns_and_does_not_evaluate_formulas(
    workspace_files: WorkspaceFiles,
) -> None:
    """Spreadsheet parsing must expose a bounded value window without executing formulas."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["a", "b", "c"])
    sheet.append([1, 2, "=1+1"])
    sheet.append([3, 4, 5])
    workbook.save(workspace_files.project_root / "Inputs" / "data.xlsx")
    workbook.close()

    dto = PreviewService(
        workspace_files,
        max_preview_bytes=64 * 1024,
        row_limit=2,
        column_limit=2,
    ).preview("Inputs/data.xlsx", content_url="/content/data.xlsx")

    assert dto["type"] == "spreadsheet"
    assert dto["sheets"] == [
        {
            "name": "Data",
            "rowCount": 3,
            "columnCount": 3,
            "rows": [["a", "b"], ["1", "2"]],
            "truncated": True,
        }
    ]


def test_docx_preview_preserves_passive_paragraph_table_image_order(
    workspace_files: WorkspaceFiles,
) -> None:
    """DOCX preview must extract passive blocks without exposing or executing package content."""
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = Document()
    document.add_paragraph("intro")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.add_picture(io.BytesIO(png))
    document.save(workspace_files.project_root / "Inputs" / "report.docx")

    dto = PreviewService(workspace_files, max_preview_bytes=128 * 1024).preview(
        "Inputs/report.docx", content_url="/content/report.docx"
    )

    assert dto["type"] == "docx"
    assert [block["type"] for block in dto["blocks"]] == ["paragraph", "table", "image"]
    assert dto["blocks"][0]["text"] == "intro"
    assert dto["blocks"][1]["rows"] == [["left", "right"]]
    assert dto["blocks"][2]["dataUrl"].startswith("data:image/png;base64,")
    assert str(workspace_files.project_root) not in repr(dto)
