"""Bounded passive preview normalization for supported workspace documents."""

import base64
import csv
import html
import io
import mimetypes
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from PIL import Image

from .workspace_files import WorkspaceFileError, WorkspaceFiles


class PreviewError(WorkspaceFileError):
    """Base class for stable preview failures."""


class PreviewTooLarge(PreviewError):  # noqa: N818 - stable application error
    code = "PREVIEW_TOO_LARGE"

    def __init__(self) -> None:
        super().__init__("File exceeds the configured preview limit")


class PreviewEncodingError(PreviewError):
    code = "PREVIEW_ENCODING_ERROR"

    def __init__(self) -> None:
        super().__init__("Preview text is not valid UTF-8")


class UnsupportedPreview(PreviewError):  # noqa: N818 - stable application error
    code = "PREVIEW_UNSUPPORTED"

    def __init__(self) -> None:
        super().__init__("Preview format is unsupported")


class InvalidPreviewDocument(PreviewError):  # noqa: N818 - stable application error
    code = "PREVIEW_INVALID_DOCUMENT"

    def __init__(self) -> None:
        super().__init__("Preview document is invalid")


class PreviewService:
    """Return bounded JSON-compatible DTOs without evaluating active content."""

    def __init__(
        self,
        files: WorkspaceFiles,
        *,
        max_preview_bytes: int = 8 * 1024 * 1024,
        max_archive_expanded_bytes: int = 32 * 1024 * 1024,
        row_limit: int = 200,
        column_limit: int = 100,
        block_limit: int = 500,
    ) -> None:
        if min(
            max_preview_bytes,
            max_archive_expanded_bytes,
            row_limit,
            column_limit,
            block_limit,
        ) <= 0:
            raise ValueError("Preview limits must be positive")
        self._files = files
        self._max_preview_bytes = max_preview_bytes
        self._max_archive_expanded_bytes = max_archive_expanded_bytes
        self._row_limit = row_limit
        self._column_limit = column_limit
        self._block_limit = block_limit

    def preview(self, relative_path: str, *, content_url: str) -> dict[str, Any]:
        path = self._files.file_path(relative_path)
        suffix = path.suffix.casefold()
        portable_path = self._files.relative(path)
        if suffix == ".pdf":
            return self._pdf(path, portable_path, content_url)
        if suffix == ".svg":
            return self._svg(path, portable_path, content_url)
        if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}:
            return self._image(path, portable_path, content_url)
        if suffix in {".xlsx", ".csv"}:
            return self._spreadsheet(path, portable_path)
        if suffix == ".docx":
            return self._docx(path, portable_path)
        if suffix == ".md":
            text = self._decode_utf8(self._read_bounded(path))
            return {
                "type": "markdown",
                "path": portable_path,
                "content": html.escape(text, quote=False),
                "truncated": False,
            }
        if suffix in {".py", ".json", ".yaml", ".yml", ".txt", ".log"}:
            return {
                "type": "text",
                "path": portable_path,
                "content": self._decode_utf8(self._read_bounded(path)),
                "truncated": False,
            }
        stat = path.stat()
        return {
            "type": "unsupported",
            "path": portable_path,
            "mimeType": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            "size": stat.st_size,
            "downloadUrl": content_url,
        }

    def _pdf(self, path: Path, portable_path: str, content_url: str) -> dict[str, Any]:
        raw = self._read_bounded(path)
        if not raw.startswith(b"%PDF-"):
            raise InvalidPreviewDocument()
        return {
            "type": "pdf",
            "path": portable_path,
            "mimeType": "application/pdf",
            "size": len(raw),
            "rangeUrl": content_url,
        }

    def _image(self, path: Path, portable_path: str, content_url: str) -> dict[str, Any]:
        raw = self._read_bounded(path)
        try:
            with Image.open(io.BytesIO(raw)) as image:
                width, height = image.size
                image_format = image.format
                image.verify()
        except Exception as error:
            raise InvalidPreviewDocument() from error
        if width <= 0 or height <= 0 or width * height > 40_000_000:
            raise PreviewTooLarge()
        mime_type = Image.MIME.get(image_format or "") or mimetypes.guess_type(path.name)[0]
        return {
            "type": "image",
            "path": portable_path,
            "mimeType": mime_type or "application/octet-stream",
            "width": width,
            "height": height,
            "contentUrl": content_url,
        }

    def _svg(self, path: Path, portable_path: str, content_url: str) -> dict[str, Any]:
        raw = self._read_bounded(path)
        try:
            root = ElementTree.fromstring(raw)
        except (ElementTree.ParseError, ValueError) as error:
            raise InvalidPreviewDocument() from error
        forbidden = {"script", "foreignobject", "iframe", "object", "embed"}
        for parent in list(root.iter()):
            for child in list(parent):
                if _local_name(child.tag).casefold() in forbidden:
                    parent.remove(child)
            for name, value in list(parent.attrib.items()):
                local_name = _local_name(name).casefold()
                if local_name.startswith("on") or (
                    local_name in {"href", "src"}
                    and value.strip().casefold().startswith(("javascript:", "data:text/html"))
                ):
                    del parent.attrib[name]
        content = ElementTree.tostring(root, encoding="unicode")
        return {
            "type": "image",
            "path": portable_path,
            "mimeType": "image/svg+xml",
            "content": content,
            "contentUrl": content_url,
        }

    def _spreadsheet(self, path: Path, portable_path: str) -> dict[str, Any]:
        raw = self._read_bounded(path)
        if path.suffix.casefold() == ".csv":
            text = self._decode_utf8(raw)
            all_rows = list(csv.reader(io.StringIO(text)))
            column_count = max((len(row) for row in all_rows), default=0)
            rows = [row[: self._column_limit] for row in all_rows[: self._row_limit]]
            return {
                "type": "spreadsheet",
                "path": portable_path,
                "sheets": [
                    {
                        "name": path.stem,
                        "rowCount": len(all_rows),
                        "columnCount": column_count,
                        "rows": rows,
                        "truncated": len(all_rows) > self._row_limit
                        or column_count > self._column_limit,
                    }
                ],
            }

        self._validate_archive(raw)
        try:
            workbook = load_workbook(
                io.BytesIO(raw),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as error:
            raise InvalidPreviewDocument() from error
        try:
            sheets = []
            for sheet in workbook.worksheets:
                rows = [
                    [_cell_text(cell.value) for cell in row[: self._column_limit]]
                    for row in sheet.iter_rows(
                        max_row=self._row_limit,
                        max_col=self._column_limit,
                    )
                ]
                sheets.append(
                    {
                        "name": sheet.title,
                        "rowCount": sheet.max_row,
                        "columnCount": sheet.max_column,
                        "rows": rows,
                        "truncated": sheet.max_row > self._row_limit
                        or sheet.max_column > self._column_limit,
                    }
                )
            return {"type": "spreadsheet", "path": portable_path, "sheets": sheets}
        finally:
            workbook.close()

    def _docx(self, path: Path, portable_path: str) -> dict[str, Any]:
        raw = self._read_bounded(path)
        self._validate_archive(raw)
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                document = ElementTree.fromstring(archive.read("word/document.xml"))
                relationships = _docx_relationships(archive)
                blocks: list[dict[str, Any]] = []
                body = next(
                    (node for node in document.iter() if _local_name(node.tag) == "body"), None
                )
                if body is None:
                    raise InvalidPreviewDocument()
                for node in body:
                    if len(blocks) >= self._block_limit:
                        break
                    kind = _local_name(node.tag)
                    if kind == "p":
                        text = "".join(
                            child.text or ""
                            for child in node.iter()
                            if _local_name(child.tag) == "t"
                        )
                        if text:
                            blocks.append({"type": "paragraph", "text": text})
                        for child in node.iter():
                            if len(blocks) >= self._block_limit:
                                break
                            if _local_name(child.tag) != "blip":
                                continue
                            relationship_id = next(
                                (
                                    value
                                    for name, value in child.attrib.items()
                                    if _local_name(name) == "embed"
                                ),
                                None,
                            )
                            target = relationships.get(relationship_id or "")
                            if target is None or not target.startswith("word/media/"):
                                continue
                            image = archive.read(target)
                            mime = mimetypes.guess_type(target)[0] or "application/octet-stream"
                            blocks.append(
                                {
                                    "type": "image",
                                    "mimeType": mime,
                                    "dataUrl": f"data:{mime};base64,{base64.b64encode(image).decode('ascii')}",
                                }
                            )
                    elif kind == "tbl":
                        rows = []
                        for row in (child for child in node if _local_name(child.tag) == "tr"):
                            cells = []
                            for cell in (child for child in row if _local_name(child.tag) == "tc"):
                                cells.append(
                                    "".join(
                                        value.text or ""
                                        for value in cell.iter()
                                        if _local_name(value.tag) == "t"
                                    )
                                )
                            rows.append(cells[: self._column_limit])
                            if len(rows) >= self._row_limit:
                                break
                        blocks.append({"type": "table", "rows": rows})
                return {
                    "type": "docx",
                    "path": portable_path,
                    "blocks": blocks,
                    "truncated": len(body) > len(blocks),
                }
        except (KeyError, OSError, zipfile.BadZipFile, ElementTree.ParseError) as error:
            raise InvalidPreviewDocument() from error

    def _read_bounded(self, path: Path) -> bytes:
        if path.stat().st_size > self._max_preview_bytes:
            raise PreviewTooLarge()
        with path.open("rb") as stream:
            raw = stream.read(self._max_preview_bytes + 1)
        if len(raw) > self._max_preview_bytes:
            raise PreviewTooLarge()
        return raw

    def _validate_archive(self, raw: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                infos = archive.infolist()
                if any(info.flag_bits & 0x1 for info in infos):
                    raise InvalidPreviewDocument()
                if sum(info.file_size for info in infos) > self._max_archive_expanded_bytes:
                    raise PreviewTooLarge()
        except zipfile.BadZipFile as error:
            raise InvalidPreviewDocument() from error

    @staticmethod
    def _decode_utf8(raw: bytes) -> str:
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError as error:
            raise PreviewEncodingError() from error


def _local_name(name: str) -> str:
    return name.rsplit("}", 1)[-1]


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _docx_relationships(archive: zipfile.ZipFile) -> dict[str, str]:
    try:
        root = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
    except KeyError:
        return {}
    relationships: dict[str, str] = {}
    for node in root:
        relationship_id = node.attrib.get("Id")
        target = node.attrib.get("Target", "")
        target_mode = node.attrib.get("TargetMode", "Internal")
        if relationship_id and target_mode == "Internal" and not target.startswith(("/", "..")):
            relationships[relationship_id] = f"word/{target}"
    return relationships
