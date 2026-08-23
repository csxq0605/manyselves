"""Deterministically render the standalone evidence/source index to DOCX."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_FIXED_DOCX_TIME = datetime(2000, 1, 1, tzinfo=UTC)
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET = re.compile(r"^-\s+(.+?)\s*$")


@dataclass(frozen=True, slots=True)
class SourceIndexRenderResult:
    output_ref: Path
    output_sha256: str
    paragraph_count: int


class SourceIndexDocxRenderer:
    """Render a compact-reference companion document from canonical Markdown."""

    BASE_FONT = "Calibri"
    CJK_FONT = "Arial Unicode MS"
    HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
    HEADING_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
    MUTED = RGBColor(0x66, 0x66, 0x66)

    @classmethod
    def render(cls, markdown: str, output: Path) -> SourceIndexRenderResult:
        if not markdown.strip():
            raise ValueError("source-index Markdown cannot be blank")
        document = Document()
        cls._configure_document(document)
        expected_visible = cls._write_content(document, markdown)
        cls._verify_visible_content(document, expected_visible)

        output = Path(output)
        output.parent.mkdir(parents=True, exist_ok=True)
        buffer = io.BytesIO()
        document.save(buffer)
        data = cls._canonical_docx(buffer.getvalue())
        output.write_bytes(data)
        cls._verify_openable(output, expected_visible)
        return SourceIndexRenderResult(
            output_ref=output,
            output_sha256=hashlib.sha256(data).hexdigest(),
            paragraph_count=len(expected_visible),
        )

    @classmethod
    def _configure_document(cls, document: Document) -> None:
        section = document.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        normal = document.styles["Normal"]
        cls._set_style_font(normal, size=11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25

        heading_one = document.styles["Heading 1"]
        cls._set_style_font(
            heading_one,
            size=16,
            color=cls.HEADING_BLUE,
            bold=True,
        )
        heading_one.paragraph_format.space_before = Pt(18)
        heading_one.paragraph_format.space_after = Pt(10)
        heading_one.paragraph_format.keep_with_next = True

        heading_two = document.styles["Heading 2"]
        cls._set_style_font(
            heading_two,
            size=13,
            color=cls.HEADING_BLUE,
            bold=True,
        )
        heading_two.paragraph_format.space_before = Pt(14)
        heading_two.paragraph_format.space_after = Pt(7)
        heading_two.paragraph_format.keep_with_next = True

        bullet = document.styles["List Bullet"]
        cls._set_style_font(bullet, size=11)
        bullet.paragraph_format.left_indent = Inches(0.375)
        bullet.paragraph_format.first_line_indent = Inches(-0.188)
        bullet.paragraph_format.space_before = Pt(0)
        bullet.paragraph_format.space_after = Pt(4)
        bullet.paragraph_format.line_spacing = 1.25

        properties = document.core_properties
        properties.title = "证据与来源索引"
        properties.subject = "配电安全专家咨询报告独立来源附件"
        properties.author = "Manyselves"
        properties.last_modified_by = "Manyselves"
        properties.created = _FIXED_DOCX_TIME
        properties.modified = _FIXED_DOCX_TIME

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        cls._set_run_font(
            header.add_run("配电安全专家咨询报告 · 来源附件"),
            size=9,
            color=cls.MUTED,
        )

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.paragraph_format.space_before = Pt(0)
        cls._set_run_font(footer.add_run("证据与来源索引  ·  "), size=9, color=cls.MUTED)
        page_run = footer.add_run()
        cls._set_run_font(page_run, size=9, color=cls.MUTED)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for element in (begin, instruction, separate, text, end):
            page_run._r.append(element)

    @classmethod
    def _write_content(cls, document: Document, markdown: str) -> list[str]:
        expected_visible: list[str] = []
        title_written = False
        for raw_line in markdown.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            heading = _HEADING.match(line)
            if heading is not None:
                level = len(heading.group(1))
                text = heading.group(2).strip()
                if not title_written:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.keep_with_next = True
                    cls._set_run_font(
                        paragraph.add_run(text),
                        size=22,
                        color=cls.HEADING_DARK_BLUE,
                        bold=True,
                    )
                    subtitle = document.add_paragraph()
                    subtitle.paragraph_format.space_before = Pt(0)
                    subtitle.paragraph_format.space_after = Pt(14)
                    subtitle.paragraph_format.keep_with_next = True
                    cls._set_run_font(
                        subtitle.add_run("独立附件 · 与报告正文脚注编号对应"),
                        size=10,
                        color=cls.MUTED,
                    )
                    title_written = True
                else:
                    document.add_paragraph(text, style="Heading 1" if level <= 3 else "Heading 2")
                expected_visible.append(text)
                continue
            bullet = _BULLET.match(line)
            if bullet is not None:
                text = bullet.group(1).strip()
                document.add_paragraph(text, style="List Bullet")
                expected_visible.append(text)
                continue
            document.add_paragraph(line, style="Normal")
            expected_visible.append(line)
        if not title_written:
            raise ValueError("source-index Markdown requires a heading")
        return expected_visible

    @classmethod
    def _set_style_font(
        cls,
        style,
        *,
        size: float,
        color: RGBColor | None = None,
        bold: bool | None = None,
    ) -> None:
        style.font.name = cls.BASE_FONT
        style.font.size = Pt(size)
        if color is not None:
            style.font.color.rgb = color
        if bold is not None:
            style.font.bold = bold
        run_properties = style.element.get_or_add_rPr()
        fonts = run_properties.get_or_add_rFonts()
        fonts.set(qn("w:ascii"), cls.BASE_FONT)
        fonts.set(qn("w:hAnsi"), cls.BASE_FONT)
        fonts.set(qn("w:eastAsia"), cls.CJK_FONT)

    @classmethod
    def _set_run_font(
        cls,
        run,
        *,
        size: float,
        color: RGBColor,
        bold: bool = False,
    ) -> None:
        run.font.name = cls.BASE_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.bold = bold
        properties = run._element.get_or_add_rPr()
        fonts = properties.get_or_add_rFonts()
        fonts.set(qn("w:ascii"), cls.BASE_FONT)
        fonts.set(qn("w:hAnsi"), cls.BASE_FONT)
        fonts.set(qn("w:eastAsia"), cls.CJK_FONT)

    @staticmethod
    def _verify_visible_content(document: Document, expected: list[str]) -> None:
        visible = [paragraph.text.strip() for paragraph in document.paragraphs]
        expected_counts = Counter(expected)
        visible_counts = Counter(visible)
        missing = [
            text
            for text, count in expected_counts.items()
            if visible_counts[text] < count
        ]
        if missing:
            raise ValueError(f"source-index DOCX omitted content: {missing[:5]}")

    @classmethod
    def _verify_openable(cls, output: Path, expected: list[str]) -> None:
        try:
            document = Document(output)
        except Exception as exc:
            raise ValueError("source-index DOCX is not Word/WPS-openable") from exc
        cls._verify_visible_content(document, expected)

    @staticmethod
    def _canonical_docx(data: bytes) -> bytes:
        source = io.BytesIO(data)
        target = io.BytesIO()
        with (
            zipfile.ZipFile(source, "r") as input_zip,
            zipfile.ZipFile(
                target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
            ) as output_zip,
        ):
            for name in sorted(input_zip.namelist()):
                info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = input_zip.getinfo(name).external_attr
                output_zip.writestr(info, input_zip.read(name))
        return target.getvalue()
