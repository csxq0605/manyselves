"""Search deterministic project and optional global knowledge references."""

from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Literal

from docx import Document

from manyselves.capabilities.distribution_reporting.runtime.state.parallel import (
    exclusive_file_lock,
)

KnowledgeNamespace = Literal["project", "global"]
_RETIRED_HISTORY_SENTINEL = "<persisted_result_part"


def _validate_knowledge_text(text: str, *, reference: str) -> None:
    """Reject self-consistent but semantically poisoned Knowledge snapshots."""

    if _RETIRED_HISTORY_SENTINEL in text.casefold():
        raise ValueError(
            f"Knowledge text contains a retired internal history token: {reference}"
        )


@dataclass(frozen=True)
class ReferenceDocument:
    title: str
    relative_path: str
    text: str
    content_sha256: str
    namespace: KnowledgeNamespace = "project"


@dataclass(frozen=True)
class ReferenceHit:
    title: str
    relative_path: str
    snippet: str
    score: int
    content_sha256: str
    namespace: KnowledgeNamespace = "project"


@dataclass(frozen=True)
class _ReferenceRoot:
    namespace: KnowledgeNamespace
    logical_prefix: str
    path: Path


class ReferenceLibrary:
    """Read logical knowledge references without exposing either physical root."""

    TEXT_SUFFIXES = {".csv", ".html", ".htm", ".json", ".md", ".txt"}
    DOCUMENT_SUFFIXES = {".docx"}
    SNAPSHOT_VERSION = 1

    def __init__(
        self,
        workspace: Path,
        *,
        knowledge_root: Path | None = None,
        index_root: Path | None = None,
        global_root: Path | None = None,
    ):
        self.workspace = Path(workspace).resolve()
        project_input = (
            Path(knowledge_root).resolve()
            if knowledge_root is not None
            else (self.workspace / "Knowledge").resolve()
        )
        self._allow_cas_views = knowledge_root is not None
        if knowledge_root is None and (self.workspace / "Knowledge").is_symlink():
            raise ValueError("project Knowledge root must not be a symlink")
        self.root = project_input
        self.index_root = (
            Path(index_root).resolve()
            if index_root is not None
            else self.workspace / "Work/indexes/knowledge"
        )
        if not self.root.is_relative_to(self.workspace):
            raise ValueError("Knowledge root must stay inside the project workspace")
        if not self.index_root.is_relative_to(self.workspace):
            raise ValueError("Knowledge index must stay inside the project workspace")
        roots = [_ReferenceRoot("project", "Knowledge", self.root)]
        self.global_root: Path | None = None
        if global_root is not None:
            global_input = Path(global_root)
            if global_input.is_symlink():
                raise ValueError("GlobalKnowledge root must not be a symlink")
            self.global_root = global_input.resolve()
            roots.append(_ReferenceRoot("global", "GlobalKnowledge", self.global_root))
        self._roots = tuple(roots)
        self._documents_cache: tuple[ReferenceDocument, ...] | None = None

    def _root_for(self, value: str | Path) -> tuple[_ReferenceRoot, Path]:
        candidate = Path(value)
        if candidate.is_absolute():
            logical = candidate.absolute()
            for root in self._roots:
                if logical.is_relative_to(root.path):
                    return root, logical
            raise ValueError("reference path must use Knowledge or GlobalKnowledge")
        logical_path = candidate.as_posix()
        if not logical_path or "\x00" in logical_path or "\\" in logical_path:
            raise ValueError("reference path must use Knowledge or GlobalKnowledge")
        portable = PurePosixPath(logical_path)
        parts = portable.parts
        if portable.is_absolute() or any(part in {"", ".", ".."} for part in parts):
            raise ValueError("reference path must use Knowledge or GlobalKnowledge")
        if parts and parts[0] in {root.logical_prefix for root in self._roots}:
            if len(parts) < 2:
                raise ValueError("reference path must name a Knowledge file")
            root = next(root for root in self._roots if root.logical_prefix == parts[0])
            return root, root.path.joinpath(*parts[1:]).absolute()
        root = self._roots[0]
        return root, root.path.joinpath(*parts).absolute()

    def _safe_path(self, value: str | Path) -> tuple[_ReferenceRoot, Path]:
        root, logical = self._root_for(value)
        if not logical.is_relative_to(root.path):
            raise ValueError(f"reference path must stay beneath {root.logical_prefix}")
        resolved = logical.resolve()
        cas_root = (self.workspace / "Work/content/sha256").resolve()
        if not resolved.is_relative_to(root.path) and not (
            root.namespace == "project"
            and self._allow_cas_views
            and resolved.is_relative_to(cas_root)
        ):
            raise ValueError(f"reference path must stay beneath {root.logical_prefix}")
        return root, logical

    def _read_text(self, path: Path) -> str:
        suffix = path.suffix.casefold()
        if suffix in self.TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="ignore")
            if suffix in {".html", ".htm"}:
                text = re.sub(r"<[^>]+>", " ", text)
            return text
        if suffix == ".docx":
            document = Document(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            return text
        raise ValueError(f"unsupported reference file type: {suffix or '<none>'}")

    def _source_inventory(self) -> list[dict]:
        supported = self.TEXT_SUFFIXES | self.DOCUMENT_SUFFIXES
        inventory: list[dict] = []
        seen_subpaths: set[str] = set()
        for root in self._roots:
            if not root.path.is_dir():
                continue
            for path in sorted(root.path.rglob("*")):
                if not path.is_file():
                    continue
                try:
                    _, safe_path = self._safe_path(path)
                except ValueError:
                    continue
                if safe_path.suffix.casefold() not in supported:
                    continue
                subpath = safe_path.relative_to(root.path).as_posix()
                subpath_key = subpath.casefold()
                if subpath_key in seen_subpaths:
                    continue
                seen_subpaths.add(subpath_key)
                stat = safe_path.stat()
                inventory.append(
                    {
                        "namespace": root.namespace,
                        "relative_path": f"{root.logical_prefix}/{subpath}",
                        "size": stat.st_size,
                        "mtime_ns": stat.st_mtime_ns,
                        "device": stat.st_dev,
                        "inode": stat.st_ino,
                    }
                )
        return inventory

    @staticmethod
    def _inventory_digest(inventory: list[dict]) -> str:
        return hashlib.sha256(
            json.dumps(
                inventory,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()

    @staticmethod
    def _write_immutable_text(path: Path, text: str) -> None:
        payload = text.encode("utf-8")
        if path.is_file():
            if path.read_bytes() != payload:
                raise ValueError("knowledge snapshot text changed for one digest")
            return
        path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            mode="wb",
            prefix=f".{path.name}-",
            suffix=".tmp",
            dir=path.parent,
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        try:
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)

    def _build_or_load_snapshot(self) -> tuple[ReferenceDocument, ...]:
        if self._documents_cache is not None:
            return self._documents_cache
        inventory = self._source_inventory()
        inventory_digest = self._inventory_digest(inventory)
        manifest_path = self.index_root / "snapshots" / f"{inventory_digest}.json"
        lock_path = self.index_root / "snapshot.lock"
        with exclusive_file_lock(lock_path):
            if not manifest_path.is_file():
                documents: list[dict] = []
                seen_content: set[str] = set()
                for item in inventory:
                    relative = item["relative_path"]
                    root, safe_path = self._safe_path(relative)
                    text = self._read_text(safe_path)
                    _validate_knowledge_text(text, reference=relative)
                    text_sha256 = hashlib.sha256(text.encode("utf-8")).hexdigest()
                    if text_sha256 in seen_content:
                        continue
                    seen_content.add(text_sha256)
                    text_ref = (
                        self.index_root.relative_to(self.workspace)
                        / "text"
                        / text_sha256[:2]
                        / f"{text_sha256}.txt"
                    )
                    self._write_immutable_text(self.workspace / text_ref, text)
                    documents.append(
                        {
                            "title": safe_path.stem,
                            "relative_path": relative,
                            "namespace": root.namespace,
                            "text_ref": text_ref.as_posix(),
                            "text_sha256": text_sha256,
                            "text_chars": len(text),
                            "source_stat": item,
                        }
                    )
                manifest = {
                    "snapshot_version": self.SNAPSHOT_VERSION,
                    "inventory_digest": inventory_digest,
                    "inventory": inventory,
                    "documents": documents,
                }
                manifest_path.parent.mkdir(parents=True, exist_ok=True)
                self._write_immutable_text(
                    manifest_path,
                    json.dumps(
                        manifest,
                        ensure_ascii=False,
                        sort_keys=True,
                        indent=2,
                    )
                    + "\n",
                )
                current_path = self.index_root / "current.json"
                current_path.parent.mkdir(parents=True, exist_ok=True)
                temporary = current_path.with_suffix(".tmp")
                temporary.write_text(
                    json.dumps(
                        {
                            "snapshot_version": self.SNAPSHOT_VERSION,
                            "inventory_digest": inventory_digest,
                            "manifest_ref": manifest_path.relative_to(
                                self.workspace
                            ).as_posix(),
                        },
                        ensure_ascii=False,
                        sort_keys=True,
                        indent=2,
                    )
                    + "\n",
                    encoding="utf-8",
                )
                os.replace(temporary, current_path)
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            current_payload = {
                "snapshot_version": self.SNAPSHOT_VERSION,
                "inventory_digest": inventory_digest,
                "manifest_ref": manifest_path.relative_to(
                    self.workspace
                ).as_posix(),
            }
            current_path = self.index_root / "current.json"
            current = (
                json.loads(current_path.read_text(encoding="utf-8"))
                if current_path.is_file()
                else None
            )
            if current != current_payload:
                current_path.parent.mkdir(parents=True, exist_ok=True)
                temporary = current_path.with_suffix(".tmp")
                temporary.write_text(
                    json.dumps(
                        current_payload,
                        ensure_ascii=False,
                        sort_keys=True,
                        indent=2,
                    )
                    + "\n",
                    encoding="utf-8",
                )
                os.replace(temporary, current_path)
        documents = []
        for item in manifest["documents"]:
            text_path = (self.workspace / item["text_ref"]).resolve()
            if not text_path.is_relative_to(self.index_root.resolve()) or not text_path.is_file():
                raise ValueError("knowledge snapshot text is missing")
            text = text_path.read_text(encoding="utf-8")
            if hashlib.sha256(text.encode("utf-8")).hexdigest() != item["text_sha256"]:
                raise ValueError("knowledge snapshot text hash mismatch")
            _validate_knowledge_text(
                text,
                reference=str(item["relative_path"]),
            )
            documents.append(
                ReferenceDocument(
                    title=item["title"],
                    relative_path=item["relative_path"],
                    text=text,
                    content_sha256=item["text_sha256"],
                    namespace=item.get("namespace", "project"),
                )
            )
        self._documents_cache = tuple(documents)
        return self._documents_cache

    def documents(self) -> tuple[ReferenceDocument, ...]:
        """Return the complete immutable parsed snapshot for this revision."""

        return self._build_or_load_snapshot()

    def snapshot_manifest_ref(self) -> Path:
        self._build_or_load_snapshot()
        current = json.loads(
            (self.index_root / "current.json").read_text(encoding="utf-8")
        )
        return Path(current["manifest_ref"])

    def open(self, relative_path: str) -> ReferenceDocument:
        root, path = self._safe_path(relative_path)
        if not path.is_file():
            raise ValueError(f"reference path is not a file in {root.logical_prefix}")
        canonical = f"{root.logical_prefix}/{path.relative_to(root.path).as_posix()}"
        for document in self._build_or_load_snapshot():
            if document.relative_path == canonical:
                return document
        raise ValueError("reference path is absent from the immutable Knowledge snapshot")

    def search(self, query: str, limit: int = 5) -> list[ReferenceHit]:
        terms = [term.casefold() for term in query.split() if term.strip()]
        if not terms or limit <= 0:
            return []

        hits: list[ReferenceHit] = []
        for document in self.documents():
            lowered = document.text.casefold()
            score = sum(lowered.count(term) for term in terms)
            if not score:
                continue
            first = min(lowered.find(term) for term in terms if term in lowered)
            start = max(0, first - 160)
            snippet = " ".join(document.text[start : first + 440].split())
            hits.append(
                ReferenceHit(
                    title=document.title,
                    relative_path=document.relative_path,
                    snippet=snippet,
                    score=score,
                    content_sha256=document.content_sha256,
                    namespace=document.namespace,
                )
            )
        priority = {"project": 0, "global": 1}
        return sorted(
            hits,
            key=lambda hit: (
                -hit.score,
                priority[hit.namespace],
                hit.relative_path.casefold(),
                hit.relative_path,
            ),
        )[:limit]
