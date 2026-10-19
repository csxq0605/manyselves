"""Suffix-dispatched intake adapters with explicit manual-review outcomes."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from manyselves.capabilities.distribution_reporting.runtime.models.preparation import (
    ManifestFile,
    ParsedArtifact,
)
from manyselves.capabilities.distribution_reporting.runtime.models.reporting import (
    SourceLocation,
)
from manyselves.runtime.artifacts.parsers import parse_artifact

Parser = Callable[[Path, ManifestFile], list[ParsedArtifact]]


def _artifact(
    manifest: ManifestFile,
    kind: str,
    index: int,
    payload: dict,
    **location,
) -> ParsedArtifact:
    return ParsedArtifact(
        id=f"{manifest.id}-{kind}-{index:04d}",
        kind=kind,
        source=SourceLocation(
            file_id=manifest.id,
            path=manifest.path,
            **location,
        ),
        payload=payload,
    )


def _parse_text(path: Path, manifest: ManifestFile) -> list[ParsedArtifact]:
    return [
        _artifact(
            manifest,
            "text",
            1,
            {"text": path.read_text(encoding="utf-8", errors="replace")},
        )
    ]


def _parse_docx(path: Path, manifest: ManifestFile) -> list[ParsedArtifact]:
    document = parse_artifact(path)
    if document.error:
        raise ValueError(document.error)
    return [
        _artifact(
            manifest, "document_paragraph", index,
            {"text": block.text, "locator": block.locator},
            cell=block.locator, row=int(block.locator.rsplit(":", 1)[-1]),
        )
        for index, block in enumerate(document.blocks, start=1)
        if block.text.strip()
    ]


def _parse_image(path: Path, manifest: ManifestFile) -> list[ParsedArtifact]:
    with Image.open(path) as image:
        payload = {
            "width": image.width,
            "height": image.height,
            "format": image.format or "unknown",
            "mode": image.mode,
        }
    return [_artifact(manifest, "image_metadata", 1, payload)]


def _parse_workbook(path: Path, manifest: ManifestFile) -> list[ParsedArtifact]:
    # Raw cells are observations, not an inferred header/value schema. Normal
    # mode exposes merge ranges. Load cached values in normal mode too: a
    # read-only stream trusts optional/stale worksheet dimension metadata and
    # can yield a different row shape or omit cells present in the formula view.
    workbook = load_workbook(path, read_only=False, data_only=False)
    artifacts: list[ParsedArtifact] = []
    try:
        values = load_workbook(path, read_only=False, data_only=True)
        try:
            for sheet in workbook.worksheets:
                cached_rows = values[sheet.title].iter_rows()
                for row_index, (row, cached) in enumerate(
                    zip(sheet.iter_rows(), cached_rows, strict=True), start=1
                ):
                    if not any(cell.value not in (None, "") for cell in row):
                        continue
                    row_values = []
                    for cell, cached_cell in zip(row, cached, strict=True):
                        value = cell.value
                        if cell.data_type == "f":
                            value = cached_cell.value
                            if value is None or cached_cell.data_type == "e":
                                reason = "formula_result_error" if cached_cell.data_type == "e" else "formula_result_unavailable"
                                artifacts.append(_artifact(
                                    manifest, "manual_required", len(artifacts) + 1,
                                    {"reason": reason, "formula": str(cell.value), "cached_value": value},
                                    sheet=sheet.title, cell=cell.coordinate, row=row_index,
                                ))
                                value = None
                        row_values.append(value)
                    merges = [str(area) for area in sheet.merged_cells.ranges
                              if area.min_row == row_index]
                    artifacts.append(_artifact(
                        manifest, "workbook_row", len(artifacts) + 1,
                        {"headers": [cell.coordinate for cell in row],
                         "values": row_values, "merged_ranges": sorted(merges),
                         "header_inferred": False},
                        sheet=sheet.title, row=row_index,
                        cell=f"A{row_index}:{get_column_letter(len(row))}{row_index}",
                    ))
        finally:
            values.close()
    finally:
        workbook.close()
    return artifacts


def _manual(
    manifest: ManifestFile, reason: str, *, detail: str | None = None
) -> list[ParsedArtifact]:
    payload = {"reason": reason}
    if detail:
        payload["detail"] = detail
    return [_artifact(manifest, "manual_required", 1, payload)]


def _parse_pdf(path: Path, manifest: ManifestFile) -> list[ParsedArtifact]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return _manual(manifest, "pdf_parser_unavailable")
    reader = PdfReader(path)
    artifacts = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        # Empty extraction can mean a scan, blank page or unsupported content;
        # it is not proof that the page contains no customer observations.
        kind = "pdf_page" if text.strip() else "manual_required"
        payload = {"text": text} if text.strip() else {"reason": "pdf_page_without_text"}
        artifacts.append(_artifact(manifest, kind, page_number, payload, page=page_number))
    return artifacts


class IntakeAdapterRegistry:
    """Parse each manifest file without hiding unsupported binary formats."""

    PARSERS: dict[str, Parser] = {
        ".xlsx": _parse_workbook,
        ".xlsm": _parse_workbook,
        ".docx": _parse_docx,
        ".md": _parse_text,
        ".txt": _parse_text,
        ".png": _parse_image,
        ".jpg": _parse_image,
        ".jpeg": _parse_image,
        ".pdf": _parse_pdf,
    }
    MANUAL_REQUIRED = {".dwg", ".mp4", ".mov", ".avi"}

    def parse(self, path: Path, manifest_file: ManifestFile) -> list[ParsedArtifact]:
        path = Path(path)
        suffix = path.suffix.casefold()
        if suffix in self.MANUAL_REQUIRED:
            return _manual(manifest_file, "unsupported_binary_format")
        parser = self.PARSERS.get(suffix)
        if parser is None:
            return _manual(manifest_file, "unsupported_format", detail=suffix or "no_suffix")
        return parser(path, manifest_file)
