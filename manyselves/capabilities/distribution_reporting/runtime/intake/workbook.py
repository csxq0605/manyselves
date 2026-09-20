"""Lossless workbook structure inspection used before business mapping."""

from pathlib import Path

from openpyxl import load_workbook

from manyselves.capabilities.distribution_reporting.runtime.models.reporting import ReportingModel


class WorkbookSheetArtifact(ReportingModel):
    title: str
    max_row: int
    max_column: int
    formulas: dict[str, str]


class WorkbookArtifact(ReportingModel):
    file_id: str
    path: Path
    sheets: dict[str, WorkbookSheetArtifact]


def inspect_workbook(path: Path, *, file_id: str) -> WorkbookArtifact:
    """Return sheet dimensions and exact formula text for one workbook."""

    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheets: dict[str, WorkbookSheetArtifact] = {}
        for sheet in workbook.worksheets:
            formulas = {
                cell.coordinate: cell.value
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f" and isinstance(cell.value, str)
            }
            sheets[sheet.title] = WorkbookSheetArtifact(
                title=sheet.title,
                max_row=sheet.max_row,
                max_column=sheet.max_column,
                formulas=formulas,
            )
        return WorkbookArtifact(file_id=file_id, path=path, sheets=sheets)
    finally:
        workbook.close()
