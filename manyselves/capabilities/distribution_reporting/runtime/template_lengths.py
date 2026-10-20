"""Section lengths of the expert report template, as a writing reference.

The template's heading outline is numbered by position (empty headings are
skipped), and each heading's Chinese characters are counted over its
paragraphs and tables up to the next heading. The numbers are a reference
for how much the expert version writes per section, never an output limit.
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

_CJK = re.compile(r"[一-鿿]")
_HEADING = re.compile(r"^(?:Heading|标题)\s*(\d)$", re.IGNORECASE)


@dataclass(frozen=True)
class TemplateSection:
    number: str
    title: str
    chars: int
    tables: int


def _level(paragraph: Paragraph) -> int | None:
    match = _HEADING.match(paragraph.style.name if paragraph.style is not None else "")
    return int(match.group(1)) if match else None


def template_sections(path: Path) -> list[TemplateSection]:
    """Every non-empty heading with the characters and tables directly under it."""

    document = Document(str(path))
    counters: list[int] = []
    sections: list[dict] = []
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            level = _level(paragraph)
            text = paragraph.text.strip()
            if level is not None:
                if not text:
                    continue
                counters = counters[:level] + [0] * max(0, level - len(counters))
                counters[level - 1] += 1
                sections.append(
                    {"number": ".".join(map(str, counters)), "title": text, "chars": 0, "tables": 0}
                )
            elif sections:
                sections[-1]["chars"] += len(_CJK.findall(text))
        elif element.tag == qn("w:tbl") and sections:
            table = Table(element, document)
            cells = {id(cell._tc): cell.text for row in table.rows for cell in row.cells}
            sections[-1]["chars"] += sum(len(_CJK.findall(text)) for text in cells.values())
            sections[-1]["tables"] += 1
    return [TemplateSection(**section) for section in sections]


def _similar(left: str, right: str) -> float:
    return SequenceMatcher(None, left, right).ratio()


def _total(sections: list[TemplateSection], number: str) -> TemplateSection:
    members = [s for s in sections if s.number == number or s.number.startswith(number + ".")]
    head = next(s for s in sections if s.number == number)
    return TemplateSection(
        number, head.title, sum(s.chars for s in members), sum(s.tables for s in members)
    )


def matched_lengths(
    sections: list[TemplateSection], wanted: Iterable[tuple[str, str]]
) -> dict[str, TemplateSection]:
    """Map report section ids to template sections by number, checked by title.

    A number whose template title is unlike the wanted title falls back to
    the most similar title; nothing is reported without a reasonable match.
    """

    by_number = {s.number: s for s in sections}
    found: dict[str, TemplateSection] = {}
    for section_id, title in wanted:
        candidate = by_number.get(section_id)
        if candidate is None or _similar(candidate.title, title) < 0.5:
            best = max(sections, key=lambda s: _similar(s.title, title), default=None)
            candidate = best if best is not None and _similar(best.title, title) >= 0.6 else None
        if candidate is not None:
            found[section_id] = _total(sections, candidate.number)
    return found


def render_length_reference(
    lengths: Mapping[str, TemplateSection], titles: Mapping[str, str]
) -> str:
    lines = [
        "# 专家版各节篇幅参考",
        "专家版模板对应小节的中文字数和表格数，只作写作参考，不是上限：同类对象一句概括，"
        "缺资料收进数据缺口表，第一、三章只引用第二章结论。本项目对象更多时可以更长，但不重复论证。",
    ]
    for section_id, title in titles.items():
        if (section := lengths.get(section_id)) is not None:
            lines.append(
                f"{section_id} {title}：约 {section.chars} 字"
                + (f"，{section.tables} 张表" if section.tables else "")
            )
    return "\n".join(lines) + "\n"


__all__ = [
    "TemplateSection",
    "matched_lengths",
    "render_length_reference",
    "template_sections",
]
