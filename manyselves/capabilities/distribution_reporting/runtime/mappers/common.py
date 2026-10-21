"""Shared mapper carriers and evidence construction helpers."""

import hashlib
import re
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from manyselves.capabilities.distribution_reporting.runtime.models.preparation import MappingGap
from manyselves.capabilities.distribution_reporting.runtime.models.reporting import (
    EvidenceItem,
    EvidenceMeasurement,
    SourceLocation,
)

_DISPIMG_PATTERN = re.compile(r'DISPIMG\("([^"]+)"', re.IGNORECASE)


def dispimg_refs(*values: Any) -> list[str]:
    refs: list[str] = []
    for value in values:
        if not isinstance(value, str):
            continue
        refs.extend(match.group(1) for match in _DISPIMG_PATTERN.finditer(value))
    return list(dict.fromkeys(refs))


def unbound_photo_gaps(workbook: Workbook, evidence: list[EvidenceItem]) -> list[MappingGap]:
    """Every in-cell photo either rides on a mapped record or becomes a gap.

    A photo whose row or checkpoint the mapper could not read must not vanish
    silently; the gap names its cell so the report can state it.
    """

    bound = {photo_id for item in evidence for photo_id in item.photo_refs}
    gaps: list[MappingGap] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                unbound = [ref for ref in dispimg_refs(cell.value) if ref not in bound]
                if unbound:
                    gaps.append(MappingGap(
                        code="photo_without_record",
                        message=f"图片未对应到可读取的检查记录：{sheet.title}!{cell.coordinate}",
                        sheet=sheet.title,
                        cell=cell.coordinate,
                    ))
    return gaps


HEADER_ROWS = 3

# A trailing unit in a header label, e.g. "剩余电流(A)" -> "剩余电流".
_TRAILING_UNIT = re.compile(r"\s*[（(][^（）()]*[）)]\s*$")


def field_name(label: Any) -> str:
    """A table header label as a field name: without a trailing unit."""

    return _TRAILING_UNIT.sub("", str(label or "").strip()).strip()


def header_label(sheet: Any, row: int, column: int, fallback: str) -> str:
    """The checkpoint name the table's own label row gives a column."""

    return field_name(sheet.cell(row, column).value) or fallback


def source_layout(
    path: Path, evidence: list[EvidenceItem], image_ids: set[str]
) -> tuple[dict[str, dict[str, str]], dict[str, tuple[str, str]]]:
    """Read the header text above each sheet's records and where each image sits.

    Header rows are the rows above the first mapped record of a sheet, at most
    ``HEADER_ROWS`` of them; merged header cells count for every column they
    span. The image location is the cell holding its DISPIMG formula.
    """

    first_row: dict[str, int] = {}
    for item in evidence:
        sheet, row = item.source.sheet, item.source.row
        if sheet and row:
            first_row[sheet] = min(first_row.get(sheet, row), row)
    headers: dict[str, dict[str, str]] = {}
    cells: dict[str, tuple[str, str]] = {}
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        for sheet in workbook.worksheets:
            if sheet.title in first_row:
                top = max(1, first_row[sheet.title] - HEADER_ROWS)
                rows = range(top, first_row[sheet.title])
                merged = {}
                for span in sheet.merged_cells.ranges:
                    value = sheet.cell(span.min_row, span.min_col).value
                    for row in range(span.min_row, span.max_row + 1):
                        for column in range(span.min_col, span.max_col + 1):
                            merged[(row, column)] = value
                columns: dict[str, str] = {}
                for column in range(1, sheet.max_column + 1):
                    texts: list[str] = []
                    for row in rows:
                        value = merged.get((row, column), sheet.cell(row, column).value)
                        text = str(value).strip() if value not in (None, "") else ""
                        if text and not dispimg_refs(text) and text not in texts:
                            texts.append(text)
                    if texts:
                        columns[get_column_letter(column)] = " / ".join(texts)
                if columns:
                    headers[sheet.title] = columns
            if image_ids:
                for row in sheet.iter_rows():
                    for cell in row:
                        for ref in dispimg_refs(cell.value):
                            if ref in image_ids:
                                cells.setdefault(ref, (sheet.title, cell.coordinate))
    finally:
        workbook.close()
    return headers, cells


@contextmanager
def cached_formula_workbook(path: Path) -> Iterator[tuple[Workbook, list[MappingGap]]]:
    """Resolve saved scalar results in memory, retaining DISPIMG for photo binding.

    Both views derive dimensions from cells. Missing/error caches remain gaps,
    even if an unresolved routing cell prevents its row from becoming evidence.
    The source workbook is never saved or recalculated.
    """
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        values = load_workbook(path, read_only=False, data_only=True)
        gaps = []
        try:
            for sheet in workbook.worksheets:
                cached_sheet = values[sheet.title]
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type != "f" or dispimg_refs(cell.value):
                            continue
                        cached = cached_sheet.cell(cell.row, cell.column)
                        if cached.value is None or cached.data_type == "e":
                            error = cached.data_type == "e"
                            detail = f"；缓存错误={cached.value}" if error else ""
                            gaps.append(MappingGap(
                                code="formula_result_error" if error else "formula_result_unavailable",
                                message=f"公式缺少可用缓存结果：{sheet.title}!{cell.coordinate}；{cell.value}{detail}",
                                sheet=sheet.title,
                                cell=cell.coordinate,
                            ))
                            cell.value = None
                        else:
                            cell.value = cached.value
        finally:
            values.close()
        yield workbook, gaps
    finally:
        workbook.close()


def evidence_id(
    file_id: str,
    sheet: str,
    cell: str,
    submodule_id: str | None,
    fact: str,
) -> str:
    raw = "\0".join((file_id, sheet, cell, submodule_id or "", fact)).encode()
    return f"ev-{hashlib.sha256(raw).hexdigest()[:20]}"


def make_evidence(
    *,
    file_id: str,
    path: Path,
    sheet: str,
    cell: str,
    row: int,
    column: str,
    subject: str,
    fact: str,
    module_id: str | None = None,
    submodule_id: str | None = None,
    photo_refs: list[str] | None = None,
    value: str | float | int | None = None,
    unit: str | None = None,
    measurements: list[EvidenceMeasurement] | None = None,
    confidence: float = 1.0,
    needs_confirmation: bool = False,
    summary_entry: bool = False,
    program_note: str | None = None,
    checkpoint_field: str | None = None,
) -> EvidenceItem:
    return EvidenceItem(
        id=evidence_id(file_id, sheet, cell, submodule_id, fact),
        subject=subject,
        fact=fact,
        source=SourceLocation(
            file_id=file_id,
            path=path,
            sheet=sheet,
            cell=cell,
            row=row,
            column=column,
        ),
        value=value,
        unit=unit,
        measurements=measurements or [],
        confidence=confidence,
        needs_confirmation=needs_confirmation,
        module_id=module_id,
        submodule_id=submodule_id,
        photo_refs=photo_refs or [],
        summary_entry=summary_entry,
        program_note=program_note,
        checkpoint_field=checkpoint_field,
    )
