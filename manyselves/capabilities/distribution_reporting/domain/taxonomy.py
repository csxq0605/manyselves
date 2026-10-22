"""Capability-owned run-scoped taxonomy for distribution reporting.

Chapter 2's section tree comes only from the run's own 评估信息汇总表 (the
S4-6 workbook).  There is no built-in tree: a step that reads the taxonomy
without one bound fails instead of falling back to some other project's
sections.  ``report-taxonomy-example.json`` illustrates contracts and serves
tests; a run never reads it.
"""

import hashlib
import json
import re
from collections.abc import Callable, Iterator, Mapping
from contextlib import contextmanager
from contextvars import ContextVar, Token
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True, slots=True)
class SubmoduleDefinition:
    id: str
    title: str
    module_id: str


@dataclass(frozen=True, slots=True)
class ModuleDefinition:
    id: str
    title: str
    sections: dict[str, SubmoduleDefinition]
    submodules: dict[str, SubmoduleDefinition]

    @property
    def groups(self) -> dict[str, SubmoduleDefinition]:
        """Return structural headings that own child sections but no prose part."""

        return {
            section_id: section
            for section_id, section in self.sections.items()
            if section_id not in self.submodules
        }


def _module(
    module_id: str,
    title: str,
    sections: tuple[tuple[str, str], ...],
) -> ModuleDefinition:
    definitions = {
        section_id: SubmoduleDefinition(
            id=section_id,
            title=section_title,
            module_id=module_id,
        )
        for section_id, section_title in sections
    }
    return ModuleDefinition(
        id=module_id,
        title=title,
        sections=definitions,
        submodules={
            section_id: definition
            for section_id, definition in definitions.items()
            if not any(
                other_id.startswith(f"{section_id}.")
                for other_id in definitions
            )
        },
    )


_ACTIVE_REPORT_TAXONOMY: ContextVar[Mapping[str, ModuleDefinition] | None] = (
    ContextVar("active_report_taxonomy", default=None)
)


# Where the tree of the run being executed can be read once it exists: set
# around a whole run, so that a step outside any Tool call (validating a
# skipped step's recorded result, or state a resumed run reloaded) still
# reads the run's own tree.
_RUN_TAXONOMY_SOURCE: ContextVar[Callable[[], Mapping[str, Any] | None] | None] = ContextVar(
    "run_taxonomy_source", default=None
)
_SOURCE_CACHE: dict[str, Mapping[str, ModuleDefinition]] = {}


EXAMPLE_TAXONOMY_FILE = Path(__file__).with_name("report-taxonomy-example.json")


class ReportTaxonomyUnbound(LookupError):
    """A step read the section tree before the run's own tree was bound."""


class _RunTaxonomyView(Mapping[str, ModuleDefinition]):
    """Mapping facade whose value follows the current async run context."""

    @staticmethod
    def _current() -> Mapping[str, ModuleDefinition]:
        modules = _ACTIVE_REPORT_TAXONOMY.get()
        if modules is None:
            modules = _from_run_source()
        if modules is None:
            raise ReportTaxonomyUnbound(
                "报告小节树未绑定：第二章小节只从本次运行的评估信息汇总表读取"
            )
        return modules

    def __getitem__(self, key: str) -> ModuleDefinition:
        return self._current()[key]

    def __iter__(self) -> Iterator[str]:
        return iter(self._current())

    def __len__(self) -> int:
        return len(self._current())


REPORT_TAXONOMY: Mapping[str, ModuleDefinition] = _RunTaxonomyView()


def _section_id(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        text = format(float(value), ".12g")
    else:
        text = str(value).strip()
    text = re.sub(r"\s+", "", text)
    return text if re.fullmatch(r"2\.\d+(?:\.\d+)*", text) else None


def _validate_taxonomy(modules: Mapping[str, ModuleDefinition]) -> None:
    """The tree is whatever the table lists, as long as it is a tree.

    Which modules and sections appear is the table's call; a project without
    a problem leaves its section out.  Ids must be chapter 2 and in order.
    """

    if not modules:
        raise ValueError("S4-6 workbook taxonomy has no chapter 2 module")
    keys = [tuple(int(part) for part in module_id.split(".")) for module_id in modules]
    if any(len(key) != 2 or key[0] != 2 for key in keys) or keys != sorted(set(keys)):
        raise ValueError(
            f"S4-6 workbook taxonomy modules must be ordered chapter 2 ids; actual={list(modules)}"
        )
    seen: set[str] = set(modules)
    for module_id, module in modules.items():
        if not module.title.strip():
            raise ValueError(f"workbook taxonomy module has no title: {module_id}")
        for section_id, section in module.sections.items():
            if section_id in seen:
                raise ValueError(f"duplicate workbook taxonomy id: {section_id}")
            seen.add(section_id)
            if section.module_id != module_id or not section.title.strip():
                raise ValueError(f"invalid workbook taxonomy section: {section_id}")
            parent_id = section_id.rsplit(".", 1)[0]
            if parent_id != module_id and parent_id not in module.sections:
                raise ValueError(
                    f"workbook taxonomy section has no parent: {section_id} -> {parent_id}"
                )


def parse_report_taxonomy_workbook(
    path: Path,
    *,
    source_ref: str | None = None,
    source_sha256: str | None = None,
) -> dict[str, Any]:
    """Parse the chapter 2 heading tree a frozen S4-6 workbook lists."""

    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if "评估信息汇总表" not in workbook.sheetnames:
            raise ValueError("S4-6 workbook is missing sheet: 评估信息汇总表")
        sheet = workbook["评估信息汇总表"]
        module_titles: dict[str, str] = {}
        section_rows: dict[str, list[tuple[str, str]]] = {}
        ordered_ids: list[str] = []
        for row in range(1, sheet.max_row + 1):
            identifier = _section_id(sheet.cell(row, 2).value)
            if identifier is None:
                continue
            title = next(
                (
                    str(sheet.cell(row, column).value).strip()
                    for column in range(3, 6)
                    if sheet.cell(row, column).value is not None
                    and str(sheet.cell(row, column).value).strip()
                ),
                "",
            )
            if not title:
                raise ValueError(
                    f"S4-6 workbook taxonomy title is empty at {sheet.title}!B{row}:E{row}"
                )
            if identifier in ordered_ids:
                raise ValueError(f"duplicate S4-6 workbook taxonomy id: {identifier}")
            ordered_ids.append(identifier)
            if identifier.count(".") == 1:
                module_titles[identifier] = title
                section_rows.setdefault(identifier, [])
            else:
                module_id = ".".join(identifier.split(".")[:2])
                if module_id not in module_titles:
                    raise ValueError(
                        f"S4-6 workbook taxonomy child precedes module: {identifier}"
                    )
                section_rows[module_id].append((identifier, title))
        modules = {
            module_id: _module(
                module_id,
                module_titles[module_id],
                tuple(section_rows[module_id]),
            )
            for module_id in module_titles
        }
        _validate_taxonomy(modules)
        digest = source_sha256 or hashlib.sha256(path.read_bytes()).hexdigest()
        payload = report_taxonomy_snapshot(
            modules,
            source_ref=source_ref or path.as_posix(),
            source_sha256=digest,
            sheet=sheet.title,
        )
        payload["source_kind"] = "xlsx"
        return payload
    finally:
        workbook.close()


def taxonomy_from_snapshot(payload: Mapping[str, Any]) -> dict[str, ModuleDefinition]:
    """Validate and materialize one persisted run taxonomy snapshot."""

    if payload.get("schema_version") != 1:
        raise ValueError("unsupported report taxonomy snapshot version")
    modules: dict[str, ModuleDefinition] = {}
    for raw_module in payload.get("modules") or ():
        if not isinstance(raw_module, Mapping):
            raise ValueError("invalid report taxonomy module payload")
        module_id = str(raw_module.get("id") or "")
        title = str(raw_module.get("title") or "").strip()
        sections = tuple(
            (str(item.get("id") or ""), str(item.get("title") or "").strip())
            for item in (raw_module.get("sections") or ())
            if isinstance(item, Mapping)
        )
        modules[module_id] = _module(module_id, title, sections)
    _validate_taxonomy(modules)
    return modules


def report_taxonomy_snapshot(
    modules: Mapping[str, ModuleDefinition],
    *,
    source_ref: str,
    source_sha256: str,
    sheet: str,
) -> dict[str, Any]:
    """Serialize one validated taxonomy without report prose or evidence."""

    _validate_taxonomy(modules)
    payload = {
        "schema_version": 1,
        "source_ref": source_ref,
        "source_sha256": source_sha256,
        "sheet": sheet,
        "modules": [
            {
                "id": module.id,
                "title": module.title,
                "sections": [
                    {"id": section.id, "title": section.title}
                    for section in module.sections.values()
                ],
            }
            for module in modules.values()
        ],
    }
    return payload


def activate_report_taxonomy(
    payload: Mapping[str, Any],
) -> Token[Mapping[str, ModuleDefinition] | None]:
    """Bind a validated immutable taxonomy to the current async run context."""

    modules = taxonomy_from_snapshot(payload)
    return _ACTIVE_REPORT_TAXONOMY.set(modules)


def reset_report_taxonomy(
    token: Token[Mapping[str, ModuleDefinition] | None],
) -> None:
    _ACTIVE_REPORT_TAXONOMY.reset(token)


@contextmanager
def bound_report_taxonomy(payload: Mapping[str, Any]) -> Iterator[None]:
    """Bind a run's taxonomy snapshot for the duration of a block."""

    token = activate_report_taxonomy(payload)
    try:
        yield
    finally:
        reset_report_taxonomy(token)


def _from_run_source() -> Mapping[str, ModuleDefinition] | None:
    source = _RUN_TAXONOMY_SOURCE.get()
    payload = source() if source is not None else None
    if not payload:
        return None
    key = json.dumps(payload, sort_keys=True, ensure_ascii=False)
    modules = _SOURCE_CACHE.get(key)
    if modules is None:
        modules = taxonomy_from_snapshot(payload)
        _SOURCE_CACHE.clear()
        _SOURCE_CACHE[key] = modules
    return modules


@contextmanager
def run_taxonomy_source(read: Callable[[], Mapping[str, Any] | None]) -> Iterator[None]:
    """Read the run's tree through ``read`` wherever no tree is bound."""

    token = _RUN_TAXONOMY_SOURCE.set(read)
    try:
        yield
    finally:
        _RUN_TAXONOMY_SOURCE.reset(token)


def report_taxonomy_bound() -> bool:
    return _ACTIVE_REPORT_TAXONOMY.get() is not None


def example_report_taxonomy() -> dict[str, Any]:
    """The illustration tree for contract examples and tests; never a run's tree."""

    payload = json.loads(EXAMPLE_TAXONOMY_FILE.read_text(encoding="utf-8"))
    payload.pop("_note", None)
    return payload


_MARKDOWN_HEADING = re.compile(r"^#{1,6}\s+(2\.\d+(?:\.\d+)*)\.?\s+(.+?)\s*$")


def taxonomy_from_markdown(markdown: str) -> dict[str, Any]:
    """The chapter 2 tree a finished report's own headings list.

    For rendering a report on its own, outside a run that has its table.
    """

    titles: dict[str, str] = {}
    for line in markdown.splitlines():
        match = _MARKDOWN_HEADING.match(line.strip())
        if match and match.group(1) not in titles:
            titles[match.group(1)] = match.group(2).strip()
    modules = {
        module_id: _module(
            module_id,
            title,
            tuple(
                (section_id, section_title)
                for section_id, section_title in titles.items()
                if section_id.startswith(module_id + ".")
            ),
        )
        for module_id, title in titles.items()
        if module_id.count(".") == 1
    }
    return report_taxonomy_snapshot(
        modules, source_ref="report-markdown", source_sha256="", sheet=""
    )


def resolve_submodule(submodule_id: str) -> SubmoduleDefinition:
    """Return a fixed submodule definition or reject an invented identifier."""

    for module in REPORT_TAXONOMY.values():
        if submodule := module.submodules.get(submodule_id):
            return submodule
    raise ValueError(f"unknown report submodule: {submodule_id}")


def compose_module_markdown(
    module_id: str, submodule_narratives: dict[str, str]
) -> str:
    """Derive the only canonical module Markdown from fixed submodule prose."""

    definition = REPORT_TAXONOMY[module_id]
    expected = set(definition.submodules)
    actual = set(submodule_narratives)
    if actual != expected:
        raise ValueError(
            "cannot compose module Markdown with incomplete taxonomy; "
            f"missing={sorted(expected - actual)}; extra={sorted(actual - expected)}"
        )

    blocks = [f"## {module_id} {definition.title}"]
    for section_id, section in definition.sections.items():
        heading_level = section_id.count(".") + 1
        if section_id not in definition.submodules:
            blocks.append(f"{'#' * heading_level} {section_id} {section.title}")
            continue
        submodule_id = section_id
        submodule = section
        narrative = submodule_narratives[submodule_id].strip()
        if not narrative:
            raise ValueError(f"cannot compose empty submodule narrative: {submodule_id}")
        lines = narrative.splitlines()
        first_content = next(
            (index for index, line in enumerate(lines) if line.strip()), None
        )
        if first_content is not None and re.match(
            rf"^#{{1,6}}\s+{re.escape(submodule_id)}(?:\.|\s|$)",
            lines[first_content].strip(),
        ):
            del lines[first_content]
            while first_content < len(lines) and not lines[first_content].strip():
                del lines[first_content]
        body_heading = "#" * min(heading_level + 1, 6)
        normalized_body = "\n".join(
            re.sub(r"^#{1,6}\s+", f"{body_heading} ", line)
            if re.match(r"^#{1,6}\s+", line)
            else line
            for line in lines
        ).strip()
        blocks.append(
            f"{'#' * heading_level} {submodule_id} {submodule.title}\n\n"
            f"{normalized_body}"
        )
    return "\n\n".join(blocks)
