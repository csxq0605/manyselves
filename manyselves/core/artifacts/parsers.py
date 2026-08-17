"""Truthful format-aware artifact extraction."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from docx import Document
from openpyxl import load_workbook
from PIL import Image

from .types import ArtifactDescriptor, descriptor_for_path


@dataclass(frozen=True)
class ArtifactBlock:
    locator: str
    text: str


@dataclass(frozen=True)
class ParsedArtifactBlocks:
    kind: Literal["text", "docx", "xlsx", "pdf", "image_metadata", "unsupported"]
    blocks: tuple[ArtifactBlock, ...]
    visual_verified: bool = False
    error: str | None = None
    descriptor: ArtifactDescriptor | None = None
    required_tool: str | None = None

    def as_dict(self) -> dict:
        payload = {
            "kind": self.kind,
            "blocks": [
                {"locator": block.locator, "text": block.text}
                for block in self.blocks
            ],
            "visual_verified": self.visual_verified,
            "error": self.error,
            "required_tool": self.required_tool,
        }
        if self.descriptor is not None:
            payload["descriptor"] = self.descriptor.as_dict()
        return payload


def parse_artifact(path: Path) -> ParsedArtifactBlocks:
    target = Path(path)
    descriptor = descriptor_for_path(target)
    suffix = target.suffix.casefold()
    if descriptor.kind == "document":
        try:
            document = Document(target)
            blocks = [
                ArtifactBlock(f"paragraph:{i + 1}", p.text)
                for i, p in enumerate(document.paragraphs)
                if p.text
            ]
            for table_index, table in enumerate(document.tables, start=1):
                for row_index, row in enumerate(table.rows, start=1):
                    blocks.append(
                        ArtifactBlock(
                            f"table:{table_index}:row:{row_index}",
                            "\t".join(cell.text for cell in row.cells),
                        )
                    )
            return ParsedArtifactBlocks(
                "docx", tuple(blocks), descriptor=descriptor, required_tool="inspect_document"
            )
        except Exception as exc:
            return ParsedArtifactBlocks(
                "unsupported",
                (),
                error=f"document extraction unavailable: {exc}",
                descriptor=descriptor,
                required_tool="inspect_document",
            )
    if descriptor.kind == "spreadsheet":
        try:
            workbook = load_workbook(target, read_only=True, data_only=True)
            blocks = []
            try:
                for sheet in workbook.worksheets:
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                        values = ["" if value is None else str(value) for value in row]
                        if any(values):
                            blocks.append(
                                ArtifactBlock(
                                    f"sheet:{sheet.title}:row:{row_index}",
                                    "\t".join(values),
                                )
                            )
            finally:
                workbook.close()
            return ParsedArtifactBlocks(
                "xlsx", tuple(blocks), descriptor=descriptor, required_tool="inspect_document"
            )
        except Exception as exc:
            return ParsedArtifactBlocks(
                "unsupported",
                (),
                error=f"spreadsheet extraction unavailable: {exc}",
                descriptor=descriptor,
                required_tool="inspect_document",
            )
    if descriptor.kind == "pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(target))
            blocks = tuple(
                ArtifactBlock(f"page:{i + 1}", page.extract_text() or "")
                for i, page in enumerate(reader.pages)
            )
            return ParsedArtifactBlocks(
                "pdf", blocks, descriptor=descriptor, required_tool="inspect_document"
            )
        except Exception as exc:
            return ParsedArtifactBlocks(
                "unsupported",
                (),
                error=f"PDF extraction unavailable: {exc}",
                descriptor=descriptor,
                required_tool="inspect_document",
            )
    if descriptor.kind == "image":
        try:
            with Image.open(target) as image:
                text = (
                    f"format={image.format}; width={image.width}; "
                    f"height={image.height}; mode={image.mode}"
                )
            return ParsedArtifactBlocks(
                "image_metadata",
                (ArtifactBlock("metadata", text),),
                visual_verified=False,
                descriptor=descriptor,
                required_tool="inspect_image",
            )
        except Exception as exc:
            return ParsedArtifactBlocks(
                "unsupported",
                (),
                error=f"image metadata unavailable: {exc}",
                descriptor=descriptor,
                required_tool="inspect_image",
            )
    if descriptor.kind == "text":
        # ``descriptor_for_path`` has already validated UTF-8.  Decode only
        # after format routing, never as a binary probe.
        return ParsedArtifactBlocks(
            "text",
            (ArtifactBlock("text", target.read_bytes().decode("utf-8")),),
            descriptor=descriptor,
        )
    required = descriptor.required_tool or "manual_review"
    return ParsedArtifactBlocks(
        "unsupported",
        (),
        error=(
            f"unsupported binary format: {descriptor.media_type}; "
            f"use {required}"
        ),
        descriptor=descriptor,
        required_tool=required,
    )
