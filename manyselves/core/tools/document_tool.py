"""Local, dependency-light document inspection tools."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from docx import Document

from ..artifacts import parse_artifact
from ..access_policy import reject_forbidden_agent_document
from .registry import Tool


class InspectDocumentTool(Tool):
    """Inspect common project documents without requiring MinerU."""

    name = "inspect_document"
    description = (
        "Inspect a project-local text, DOCX, XLSX, or text-based PDF without "
        "changing it or requiring MinerU. DOCX results include headings, tables, "
        "and image positions."
    )

    def __init__(
        self,
        workspace: Path,
        *,
        one_shot: bool = False,
        allow_template_distiller_source: bool = False,
        required_path: str | None = None,
        required_max_chars: int | None = None,
        cache_ref: str | None = None,
    ):
        self.workspace = Path(workspace).resolve()
        self.one_shot = one_shot
        self.allow_template_distiller_source = allow_template_distiller_source
        self.required_path = required_path
        self.required_max_chars = required_max_chars
        self.cache_ref = cache_ref
        self._used = False

    async def __call__(self, path: str, max_chars: int = 30000) -> dict:
        """Inspect one project document.

        Args:
            path: Project-relative file path.
            max_chars: Maximum returned text characters.
        """
        target = (self.workspace / path).resolve()
        if not target.is_relative_to(self.workspace) or not target.is_file():
            raise ValueError("document must be a file inside the project")
        relative = target.relative_to(self.workspace).as_posix()
        if self.required_path is not None and relative != self.required_path:
            raise RuntimeError(
                "TEMPLATE_INSPECTION_PATH_MISMATCH: "
                f"inspect exactly {self.required_path}; got {relative}"
            )
        if (
            self.required_max_chars is not None
            and max_chars != self.required_max_chars
        ):
            raise RuntimeError(
                "TEMPLATE_INSPECTION_LIMIT_MISMATCH: "
                f"max_chars must equal {self.required_max_chars}; got {max_chars}"
            )
        reject_forbidden_agent_document(
            target, allow_template_distiller=self.allow_template_distiller_source
        )
        if self.one_shot and self._used:
            raise RuntimeError(
                "TEMPLATE_ALREADY_INSPECTED: reuse the first inspection result and submit the Skill"
            )
        source_sha256 = hashlib.sha256(target.read_bytes()).hexdigest()
        cached = self._load_cache(
            relative=relative,
            source_sha256=source_sha256,
            max_chars=max_chars,
        )
        if cached is not None:
            self._used = True
            return cached
        parsed = parse_artifact(target)
        if self.required_path is not None and parsed.error:
            raise RuntimeError(
                "TEMPLATE_INSPECTION_FAILED: "
                f"the contracted template could not be parsed: {parsed.error}"
            )
        text = "\n".join(f"[{block.locator}] {block.text}" for block in parsed.blocks)
        result = {
            "path": relative,
            "kind": parsed.kind,
            "text": text[:max_chars],
            "truncated": len(text) > max_chars,
            "visual_verified": parsed.visual_verified,
            "error": parsed.error,
        }
        if target.suffix.casefold() == ".docx":
            document = Document(target)
            headings = [
                {
                    "paragraph": index + 1,
                    "style": paragraph.style.name,
                    "text": paragraph.text[:300],
                }
                for index, paragraph in enumerate(document.paragraphs)
                if paragraph.text.strip()
                and (
                    paragraph.style.name.casefold().startswith("heading")
                    or paragraph.style.name.startswith("标题")
                )
            ]
            image_paragraphs = [
                {
                    "paragraph": index + 1,
                    "style": paragraph.style.name,
                    "previous_text": document.paragraphs[index - 1].text[:300]
                    if index
                    else "",
                    "paragraph_text": paragraph.text[:300],
                    "next_text": document.paragraphs[index + 1].text[:300]
                    if index + 1 < len(document.paragraphs)
                    else "",
                    "image_count": len(paragraph._p.xpath(".//w:drawing | .//w:pict")),
                }
                for index, paragraph in enumerate(document.paragraphs)
                if paragraph._p.xpath(".//w:drawing | .//w:pict")
            ]
            tables = []
            table_image_count = 0
            for index, table in enumerate(document.tables, start=1):
                image_count = len(table._tbl.xpath(".//w:drawing | .//w:pict"))
                table_image_count += image_count
                tables.append(
                    {
                        "table": index,
                        "rows": len(table.rows),
                        "columns": len(table.columns),
                        "headers": [cell.text[:160] for cell in table.rows[0].cells]
                        if table.rows
                        else [],
                        "sample_rows": [
                            [cell.text[:240] for cell in row.cells]
                            for row in table.rows[:5]
                        ],
                        "image_count": image_count,
                    }
                )
            result["structure"] = {
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
                "image_count": sum(item["image_count"] for item in image_paragraphs)
                + table_image_count,
                "images_in_tables": table_image_count,
                "headings": headings,
                "image_paragraphs": image_paragraphs,
                "tables": tables,
            }
        self._write_cache(
            relative=relative,
            source_sha256=source_sha256,
            max_chars=max_chars,
            result=result,
        )
        self._used = True
        return result

    def _cache_path(self) -> Path | None:
        if self.cache_ref is None:
            return None
        target = (self.workspace / self.cache_ref).resolve()
        if not target.is_relative_to(self.workspace):
            raise ValueError("document inspection cache must stay inside the project")
        return target

    def load_cached_result(
        self,
        path: str,
        *,
        max_chars: int,
    ) -> dict[str, Any] | None:
        """Validate and return the durable result for one exact source."""

        target = (self.workspace / path).resolve()
        if not target.is_relative_to(self.workspace) or not target.is_file():
            raise ValueError("document must be a file inside the project")
        relative = target.relative_to(self.workspace).as_posix()
        return self._load_cache(
            relative=relative,
            source_sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
            max_chars=max_chars,
        )

    def _load_cache(
        self,
        *,
        relative: str,
        source_sha256: str,
        max_chars: int,
    ) -> dict[str, Any] | None:
        cache_path = self._cache_path()
        if cache_path is None or not cache_path.is_file():
            return None
        try:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuntimeError(
                "TEMPLATE_INSPECTION_CACHE_INVALID: cached inspection is unreadable"
            ) from exc
        expected = {
            "source_path": relative,
            "source_sha256": source_sha256,
            "max_chars": max_chars,
        }
        actual = {key: cached.get(key) for key in expected}
        result = cached.get("result")
        result_sha256 = (
            self._result_sha256(result)
            if isinstance(result, dict)
            else None
        )
        if (
            actual != expected
            or not isinstance(result, dict)
            or cached.get("result_sha256") != result_sha256
        ):
            raise RuntimeError(
                "TEMPLATE_INSPECTION_CACHE_MISMATCH: "
                "cached inspection does not match the current template contract"
            )
        return result

    def _write_cache(
        self,
        *,
        relative: str,
        source_sha256: str,
        max_chars: int,
        result: dict[str, Any],
    ) -> None:
        cache_path = self._cache_path()
        if cache_path is None:
            return
        payload = {
            "source_path": relative,
            "source_sha256": source_sha256,
            "max_chars": max_chars,
            "result_sha256": self._result_sha256(result),
            "result": result,
        }
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = cache_path.with_name(f".{cache_path.name}.tmp")
        temporary.write_text(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2)
            + "\n",
            encoding="utf-8",
        )
        temporary.replace(cache_path)

    @staticmethod
    def _result_sha256(result: dict[str, Any]) -> str:
        canonical = json.dumps(
            result,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        return hashlib.sha256(canonical).hexdigest()
