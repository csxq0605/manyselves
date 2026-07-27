"""Repository-packaged approved-Markdown to DOCX core."""

import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

DEFAULT_FILENAME = "配电安全专家咨询报告.docx"
_HEADING = re.compile(r"^(#{1,6})\s+(.+)$")


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_separator(line: str) -> bool:
    cells = _cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _append_markdown(document: Document, report_text: str) -> None:
    style_names = {style.name for style in document.styles}
    lines = report_text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line in {"---", "***", "___"}:
            index += 1
            continue
        heading = _HEADING.match(line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and _is_separator(lines[index + 1]):
            headers = _cells(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = _cells(lines[index])
                if len(row) == len(headers):
                    rows.append(row)
                index += 1
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for column, value in enumerate(headers):
                table.cell(0, column).text = value
            for row in rows:
                cells = table.add_row().cells
                for column, value in enumerate(row):
                    cells[column].text = value
            continue
        if line.startswith("- "):
            text = line[2:].strip()
            if "List Bullet" in style_names:
                document.add_paragraph(text, style="List Bullet")
            else:
                document.add_paragraph(f"• {text}")
        elif re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            if "List Number" in style_names:
                document.add_paragraph(text, style="List Number")
            else:
                document.add_paragraph(line)
        else:
            document.add_paragraph(line)
        index += 1


def _clear_template_body(document: Document) -> None:
    """Keep template styles/sections/headers while removing its sample report body."""

    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _expected_markdown_fragments(report_text: str) -> list[str]:
    """Return the visible text fragments the narrow Markdown renderer must preserve."""

    fragments: list[str] = []
    lines = report_text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line in {"---", "***", "___"}:
            index += 1
            continue
        heading = _HEADING.match(line)
        if heading:
            heading_text = heading.group(2).strip()
            numbered = re.match(r"^(\d+(?:\.\d+)*)(?:\.)?\s+", heading_text)
            # V2 owns titles for known numbered sections. Presence of the
            # section number proves the structural slot survived; body prose is
            # still checked verbatim by the remaining fragments.
            fragments.append(numbered.group(1) if numbered else heading_text)
            index += 1
            continue
        numbered_heading = re.match(r"^(\d+(?:\.\d+)+)(?:\.)?\s+.+$", line)
        if numbered_heading:
            fragments.append(numbered_heading.group(1))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and _is_separator(lines[index + 1]):
            fragments.extend(cell for cell in _cells(line) if cell)
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                fragments.extend(cell for cell in _cells(lines[index]) if cell)
                index += 1
            continue
        if line.startswith("- "):
            fragments.append(line[2:].strip())
        elif re.match(r"^\d+\.\s+", line):
            fragments.append(re.sub(r"^\d+\.\s+", "", line))
        else:
            fragments.append(line)
        index += 1
    return fragments


def verify_rendered_markdown(output_path: Path, report_text: str) -> None:
    """Fail unless every source Markdown fragment survives as visible DOCX text.

    Word may split one Markdown line across paragraphs, list items, runs, or
    table cells, and the V2 template may retain a leading heading marker for an
    unrecognised custom heading.  Compare normalized visible text as one corpus
    so those presentation-only differences do not become false omissions.  We
    deliberately do not remove citations, punctuation, or prose: content loss
    must still fail the workflow.
    """

    try:
        rendered = Document(output_path)
    except Exception as exc:
        raise ValueError("rendered file is not Word/WPS-openable") from exc
    def semantic_text(value: str) -> str:
        value = value.strip()
        # Strip inline Markdown before recognizing block/list prefixes.  Source
        # lines commonly start with bold numbered labels (``**1. ...**``),
        # while their DOCX paragraphs contain only ``1. ...``.  Doing this
        # later leaves the source-side number behind and falsely reports that
        # faithfully rendered prose was omitted.
        value = value.replace("**", "").replace("__", "").replace("`", "")
        value = re.sub(r"^#{1,6}\s*", "", value)
        value = re.sub(r"^>\s*", "", value)
        value = re.sub(r"^[•·]\s*", "", value)
        value = re.sub(r"^\d+\.\s+", "", value)
        value = re.sub(r"^(\d+(?:\.\d+)+)\.\s+", r"\1 ", value)
        value = re.sub(r"【([^】]+)】[:：]?", r"\1：", value)
        value = re.sub(r"\s*([：:])\s*", r"\1", value)
        return re.sub(r"\s+", "", value).strip()

    visible = [semantic_text(paragraph.text) for paragraph in rendered.paragraphs]
    visible.extend(
        semantic_text(cell.text)
        for table in rendered.tables
        for row in table.rows
        for cell in row.cells
    )
    visible_corpus = semantic_text("\n".join(value for value in visible if value))
    missing = [
        fragment
        for fragment in _expected_markdown_fragments(report_text)
        if semantic_text(fragment) not in visible_corpus
    ]
    if missing:
        preview = ", ".join(repr(fragment[:80]) for fragment in missing[:5])
        raise ValueError(f"rendered DOCX omitted approved Markdown content: {preview}")


class PackagedDocxCore:
    """Render approved report prose without external runtime files."""

    def __init__(self, template_path: Path):
        self.template_path = Path(template_path)

    def render_approved_prose(
        self,
        report_text: str,
        *,
        filename: str | None = None,
        report_model: Any = None,
    ) -> tuple[str, bytes]:
        if report_model is not None:
            raise ValueError(
                "structured-model prose generation is disabled; pass approved report prose only"
            )
        if not self.template_path.is_file():
            raise FileNotFoundError(f"packaged DOCX template not found: {self.template_path}")
        document = Document(self.template_path)
        _clear_template_body(document)
        _append_markdown(document, report_text)
        buffer = io.BytesIO()
        document.save(buffer)
        return filename or DEFAULT_FILENAME, buffer.getvalue()
