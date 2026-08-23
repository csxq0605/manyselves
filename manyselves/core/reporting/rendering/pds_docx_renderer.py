"""Render an approved five-module agent report through the handoff DOCX core."""

from __future__ import annotations

import hashlib
import io
import re
import tempfile
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image, ImageOps
from pydantic import Field, field_validator, model_validator

from manyselves.capabilities.distribution_reporting.domain.taxonomy import (
    REPORT_TAXONOMY,
    resolve_submodule,
)

from ..agentic_models import StrictModel
from ..claim_ledger import ClaimLedger
from ..models import REPORT_MODULE_IDS, SpecialTopicPlan
from ..report_markdown import (
    CanonicalMarkdownTable,
    CanonicalReportContent,
    compose_canonical_markdown,
    markdown_table,
    strip_leading_module_heading,
)
from .handoff_docx import HandoffDocxCore
from .packaged_docx import _expected_markdown_fragments

_CITATION_TOKEN = re.compile(r"\[\[CITE:(\d+)\]\]")
_PHOTO_TOKEN = re.compile(r"^\[\[PHOTO:([^]]+)\]\]$")
_FIXED_DOCX_TIME = datetime(2000, 1, 1, tzinfo=UTC)
_PHOTOS_PER_SUMMARY_TABLE = 2
_PHOTO_MAX_WIDTH_CM = 8.2
_PHOTO_MAX_HEIGHT_CM = 7.0


class ReportTable(StrictModel):
    title: str = Field(min_length=1)
    headers: list[str] = Field(min_length=1)
    rows: list[list[str]] = Field(default_factory=list)
    source_ids: list[str] = Field(default_factory=list)
    claim_ids: list[str] = Field(default_factory=list)

    @model_validator(mode="after")
    def rows_match_headers(self) -> "ReportTable":
        if not self.source_ids:
            raise ValueError("table requires at least one source id")
        invalid = [index for index, row in enumerate(self.rows) if len(row) != len(self.headers)]
        if invalid:
            raise ValueError(f"table rows do not match header width: {invalid}")
        return self


class ReportPhoto(StrictModel):
    id: str = Field(min_length=1)
    path: Path
    caption: str = Field(min_length=1)
    source_id: str = Field(pattern=r"^E-")
    claim_ids: list[str] = Field(default_factory=list)
    submodule_id: str | None = None


class ApprovedReport(StrictModel):
    title: str = Field(min_length=1)
    assessment_background: str = Field(min_length=1)
    findings_overview: str = Field(min_length=1)
    regional_executive_summary: str = Field(min_length=1)
    module_narratives: dict[str, str]
    risk_panorama: str = Field(min_length=1)
    dimension_risk_analysis: str = Field(min_length=1)
    data_gap_analysis: str = Field(min_length=1)
    improvement_action_plan: str = Field(min_length=1)
    special_topic_plan: SpecialTopicPlan | None = None
    special_topic_analysis: str | None = Field(default=None, min_length=1)
    ledger: ClaimLedger
    tables: list[ReportTable] = Field(default_factory=list)
    photos: list[ReportPhoto] = Field(default_factory=list)

    @field_validator("module_narratives")
    @classmethod
    def has_complete_five_modules(cls, narratives: dict[str, str]) -> dict[str, str]:
        if set(narratives) != set(REPORT_MODULE_IDS):
            raise ValueError("approved report requires exactly modules 2.1-2.5")
        if any(not text.strip() for text in narratives.values()):
            raise ValueError("module narratives cannot be blank")
        return narratives

    @model_validator(mode="after")
    def assets_are_traceable(self) -> "ApprovedReport":
        if (self.special_topic_plan is None) != (self.special_topic_analysis is None):
            raise ValueError(
                "special_topic_plan and special_topic_analysis must either both be present "
                "or both be absent"
            )
        if self.special_topic_plan is not None and self.special_topic_analysis is not None:
            self.special_topic_plan.validate_analysis(self.special_topic_analysis)
        claim_modules = {claim.module_id for claim in self.ledger.claims}
        if claim_modules != set(REPORT_MODULE_IDS):
            raise ValueError("approved report claim ledger requires exactly modules 2.1-2.5")
        source_ids = {source.id for source in self.ledger.sources}
        claim_ids = {claim.id for claim in self.ledger.claims}
        claims_by_id = {claim.id: claim for claim in self.ledger.claims}
        for photo in self.photos:
            if photo.source_id not in source_ids:
                raise ValueError(f"photo {photo.id} references unknown project source")
            unknown_claims = sorted(set(photo.claim_ids) - claim_ids)
            if unknown_claims:
                raise ValueError(f"photo {photo.id} references unknown claims: {unknown_claims}")
            if photo.submodule_id is not None:
                if photo.submodule_id not in {
                    submodule_id
                    for module in REPORT_TAXONOMY.values()
                    for submodule_id in module.submodules
                }:
                    raise ValueError(
                        f"photo {photo.id} references unknown submodule {photo.submodule_id}"
                    )
                if photo.claim_ids and not any(
                    claims_by_id[claim_id].submodule_id == photo.submodule_id
                    for claim_id in photo.claim_ids
                ):
                    raise ValueError(
                        f"photo {photo.id} submodule does not match its linked claims"
                    )
        for table in self.tables:
            unknown_sources = sorted(set(table.source_ids) - source_ids)
            if unknown_sources:
                raise ValueError(
                    f"table {table.title} references unknown sources: {unknown_sources}"
                )
            unknown_claims = sorted(set(table.claim_ids) - claim_ids)
            if unknown_claims:
                raise ValueError(f"table {table.title} references unknown claims: {unknown_claims}")
        return self


class PdsRenderResult(StrictModel):
    output_path: Path
    output_sha256: str
    citation_count: int
    table_count: int
    photo_count: int
    handoff_core_used: bool
    protected_prose_verified: bool


class PdsDocxRenderer:
    def __init__(self, handoff_core: HandoffDocxCore):
        self.handoff_core = handoff_core

    def render(
        self,
        report: ApprovedReport,
        output_path: Path,
        *,
        approved_markdown: str | None = None,
        before_publish: Callable[[], None] | None = None,
    ) -> PdsRenderResult:
        """Render approved prose only and validate semantic preservation."""

        cited_markdown = (
            approved_markdown
            if approved_markdown is not None
            else report.ledger.bind_citations(self._compose_markdown(report))
        )
        _, raw_docx = self.handoff_core.render_approved_prose(
            cited_markdown,
            filename=Path(output_path).name,
            report_model=None,
        )
        document = Document(io.BytesIO(raw_docx))
        self._ensure_title(document, report.title)
        self._materialize_citations(document)
        self._materialize_photos(document, report)
        document.core_properties.created = _FIXED_DOCX_TIME
        document.core_properties.modified = _FIXED_DOCX_TIME
        document.core_properties.last_printed = _FIXED_DOCX_TIME

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        buffer = io.BytesIO()
        document.save(buffer)
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{output_path.stem}-",
                suffix=".docx",
                dir=output_path.parent,
                delete=False,
            ) as temporary:
                temporary.write(self._canonical_docx(buffer.getvalue()))
                temporary_path = Path(temporary.name)
            self._verify_output(temporary_path, report)
            if before_publish is not None:
                before_publish()
            temporary_path.replace(output_path)
            temporary_path = None
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        return PdsRenderResult(
            output_path=output_path,
            output_sha256=hashlib.sha256(output_path.read_bytes()).hexdigest(),
            citation_count=len(report.ledger.build_citation_plan().entries),
            table_count=len(report.tables),
            photo_count=len(report.photos),
            handoff_core_used=True,
            protected_prose_verified=True,
        )

    @staticmethod
    def _compose_markdown(report: ApprovedReport) -> str:
        unresolved_claims = [
            claim.text for claim in report.ledger.claims if claim.unresolved
        ]
        module_narratives: dict[str, str] = {}
        for module_id in REPORT_MODULE_IDS:
            module_photos = [
                photo
                for photo in report.photos
                if (
                    photo.submodule_id is not None
                    and resolve_submodule(photo.submodule_id).module_id == module_id
                )
            ]
            module_narratives[module_id] = PdsDocxRenderer._place_photo_tokens(
                strip_leading_module_heading(
                    report.module_narratives[module_id],
                    module_id,
                ),
                module_photos,
            )
        return compose_canonical_markdown(
            CanonicalReportContent(
                title=report.title,
                assessment_background=report.assessment_background,
                findings_overview=report.findings_overview,
                regional_executive_summary=report.regional_executive_summary,
                module_narratives=module_narratives,
                risk_panorama=report.risk_panorama,
                dimension_risk_analysis=report.dimension_risk_analysis,
                data_gap_analysis=report.data_gap_analysis
                or "\n".join(f"- {text}" for text in unresolved_claims),
                improvement_action_plan=report.improvement_action_plan,
                special_topic_plan=report.special_topic_plan,
                special_topic_analysis=report.special_topic_analysis,
                tables=[
                    CanonicalMarkdownTable(
                        title=table.title,
                        headers=table.headers,
                        rows=table.rows,
                        source_ids=table.source_ids,
                    )
                    for table in report.tables
                ],
            )
        )

    @staticmethod
    def _strip_leading_module_heading(narrative: str, module_id: str) -> str:
        """Remove the canonical module heading even after an editor transition."""

        return strip_leading_module_heading(narrative, module_id)

    @staticmethod
    def _place_photo_tokens(narrative: str, photos: list[ReportPhoto]) -> str:
        """Place photo evidence beside its linked submodule, not at module end."""

        if not photos:
            return narrative
        by_submodule: dict[str, list[ReportPhoto]] = {}
        fallback: list[ReportPhoto] = []
        for photo in photos:
            if photo.submodule_id:
                by_submodule.setdefault(photo.submodule_id, []).append(photo)
            else:
                fallback.append(photo)
        lines = narrative.splitlines()
        output: list[str] = []
        active_submodule: str | None = None
        inserted: set[str] = set()

        def append_tokens(submodule_id: str | None) -> None:
            if not submodule_id or submodule_id in inserted:
                return
            selected = by_submodule.get(submodule_id, [])
            if selected:
                output.extend(["", *[f"[[PHOTO:{photo.id}]]" for photo in selected], ""])
                inserted.add(submodule_id)

        for line in lines:
            heading = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
            if heading:
                append_tokens(active_submodule)
                active_submodule = next(
                    (
                        submodule_id
                        for submodule_id in by_submodule
                        if submodule_id in heading.group(1)
                    ),
                    None,
                )
            output.append(line)
        append_tokens(active_submodule)
        for submodule_id, selected in by_submodule.items():
            if submodule_id not in inserted:
                output.extend(["", *[f"[[PHOTO:{photo.id}]]" for photo in selected]])
        output.extend(["", *[f"[[PHOTO:{photo.id}]]" for photo in fallback]])
        return "\n".join(output).strip()

    @staticmethod
    def _markdown_table(table: ReportTable) -> str:
        return markdown_table(
            CanonicalMarkdownTable(
                title=table.title,
                headers=table.headers,
                rows=table.rows,
                source_ids=table.source_ids,
            )
        )

    @staticmethod
    def _materialize_citations(document: Document) -> None:
        paragraphs = list(document.paragraphs)
        paragraphs.extend(
            p
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        )
        for paragraph in paragraphs:
            text = paragraph.text
            if not _CITATION_TOKEN.search(text):
                continue
            pieces = _CITATION_TOKEN.split(text)
            paragraph.clear()
            for index, piece in enumerate(pieces):
                if not piece:
                    continue
                run = paragraph.add_run(piece)
                if index % 2 == 1:
                    run.font.superscript = True

    @staticmethod
    def _ensure_title(document: Document, title: str) -> None:
        if any(paragraph.text.strip() == title for paragraph in document.paragraphs):
            return
        paragraph = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
        style_names = {style.name for style in document.styles}
        style = "Title" if "Title" in style_names else "Normal"
        title_paragraph = paragraph.insert_paragraph_before(title, style=style)
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title_paragraph.runs:
            run.bold = True
            run.font.size = Pt(18)

    @staticmethod
    def _materialize_photos(document: Document, report: ApprovedReport) -> None:
        photos = {photo.id: photo for photo in report.photos}
        paragraphs = list(document.paragraphs)
        index = 0
        while index < len(paragraphs):
            match = _PHOTO_TOKEN.match(paragraphs[index].text.strip())
            if match is None:
                index += 1
                continue
            marker_paragraphs = [paragraphs[index]]
            selected = [photos[match.group(1)]]
            scan = index + 1
            while scan < len(paragraphs):
                text = paragraphs[scan].text.strip()
                next_match = _PHOTO_TOKEN.match(text)
                if next_match is not None:
                    marker_paragraphs.append(paragraphs[scan])
                    selected.append(photos[next_match.group(1)])
                    scan += 1
                    continue
                if not text:
                    marker_paragraphs.append(paragraphs[scan])
                    scan += 1
                    continue
                break
            for photo in selected:
                if not photo.path.is_file():
                    raise FileNotFoundError(f"report photo not found: {photo.path}")
            for batch_start in range(0, len(selected), _PHOTOS_PER_SUMMARY_TABLE):
                batch = selected[
                    batch_start : batch_start + _PHOTOS_PER_SUMMARY_TABLE
                ]
                if batch_start:
                    continuation = document.add_paragraph()
                    continuation.add_run("原表图证汇总（续）").bold = True
                    marker_paragraphs[0]._p.addprevious(continuation._p)
                table = document.add_table(rows=len(batch) + 1, cols=2)
                table.autofit = False
                if "Table Grid" in [style.name for style in document.styles]:
                    table.style = "Table Grid"
                headers = table.rows[0].cells
                headers[0].text = "原表对应内容"
                headers[1].text = "图证"
                for header in headers:
                    header.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for run in header.paragraphs[0].runs:
                        run.bold = True
                    header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for photo_index, photo in enumerate(batch, start=1):
                    detail_cell, image_cell = table.rows[photo_index].cells
                    detail_cell.width = Cm(6.0)
                    image_cell.width = Cm(9.0)
                    detail_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    image_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    detail = detail_cell.paragraphs[0]
                    detail.add_run(photo.caption).bold = True
                    detail.add_run(f"\n来源：{photo.source_id}")
                    image = image_cell.paragraphs[0]
                    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    image.add_run().add_picture(
                        str(photo.path),
                        **PdsDocxRenderer._photo_fit_dimensions(photo.path),
                    )
                    image.add_run("\n原表图证")
                marker_paragraphs[0]._p.addprevious(table._tbl)
            for paragraph in marker_paragraphs:
                parent = paragraph._p.getparent()
                if parent is not None:
                    parent.remove(paragraph._p)
            index = scan

    @staticmethod
    def _photo_fit_dimensions(path: Path) -> dict[str, object]:
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source)
            pixel_width, pixel_height = image.size
        if pixel_width <= 0 or pixel_height <= 0:
            return {"width": Cm(_PHOTO_MAX_WIDTH_CM)}
        aspect = pixel_width / pixel_height
        box_aspect = _PHOTO_MAX_WIDTH_CM / _PHOTO_MAX_HEIGHT_CM
        if aspect >= box_aspect:
            return {"width": Cm(_PHOTO_MAX_WIDTH_CM)}
        return {"height": Cm(_PHOTO_MAX_HEIGHT_CM)}

    @staticmethod
    def _canonical_docx(data: bytes) -> bytes:
        source = io.BytesIO(data)
        target = io.BytesIO()
        with (
            zipfile.ZipFile(source, "r") as input_zip,
            zipfile.ZipFile(
                target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
            ) as output_zip,
        ):
            for name in sorted(input_zip.namelist()):
                info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = input_zip.getinfo(name).external_attr
                output_zip.writestr(info, input_zip.read(name))
        return target.getvalue()

    @staticmethod
    def _verify_output(output_path: Path, report: ApprovedReport) -> None:
        try:
            rendered = Document(output_path)
        except Exception as exc:
            raise ValueError("rendered file is not Word/WPS-openable") from exc
        def semantic_text(value: str) -> str:
            value = value.strip()
            # Remove inline Markdown before recognizing numbered/list prefixes.
            # A protected line such as ``**1. 风险标题**`` is rendered as a
            # numbered paragraph without the literal ``1.``.  Stripping the
            # Markdown later leaves the source-side number behind and produces
            # a false omission.
            value = value.replace("**", "").replace("__", "").replace("`", "")
            value = re.sub(r"^#{1,6}\s*", "", value)
            value = re.sub(r"^>\s*", "", value)
            value = re.sub(r"^[-*+•·]\s+", "", value)
            value = re.sub(r"^\d+\.\s+", "", value)
            value = re.sub(r"^(\d+(?:\.\d+)+)\.\s+", r"\1 ", value)
            value = re.sub(r"\[\[CLAIM:C-[^\]\s]+\]\]", "", value)
            value = re.sub(r"【([^】]+)】[:：]?", r"\1：", value)
            value = re.sub(r"\s*([：:])\s*", r"\1", value)
            return re.sub(r"\s+", "", value)

        visible = [
            "".join(run.text for run in paragraph.runs if run.font.superscript is not True)
            for paragraph in rendered.paragraphs
        ]
        visible.extend(
            cell.text
            for table in rendered.tables
            for row in table.rows
            for cell in row.cells
        )
        semantic_visible = semantic_text("\n".join(visible))
        text = "\n".join(visible)
        protected = [
            report.assessment_background,
            report.findings_overview,
            report.risk_panorama,
            report.dimension_risk_analysis,
            report.data_gap_analysis,
            report.improvement_action_plan,
            *report.module_narratives.values(),
        ]
        if report.special_topic_analysis is not None:
            protected.append(report.special_topic_analysis)
        missing = [
            fragment
            for value in protected
            for fragment in _expected_markdown_fragments(value)
            if semantic_text(fragment) not in semantic_visible
        ]
        if missing:
            preview = ", ".join(repr(fragment[:120]) for fragment in missing[:8])
            raise ValueError(
                "rendered DOCX changed or omitted protected Agent/Chief Editor prose: "
                + preview
            )
        if report.title not in text:
            raise ValueError("rendered DOCX is missing the approved title")
        if _CITATION_TOKEN.search(text):
            raise ValueError("rendered DOCX contains unresolved citation tokens")
        if any(_PHOTO_TOKEN.match(paragraph.text.strip()) for paragraph in rendered.paragraphs):
            raise ValueError("rendered DOCX contains unresolved photo tokens")
