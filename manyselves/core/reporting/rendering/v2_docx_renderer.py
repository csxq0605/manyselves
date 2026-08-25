from __future__ import annotations

import hashlib
import io
import json
import re
import tempfile
from contextvars import ContextVar
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from manyselves.core.reporting.rendering.markdown_content import (
    clean_bullet_item_text as shared_clean_bullet_item_text,
)
from manyselves.core.reporting.rendering.markdown_content import (
    clean_numbered_item_text as shared_clean_numbered_item_text,
)
from manyselves.core.reporting.rendering.markdown_content import (
    is_bullet_line as shared_is_bullet_line,
)
from manyselves.core.reporting.rendering.markdown_content import (
    is_numbered_list_line as shared_is_numbered_list_line,
)
from manyselves.core.reporting.rendering.markdown_content import (
    is_report_title as shared_is_report_title,
)
from manyselves.core.reporting.rendering.markdown_content import (
    is_table_line as shared_is_table_line,
)
from manyselves.core.reporting.rendering.markdown_content import (
    parse_markdown_blocks,
)
from manyselves.core.reporting.taxonomy import REPORT_TAXONOMY

try:
    from PIL import Image, ImageOps
except Exception:
    Image = None
    ImageOps = None


DEFAULT_FILENAME = "配电安全专家咨询报告.docx"
TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "report_template.docx"

PHOTO_MAX_PIXELS = (900, 900)
PHOTO_JPEG_QUALITY = 68
PHOTO_INLINE_WIDTH_CM = 3.0
PHOTO_CELL_MAX_WIDTH_CM = 3.0
PHOTO_CELL_MAX_HEIGHT_CM = 3.0
MAX_REPORT_PHOTO_INSERTIONS = 24
PHOTO_CACHE_DIR = Path(tempfile.gettempdir()) / "pds_report_docx_photo_cache"
PHOTO_INSERTION_COUNT = 0

DEFAULT_REPORT_POLICY: dict[str, Any] = {
    "table_policy": {
        "missing_photo_text": "待补充照片",
        "broken_photo_ref_text": "图片引用缺失，待补充照片",
        "photo_width_cm": PHOTO_INLINE_WIDTH_CM,
        "photo_height_cm": 3.0,
        "show_only_problem_rows": False,
        "max_rows_per_table": 30,
    },
    "chapter_3_policy": {
        "top_risk_count": 3,
        "include_data_gap_analysis": True,
    },
    "suggestion_policy": {
        "group_by_time": True,
        "include_expected_effect": True,
    },
    "risk_policy": {
        "max_words": 0,
    },
    "conclusion_policy": {
        "max_words_per_item": 0,
    },
}
REPORT_POLICY_CONTEXT: ContextVar[dict[str, Any]] = ContextVar(
    "REPORT_POLICY_CONTEXT",
    default=DEFAULT_REPORT_POLICY,
)

BODY_FONT_SIZE_PT = 11
TABLE_FONT_SIZE_PT = 10
HEADING_STYLE_TOKENS: dict[str, tuple[float, float, float, float]] = {
    "Title": (18, 0, 8, 1.15),
    "Heading 1": (18, 12, 6, 1.2),
    "Heading 2": (16, 10, 5, 1.2),
    "Heading 3": (14, 8, 4, 1.2),
    "Heading 4": (12, 6, 3, 1.15),
}
BODY_LINE_SPACING = 1.3
BODY_SPACE_AFTER_PT = 4
LIST_LEFT_INDENT_TWIPS = 420
LIST_HANGING_INDENT_TWIPS = 240


CHAPTERS = [
    ("1", "配电评估概述"),
    ("2", "评估内容描述"),
    ("3", "结论与建议"),
    ("4", "专项问题分析"),
]

_HEADINGS_BEFORE_MODULES: list[tuple[str, str, int]] = [
    ("1", "配电评估概述", 1),
    ("1.1", "评估背景", 2),
    ("1.2", "健康度总览", 2),
    ("1.3", "各区域执行摘要", 2),
    ("2", "评估内容描述", 1),
]
_MODULE_HEADINGS = [
    heading
    for module in REPORT_TAXONOMY.values()
    for heading in (
        (module.id, module.title, 2),
        *(
            (
                section.id,
                section.title,
                section.id.count(".") + 1,
            )
            for section in module.sections.values()
        ),
    )
]
_HEADINGS_AFTER_MODULES: list[tuple[str, str, int]] = [
    ("3", "结论与建议", 1),
    ("3.1", "风险/问题汇总与概览", 2),
    ("3.1.1", "风险全景图", 3),
    ("3.1.2", "各维度风险分析", 3),
    ("3.1.3", "数据缺口分析", 3),
    ("3.2", "改善行动速查表", 2),
    ("4", "专项问题分析", 1),
]
HEADINGS: list[tuple[str, str, int]] = [
    *_HEADINGS_BEFORE_MODULES,
    *_MODULE_HEADINGS,
    *_HEADINGS_AFTER_MODULES,
]


def _heading_key(text: str) -> str:
    text = str(text).replace("**", "").replace("__", "").strip()
    text = re.sub(r"^#{1,6}\s*", "", text)
    text = re.sub(r"\s+", "", text)
    return re.sub(r"[。．.：:]+", "", text)


HEADING_BY_NUMBER = {number: (title, level) for number, title, level in HEADINGS}
HEADING_BY_TITLE = {_heading_key(title): (number, title, level) for number, title, level in HEADINGS}
HEADING_ORDER = {number: index for index, (number, _, _) in enumerate(HEADINGS)}
DETAIL_SECTION_IDS = [number for number, _, level in HEADINGS if number.startswith("2.") and level >= 3]
INLINE_LABELS = (
    "现状描述",
    "现状分析",
    "问题描述",
    "评估结果",
    "评估发现",
    "证据限定",
    "判断",
    "判定",
    "原因分析",
    "风险等级",
    "资料来源",
    "方法论依据",
    "方法论核查要求",
    "核查方法",
    "判定依据",
    "风险分析",
    "风险与影响",
    "潜在风险",
    "整改建议",
    "建议",
    "后续建议",
    "参考标准",
    "相关标准",
    "依据",
    "结论",
)
EXPANDABLE_LABELS = (
    "方法论核查要求",
    "核查方法",
    "判定依据",
    "风险分析",
    "潜在风险",
    "整改建议",
    "建议",
    "后续建议",
)
TOC_PAGE_HINTS = {
    "1": "4",
    "1.1": "4",
    "1.2": "4",
    "1.3": "5",
    "2": "6",
    "2.1": "6",
    "2.1.1": "6",
    "2.1.2": "8",
    "2.1.3": "9",
    "2.1.4": "10",
    "2.1.5": "10",
    "2.2": "11",
    "2.2.1": "11",
    "2.2.2": "14",
    "2.3": "20",
    "2.3.1": "20",
    "2.3.2": "21",
    "2.3.3": "22",
    "2.4": "24",
    "2.4.1": "24",
    "2.4.2": "25",
    "2.4.3": "29",
    "2.4.4": "32",
    "2.5": "33",
    "2.5.1": "33",
    "2.5.2": "33",
    "2.5.3": "33",
    "2.5.4": "38",
    "2.5.5": "38",
    "2.5.6": "39",
    "2.5.7": "39",
    "3": "40",
    "3.1": "40",
    "3.1.1": "40",
    "3.1.2": "40",
    "3.1.3": "41",
    "3.2": "42",
    "4": "43",
    "4.1": "43",
    "4.2": "45",
    "4.3": "47",
    "4.4": "49",
}


def parse_json_object(value: Any) -> dict[str, Any]:
    if not value:
        return {}
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return {}
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            return {}
        return parsed if isinstance(parsed, dict) else {}
    return {}


def deep_merge_policy(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    merged = dict(base)
    for key, value in (override or {}).items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = deep_merge_policy(merged[key], value)
        else:
            merged[key] = value
    return merged


def build_effective_report_policy(
    active_skill_pack_json: Any = None,
    session_policy_patch_json: Any = None,
    effective_report_policy_json: Any = None,
) -> dict[str, Any]:
    explicit_policy = parse_json_object(effective_report_policy_json)
    if explicit_policy:
        return deep_merge_policy(DEFAULT_REPORT_POLICY, explicit_policy)

    active_pack = parse_json_object(active_skill_pack_json)
    active_policy = parse_json_object(active_pack.get("report_policy") or active_pack.get("report_policy_json"))
    if not active_policy and "report_policy" in active_pack:
        active_policy = active_pack["report_policy"] if isinstance(active_pack["report_policy"], dict) else {}
    session_patch = parse_json_object(session_policy_patch_json)

    policy = deep_merge_policy(DEFAULT_REPORT_POLICY, active_policy)
    return deep_merge_policy(policy, session_patch)


def current_report_policy() -> dict[str, Any]:
    return REPORT_POLICY_CONTEXT.get()


def policy_section(name: str) -> dict[str, Any]:
    value = current_report_policy().get(name) or {}
    return value if isinstance(value, dict) else {}


def policy_int(section: str, key: str, default: int) -> int:
    value = policy_section(section).get(key, default)
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def policy_float(section: str, key: str, default: float) -> float:
    value = policy_section(section).get(key, default)
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def policy_bool(section: str, key: str, default: bool) -> bool:
    value = policy_section(section).get(key, default)
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "y"}
    return bool(value)


def missing_photo_text() -> str:
    return clean_text(policy_section("table_policy").get("missing_photo_text")) or "待补充照片"


def broken_photo_ref_text() -> str:
    return clean_text(policy_section("table_policy").get("broken_photo_ref_text")) or "图片引用缺失，待补充照片"


def truncate_policy_text(text: str, section: str, key: str) -> str:
    limit = policy_int(section, key, 0)
    text = clean_text(text)
    if limit <= 0 or len(text) <= limit:
        return text
    return text[:limit].rstrip("，。；;、 ") + "..."


def render_report_docx(
    report_text: str,
    filename: str | None = None,
    report_model: Any = None,
    active_skill_pack_json: Any = None,
    session_policy_patch_json: Any = None,
    effective_report_policy_json: Any = None,
    section_reasoning_result_json: Any = None,
) -> tuple[str, bytes]:
    reset_photo_insertion_count()
    policy = build_effective_report_policy(
        active_skill_pack_json=active_skill_pack_json,
        session_policy_patch_json=session_policy_patch_json,
        effective_report_policy_json=effective_report_policy_json,
    )
    token = REPORT_POLICY_CONTEXT.set(policy)
    output_filename = filename or DEFAULT_FILENAME
    try:
        if not output_filename.lower().endswith(".docx"):
            output_filename += ".docx"

        model = parse_report_model(report_model)
        if model is not None and not is_structured_report_model(model):
            raise ValueError("report_model_json 不是有效的结构化报告模型，已阻止生成不完整报告")
        has_report_model = model is not None
        if model is None:
            model = {}
        model["_effective_report_policy"] = policy
        model["_section_reasoning_results"] = parse_section_reasoning_results(section_reasoning_result_json)
        report_text = normalize_report_text(report_text)
        if has_report_model:
            report_text = build_report_text_from_model(model, report_text)

        if TEMPLATE_PATH.exists():
            document = Document(TEMPLATE_PATH)
            configure_styles(document)
            replace_template_content(document, report_text, model)
            normalize_static_toc(document)
            enforce_all_table_styles(document)
            compact_report_flow(document)
            mark_fields_for_update(document)
        else:
            document = Document()
            configure_styles(document)
            add_report_content(document, report_text, model)
            enforce_all_table_styles(document)
            compact_report_flow(document)

        enforce_cjk_font_mapping(document)

        buffer = io.BytesIO()
        document.save(buffer)
        return output_filename, buffer.getvalue()
    finally:
        REPORT_POLICY_CONTEXT.reset(token)


def parse_report_model(value: Any) -> dict[str, Any] | None:
    if not value:
        return None
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            return None
        return parsed if isinstance(parsed, dict) else None
    return None


def is_structured_report_model(model: dict[str, Any]) -> bool:
    return (
        isinstance(model.get("issues"), list)
        or isinstance(model.get("sections"), dict)
        or isinstance(model.get("rules_by_section"), dict)
        or isinstance(model.get("status_counts"), dict)
        or isinstance(model.get("action_rows"), list)
    )


def build_report_text_from_model(model: dict[str, Any], fallback_text: str) -> str:
    status_counts = model.get("status_counts") or {}
    issues = [issue for issue in model.get("issues") or [] if isinstance(issue, dict)]
    sections = model.get("sections") or {}
    rules_by_section = model.get("rules_by_section") or {}
    action_rows = model.get("action_rows") or []
    cross_link_map = v2_cross_link_map(issues)
    section_reasoning_results = model.get("_section_reasoning_results") or {}

    lines: list[str] = [
        "# 配电安全专家咨询报告",
        "",
        "## 1. 配电评估概述",
        "",
        "### 1.1 评估背景",
        "",
        chapter1_background_text(model),
        "",
        "### 1.2 健康度总览",
        "",
        chapter1_overview_text(model),
        "",
        "### 1.3 各区域执行摘要",
        "",
        clean_text(model.get("regional_executive_summary"))
        or "现有结构化资料未提供可验证的区域拆分字段；本节保留该证据边界，区域级行动应在补齐对象、位置、责任界面和验证状态后形成。",
        "",
        "评估信息汇总表",
        "",
    ]
    ng_count = int(status_counts.get("ng") or 0)
    attention_count = int(status_counts.get("attention") or status_counts.get("一般") or 0)
    ok_count = int(status_counts.get("ok") or 0)
    overall_evaluation = chapter1_overall_evaluation(ng_count, attention_count)
    lines.extend(markdown_table(
        ["统计项", "结果"],
        [
            ["NG数量", ng_count],
            ["一般数量", attention_count],
            ["OK数量", ok_count],
            ["整体评价", overall_evaluation],
        ],
    ))
    lines.extend([
        "",
        f"整体评价：{overall_evaluation}。",
    ])

    lines.extend([
        "",
        "## 2. 评估内容描述",
        "",
    ])

    for number, title, level in HEADINGS:
        if number in {"1", "1.1", "1.2", "1.3", "2", "3", "3.1", "3.2", "4", "4.1", "4.2", "4.3", "4.4"}:
            continue
        if not number.startswith("2."):
            continue
        lines.extend(["#" * (level + 1) + f" {number} {title}", ""])
        if level < 3:
            count = count_issues_under_prefix(sections, number)
            if count:
                intro = section_intro(number)
                if intro:
                    lines.append(intro)
                lines.append(f"本项涉及 {count} 项相关记录，以下按检测项逐项说明。")
            else:
                lines.append("本次上传资料中未识别到该类明确异常，建议结合现场资料进一步核实。")
            lines.append("")
            continue

        section_issues = get_section_issues(sections, number)
        section_rule = primary_section_rule(rules_by_section, sections, number)
        if not section_issues:
            reasoning_lines = section_reasoning_block_lines(section_reasoning_results.get(number) or {})
            if reasoning_lines:
                lines.extend(reasoning_lines)
                continue
            lines.extend(status_block_lines(
                situation="",
                conclusion=f"{title}因数据缺失，本次评估未能给出明确结论。建议补充{missing_fields_hint(number)}后重新评估。",
                risk="",
                suggestions=["后续结合现场记录、检测数据或设备台账补充核实"],
            ))
            continue
        section_detail_rows = detail_rows_for_section(model, number)
        if section_detail_rows:
            lines.extend(["现场问题明细表", ""])
            lines.extend(markdown_table(
                issue_table_headers(),
                [detail_markdown_table_row(detail) for detail in section_detail_rows],
            ))
            lines.append("")
        reasoning_lines = section_reasoning_block_lines(section_reasoning_results.get(number) or {})
        if reasoning_lines:
            lines.extend(reasoning_lines)
            section_cross_lines = v2_section_cross_link_lines(cross_link_map.get(number, []))
            if section_cross_lines:
                lines.extend(section_cross_lines)
            continue
        for issue in section_issues[:8]:
            rule = issue_rule(issue) or section_rule or {}
            problem = clean_text(issue.get("problem") or issue.get("description") or "待核实问题")
            recommendation = combine_recommendations(
                issue.get("recommendation"),
                rule.get("recommendation"),
            )
            risk_analysis = clean_text(issue.get("risk_description") or rule.get("risk_description") or "")
            ku_title = clean_text(issue.get("ku_title") or rule.get("ku_title") or "")
            status_label = issue_status_label(issue, rule)
            conclusion_text = status_conclusion_text(number, ku_title or title, status_label, problem)
            risk_text = matched_risk_analysis(number, problem, risk_analysis, status_label)
            suggestion_lines = [] if status_label == "OK" else split_recommendations(recommendation, limit=5 if status_label == "NG" else 3)
            lines.extend(status_block_lines(
                situation=problem if status_label in {"一般", "NG", "数据冲突"} else "",
                conclusion=conclusion_text,
                risk=risk_text,
                suggestions=suggestion_lines if status_label == "OK" else (suggestion_lines or ["结合现场资料制定整改措施"]),
            ))
            for table in issue.get("tables") or []:
                columns = table.get("columns") or []
                rows = table.get("rows") or []
                if not columns or not rows:
                    continue
                title = clean_text(table.get("title") or "")
                if title:
                    lines.append(title)
                normalized_rows = []
                for row in rows:
                    if isinstance(row, dict):
                        normalized_rows.append(row)
                    else:
                        normalized_rows.append(list(row))
                output_columns, output_rows = normalize_issue_markdown_table(columns, normalized_rows, issue)
                lines.extend(markdown_table(output_columns, output_rows))
                lines.append("")
        section_cross_lines = v2_section_cross_link_lines(cross_link_map.get(number, []))
        if section_cross_lines:
            lines.extend(section_cross_lines)

    lines.extend([
        "## 3. 结论与建议",
        "",
        "### 3.1 风险/问题汇总与概览",
        "",
        chapter3_summary_text(model, issues, action_rows),
        "",
    ])
    lines.extend(visual_environment_lines(model))
    lines.extend([
        "### 3.2 改善行动速查表",
        "",
        "改善行动列表见下表。建议按高、中、低优先级逐项推进，并在实施后复核整改效果。",
        "",
    ])
    if action_rows:
        lines.extend(markdown_table(
            ["类别", "问题/主题", "优先级", "建议措施", "备注"],
            [
                [
                    row.get("类别") or row.get("category") or "待归类",
                    row.get("问题/主题") or row.get("problem") or "",
                    row.get("优先级") or row.get("priority") or "待核实",
                    row.get("建议措施") or row.get("recommendation") or "结合现场资料制定整改措施",
                    clean_action_remark(row.get("备注") or row.get("ku_id") or ""),
                ]
                for row in action_rows[:40]
            ],
        ))
    return "\n".join(lines)


def replace_template_content(document: Document, report_text: str, model: dict[str, Any] | None = None) -> None:
    start_index = find_content_start_paragraph_index(document)
    if start_index is None:
        add_report_content(document, report_text, model)
        return

    # Template tables are fixed in the original DOCX body. Once generated text
    # length changes, those tables drift under the wrong headings. Remove them
    # and insert generated Markdown tables inline so table position follows the
    # corresponding KU section.
    remove_body_tables_after_start(document, start_index)
    target_paragraphs = document.paragraphs[start_index:]
    for paragraph in target_paragraphs:
        paragraph.clear()
        strip_template_breaks(paragraph)
        paragraph.style = "Normal"
        remove_paragraph_numbering(paragraph)

    write_report_into_paragraphs(document, target_paragraphs, report_text, model)


def find_content_start_paragraph_index(document: Document) -> int | None:
    for index, paragraph in enumerate(document.paragraphs):
        text = normalize_heading_title(paragraph.text)
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.startswith("Heading") and text in {"配电评估概述", "1配电评估概述"}:
            return index
    for index, paragraph in enumerate(document.paragraphs):
        text = normalize_heading_title(paragraph.text)
        if text in {"配电评估概述", "1配电评估概述"}:
            return index
    return None


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.size = Pt(BODY_FONT_SIZE_PT)
    normal_pf = normal.paragraph_format
    normal_pf.space_before = Pt(0)
    normal_pf.space_after = Pt(BODY_SPACE_AFTER_PT)
    normal_pf.line_spacing = BODY_LINE_SPACING
    for style_name, (size, before, after, line_spacing) in HEADING_STYLE_TOKENS.items():
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.size = Pt(size)
            style.font.bold = True
            style_pf = style.paragraph_format
            style_pf.space_before = Pt(before)
            style_pf.space_after = Pt(after)
            style_pf.line_spacing = line_spacing
            remove_style_numbering(style)
    ensure_preferred_table_style_name(document)
    ensure_preferred_table_paragraph_style(document)


def enforce_cjk_font_mapping(document: Document) -> None:
    """Preserve the template's per-role East Asian font mappings.

    Generated runs inherit their paragraph styles, so a global rewrite is both
    unnecessary and destructive: it previously flattened the template's
    Songti/YaHei/Heiti hierarchy into one font across body, headings, tables,
    headers, and footers.
    """

    return None


def remove_style_numbering(style) -> None:
    """Prevent template Heading styles from adding automatic outline numbers."""
    p_pr = style.element.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def strip_template_breaks(paragraph) -> None:
    """Remove hidden template breaks that can leave blank pages in WPS/Word."""
    p = paragraph._p
    for node in list(p.iter()):
        parent = node.getparent()
        if parent is None:
            continue
        if node.tag == qn("w:br") and node.get(qn("w:type")) == "page":
            parent.remove(node)
        elif node.tag == qn("w:lastRenderedPageBreak"):
            parent.remove(node)

    p_pr = p.get_or_add_pPr()
    for tag in ("w:pageBreakBefore", "w:sectPr"):
        child = p_pr.find(qn(tag))
        if child is not None:
            p_pr.remove(child)


def remove_paragraph_tabs(paragraph) -> None:
    """Clear inherited tab stops that can make advice bullets drift in WPS/Word."""
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)


def remove_paragraph_indentation(paragraph) -> None:
    """Clear all inherited indentation, including WPS character-based indents."""
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is not None:
        p_pr.remove(ind)


def write_report_into_paragraphs(
    document: Document,
    paragraphs: list,
    report_text: str,
    model: dict[str, Any] | None = None,
) -> None:
    paragraph_index = 0
    last_written_paragraph = None

    def next_paragraph():
        nonlocal paragraph_index
        if paragraph_index < len(paragraphs):
            paragraph = paragraphs[paragraph_index]
            paragraph_index += 1
            return paragraph
        paragraph = document.add_paragraph()
        paragraphs.append(paragraph)
        paragraph_index += 1
        return paragraph

    for block in parse_markdown_blocks(report_text):
        if block.kind == "title":
            continue
        if block.kind == "table":
            if last_written_paragraph is not None:
                format_table_title_paragraph(last_written_paragraph)
            write_markdown_table(document, next_paragraph(), list(block.raw_lines), model)
            continue

        written = write_line_or_expanded(next_paragraph, block.text)
        if written:
            last_written_paragraph = written[-1]

    remove_unused_template_paragraphs(paragraphs[paragraph_index:])


def write_line_or_expanded(next_paragraph, line: str) -> list:
    expanded = split_labeled_numbered_line(line)
    if expanded:
        label, items = expanded
        written = []
        paragraph = next_paragraph()
        write_label_paragraph(paragraph, label)
        written.append(paragraph)
        for item in items:
            paragraph = next_paragraph()
            write_bullet_paragraph(paragraph, item)
            written.append(paragraph)
        return written
    paragraph = next_paragraph()
    write_line_to_paragraph(paragraph, line)
    return [paragraph]


def split_labeled_numbered_line(line: str) -> tuple[str, list[str]] | None:
    clean = clean_inline_marks(line)
    labels = "|".join(re.escape(label) for label in EXPANDABLE_LABELS)
    match = re.match(rf"^({labels})[:：]\s*(.+)$", clean)
    if not match:
        return None
    label, body = match.group(1), match.group(2).strip()
    if not re.search(r"(?:^|[；;。:：]\s*)1[.．、]\s*", body) or not re.search(r"(?:^|[；;。:：]\s*)2[.．、]\s*", body):
        return None
    parts = [
        part.strip(" ；;。")
        for part in re.split(r"(?:^|[；;。:：]\s*)\d+[.．、]\s*", body)
        if part.strip(" ；;。")
    ]
    return (f"{label}：", parts) if len(parts) >= 2 else None


def write_label_paragraph(paragraph, label: str) -> None:
    paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    reset_direct_paragraph_format(paragraph)
    paragraph.style = "Normal"
    pf = paragraph.paragraph_format
    pf.first_line_indent = None
    pf.space_before = Pt(3)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.2
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False
    run = paragraph.add_run(label)
    run.bold = True
    run.font.size = Pt(BODY_FONT_SIZE_PT)


def write_bullet_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    remove_paragraph_tabs(paragraph)
    reset_direct_paragraph_format(paragraph)
    paragraph.style = "Normal"
    pf = paragraph.paragraph_format
    pf.left_indent = Cm(LIST_LEFT_INDENT_TWIPS / 567)
    pf.first_line_indent = Cm(-LIST_HANGING_INDENT_TWIPS / 567)
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.25
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False
    apply_list_numbering(
        paragraph,
        _bullet_numbering_id(paragraph.part.document),
    )
    body = paragraph.add_run(clean_bullet_item_text(text))
    body.font.size = Pt(BODY_FONT_SIZE_PT)


def write_numbered_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    remove_paragraph_tabs(paragraph)
    reset_direct_paragraph_format(paragraph)
    paragraph.style = "Normal"
    pf = paragraph.paragraph_format
    pf.left_indent = Cm(LIST_LEFT_INDENT_TWIPS / 567)
    pf.first_line_indent = Cm(-LIST_HANGING_INDENT_TWIPS / 567)
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.25
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False
    apply_list_numbering(
        paragraph,
        _numbered_list_id(
            paragraph.part.document,
            numbered_list_start(text),
        ),
    )
    body = paragraph.add_run(clean_numbered_item_text(text))
    body.font.size = Pt(BODY_FONT_SIZE_PT)


def write_line_to_paragraph(paragraph, line: str) -> None:
    heading = parse_heading(line)
    paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    reset_direct_paragraph_format(paragraph)
    if heading:
        number, title, level = heading
        paragraph.style = f"Heading {level}"
        remove_paragraph_numbering(paragraph)
        apply_heading_format(paragraph, level)
        set_paragraph_outline_level(paragraph, level - 1)
        remove_paragraph_numbering(paragraph)
        if level == 1 and number in {"1", "2", "3", "4"}:
            paragraph.paragraph_format.page_break_before = number != "1"
        run = paragraph.add_run(f"{number}. {title}")
        run.bold = True
        return

    paragraph.style = "Normal"
    display_line = clean_block_markup(line)
    if is_exact_label_line(display_line):
        write_label_paragraph(paragraph, normalized_label_text(display_line))
        return
    if re.match(r"^#{1,6}\s+", line.strip()):
        write_markdown_subheading_paragraph(paragraph, clean_markdown_heading(line))
        return
    if is_time_group_heading_line(display_line):
        write_time_group_heading_paragraph(paragraph, clean_inline_marks(display_line))
        return
    if is_bullet_line(display_line):
        display_line = re.sub(r"^[-*•➢]\s*", "", display_line).strip()
        write_bullet_paragraph(paragraph, display_line)
        return
    if is_numbered_list_line(display_line):
        write_numbered_paragraph(paragraph, display_line)
        return
    apply_body_format(paragraph, display_line)
    add_inline_formatted_runs(paragraph, clean_inline_marks(display_line))


def add_inline_formatted_runs(paragraph, line: str) -> None:
    labels = "|".join(re.escape(label) for label in INLINE_LABELS)
    match = re.match(r"^【([^】]{1,24})】[:：]?\s*(.*)$", line)
    if not match:
        match = re.match(rf"^({labels})[:：]\s*(.*)$", line)
    if not match:
        paragraph.add_run(line)
        return
    label, body = match.group(1), match.group(2)
    label_run = paragraph.add_run(f"{label}：")
    label_run.bold = True
    if body:
        paragraph.add_run(body)


def clean_block_markup(line: str) -> str:
    """Remove Markdown-only block markers before Word formatting."""

    return re.sub(r"^>\s*", "", str(line).strip())


def normalized_label_text(line: str) -> str:
    """Return one consistent, bracket-free label for the expert report style."""

    clean = clean_inline_marks(clean_markdown_heading(clean_block_markup(line)))
    clean = clean.rstrip("：:").strip()
    clean = re.sub(r"^【\s*|\s*】$", "", clean).strip()
    return clean + "："


def write_markdown_subheading_paragraph(paragraph, text: str) -> None:
    """Render non-taxonomy Markdown headings as the template's small label role."""

    paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    reset_direct_paragraph_format(paragraph)
    paragraph.style = "小标" if "小标" in {style.name for style in paragraph.part.document.styles} else "Normal"
    pf = paragraph.paragraph_format
    pf.first_line_indent = None
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.2
    run = paragraph.add_run(normalized_label_text(text))
    run.bold = True


def reset_direct_paragraph_format(paragraph) -> None:
    """Remove direct formatting inherited from whatever template paragraph is reused."""
    strip_template_breaks(paragraph)
    remove_paragraph_tabs(paragraph)
    remove_paragraph_indentation(paragraph)
    pf = paragraph.paragraph_format
    pf.alignment = None
    pf.left_indent = None
    pf.right_indent = None
    pf.first_line_indent = None
    pf.space_before = None
    pf.space_after = None
    pf.line_spacing = None
    pf.line_spacing_rule = None
    pf.keep_with_next = None
    pf.keep_together = None
    pf.page_break_before = None


def apply_heading_format(paragraph, level: int) -> None:
    # Use one explicit heading rhythm after the template font family is retained.
    pf = paragraph.paragraph_format
    pf.left_indent = None
    pf.right_indent = None
    pf.first_line_indent = None
    pf.alignment = None
    _size, before, after, line_spacing = HEADING_STYLE_TOKENS[f"Heading {level}"]
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    pf.keep_with_next = level >= 3


def apply_body_format(paragraph, line: str) -> None:
    pf = paragraph.paragraph_format
    pf.first_line_indent = Cm(0.74)
    pf.left_indent = None
    pf.right_indent = None
    pf.space_before = Pt(0)
    pf.space_after = Pt(BODY_SPACE_AFTER_PT)
    pf.line_spacing = BODY_LINE_SPACING
    pf.keep_together = False
    pf.keep_with_next = False
    pf.page_break_before = False


def write_time_group_heading_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    reset_direct_paragraph_format(paragraph)
    paragraph.style = "Normal"
    pf = paragraph.paragraph_format
    pf.left_indent = None
    pf.right_indent = None
    pf.first_line_indent = None
    pf.space_before = Pt(4)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.2
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False
    display = re.sub(r"^【\s*|\s*】$", "", text).strip().rstrip("：:") + "："
    run = paragraph.add_run(display)
    run.bold = True
    run.font.size = Pt(BODY_FONT_SIZE_PT)


def is_inline_analysis_line(line: str) -> bool:
    clean = clean_inline_marks(line)
    return bool(re.match(r"^【(根因|直接影响|恶化条件|建议|风险分析|结论|现状描述)】", clean))


def is_label_line(line: str) -> bool:
    return bool(re.match(
        r"^(?:【)?(现状描述|风险等级|资料来源|方法论依据|核查方法|判定依据|风险分析|整改建议|参考标准|潜在风险|后续建议|结论|建议)(?:】)?[:：]",
        clean_inline_marks(line),
    ))


def is_exact_label_line(line: str) -> bool:
    clean = clean_inline_marks(clean_markdown_heading(clean_block_markup(line)))
    clean = clean.rstrip("：:").strip()
    labels = "|".join(re.escape(label) for label in INLINE_LABELS)
    return bool(re.fullmatch(rf"(?:【)?(?:{labels})(?:】)?", clean))


def is_time_group_heading_line(line: str) -> bool:
    return bool(re.match(
        r"^【(立即执行|短期整改|中期规划|长期关注|优先补充|建议补充).+】$",
        clean_inline_marks(line),
    ))


def parse_heading(line: str) -> tuple[str, str, int] | None:
    # Only explicit Markdown headings own the report taxonomy.  Chief-editor
    # prose may contain local numbered labels such as ``**1.1 负荷均衡调整**``;
    # treating those body lines as canonical headings silently replaces their
    # approved titles when the local number collides with HEADINGS.
    if not re.match(r"^#{1,6}\s+", line.strip()):
        return None
    clean = clean_inline_marks(clean_markdown_heading(line))
    # Only treat real report headings as headings. Lines such as
    # "1)评估所见..." or "1、..." are body numbered lists and must not be
    # mapped to chapter "1. 配电评估概述".
    match = re.match(r"^(\d+(?:\.\d+)+)(?:\.)?\s+(.+)$", clean)
    if match:
        number = match.group(1)
        if number in HEADING_BY_NUMBER:
            title, level = HEADING_BY_NUMBER[number]
            return number, title, level
        if number.startswith("4."):
            return number, match.group(2).strip(), 2
    top_level_match = re.match(r"^([1234])\.\s+(.+)$", clean)
    if top_level_match:
        number = top_level_match.group(1)
        expected = HEADING_BY_NUMBER.get(number)
        if expected and normalize_heading_title(top_level_match.group(2)) == normalize_heading_title(expected[0]):
            return number, expected[0], expected[1]
    normalized = normalize_heading_title(strip_number_prefix(clean))
    by_title = HEADING_BY_TITLE.get(normalized)
    return by_title


def remove_unused_template_paragraphs(paragraphs: list) -> None:
    for paragraph in paragraphs:
        if paragraph_has_section_properties(paragraph):
            paragraph.clear()
            strip_template_breaks(paragraph)
            paragraph.style = "Normal"
            remove_paragraph_numbering(paragraph)
            continue
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def normalize_static_toc(document: Document) -> None:
    """Rewrite the visible TOC lines to match the generated fixed report outline."""
    title_index = next(
        (
            index
            for index, paragraph in enumerate(document.paragraphs)
            if normalize_heading_title(paragraph.text) == "目录"
        ),
        -1,
    )
    if title_index >= 0:
        title_paragraph = document.paragraphs[title_index]
        title_paragraph.style = "Normal"
        reset_direct_paragraph_format(title_paragraph)
        remove_paragraph_numbering(title_paragraph)
        remove_paragraph_outline_level(title_paragraph)
        for run in title_paragraph.runs:
            run.bold = True
            run.font.size = Pt(14)
    toc_paragraphs = [
        paragraph
        for index, paragraph in enumerate(document.paragraphs)
        if index > title_index
        and paragraph.style is not None
        and paragraph.style.name.startswith("toc ")
        and normalize_heading_title(paragraph.text)
    ]
    entries = [
        (number, title, level)
        for number, title, level in HEADINGS
        if level <= 3
    ]
    if len(toc_paragraphs) < len(entries):
        content_start = find_content_start_paragraph_index(document)
        if content_start is not None:
            anchor = document.paragraphs[content_start]
            for _ in range(len(entries) - len(toc_paragraphs)):
                toc_paragraphs.append(anchor.insert_paragraph_before())
    for paragraph, (number, title, level) in zip(toc_paragraphs, entries):
        paragraph.clear()
        paragraph.style = "Normal"
        reset_direct_paragraph_format(paragraph)
        remove_paragraph_numbering(paragraph)
        remove_paragraph_outline_level(paragraph)
        pf = paragraph.paragraph_format
        pf.left_indent = Cm(0.45 * max(level - 1, 0))
        pf.first_line_indent = None
        pf.space_after = Pt(2)
        pf.line_spacing = 1.1
        try:
            pf.tab_stops.clear_all()
            pf.tab_stops.add_tab_stop(Cm(15.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        except Exception:
            pass
        paragraph.add_run(f"{number}. {title}")
        paragraph.add_run().add_tab()
        paragraph.add_run(TOC_PAGE_HINTS.get(number, ""))
    for paragraph in toc_paragraphs[len(entries):]:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def paragraph_has_section_properties(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.sectPr is not None


def remove_paragraph_numbering(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def _next_numbering_id(numbering, element_name: str, attribute_name: str) -> int:
    values: list[int] = []
    for element in numbering.findall(qn(f"w:{element_name}")):
        raw = element.get(qn(f"w:{attribute_name}"))
        try:
            values.append(int(raw or 0))
        except (TypeError, ValueError):
            continue
    return max(values, default=0) + 1


def _append_abstract_numbering(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    numbering_name = f"manyselves-{kind}"
    for abstract_num in numbering.findall(qn("w:abstractNum")):
        name = abstract_num.find(qn("w:name"))
        if name is None or name.get(qn("w:val")) != numbering_name:
            continue
        try:
            return int(abstract_num.get(qn("w:abstractNumId")) or 0)
        except (TypeError, ValueError):
            continue

    abstract_num_id = _next_numbering_id(numbering, "abstractNum", "abstractNumId")
    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_num_id))
    name = OxmlElement("w:name")
    name.set(qn("w:val"), numbering_name)
    abstract_num.append(name)
    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi_level_type)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.append(level_justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(LIST_LEFT_INDENT_TWIPS))
    tabs.append(tab)
    p_pr.append(tabs)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), str(LIST_LEFT_INDENT_TWIPS))
    indentation.set(qn("w:hanging"), str(LIST_HANGING_INDENT_TWIPS))
    p_pr.append(indentation)
    level.append(p_pr)
    abstract_num.append(level)

    first_num_index = next(
        (
            index
            for index, child in enumerate(numbering)
            if child.tag == qn("w:num")
        ),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract_num)
    return abstract_num_id


def _append_numbering_instance(
    document: Document,
    abstract_num_id: int,
    *,
    start_override: int | None = None,
) -> int:
    numbering = document.part.numbering_part.element
    num_id = _next_numbering_id(numbering, "num", "numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id_element = OxmlElement("w:abstractNumId")
    abstract_num_id_element.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract_num_id_element)
    if start_override is not None:
        level_override = OxmlElement("w:lvlOverride")
        level_override.set(qn("w:ilvl"), "0")
        start_override_element = OxmlElement("w:startOverride")
        start_override_element.set(qn("w:val"), str(max(1, start_override)))
        level_override.append(start_override_element)
        num.append(level_override)
    numbering.append(num)
    return num_id


def _bullet_numbering_id(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_num_id = _append_abstract_numbering(document, "bullet")
    for num in numbering.findall(qn("w:num")):
        abstract_ref = num.find(qn("w:abstractNumId"))
        if (
            abstract_ref is not None
            and abstract_ref.get(qn("w:val")) == str(abstract_num_id)
            and num.find(qn("w:lvlOverride")) is None
        ):
            try:
                return int(num.get(qn("w:numId")) or 0)
            except (TypeError, ValueError):
                continue
    num_id = _append_numbering_instance(document, abstract_num_id)
    return num_id


def _numbered_list_id(document: Document, start: int) -> int:
    abstract_num_id = _append_abstract_numbering(document, "decimal")
    return _append_numbering_instance(
        document,
        abstract_num_id,
        start_override=start,
    )


def apply_list_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_element)
    p_pr.append(num_pr)


def remove_paragraph_outline_level(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is not None:
        p_pr.remove(outline)


def set_paragraph_outline_level(paragraph, outline_level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(max(0, min(outline_level, 8))))


def restore_template_heading_numbering(paragraph, level: int) -> None:
    """Restore the outline numbering used by the bundled Word template."""
    if level < 1 or level > 4:
        return
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)

    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level - 1))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)


def format_table_cell_paragraph(paragraph, header: bool = False) -> None:
    try:
        paragraph.style = find_preferred_table_paragraph_style(paragraph) or "Normal"
    except Exception:
        pass
    reset_direct_paragraph_format(paragraph)
    paragraph.alignment = None
    pf = paragraph.paragraph_format
    pf.left_indent = None
    pf.right_indent = None
    pf.first_line_indent = None
    pf.line_spacing = None
    pf.space_before = None
    pf.space_after = None
    for run in paragraph.runs:
        run.font.size = None
        run.bold = None


def find_preferred_table_paragraph_style(paragraph):
    try:
        styles = paragraph.part.document.styles
    except Exception:
        return None
    for style in styles:
        if style.element.get(qn("w:type")) != "paragraph":
            continue
        if style.style_id == "24" or style.name == "表格2":
            return style
    return None


def ensure_preferred_table_paragraph_style(document: Document):
    for style in document.styles:
        if style.element.get(qn("w:type")) == "paragraph" and style.name == "表格2":
            configure_table_paragraph_style(style)
            return style
    try:
        style = document.styles.add_style("表格2", WD_STYLE_TYPE.PARAGRAPH)
    except Exception:
        return None
    configure_table_paragraph_style(style)
    return style


def configure_table_paragraph_style(style) -> None:
    # Preserve the template font family while keeping dense technical tables
    # readable and independent from the Normal style's first-line indent.
    style.font.size = Pt(TABLE_FONT_SIZE_PT)
    p_pr = style.element.get_or_add_pPr()
    for tag in ("w:numPr", "w:tabs", "w:ind"):
        child = p_pr.find(qn(tag))
        if child is not None:
            p_pr.remove(child)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    p_pr.append(ind)


def parse_section_reasoning_results(value: Any) -> dict[str, dict[str, Any]]:
    payload = parse_json_object(value)
    if not payload:
        return {}

    sections = payload.get("sections")
    if sections is None and any(key in payload for key in ("section_id", "section_title", "conclusion")):
        sections = [payload]
    if not isinstance(sections, list):
        return {}

    results: dict[str, dict[str, Any]] = {}
    for item in sections:
        if not isinstance(item, dict):
            continue
        section_id = clean_text(item.get("section_id"))
        if not section_id:
            section_title = clean_text(item.get("section_title"))
            if section_title:
                section_id = HEADING_BY_TITLE.get(_heading_key(section_title), ("", "", 0))[0]
        if not section_id:
            continue
        results[section_id] = item
    return results


def section_reasoning_block_lines(result: dict[str, Any]) -> list[str]:
    if not isinstance(result, dict):
        return []

    situation = clean_text(
        result.get("current_status")
        or result.get("situation")
        or result.get("status")
        or result.get("现状描述")
        or ""
    )
    conclusion = clean_text(result.get("conclusion") or result.get("结论") or "")
    risk = clean_text(
        result.get("risk_analysis")
        or result.get("risk")
        or result.get("风险分析")
        or ""
    )
    recommendations = result.get("recommendations")
    if recommendations is None:
        recommendations = result.get("suggestions") or result.get("建议") or []
    if isinstance(recommendations, str):
        suggestion_lines = split_recommendations(recommendations, limit=8)
    elif isinstance(recommendations, list):
        suggestion_lines = [
            clean_text(item.get("text") if isinstance(item, dict) else item)
            for item in recommendations
        ]
        suggestion_lines = [item for item in suggestion_lines if item]
    else:
        suggestion_lines = []

    evidence_refs = result.get("evidence_refs") or result.get("evidence") or []
    if isinstance(evidence_refs, str):
        evidence_lines = split_recommendations(evidence_refs, limit=6)
    elif isinstance(evidence_refs, list):
        evidence_lines = [
            clean_text(item.get("text") if isinstance(item, dict) else item)
            for item in evidence_refs
        ]
        evidence_lines = [item for item in evidence_lines if item]
    else:
        evidence_lines = []

    if evidence_lines:
        evidence_text = "；".join(evidence_lines[:5])
        if situation and evidence_text not in situation:
            situation = f"{situation} 证据依据：{evidence_text}。"
        elif not situation:
            situation = f"证据依据：{evidence_text}。"

    if not any([situation, conclusion, risk, suggestion_lines]):
        return []

    return status_block_lines(
        situation=situation,
        conclusion=conclusion or "该章节已按专项推理规则完成复核，建议结合现场证据持续跟踪。",
        risk=risk,
        suggestions=suggestion_lines or ["结合章节专项推理结果完善整改闭环。"],
    )


def format_table_title_paragraph(paragraph) -> None:
    reset_direct_paragraph_format(paragraph)
    paragraph.alignment = None
    pf = paragraph.paragraph_format
    pf.first_line_indent = None
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15


def write_markdown_table(
    document: Document,
    anchor_paragraph,
    table_buffer: list[str],
    model: dict[str, Any] | None = None,
) -> None:
    rows: list[list[str]] = []
    for line in table_buffer:
        cells = [clean_inline_marks(cell.strip()) for cell in line.strip("|").split("|")]
        if cells and all(set(cell) <= {"-", ":", " "} for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    max_rows = policy_int("table_policy", "max_rows_per_table", 30)
    if max_rows > 0 and len(rows) > max_rows + 1:
        rows = rows[: max_rows + 1]

    anchor_paragraph.clear()
    anchor_paragraph.style = "Normal"
    reset_direct_paragraph_format(anchor_paragraph)
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    apply_preferred_table_style(document, table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_layout(table)
    set_table_widths(table, column_count)

    for row_index, row in enumerate(rows):
        for col_index in range(column_count):
            header = rows[0][col_index] if rows and col_index < len(rows[0]) else ""
            value = normalize_table_cell_value(row[col_index] if col_index < len(row) else "", header, row_index == 0)
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=80, start=80, bottom=80, end=80)
            if row_index == 0:
                shade_cell(cell, "D9F2F0")
            photo_payload = photo_payload_from_table_cell(value, model)
            if photo_payload:
                write_photo_cell(cell, photo_payload)
                continue
            else:
                cell.text = str(value)
            for paragraph in cell.paragraphs:
                format_table_cell_paragraph(paragraph, header=row_index == 0)
    set_table_widths(table, column_count)

    table_element = table._tbl
    table_element.getparent().remove(table_element)
    anchor_p = anchor_paragraph._p
    anchor_p.addnext(table_element)
    insert_table_bottom_spacer(table_element)
    remove_paragraph_element(anchor_p)


def insert_table_bottom_spacer(table_element) -> None:
    spacer = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "160")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "160")
    spacing.set(qn("w:lineRule"), "exact")
    p_pr.append(spacing)
    spacer.append(p_pr)
    table_element.addnext(spacer)


def remove_paragraph_element(paragraph_element) -> None:
    parent = paragraph_element.getparent()
    if parent is not None:
        parent.remove(paragraph_element)


def ensure_preferred_table_style_name(document: Document) -> str:
    """Return the bundled template table style id without renaming styles."""
    preferred = find_preferred_table_style(document)
    if preferred is None:
        return "21"
    return preferred.style_id


def find_preferred_table_style(document: Document):
    for style in document.styles:
        style_type = style.element.get(qn("w:type"))
        if style_type != "table":
            continue
        if style.style_id == "21" or style.name in {"表格2", "表格 2", "Doc Table Band 2nd"}:
            return style
    for style in document.styles:
        if style.element.get(qn("w:type")) == "table" and style.name == "Table Grid":
            return style
    return None


def apply_preferred_table_style(document: Document, table) -> None:
    """Use the template table style shown as "表格2" in WPS when available."""
    style_id = ensure_preferred_table_style_name(document)
    preferred = find_preferred_table_style(document)
    if preferred is not None:
        try:
            table.style = preferred
        except Exception:
            table.style = preferred.name
        force_table_style_id(table, preferred.style_id)
        return
    force_table_style_id(table, style_id)


def force_table_style_id(table, style_id: str | None = None) -> None:
    style_id = style_id or "21"
    tbl_pr = table._tbl.tblPr
    tbl_style = tbl_pr.find(qn("w:tblStyle"))
    if tbl_style is None:
        tbl_style = OxmlElement("w:tblStyle")
        tbl_pr.insert(0, tbl_style)
    tbl_style.set(qn("w:val"), style_id)


def enforce_all_table_styles(document: Document) -> None:
    for table in document.tables:
        apply_preferred_table_style(document, table)


def normalize_table_cell_value(value: Any, header: Any = "", is_header: bool = False) -> str:
    text = clean_text(value)
    header_text = clean_text(header)
    if is_header:
        return text or "字段"
    if text:
        return text
    if "图片" in header_text or "照片" in header_text:
        return missing_photo_text()
    if "时间" in header_text:
        return "未测试"
    if "状态" in header_text or "判定" in header_text:
        return "待评估"
    if any(token in header_text for token in ["位置", "区域", "配电"]):
        return "待核实位置"
    if any(token in header_text for token in ["柜号", "回路", "设备", "负荷", "名称"]):
        return "待核实"
    if any(token in header_text for token in ["数值", "电流", "容量", "温度", "THD", "限值", "超标率"]):
        return "未测量"
    if "建议" in header_text or "措施" in header_text:
        return "结合现场资料进一步核实。"
    return "待补充"


def photo_payload_from_table_cell(value: Any, model: dict[str, Any] | None) -> dict[str, Any] | None:
    text = clean_text(value)
    match = re.fullmatch(r"\[\[PHOTO_REFS:([A-Za-z0-9_,\-\s]+)\]\]", text)
    if not match:
        return None
    refs = [item.strip() for item in match.group(1).split(",") if item.strip()]
    return {
        "photo_refs": refs,
        "photo_status": "matched" if refs else "missing_photo",
        "photo_manifest": (model or {}).get("photo_manifest") or [],
    }


def set_table_widths(table, column_count: int) -> None:
    header_texts = []
    if table.rows:
        header_texts = [clean_text(cell.text) for cell in table.rows[0].cells]
    width_maps = {
        3: [4.2, 6.4, 6.4],
        4: [3.2, 4.6, 5.8, 3.8],
        5: [1.5, 2.6, 7.0, 2.2, 4.0],
        6: [2.3, 2.3, 2.7, 4.2, 4.0, 2.8],
        7: [2.3, 3.0, 1.5, 1.6, 2.2, 1.8, 1.6],
        8: [2.0, 2.0, 1.8, 1.5, 1.7, 2.4, 2.2, 3.0],
        9: [2.0, 2.0, 2.4, 2.3, 2.2, 2.2, 2.4, 2.4, 2.2],
        11: [1.4, 1.8, 1.5, 2.0, 2.8, 1.4, 1.4, 3.0, 2.0, 1.7, 2.2],
    }
    if column_count == 6 and {"序号", "章节", "问题", "优先级", "建议", "预期收益"}.issubset(set(header_texts)):
        weights = [1.1, 3.0, 4.5, 1.5, 4.2, 3.2]
    else:
        weights = width_maps.get(column_count)
    if not weights:
        weights = [
            max(
                1.0,
                min(
                    6.0,
                    max(
                        (
                            len(clean_text(row.cells[index].text))
                            for row in table.rows
                            if index < len(row.cells)
                        ),
                        default=1,
                    )
                    ** 0.5,
                ),
            )
            for index in range(column_count)
        ]
    usable_width = table_usable_width_cm(table)
    weight_total = sum(weights)
    total_twips = cm_to_twips(usable_width)
    width_twips = [
        int(total_twips * (weight / weight_total))
        for weight in weights[:-1]
    ]
    width_twips.append(total_twips - sum(width_twips))
    set_fixed_table_layout(table)
    set_table_total_width(table, total_twips)
    set_table_grid_widths(table, width_twips)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index < len(width_twips):
                cell.width = Cm(width_twips[index] / 567)
                set_cell_width(cell, width_twips[index])


def table_usable_width_cm(table) -> float:
    try:
        sections = table.part.document.sections
        widths = [
            (
                section.page_width.cm
                - section.left_margin.cm
                - section.right_margin.cm
            )
            for section in sections
            if section.page_width is not None
            and section.left_margin is not None
            and section.right_margin is not None
        ]
    except Exception:
        widths = []
    # Use the narrowest retained section so one table geometry remains safe
    # after Word repaginates the document or updates fields.
    return max(1.0, min(widths, default=15.0))


def set_table_total_width(table, width_twips: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")


def set_fixed_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_table_grid_widths(table, widths_twips: list[int]) -> None:
    tbl = table._tbl
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width_twips in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width_twips))
        grid.append(col)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def cm_to_twips(value: float) -> int:
    return int(float(value) * 567)


TABLE_SECTION_MAP = {
    1: "2.1.1",
    2: "2.1.2",
    3: "2.2.1.1",
    4: "2.2.2.1",
    5: "2.2.2.1",
    6: "2.2.2.1",
    7: "2.2.2.3",
    8: "2.3.1",
    9: "2.3.3",
    10: "2.4.1.2",
    11: "2.4.2.1",
    12: "2.4.2.2",
    13: "2.4.2.4",
    14: "2.4.2.5",
    15: "2.4.3.1",
    16: "2.4.3.2",
    17: "2.4.3.3",
    18: "2.5.2",
    19: "2.5.3.2",
    20: "2.5.3.3",
    21: "2.5.5",
}


def fill_template_tables(document: Document, model: dict[str, Any]) -> None:
    sections = model.get("sections") or {}
    action_rows = model.get("action_rows") or []
    for table_index, table in enumerate(document.tables, start=1):
        if table_index == 22:
            fill_action_table(table, action_rows)
            continue
        section_id = TABLE_SECTION_MAP.get(table_index)
        if not section_id:
            continue
        fill_issue_table(table, get_section_issues(sections, section_id), section_id, model)


def fill_issue_table(table, issues: list[dict[str, Any]], section_id: str, model: dict[str, Any] | None = None) -> None:
    col_count = len(table.columns)
    headers = headers_for_table(col_count)
    reset_table(table, headers)
    detail_rows = detail_rows_for_section(model or {}, section_id) if col_count == 6 else []
    if detail_rows:
        for detail in detail_rows[:12]:
            add_table_row(table, detail_issue_row(detail, model or {}))
        set_table_widths(table, col_count)
        return
    if not issues:
        add_table_row(table, placeholder_row(col_count, section_id))
        set_table_widths(table, col_count)
        return
    for issue in issues[:12]:
        add_table_row(table, issue_row(issue, col_count))
    set_table_widths(table, col_count)


def fill_action_table(table, rows: list[dict[str, Any]]) -> None:
    headers = ["类别 / Category", "话题 / Topic", "优先级 / Priority", "收益 / Profit", "备注 / Comments"]
    reset_table(table, headers)
    if not rows:
        add_table_row(table, ["待核实", "上传资料中未识别到明确改善项", "待核实", "补充资料后重新生成", ""])
        set_table_widths(table, len(table.columns))
        return
    for row in rows[:40]:
        add_table_row(table, [
            row.get("类别") or row.get("category") or "待归类",
            row.get("问题/主题") or row.get("problem") or "",
            row.get("优先级") or row.get("priority") or "待核实",
            row.get("建议措施") or row.get("recommendation") or "结合现场资料制定整改措施",
            row.get("备注") or row.get("evidence") or "",
        ])
    set_table_widths(table, len(table.columns))


def reset_table(table, headers: list[str]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    if not table.rows:
        table.add_row()
    write_cells(table.rows[0].cells, headers)
    format_row(table.rows[0], header=True)


def add_table_row(table, values: list[Any]) -> None:
    row = table.add_row()
    write_cells(row.cells, values)
    format_row(row, header=False)


def write_cells(cells, values: list[Any]) -> None:
    for index, cell in enumerate(cells):
        header = ""
        value = values[index] if index < len(values) else ""
        if isinstance(value, dict) and "photo_refs" in value:
            write_photo_cell(cell, value)
        else:
            cell.text = normalize_table_cell_value(value, header)


def reset_photo_insertion_count() -> None:
    global PHOTO_INSERTION_COUNT
    PHOTO_INSERTION_COUNT = 0


def can_insert_report_photo() -> bool:
    return PHOTO_INSERTION_COUNT < MAX_REPORT_PHOTO_INSERTIONS


def mark_report_photo_inserted() -> None:
    global PHOTO_INSERTION_COUNT
    PHOTO_INSERTION_COUNT += 1


def compressed_photo_path(path: Path) -> Path:
    if Image is None or ImageOps is None:
        return path
    try:
        stat = path.stat()
    except OSError:
        return path

    cache_key = "|".join(
        [
            str(path.resolve()),
            str(stat.st_size),
            str(int(stat.st_mtime)),
            str(PHOTO_MAX_PIXELS),
            str(PHOTO_JPEG_QUALITY),
        ]
    )
    digest = hashlib.sha1(cache_key.encode("utf-8")).hexdigest()
    cached = PHOTO_CACHE_DIR / f"{digest}.jpg"
    if cached.exists() and cached.stat().st_size > 0:
        return cached

    try:
        PHOTO_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            image.thumbnail(PHOTO_MAX_PIXELS)
            if image.mode in {"RGBA", "LA"}:
                background = Image.new("RGB", image.size, (255, 255, 255))
                alpha = image.getchannel("A") if "A" in image.getbands() else None
                background.paste(image, mask=alpha)
                image = background
            elif image.mode != "RGB":
                image = image.convert("RGB")
            image.save(cached, format="JPEG", quality=PHOTO_JPEG_QUALITY, optimize=True)
    except Exception:
        return path
    return cached if cached.exists() and cached.stat().st_size > 0 else path


def write_photo_cell(cell, payload: dict[str, Any]) -> None:
    refs = payload.get("photo_refs") or []
    manifest = {
        item.get("photo_id"): item
        for item in (payload.get("photo_manifest") or [])
        if isinstance(item, dict) and item.get("photo_id")
    }
    cell.text = ""
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    format_photo_cell_paragraph(paragraph, clear=True)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
    inserted = False
    for ref in refs[:1]:
        if not can_insert_report_photo():
            paragraph.add_run("图片数量较多，详见原始资料")
            break
        item = manifest.get(ref)
        path = Path(str(item.get("local_path") or "")) if item else None
        if not path or not path.exists():
            continue
        run = paragraph.add_run()
        try:
            width_cm = policy_float("table_policy", "photo_width_cm", PHOTO_INLINE_WIDTH_CM)
            height_cm = policy_float("table_policy", "photo_height_cm", 0)
            cell_width = cell_width_cm(cell)
            if cell_width:
                width_cm = min(width_cm, max(cell_width - 0.2, 0.8))
            run.add_picture(
                str(compressed_photo_path(path)),
                **photo_fit_dimensions(path, width_cm, height_cm),
            )
            inserted = True
            mark_report_photo_inserted()
        except Exception:
            continue
    if inserted:
        return
    status = clean_text(payload.get("photo_status") or "")
    if status == "missing_media":
        paragraph.add_run(broken_photo_ref_text())
    else:
        paragraph.add_run(missing_photo_text())


def format_photo_cell_paragraph(paragraph, clear: bool = False) -> None:
    if clear:
        paragraph.clear()
    strip_template_breaks(paragraph)
    remove_paragraph_numbering(paragraph)
    remove_paragraph_tabs(paragraph)
    remove_paragraph_indentation(paragraph)
    try:
        document = paragraph.part.document
        paragraph.style = ensure_preferred_table_paragraph_style(document) or find_preferred_table_paragraph_style(paragraph) or "Normal"
    except Exception:
        try:
            paragraph.style = "Normal"
        except Exception:
            pass
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.left_indent = None
    pf.right_indent = None
    pf.first_line_indent = None
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False


def photo_fit_dimensions(path: Path, width_cm: float, height_cm: float) -> dict[str, Any]:
    width_cm = min(max(float(width_cm or PHOTO_INLINE_WIDTH_CM), 0.1), PHOTO_CELL_MAX_WIDTH_CM)
    height_cm = min(max(float(height_cm or width_cm), 0.1), PHOTO_CELL_MAX_HEIGHT_CM)
    if Image is None:
        return {"width": Cm(width_cm)}
    try:
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image) if ImageOps is not None else image
            pixel_width, pixel_height = image.size
    except Exception:
        return {"width": Cm(width_cm)}
    if pixel_width <= 0 or pixel_height <= 0:
        return {"width": Cm(width_cm)}
    aspect = pixel_width / pixel_height
    box_aspect = width_cm / height_cm
    if aspect >= box_aspect:
        return {"width": Cm(width_cm)}
    return {"height": Cm(height_cm)}


def cell_width_cm(cell) -> float:
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return 0.0
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        return 0.0
    try:
        value = int(tc_w.get(qn("w:w")) or 0)
    except (TypeError, ValueError):
        return 0.0
    if value <= 0:
        return 0.0
    return value / 567.0


def format_row(row, header: bool) -> None:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=100, start=80, bottom=100, end=80)
        if header:
            shade_cell(cell, "D9F2F0")
        if cell_has_drawing(cell):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
            for paragraph in cell.paragraphs:
                format_photo_cell_paragraph(paragraph)
            continue
        for paragraph in cell.paragraphs:
            format_table_cell_paragraph(paragraph, header=header)


def cell_has_drawing(cell) -> bool:
    return bool(cell._tc.xpath(".//w:drawing"))


def set_cell_margins(cell, top: int = 80, start: int = 80, bottom: int = 80, end: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "left": start, "bottom": bottom, "right": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def headers_for_table(col_count: int) -> list[str]:
    if col_count <= 2:
        return ["项目", "内容"][:col_count]
    if col_count == 3:
        return ["位置/范围", "问题描述", "整改建议"]
    if col_count == 4:
        return ["位置/范围", "问题描述", "整改建议", "备注"]
    if col_count == 5:
        return ["类别", "问题/主题", "优先级", "建议措施", "备注"]
    if col_count == 6:
        return ["位置", "设备类型", "柜号及回路号", "问题描述", "建议", "现场图片"]
    if col_count == 7:
        return ["设备/区域", "绝缘检测", "局放/温度", "导通检测", "保护性能", "动作验证", "备注"]
    if col_count == 8:
        return ["配电室", "设备", "容量/参数", "负荷分配", "额定容量", "运行数据", "负荷率", "备注"]
    return ["配电室", "设备", "容量/参数", "负荷分配", "额定容量", "运行数据", "负荷率", "风险", "备注"][:col_count]


def placeholder_row(col_count: int, section_id: str) -> list[str]:
    title = HEADING_BY_NUMBER.get(section_id, ("待核实", 4))[0]
    base = ["待核实", "待核实", "待补充", f"{title}未在上传资料中识别到明确异常", "结合现场资料进一步核实", ""]
    if col_count == 2:
        return [title, "上传资料中未识别到明确异常，建议现场核实。"]
    if col_count == 3:
        return ["待核实", f"{title}未识别到明确异常", "结合现场资料进一步核实"]
    if col_count == 4:
        return ["待核实", f"{title}未识别到明确异常", "结合现场资料进一步核实", "待补充"]
    return (base + ["待核实"] * col_count)[:col_count]


def detail_rows_for_section(model: dict[str, Any], section_id: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in model.get("issue_detail_rows") or []:
        if isinstance(row, dict) and row.get("section_id") == section_id:
            rows.append(row)
    if rows:
        return filter_policy_issue_rows(rows)
    section = (model.get("sections") or {}).get(section_id) or {}
    return filter_policy_issue_rows([row for row in section.get("issue_detail_rows") or [] if isinstance(row, dict)])


def filter_policy_issue_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    filtered = rows
    if policy_bool("table_policy", "show_only_problem_rows", False):
        filtered = [row for row in filtered if row_policy_status(row) != "OK"]
    max_rows = policy_int("table_policy", "max_rows_per_table", 30)
    if max_rows > 0:
        filtered = filtered[:max_rows]
    return filtered


def row_policy_status(row: dict[str, Any]) -> str:
    text = " ".join(
        clean_text(row.get(key))
        for key in ("status_label", "evaluation_status", "status", "conclusion", "risk_level", "priority")
    )
    if re.search(r"\bOK\b|合格|正常|通过", text, flags=re.IGNORECASE):
        return "OK"
    return text or "unknown"


def detail_issue_row(detail: dict[str, Any], model: dict[str, Any]) -> list[Any]:
    return [
        clean_text(detail.get("position") or "待核实位置"),
        clean_text(detail.get("device_type") or "待核实设备"),
        clean_text(detail.get("cabinet_or_loop_no") or "待核实"),
        clean_text(detail.get("issue_description") or detail.get("problem") or "待核实问题"),
        clean_text(detail.get("suggestion") or detail.get("recommendation") or "结合现场资料进一步核实。"),
        {
            "photo_refs": detail.get("photo_refs") or [],
            "photo_status": detail.get("photo_status") or "",
            "photo_manifest": model.get("photo_manifest") or [],
        },
    ]


def detail_markdown_table_row(detail: dict[str, Any]) -> list[Any]:
    photo_refs = [
        clean_text(ref)
        for ref in (detail.get("photo_refs") or [])
        if clean_text(ref)
    ]
    photo_cell = f"[[PHOTO_REFS:{','.join(photo_refs)}]]" if photo_refs else missing_photo_text()
    return [
        clean_text(detail.get("position") or "待核实位置"),
        clean_text(detail.get("device_type") or "待核实设备"),
        clean_text(detail.get("cabinet_or_loop_no") or "待核实"),
        clean_text(detail.get("issue_description") or detail.get("problem") or "待核实问题"),
        clean_text(detail.get("suggestion") or detail.get("recommendation") or "结合现场资料进一步核实。"),
        photo_cell,
    ]


def issue_row(issue: dict[str, Any], col_count: int) -> list[str]:
    problem = clean_text(issue.get("problem") or issue.get("description") or "待核实问题")
    rule = issue_rule(issue) or {}
    recommendation = combine_recommendations(issue.get("recommendation"), rule.get("recommendation"))
    priority = clean_text(issue.get("priority") or issue.get("risk_level") or "待核实")
    location = safe_location(
        issue.get("location"),
        issue.get("position"),
        issue.get("位置"),
        issue.get("位置/范围"),
        extract_location_from_text(problem),
    )
    source_note = clean_text(issue.get("evidence") or issue.get("source") or "上传资料")
    category = safe_device_type(
        issue.get("equipment_type"),
        issue.get("device_type"),
        issue.get("设备类型"),
        infer_device_type(str(issue.get("section_id") or ""), problem),
    )
    if col_count == 2:
        return [category, problem]
    if col_count == 3:
        return [location, problem, recommendation]
    if col_count == 4:
        return [location, problem, recommendation, priority]
    if col_count == 5:
        return [category, problem, priority, recommendation, source_note]
    if col_count == 6:
        return [location, category, "待核实", problem, recommendation, missing_photo_text()]
    if col_count == 7:
        return [category, "待核实", problem, "待核实", priority, recommendation, source_note]
    if col_count == 8:
        return [location, category, "待核实", problem, "待核实", "待核实", priority, recommendation]
    return [location, category, "待核实", problem, "待核实", "待核实", priority, recommendation, "待核实"][:col_count]


def issue_table_headers() -> list[str]:
    return ["位置", "设备类型", "柜号及回路号", "问题描述", "建议", "现场图片"]


def issue_problem_table_row(
    issue: dict[str, Any],
    problem: str,
    recommendation: str,
    suggestion_lines: list[str] | None = None,
) -> list[str]:
    location = safe_location(
        issue.get("location"),
        issue.get("position"),
        issue.get("位置"),
        issue.get("位置/范围"),
        issue.get("room"),
        issue.get("配电室"),
        extract_location_from_text(problem),
    )
    equipment_type = safe_device_type(
        issue.get("equipment_type"),
        issue.get("device_type"),
        issue.get("设备类型"),
        infer_device_type(str(issue.get("section_id") or ""), problem),
    )
    cabinet = first_nonempty(
        issue.get("cabinet"),
        issue.get("loop"),
        issue.get("柜号及回路号"),
        issue.get("柜号"),
        issue.get("回路号"),
        issue.get("circuit"),
    )
    advice = "；".join(suggestion_lines or []) if suggestion_lines else recommendation
    photo = first_nonempty(issue.get("现场图片"), issue.get("photo"), issue.get("image"), issue.get("照片编号"), missing_photo_text())
    return [
        normalize_table_cell_value(location, "位置"),
        normalize_table_cell_value(equipment_type, "设备类型"),
        normalize_table_cell_value(cabinet, "柜号及回路号"),
        clean_text(problem),
        clean_text(advice),
        normalize_table_cell_value(photo, "现场图片"),
    ]


def normalize_issue_markdown_table(
    columns: list[str],
    rows: list[Any],
    issue: dict[str, Any],
) -> tuple[list[str], list[list[Any]]]:
    # Issue-list tables in the report use one fixed schema. Wider statistical
    # tables, such as load-rate summaries, keep their original structure.
    if len(columns) >= 6:
        return columns, [
            [row.get(column, "") for column in columns] if isinstance(row, dict) else list(row)
            for row in rows
        ]

    return issue_table_headers(), [
        normalize_issue_table_row(columns, row, issue)
        for row in rows
    ]


def normalize_issue_table_row(columns: list[str], row: Any, issue: dict[str, Any]) -> list[str]:
    if isinstance(row, dict):
        values_by_column = {clean_text(key): value for key, value in row.items()}
    else:
        raw_values = list(row)
        values_by_column = {
            clean_text(column): raw_values[index] if index < len(raw_values) else ""
            for index, column in enumerate(columns)
        }

    location = safe_location(
        pick_column(values_by_column, "位置/范围", "位置", "范围", "配电室", "区域"),
        issue.get("location"),
        issue.get("position"),
        issue.get("位置"),
        extract_location_from_text(clean_text(issue.get("problem") or issue.get("description"))),
    )
    equipment_type = safe_device_type(
        pick_column(values_by_column, "设备类型", "设备类别", "类型"),
        issue.get("equipment_type"),
        issue.get("device_type"),
        issue.get("设备类型"),
        infer_device_type(str(issue.get("section_id") or ""), clean_text(issue.get("problem") or issue.get("description"))),
    )
    cabinet = pick_column(values_by_column, "柜号及回路号", "柜号", "回路号", "回路", "设备编号")
    problem = pick_column(values_by_column, "问题描述", "问题", "现象描述", "描述")
    check_item = pick_column(values_by_column, "检查项", "检测项", "项目")
    if check_item and problem:
        problem = f"{clean_text(check_item)}：{clean_text(problem)}"
    elif check_item and not problem:
        problem = clean_text(check_item)
    if not problem:
        problem = clean_text(issue.get("problem") or issue.get("description") or "")
    advice = pick_column(values_by_column, "建议", "整改建议", "措施", "建议措施") or first_nonempty(
        issue.get("recommendation"), issue.get("建议")
    )
    photo = pick_column(values_by_column, "现场图片", "图片", "照片", "照片编号") or missing_photo_text()

    return [
        normalize_table_cell_value(location, "位置"),
        normalize_table_cell_value(equipment_type, "设备类型"),
        normalize_table_cell_value(cabinet, "柜号及回路号"),
        normalize_table_cell_value(problem, "问题描述"),
        normalize_table_cell_value(advice, "建议"),
        normalize_table_cell_value(photo, "现场图片"),
    ]


FORBIDDEN_LOCATION_TOKENS = [
    "S4-6",
    "S4-4",
    "S2-1",
    "评估总表",
    "诊断工作用表",
    "收资表",
    "问题汇总表",
    "评估信息汇总表",
    "结论建议汇总表",
    "附件",
    ".xlsx",
    ".xls",
    ".zip",
    "KU-",
    "检测项",
    "知识库",
]

FORBIDDEN_DEVICE_TOPICS = [
    "裸露导体防护",
    "等电位连接与接地问题",
    "标牌标识",
    "谐波风险情况",
    "低压回路剩余电流过大",
    "系统无功补偿与电容柜问题",
    "关键负荷供电路径与应急/备用供电的问题",
    "电压扰动情况",
    "运维组织架构与人员配备",
    "SOP/EOP",
]


def is_invalid_location(value: Any) -> bool:
    text = clean_text(value)
    if not text:
        return True
    if len(text) > 60:
        return True
    return any(token in text for token in FORBIDDEN_LOCATION_TOKENS)


def safe_location(*values: Any) -> str:
    for value in values:
        text = clean_text(value)
        if text and not is_invalid_location(text):
            return text
    return "待核实位置"


def extract_location_from_text(text: Any) -> str:
    clean = clean_text(text)
    patterns = [
        r"([\u4e00-\u9fa5A-Za-z0-9#]+工厂(?:[\u4e00-\u9fa5A-Za-z0-9#]+)?(?:配电房|配电室|车间|动力站|总配|屋顶光伏区域|控制柜间|总电箱)?)",
        r"([\u4e00-\u9fa5A-Za-z0-9#]+(?:配电房|配电室|总配|动力站|车间|电站|总电箱|屋顶光伏区域|控制柜间))",
        r"([A-Za-z0-9#-]+(?:柜|回路|变压器|母线|电容柜|低压柜|高压柜))",
    ]
    for pattern in patterns:
        match = re.search(pattern, clean)
        if match:
            location = re.sub(r"(关键负荷|重要负荷)$", "", match.group(1)).strip()
            if not is_invalid_location(location):
                return location
    return ""


def is_invalid_device_type(value: Any) -> bool:
    text = clean_text(value)
    if not text:
        return True
    if len(text) > 30:
        return True
    return any(token in text for token in FORBIDDEN_DEVICE_TOPICS + FORBIDDEN_LOCATION_TOKENS)


def safe_device_type(*values: Any) -> str:
    for value in values:
        text = clean_text(value)
        if text and not is_invalid_device_type(text):
            return text
    return "待核实设备"


def infer_device_type(section_id: str, problem: str) -> str:
    text = clean_text(problem)
    if section_id == "2.1.2":
        return "关键负荷/供电路径"
    if section_id == "2.1.5" or any(token in text for token in ["电容柜", "无功补偿"]):
        return "电容柜/无功补偿"
    if section_id.startswith("2.2.1") or any(token in text for token in ["谐波", "THD", "电压", "晃电", "暂降"]):
        return "电能质量/敏感负荷"
    if section_id == "2.2.2.1" or any(token in text for token in ["温度", "发热", "过热"]):
        return "低压配电"
    if section_id.startswith("2.3"):
        return "保护装置/低压配电"
    if any(token in text for token in ["桥架", "母线", "电缆"]):
        return "电缆桥架/母线"
    if any(token in text for token in ["接地", "等电位", "PE"]):
        return "接地系统"
    return "低压配电"


def pick_column(values_by_column: dict[str, Any], *names: str) -> Any:
    compact_lookup = {re.sub(r"\s+", "", key): value for key, value in values_by_column.items()}
    for name in names:
        compact_name = re.sub(r"\s+", "", name)
        value = compact_lookup.get(compact_name)
        if clean_text(value):
            return value
    return ""


def first_nonempty(*values: Any) -> Any:
    for value in values:
        if clean_text(value):
            return value
    return ""


def remove_body_tables_after_start(document: Document, start_paragraph_index: int) -> None:
    body = document._body._element
    start_element = document.paragraphs[start_paragraph_index]._element
    start_body_index = list(body).index(start_element)
    for table in list(document.tables):
        if list(body).index(table._element) > start_body_index:
            table._element.getparent().remove(table._element)


def mark_fields_for_update(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def compact_report_flow(document: Document) -> None:
    """Remove pagination leftovers around generated analysis/advice blocks."""
    for paragraph in document.paragraphs:
        text = clean_inline_marks(paragraph.text)
        if not text:
            continue
        if not is_flow_sensitive_paragraph(text):
            continue
        strip_template_breaks(paragraph)
        remove_paragraph_numbering(paragraph)
        remove_paragraph_tabs(paragraph)
        if is_bullet_line(text):
            body = re.sub(r"^[-*•➢]\s*", "", text).strip()
            write_bullet_paragraph(paragraph, body)
            continue
        if is_time_group_heading_line(text):
            write_time_group_heading_paragraph(paragraph, text)
            continue
        if is_exact_label_line(text):
            write_label_paragraph(paragraph, text)
            continue
        pf = paragraph.paragraph_format
        pf.keep_with_next = False
        pf.keep_together = False
        pf.page_break_before = False
        if is_inline_analysis_line(text):
            pf.space_before = Pt(0)
            pf.space_after = Pt(BODY_SPACE_AFTER_PT)
            pf.line_spacing = BODY_LINE_SPACING


def is_flow_sensitive_paragraph(text: str) -> bool:
    return (
        is_exact_label_line(text)
        or is_time_group_heading_line(text)
        or is_inline_analysis_line(text)
        or text.startswith("•")
        or text.startswith("鈥?")
    )


def add_report_content(document: Document, report_text: str, model: dict[str, Any] | None = None) -> None:
    configure_styles(document)
    for block in parse_markdown_blocks(report_text):
        if block.kind == "title":
            continue
        if block.kind == "table":
            write_markdown_table(
                document,
                document.add_paragraph(),
                list(block.raw_lines),
                model,
            )
            continue
        paragraph = document.add_paragraph()
        write_line_to_paragraph(paragraph, block.text)


def normalize_report_text(report_text: str) -> str:
    text = re.sub(r"<think>.*?</think>", "", report_text or "", flags=re.DOTALL | re.IGNORECASE).strip()
    return trim_to_required_report(text)


def trim_to_required_report(report_text: str) -> str:
    stop_patterns = [
        r"资料缺失/待补充/待核实.*声明",
        r"报告编制说明",
        r"配电安全专家咨询报告\s*[·\-]\s*草稿",
        r"如有疑问",
    ]
    kept: list[str] = []
    for raw_line in report_text.splitlines():
        line = clean_inline_marks(raw_line.strip())
        if any(re.search(pattern, line) for pattern in stop_patterns):
            break
        kept.append(raw_line)
    return "\n".join(kept).strip()


def markdown_table(headers: list[str], rows: list[list[Any]]) -> list[str]:
    result = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        result.append("| " + " | ".join(clean_table_cell(value) for value in row) + " |")
    return result


def status_block_lines(
    situation: str,
    conclusion: str,
    risk: str,
    suggestions: list[str],
) -> list[str]:
    """Render every KU paragraph with the required fixed label structure."""
    lines: list[str] = []
    situation = clean_text(situation)
    conclusion = clean_text(conclusion)
    risk = clean_text(risk)
    suggestions = [clean_text(item).strip("。；;") for item in suggestions if clean_text(item)]
    if situation:
        lines.extend(["【现状描述】", situation, ""])
    lines.extend(["【结论】", conclusion or "该检测项因资料不足，本次评估暂按待核实处理。", ""])
    if risk:
        lines.extend(["【风险分析】", risk, ""])
    lines.append("【建议】")
    if suggestions:
        for suggestion in suggestions:
            lines.append(f"• {suggestion}。")
    else:
        lines.append("• 结合现场资料制定整改措施。")
    lines.append("")
    return lines


def chapter3_summary_text(model: dict[str, Any], issues: list[dict[str, Any]], action_rows: list[dict[str, Any]]) -> str:
    """Build chapter 3.1 from structured KU status results, not generic prose."""
    ng_count = attention_count = ok_count = 0
    for issue in issues:
        label = issue_status_label(issue, issue_rule(issue) or {})
        if label == "NG":
            ng_count += 1
        elif label == "一般":
            attention_count += 1
        elif label == "OK":
            ok_count += 1
    section_counts = model.get("issue_counts_by_section") or []
    section_parts = []
    for item in section_counts:
        if not isinstance(item, dict):
            continue
        title = clean_text(item.get("section_title"))
        count = int(item.get("issue_count") or 0)
        if title and count:
            section_parts.append(f"{title}{count}项")
    distribution = "、".join(section_parts) if section_parts else "暂无明确分布"
    action_count = len(action_rows)
    return (
        f"本次评估共形成 NG 项 {ng_count} 项、一般/关注项 {attention_count} 项、OK 项 {ok_count} 项。"
        f"问题分布为：{distribution}。"
        f"其中需纳入改善行动列表的问题共 {action_count} 项，应按照第 3.2 节优先级安排整改、复核与闭环。"
    )


def chapter3_summary_text(model: dict[str, Any], issues: list[dict[str, Any]], action_rows: list[dict[str, Any]]) -> str:
    """Build chapter 3.1 according to the rulebook's five-dimension summary template."""
    dimensions = [
        ("2.1", "系统架构"),
        ("2.2", "环境工况"),
        ("2.3", "故障保护"),
        ("2.4", "设备/元件状态"),
        ("2.5", "运维管理"),
    ]
    grouped: dict[str, list[dict[str, Any]]] = {prefix: [] for prefix, _ in dimensions}
    for issue in issues:
        section_id = str(issue.get("section_id") or "")
        for prefix, _title in dimensions:
            if section_id == prefix or section_id.startswith(prefix + "."):
                grouped[prefix].append(issue)
                break

    risk_lines: list[str] = []
    dimension_blocks: list[str] = []
    for index, (prefix, title) in enumerate(dimensions, 1):
        items = grouped[prefix]
        ng_items = [item for item in items if normalized_status_label(item) == "NG"]
        attention_items = [item for item in items if normalized_status_label(item) == "一般"]
        evaluation = dimension_evaluation_word(len(ng_items), len(attention_items))
        if ng_items or attention_items:
            risk_lines.append(f"{index}）{dimension_risk_brief(title, ng_items, attention_items)}")
        dimension_blocks.extend([
            title,
            dimension_summary_sentence(title, evaluation, items, ng_items, attention_items),
            "",
        ])

    if not risk_lines:
        risk_lines.append("1）本次资料未识别到明确 NG 或一般项，建议继续按既定巡检周期跟踪复核。")

    lines = ["风险汇总如下："]
    lines.extend(risk_lines)
    lines.extend(["", "详述见下：", ""])
    lines.extend(dimension_blocks)
    return "\n".join(lines).strip()


def visual_environment_lines(model: dict[str, Any]) -> list[str]:
    visual_environment = model.get("visual_environment") or {}
    chart_models = model.get("chart_models") or visual_environment.get("chart_models") or []
    if not isinstance(chart_models, list):
        return []
    lines: list[str] = []
    allowed_chart_ids = {
        "risk_panorama",
        "urgent_issue_top3",
        "dimension_status_distribution",
        "priority_matrix",
        "evidence_completeness",
        "data_gap_analysis",
    }
    if not policy_bool("chapter_3_policy", "include_data_gap_analysis", True):
        allowed_chart_ids.remove("data_gap_analysis")
    for chart in chart_models:
        if not isinstance(chart, dict):
            continue
        chart_id = clean_text(chart.get("chart_id"))
        if chart_id not in allowed_chart_ids:
            continue
        columns = [clean_text(column) for column in (chart.get("columns") or []) if clean_text(column)]
        rows = chart.get("rows") or []
        if not columns or not rows:
            continue
        title = clean_text(chart.get("title") or "图表")
        lines.append(title)
        table_rows: list[list[Any]] = []
        for row in rows[:30]:
            if isinstance(row, dict):
                table_rows.append([clean_chart_cell_value(column, row.get(column, "")) for column in columns])
            elif isinstance(row, list):
                table_rows.append([
                    clean_chart_cell_value(columns[index] if index < len(columns) else "", value)
                    for index, value in enumerate(row[:len(columns)])
                ])
        if table_rows:
            lines.extend(markdown_table(columns, table_rows))
            lines.append("")
    return lines


def clean_chart_cell_value(column: Any, value: Any) -> str:
    column_text = clean_text(column)
    text = clean_text(value)
    if "建议" in column_text or "措施" in column_text:
        return concise_recommendation_text(text)
    if "问题" in column_text or "主要问题" in column_text:
        return shorten_table_text(text, 95)
    return shorten_table_text(text, 80 if "章节" not in column_text else 45)


def shorten_table_text(text: str, max_chars: int) -> str:
    text = clean_text(text)
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rstrip("，；;、。 ") + "…"


def concise_recommendation_text(value: Any, max_points: int = 2, max_point_chars: int = 34) -> str:
    text = clean_text(value)
    if not text:
        return "补充现场资料并制定闭环整改措施。"
    text = re.sub(r"依据[:：].*?(?=；|;|。|$)", "", text)
    text = re.sub(r"预期.*$", "", text)
    if "方法论建议：" in text:
        text = text.split("方法论建议：", 1)[1]
    text = text.replace("当前资料建议：", "").replace("方法论建议：", "")
    text = re.sub(r"\s+", " ", text).strip("；;。 ")

    numbered = [
        item.strip("；;。 ，,")
        for item in re.findall(r"(?:^|[；;。\n])\s*\d+\s*[\.、）)]\s*([^；;。\n]+)", text)
        if item.strip("；;。 ，,")
    ]
    if not numbered:
        numbered = [
            item.strip("；;。 ，,")
            for item in re.split(r"[；;。\n]+", text)
            if item.strip("；;。 ，,")
        ]
    points = []
    for item in numbered:
        item = re.sub(r"^(建议|应|需|当前资料建议)\s*", "", item).strip()
        if not item:
            continue
        points.append(shorten_table_text(item, max_point_chars))
        if len(points) >= max_points:
            break
    if not points:
        return shorten_table_text(text, max_point_chars * max_points)
    return "；".join(f"{index}. {point}" for index, point in enumerate(points, 1))


def normalized_status_label(issue: dict[str, Any]) -> str:
    label = issue_status_label(issue, issue_rule(issue) or {})
    if label == "NG" or "NG" in str(label).upper():
        return "NG"
    if any(token in str(label) for token in ["一般", "关注"]):
        return "一般"
    if label == "OK" or "OK" in str(label).upper():
        return "OK"
    return str(label)


def dimension_evaluation_word(ng_count: int, attention_count: int) -> str:
    if ng_count == 0 and attention_count == 0:
        return "完善"
    if ng_count == 0 and 1 <= attention_count <= 2:
        return "较为完善"
    if 1 <= ng_count <= 2 and attention_count <= 2:
        return "存在不足"
    return "问题较多，亟需整改"


def issue_short_topic(issue: dict[str, Any]) -> str:
    text = clean_text(issue.get("problem") or issue.get("description") or issue.get("section_title") or "")
    text = re.sub(r"^(现象|问题|结论|风险)[:：]\s*", "", text)
    text = re.sub(r"[。；;].*$", "", text).strip()
    if len(text) > 34:
        text = text[:34].rstrip() + "…"
    return text


def dimension_risk_brief(title: str, ng_items: list[dict[str, Any]], attention_items: list[dict[str, Any]]) -> str:
    focus_items = (ng_items + attention_items)[:3]
    topics = [topic for item in focus_items if (topic := issue_short_topic(item))]
    if topics:
        return f"{title}方面存在{'、'.join(topics)}等问题；"
    return f"{title}方面存在需跟踪复核的问题；"


def dimension_summary_sentence(
    title: str,
    evaluation: str,
    items: list[dict[str, Any]],
    ng_items: list[dict[str, Any]],
    attention_items: list[dict[str, Any]],
) -> str:
    if not items:
        return f"{title}本次未识别到明确异常，整体评价为{evaluation}，建议按既定周期保持巡检、检测和台账复核。"
    topics = [topic for item in (ng_items + attention_items)[:4] if (topic := issue_short_topic(item))]
    topic_text = "、".join(topics) if topics else "相关核查项"
    prefix = f"{title}{evaluation}，本次识别 NG 项 {len(ng_items)} 项、一般/关注项 {len(attention_items)} 项。"
    if title == "系统架构":
        extra = "若涉及供电路径单一或自动切换能力不足，应重点提升供电连续性和故障恢复能力；若涉及负荷分配问题，应同步复核容量裕度和回路负载匹配。"
    elif title == "环境工况":
        extra = "应重点关注电能质量、设备发热、局放和物理环境条件对设备寿命及故障概率的影响。"
    elif title == "故障保护":
        extra = "应围绕保护方案、定值配合、零序/漏电及过压欠压防护能力进行专项复核。"
    elif title == "设备/元件状态":
        extra = "应重点处理设备品质、安装规范、标识接地、封堵和带病运行等影响本质安全的问题。"
    else:
        extra = "应完善 SOP/EOP、图纸资料、巡检维护、智能化手段、LOTO、生命周期和备件管理，推动问题闭环。"
    return f"{prefix}主要集中在{topic_text}。{extra}"


def clean_action_remark(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if text.startswith("KU-"):
        return text
    if re.search(r"\.(xlsx|xls|zip|docx|png|jpg|jpeg)\b", text, flags=re.IGNORECASE):
        return ""
    if len(text) > 40:
        return ""
    return text


def chapter1_background_text(model: dict[str, Any]) -> str:
    """Build 1.1 from the fixed report-generation rule template."""
    year = model_text_value(model, ["评估年份", "assessment_year", "year"]) or "X"
    month = model_text_value(model, ["评估月份", "assessment_month", "month"]) or "X"
    evaluator = model_text_value(model, ["评估方名称", "evaluator_name", "assessor_name"]) or "X"
    client = model_text_value(model, ["委托方名称", "client_name", "customer_name"]) or "X"
    location = model_text_value(model, ["项目地点", "project_location", "location"]) or "X"
    return (
        f"{year}年{month}月，我方（{evaluator}），受{client}委托，对位于{location}进行配电系统安全性评估，"
        "旨在提升现有配电系统安全性及供电力可靠性，确保电力设备长期稳定运行。\n\n"
        "此电力系统评估报告综合了我方专家及贵方配电运维人员对现场的电气设备、电力网络及维护工作组织"
        "共同进行的调研分析与评价建议。报告的内容是基于在现场进行电力安全咨询服务时所收集的信息撰写的。"
    )


def chapter1_overview_text(model: dict[str, Any]) -> str:
    """Build 1.2 from the fixed report-generation rule template."""
    highest_voltage = model_text_value(model, ["最高电压等级", "highest_voltage_level", "max_voltage_level"]) or "X"
    lowest_voltage = model_text_value(model, ["最低电压等级", "lowest_voltage_level", "min_voltage_level"]) or "X"
    kpi_count = model_text_value(model, ["KPI数量", "kpi_count", "ku_count"]) or "X"
    return (
        f"评估期间，我方评估人员与贵方运维人员共同合作，对厂房内的，从{highest_voltage}至{lowest_voltage}的整个电力系统进行了全面审视，包含了以下五个维度，共{kpi_count}个KPI细项。\n"
        "• 配电系统架构问题\n"
        "• 环境工况风险\n"
        "• 针对故障的保护\n"
        "• 配电设备/元件内在风险\n"
        "• 运维管理与风险管控机制\n\n"
        "全面细致地进行了风险排查评估，评估细节请参见如下第2章内容。\n"
        "建议纵览可参见第3章。"
    )


def chapter1_overall_evaluation(ng_count: int, attention_count: int) -> str:
    """Use the rule-file thresholds for the automatic overall evaluation."""
    if ng_count == 0 and attention_count <= 2:
        return "系统整体水平良好，系统架构完备性较高、运维管理规范"
    if ng_count == 0 and 3 <= attention_count <= 5:
        return "系统整体水平较好，系统架构完备性较高，但仍存在一些可提升空间"
    if 1 <= ng_count <= 3 and attention_count <= 5:
        return "系统整体水平尚可，部分领域存在安全风险，建议按计划整改"
    if 4 <= ng_count <= 6:
        return "系统存在较多安全隐患，建议尽快制定整改计划并落实"
    if ng_count > 6:
        return "系统安全风险较高，建议立即采取整改措施"
    return "系统整体水平尚可，部分领域存在安全风险，建议按计划整改"


def model_text_value(model: dict[str, Any], keys: list[str]) -> str:
    sources: list[Any] = [
        model,
        model.get("project_info"),
        model.get("basic_info"),
        model.get("assessment_info"),
        model.get("metadata"),
        model.get("summary"),
    ]
    for source in sources:
        if not isinstance(source, dict):
            continue
        for key in keys:
            value = clean_text(source.get(key))
            if value:
                return value
    return ""


def get_section_issues(sections: dict[str, Any], section_id: str) -> list[dict[str, Any]]:
    section = sections.get(section_id) if isinstance(sections, dict) else None
    if isinstance(section, dict):
        return [item for item in section.get("issues") or [] if isinstance(item, dict)]
    return []


def count_issues_under_prefix(sections: dict[str, Any], prefix: str) -> int:
    count = 0
    for section_id, section in (sections or {}).items():
        if str(section_id).startswith(prefix + ".") and isinstance(section, dict):
            count += int(section.get("issue_count") or len(section.get("issues") or []))
    return count


def primary_section_rule(
    rules_by_section: dict[str, Any],
    sections: dict[str, Any],
    section_id: str,
) -> dict[str, Any] | None:
    rules = rules_by_section.get(section_id) if isinstance(rules_by_section, dict) else None
    if isinstance(rules, list) and rules:
        return rules[0] if isinstance(rules[0], dict) else None
    section = sections.get(section_id) if isinstance(sections, dict) else None
    if isinstance(section, dict):
        rules = section.get("knowledge_rules")
        if isinstance(rules, list) and rules:
            return rules[0] if isinstance(rules[0], dict) else None
    return None


def issue_rule(issue: dict[str, Any]) -> dict[str, Any] | None:
    rule = issue.get("knowledge_rule")
    return rule if isinstance(rule, dict) else None


def combine_recommendations(issue_recommendation: Any, rule_recommendation: Any) -> str:
    issue_text = clean_text(issue_recommendation)
    rule_text = clean_text(rule_recommendation)
    if issue_text and rule_text and issue_text not in rule_text:
        return f"当前资料建议：{issue_text}。方法论建议：{rule_text}"
    if rule_text:
        return rule_text
    if issue_text:
        return issue_text
    return "结合现场资料制定整改措施。"


def section_intro(section_number: str) -> str:
    intros = {
        "2.1": "系统架构问题是关乎系统可靠性和稳定性，即是否有备用供电方案；是否能正确应对上下游的故障或扰动，如切除/保护/替代，而不会影响大局。这方面的问题会带来隐性风险。",
        "2.2": "环境工况是配电系统/设备老化、故障、失效的主要外在因素，评估中我们将其分为电气环境工况和物理环境工况，现逐一审视如下。",
        "2.3": "可靠完善的保护是电力系统的主要功能和关键任务，我们应关注和审视以下几点：过流/速断保护方案；零序/接地故障的保护；过电压（欠电压）的保护。现将保护方面的评估结果陈述如下。",
        "2.4": "对范围内高低压配电设备进行审视后发现，配电设备/元件内在风险主要集中在设备品质、安装规范性、选型正确性和带病运行问题，现结合配置与选型、安装规范性、带病运行和末端配电抽查逐项核实。",
        "2.5": "运维管理与风险管控机制用于评价配电系统运行维护体系、图纸资料、巡检维护、智能化手段、配电室装备与LOTO流程、设备生命周期和备件管理的完整性，现结合资料逐项核实。",
    }
    return intros.get(section_number, "")


def issue_status_label(issue: dict[str, Any], rule: dict[str, Any]) -> str:
    explicit = clean_text(
        issue.get("status_label")
        or issue.get("evaluation_status")
        or issue.get("status")
    )
    explicit_lower = explicit.lower()
    if explicit:
        if "数据冲突" in explicit:
            return "数据冲突"
        if "数据缺失" in explicit:
            return "数据缺失"
        if explicit_lower in {"ng", "不合格"} or "ng" in explicit_lower:
            return "NG"
        if any(token in explicit for token in ["一般", "关注"]):
            return "一般"
        if explicit_lower == "ok" or any(token in explicit for token in ["正常", "合格"]):
            return "OK"

    problem_text = clean_text(issue.get("problem") or issue.get("description"))
    problem_lower = problem_text.lower()
    if problem_lower == "ok" or problem_text in {"正常", "合格"}:
        return "OK"

    evidence_text = " ".join(
        clean_text(value)
        for value in [
            issue.get("priority"),
            issue.get("risk_level"),
            problem_text,
        ]
    ).lower()
    if "数据冲突" in evidence_text or "矛盾" in evidence_text:
        return "数据冲突"
    if any(token in evidence_text for token in ["ng", "不合格", "异常", "超标", "无备用", "路径单一", "未投入", "过载"]):
        return "NG"
    if any(token in evidence_text for token in ["一般", "关注", "中", "medium", "偏高", "偏低"]):
        return "一般"
    if any(token in evidence_text for token in ["ok", "正常", "合格", "low", "低"]):
        return "OK"
    return "待核实"


def status_conclusion_text(section_id: str, item_name: str, status_label: str, problem: str) -> str:
    """Build KU conclusion text using the fixed acceptance template."""
    item_name = clean_text(item_name) or HEADING_BY_NUMBER.get(section_id, ("该检测项", 3))[0]
    problem = clean_text(problem)
    detail = strip_status_tail(problem)
    if status_label == "OK":
        return f"经评估，{item_name}满足要求/未发现异常。{detail or '当前上传资料未显示该项存在明确异常。'}。"
    if status_label == "一般":
        risk_hint = attention_risk_hint(section_id, detail)
        return f"经评估，{item_name}存在关注项/部分指标接近限值。{detail}。{risk_hint}。"
    if status_label == "NG":
        severity = severity_hint(section_id, detail)
        return f"经评估，{item_name}存在不合格项/安全隐患。{detail}。{severity}。"
    if status_label == "数据冲突":
        return f"经评估，{item_name}存在数据异常，当前资料口径不一致。{detail}。建议核实原始记录后重新确认判定。"
    if status_label == "数据缺失":
        return f"{item_name}因数据缺失，本次评估未能给出明确结论。建议补充{missing_fields_hint(section_id)}后重新评估。"
    return f"{item_name}因资料不足，本次评估暂按待核实处理。建议补充{missing_fields_hint(section_id)}后重新评估。"


def strip_status_tail(text: str) -> str:
    text = clean_text(text)
    text = re.sub(
        r"[，,；;]?\s*(存在负荷率\s*80%-100%\s*的关注项|存在负荷率超过\s*100%\s*的过载风险|需持续监控并预留负荷调整方案|当前可识别变压器负荷率均不超过\s*80%)\s*[。.]?",
        "",
        text,
    )
    return text.strip(" ，。") or clean_text(text)


def attention_risk_hint(section_id: str, detail: str) -> str:
    if section_id == "2.1.1":
        if "80" in detail or "负荷率" in detail:
            return "后续产能提升或负荷增长后可能面临过载风险"
    if section_id == "2.1.5" and any(token in detail for token in ["谐波", "电容柜", "无功"]):
        return "谐波环境下继续使用普通电容柜可能带来运行可靠性风险"
    if section_id == "2.2.1.1" and any(token in detail for token in ["谐波", "THD", "THDu"]):
        return "谐波水平接近限值时可能加剧设备发热，并增加与电容柜共振的风险"
    if section_id == "2.2.1.2" and any(token in detail for token in ["晃电", "暂降", "扰动", "失电"]):
        return "后续敏感负荷增加或扰动频次上升后可能造成设备停机和生产中断"
    if section_id == "2.2.1.3" and any(token in detail for token in ["冲击", "启动"]):
        return "频繁启动或冲击负荷可能引发母线电压波动并影响同母线设备稳定运行"
    if section_id == "2.2.2.1" and any(token in detail for token in ["温度", "温升", "发热", "过热"]):
        return "温升异常若持续存在，可能造成连接点劣化并进一步放大接触电阻"
    if section_id == "2.2.2.2" and "局放" in detail:
        return "局放检测和监测不足时，绝缘缺陷可能无法及时发现"
    if section_id == "2.2.2.3" and any(token in detail for token in ["温湿度", "水", "防鼠", "粉尘", "凝露", "环境"]):
        return "物理环境条件变化可能降低绝缘水平并影响配电设备长期可靠运行"
    if section_id == "2.3.1" and any(token in detail for token in ["保护", "定值", "越级", "拒动", "误动", "选择性"]):
        return "保护配合或定值缺陷可能在故障时放大停电范围"
    if section_id == "2.3.2" and any(token in detail for token in ["剩余电流", "漏电", "零序", "监测"]):
        return "剩余电流或漏电监测不足可能使接地故障和漏电隐患难以及时发现"
    if section_id == "2.3.3" and any(token in detail for token in ["电涌", "SPD", "避雷器", "过电压", "过压"]):
        return "电涌保护投运或试验记录不足可能削弱过电压防护能力"
    return "后续运行条件变化后可能进一步放大该项风险"


def severity_hint(section_id: str, detail: str) -> str:
    if section_id == "2.1.1" and any(token in detail for token in [">100", "超过 100", "100%"]):
        return "治理紧迫性：高"
    if section_id in {"2.1.2", "2.1.3", "2.1.4"}:
        return "治理紧迫性：高"
    if section_id in {"2.2.1.1", "2.2.1.2", "2.2.2.1", "2.2.2.2"}:
        return "治理紧迫性：高"
    if section_id in {"2.3.1", "2.3.2", "2.3.3"}:
        return "治理紧迫性：高"
    if section_id in {"2.2.1.3", "2.2.2.3"}:
        return "治理紧迫性：中"
    return "治理紧迫性：需结合风险等级和现场条件确定"


def missing_fields_hint(section_id: str) -> str:
    hints = {
        "2.1.1": "变压器容量、运行电流、负荷率曲线及上下级回路容量匹配资料",
        "2.1.2": "关键负荷清单、供电路径图、备用电源配置和倒闸方案",
        "2.1.3": "ATS/双投开关配置、切换逻辑、切换时间和年度试验记录",
        "2.1.4": "两进线一母联连锁配置、倒闸防误操作记录和连锁试验记录",
        "2.1.5": "功率因数、电容柜投运状态、消谐措施、谐波水平和电容容量检测记录",
        "2.2.1.1": "10kV/0.4kV 母线 THDu 数据、电能质量分析报告和谐波治理设备运行记录",
        "2.2.1.2": "电压暂升、暂降、波动、晃电事件记录及敏感负荷抗扰动措施",
        "2.2.1.3": "大功率电机启动频次、启动电流、冲击负荷功率和电压波动记录",
        "2.2.2.1": "低压配电设备红外测温记录、连接点温升数据和热点缺陷复测记录",
        "2.2.2.2": "高压配电设备局放试验、在线监测、放电痕迹和绝缘状态记录",
        "2.2.2.3": "配电室温湿度、粉尘、水浸、防鼠、防火、安全用具和通道合规记录",
        "2.3.1": "保护定值表、短路电流计算、TCC选择性配合分析和保护传动试验记录",
        "2.3.2": "10kV零序保护启用情况、低压总剩余电流、剩余电流占比和漏保投运记录",
        "2.3.3": "避雷器状态、0.4kV电涌保护器配置投运状态和SPD试验记录",
    }
    return hints.get(section_id, "现场记录、检测数据或设备台账")


def matched_risk_analysis(section_id: str, problem: str, risk_analysis: str, status_label: str) -> str:
    if status_label not in {"一般", "NG", "数据冲突"}:
        return ""
    problem = clean_text(problem)
    matched = section_risk_matches(section_id, problem)
    if matched:
        return "；".join(matched)
    # Fall back to filtered knowledge-base risk clauses only when they directly mention
    # keywords present in the actual finding. This avoids copying unrelated template risks.
    clauses = split_recommendations(risk_analysis, limit=8)
    filtered: list[str] = []
    problem_tokens = meaningful_tokens(problem)
    for clause in clauses:
        if any(token in clause for token in problem_tokens):
            filtered.append(clause)
    return "；".join(filtered[:3])


def section_risk_matches(section_id: str, problem: str) -> list[str]:
    rules: dict[str, list[tuple[list[str], str]]] = {
        "2.1.1": [
            (["负荷率", "80"], "变压器长期高负载运行会加速绝缘老化、缩短使用寿命"),
            ([">100", "超过 100", "大于100"], "变压器过热、绝缘失效，存在火灾风险"),
            (["下级容量", "上级容量"], "上下级容量不匹配时存在过载跳闸或设备损坏风险"),
            (["重要负荷", "冲击负荷", "混接"], "重要负荷与冲击负荷混接会影响重要负荷供电可靠性，可能导致生产中断"),
        ],
        "2.1.2": [
            (["路径单一", "单电源"], "关键负荷供电路径单一时，某一电力元件失效即可能造成生产中断"),
            (["无备用", "缺乏备用", "备用供电"], "无备用供电方式会降低局部维护检修的可实施性"),
            (["恢复时间", "人工恢复"], "供电路径沿线设备故障后依赖人工恢复，会拉长应急恢复时间"),
        ],
        "2.1.3": [
            (["无自动切换", "无法自动切换", "未配置"], "无自动切换功能时，供电路径故障后需人工恢复，故障恢复时间长"),
            (["未验证", "试验"], "切换功能未定期验证，实际需要时可能切换失败"),
            (["切换时间"], "切换时间不满足负荷要求时，敏感设备存在停机风险"),
            (["切换逻辑"], "切换逻辑设置不当可能导致非计划停电范围扩大"),
        ],
        "2.1.4": [
            (["环流", "并联"], "两路电源并联可能产生大环流并导致保护跳闸"),
            (["返送"], "返送电可能导致检修线路带电，造成人身安全风险"),
            (["连锁失效", "联锁失效", "未配置", "闭锁"], "连锁失效或未配置会增加误操作导致重大事故的风险"),
        ],
        "2.1.5": [
            (["功率因数"], "功率因数偏低会导致力调电费罚款并增加运营成本"),
            (["无功补偿", "电容柜失效"], "无功补偿不足会降低系统能效并增加变压器满载率"),
            (["谐波", "THD", "THDu", "电容柜"], "谐波环境下使用普通电容柜可能导致电容过载发热，存在爆炸或火灾风险"),
            (["容量衰减"], "电容容量衰减会造成补偿效果不足，并可能因频繁投切损坏控制器"),
        ],
        "2.2.1.1": [
            (["THDu", "谐波", "超限", "超标"], "THDu超标会加速电子装置老化，并加剧变压器、电机和电缆发热"),
            (["谐波", "电容"], "谐波与电容柜共振时，可能引发电容过载、爆炸或火灾风险"),
            (["保护", "误动"], "高次谐波可能引起保护装置误动作，导致非计划停电"),
            (["零线", "过流"], "高次谐波可能导致零线过流，增加电气火灾风险"),
        ],
        "2.2.1.2": [
            (["晃电", "暂降", "电压扰动"], "电压暂降或晃电可能导致变频器、PLC、DCS等敏感设备停机"),
            (["生产中断", "停机"], "频繁电压扰动会造成生产中断，并增加设备异常停机风险"),
            (["过压", "暂升"], "暂态过电压可能损坏设备绝缘和电子元件"),
        ],
        "2.2.1.3": [
            (["频繁启动", "启动电流"], "频繁启动会造成电压波动和设备热冲击，缩短设备寿命"),
            (["冲击负荷"], "冲击负荷可能导致母线电压暂降，影响同母线敏感设备正常运行"),
            (["保护误动"], "启动电流过大可能引起保护误动作或电压暂降"),
        ],
        "2.2.2.1": [
            (["温升", "发热", "过热"], "低压配电设备温升异常会加速绝缘老化并缩短设备寿命"),
            (["接触电阻", "连接"], "连接点温升异常可能导致接触电阻增大，形成发热恶化循环"),
            (["火灾", "短路"], "严重发热可能引发短路、火灾或母排熔断等事故"),
        ],
        "2.2.2.2": [
            (["局放", "放电"], "局部放电会逐步劣化绝缘，严重时可能导致绝缘击穿"),
            (["电弧", "击穿"], "高压设备绝缘击穿可能引发电弧故障，造成人员伤害和设备损坏"),
            (["积水", "潮湿"], "电缆沟积水或高湿环境会降低绝缘裕度，并放大局放风险"),
        ],
        "2.2.2.3": [
            (["温湿度", "高温"], "温湿度异常会加速绝缘老化，并降低设备载流能力"),
            (["凝露", "潮湿"], "湿度过高或凝露可能引发沿面放电、闪络和短路风险"),
            (["粉尘"], "粉尘积聚会降低绝缘水平，并可能诱发闪络或爆燃风险"),
            (["水浸", "积水", "防水"], "水浸和积水会造成电缆绝缘下降，增加短路和接地故障风险"),
            (["防鼠", "小动物"], "防鼠措施不足可能导致小动物进入柜体造成短路故障"),
        ],
        "2.3.1": [
            (["越级跳闸"], "越级跳闸会扩大停电范围，影响配电系统供电连续性"),
            (["拒动", "不动作"], "保护拒动会导致故障不能及时切除，并可能加剧设备损坏"),
            (["误动"], "保护误动会造成正常运行时非计划停电"),
            (["保护漏洞", "定值", "选择性", "保护配合"], "保护方案、定值或选择性配合不合理会降低故障切除可靠性"),
        ],
        "2.3.2": [
            (["零序", "未启用"], "零序保护未启用时接地故障无法快速切除，可能扩大事故影响"),
            (["剩余电流", "漏电"], "剩余电流异常通常提示下游线路或设备存在零地混接、漏电等隐患"),
            (["漏保", "未投入"], "漏保未投入会削弱人身触电和电气火灾防护"),
            (["监测"], "剩余电流监测未投入会降低漏电隐患的及时发现能力"),
        ],
        "2.3.3": [
            (["雷击", "避雷器"], "避雷器缺失或失效会削弱雷击过电压防护，可能造成绝缘击穿"),
            (["操作过电压", "过电压", "过压"], "操作过电压可能损坏敏感设备和电子元件"),
            (["SPD", "电涌", "未投入"], "电涌保护器未投入时，雷击或操作过电压可能直接作用于下游设备"),
            (["老化", "失效"], "SPD老化失效会降低过电压泄放能力"),
        ],
    }
    matched: list[str] = []
    for tokens, text in rules.get(section_id, []):
        if any(token and token in problem for token in tokens) and text not in matched:
            matched.append(text)
    return matched


def meaningful_tokens(text: str) -> list[str]:
    candidates = [
        "负荷率", "路径单一", "单电源", "备用", "自动切换", "ATS", "未验证", "环流", "返送", "连锁", "闭锁",
        "功率因数", "无功", "电容柜", "谐波", "THD", "THDu", "晃电", "暂降", "电压扰动", "冲击负荷",
        "频繁启动", "启动电流", "温升", "发热", "过热", "红外", "局放", "放电", "温湿度", "凝露",
        "粉尘", "水浸", "积水", "防鼠", "小动物", "保护定值", "越级跳闸", "拒动", "误动", "选择性",
        "保护配合", "剩余电流", "漏电", "零序", "漏保", "电涌", "SPD", "避雷器", "过电压", "雷击",
    ]
    return [token for token in candidates if token in text]


def concise_risk_analysis(problem: str, risk_analysis: str, status_label: str) -> str:
    problem = clean_text(problem)
    if status_label == "OK":
        return "当前上传资料未显示该项存在明确异常，建议保留周期性检查和验证记录。"
    if status_label == "数据缺失":
        return "资料尚不足以形成确定风险结论，需结合现场记录和检测数据复核。"
    candidates = split_recommendations(risk_analysis, limit=2)
    if candidates:
        risk = "；".join(candidates)
    elif status_label == "待核实":
        risk = "资料尚不足以形成确定风险结论，需结合现场记录和检测数据复核。"
    else:
        risk = "该问题可能影响配电系统安全性、供电可靠性或运维可控性。"
    if problem and problem not in risk:
        return f"结合资料中“{problem}”的表现，{risk}"
    return risk


def split_recommendations(text: Any, limit: int = 4) -> list[str]:
    value = clean_text(text)
    if not value:
        return []
    value = re.sub(r"^(当前资料建议|方法论建议|整改建议|建议)[:：]", "", value).strip()
    parts = re.split(r"[；;\n]+|(?:(?<=。)\s*)|(?:\d+\s*[.．、）)]\s*)", value)
    cleaned: list[str] = []
    for part in parts:
        item = clean_bullet_item_text(part).strip(" ：:。；;")
        if not item or item in {"当前资料建议", "方法论建议"}:
            continue
        if item.startswith("当前资料建议："):
            item = item.replace("当前资料建议：", "", 1).strip()
        if item.startswith("方法论建议："):
            item = item.replace("方法论建议：", "", 1).strip()
        if item and item not in cleaned:
            cleaned.append(item)
        if len(cleaned) >= limit:
            break
    return cleaned


def format_rule_reference(ku_id: str, ku_title: str) -> str:
    if ku_id and ku_title:
        return f"{ku_id}《{ku_title}》"
    if ku_id:
        return ku_id
    if ku_title:
        return ku_title
    return "配电安全专家知识库对应检测项"


def risk_label(value: Any) -> str:
    text = str(value or "").lower()
    if text == "high":
        return "高风险"
    if text == "medium":
        return "中风险"
    if text == "low":
        return "低风险"
    return "待核实"


def is_table_line(line: str) -> bool:
    return shared_is_table_line(line)


def is_bullet_line(line: str) -> bool:
    return shared_is_bullet_line(line)


def is_numbered_list_line(line: str) -> bool:
    return shared_is_numbered_list_line(line)


def numbered_list_start(line: str) -> int:
    match = re.match(r"^(\d+)\s*[.．、）)]\s+", clean_inline_marks(line))
    return int(match.group(1)) if match else 1


def clean_numbered_item_text(value: Any) -> str:
    return shared_clean_numbered_item_text(value)


def is_report_title(line: str) -> bool:
    return shared_is_report_title(line)


def clean_markdown_heading(line: str) -> str:
    return re.sub(r"^#{1,6}\s*", "", line).strip()


def clean_inline_marks(line: Any) -> str:
    if line is None:
        return ""
    return str(line).replace("**", "").replace("__", "").strip()


def clean_bullet_item_text(value: Any) -> str:
    return shared_clean_bullet_item_text(value)


def clean_text(value: Any) -> str:
    text = clean_inline_marks(value)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def clean_table_cell(value: Any) -> str:
    text = clean_text(value)
    return text.replace("|", "｜").replace("\n", " ")


def normalize_heading_title(text: str) -> str:
    text = clean_inline_marks(text)
    text = re.sub(r"^#{1,6}\s*", "", text)
    text = re.sub(r"\s+", "", text)
    return re.sub(r"[。．.：:]+", "", text)


def strip_number_prefix(text: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", text).strip()


# ---------------------------------------------------------------------------
# V2 report-rule overrides.
#
# The original renderer is kept for compatibility. These late definitions
# override selected functions so generated DOCX output follows the V2 rule
# files: data-driven conclusions, three-layer risk analysis, time-grouped
# recommendations, typed tables, and chapter 3 diagnostic summaries.
# ---------------------------------------------------------------------------


def _v2_raw_text(value: Any) -> str:
    return clean_inline_marks(value).strip()


def _v2_line_items(value: Any) -> list[str]:
    raw = _v2_raw_text(value)
    if not raw:
        return []
    return [clean_text(part) for part in raw.splitlines() if clean_text(part)]


def _v2_status(issue: dict[str, Any]) -> str:
    explicit = " ".join(
        clean_text(issue.get(key))
        for key in ("status_label", "evaluation_status", "status")
    ).strip()
    explicit_upper = explicit.upper()
    if explicit:
        if "数据冲突" in explicit or "数据异常" in explicit or "矛盾" in explicit:
            return "数据冲突"
        if "数据缺失" in explicit or "待核实" in explicit:
            return "数据缺失"
        if "NG" in explicit_upper or "不合格" in explicit:
            return "NG"
        if "一般" in explicit or "关注" in explicit:
            return "一般"
        if "OK" in explicit_upper or "正常" in explicit or "合格" in explicit or "满足" in explicit:
            return "OK"

    evidence = " ".join(
        clean_text(issue.get(key))
        for key in ("problem", "priority", "risk_level")
    )
    upper = evidence.upper()
    if "数据冲突" in evidence or "数据异常" in evidence or "矛盾" in evidence:
        return "数据冲突"
    if "数据缺失" in evidence or "待核实" in evidence:
        return "数据缺失"
    if "NG" in upper or "不合格" in evidence or "超标" in evidence or "未投入" in evidence or "路径单一" in evidence:
        return "NG"
    if "一般" in evidence or "关注" in evidence or "风险" in evidence:
        return "一般"
    if "OK" in upper or "正常" in evidence or "合格" in evidence:
        return "OK"
    return "待核实"


def status_conclusion_text(section_id: str, item_name: str, status_label: str, problem: str) -> str:
    item_name = clean_text(item_name) or HEADING_BY_NUMBER.get(section_id, ("该检测项", 3))[0]
    status = clean_text(status_label)
    detail = strip_status_tail(problem)
    detail = detail or "本次上传资料未提供可直接引用的定量描述"
    if status == "OK":
        return (
            f"经评估，{item_name}满足要求。\n"
            "【趋势】当前运行状态良好。"
        )
    if status == "一般":
        load_rate = extract_load_rate(problem) if section_id == "2.1.1" else None
        if load_rate is not None and 80 < load_rate <= 100:
            return (
                f"经评估，{item_name}存在负荷率关注项（{format_percent(load_rate)}，处于80%-100%区间），尚未达到过载判定。\n"
                f"【趋势】{v2_trend_hint(section_id, detail)}"
            )
        return (
            f"经评估，{item_name}存在关注项。\n"
            f"【趋势】{v2_trend_hint(section_id, detail)}"
        )
    if status == "NG":
        load_rate = extract_load_rate(problem) if section_id == "2.1.1" else None
        if load_rate is not None:
            return (
                f"经评估，{item_name}存在过载风险（{format_percent(load_rate)}，超过100%判定线）。\n"
                f"【趋势】{v2_trend_hint(section_id, detail)}\n"
                f"治理紧迫性：{v2_urgency(section_id, detail)}。"
            )
        return (
            f"经评估，{item_name}存在不合格项/安全隐患。\n"
            f"【趋势】{v2_trend_hint(section_id, detail)}\n"
            f"治理紧迫性：{v2_urgency(section_id, detail)}。"
        )
    if status in {"数据冲突", "数据异常"}:
        return (
            f"{item_name}的数据存在矛盾/异常，需核实。\n"
            f"【矛盾点】{detail}。\n"
            "【影响】在数据核实前，按最不利原则暂判为存在风险。\n"
            f"【建议】请确认：①{v2_missing_fields_hint(section_id)} ②相关原始检测记录是否为当前运行状态。"
        )
    if status == "数据缺失":
        return (
            f"{item_name}本次未收到完整检测数据。\n"
            f"基于同类项目经验，{item_name}通常需要结合运行记录、现场检测和设备台账共同判断。\n"
            f"建议优先补充：{v2_missing_fields_hint(section_id)}；补充后可通过专项检测或台账复核完成评估。"
        )
    return (
        f"{item_name}本次资料尚不足以形成精确判定。\n"
        f"建议补充：{v2_missing_fields_hint(section_id)}，并在补充后复核。"
    )


def v2_trend_hint(section_id: str, detail: str) -> str:
    if section_id == "2.1.1":
        return "当前容量裕度已经收窄，后续产能提升或夏季高温期间可能进一步推高负荷率，应提前规划负荷转移或增容方案。"
    if section_id == "2.1.2":
        return "任一供电元件故障都可能造成关键负荷中断，且维护检修弹性不足，建议尽快完善备用路径。"
    if section_id == "2.1.3":
        return "故障时依赖人工恢复会延长停电时间，关键负荷恢复时间可能从秒级扩大到分钟或小时级。"
    if section_id == "2.1.5":
        return "若THDu持续超限且普通电容柜继续运行，后续可能出现谐波共振、过热和电容器损坏风险。"
    if section_id == "2.2.1.1":
        return "当前已在低负载或常规工况下出现谐波超标，后续新增变频设备或产能提升后THDu可能继续升高，并放大电容柜和温升风险。"
    if section_id == "2.2.2.1":
        return "温升异常若持续存在，会加速连接点氧化和绝缘老化；在高负荷或高温季节可能突破缺陷阈值。"
    if section_id == "2.3.2":
        return "剩余电流异常叠加漏保或监测不足时，漏电故障可能无法及时切除，电气火灾和触电风险会升高。"
    if section_id.startswith("2.5"):
        return "运维制度和执行闭环不足会降低问题发现与处置效率，使设备缺陷从一般隐患演化为停电或安全事件。"
    return "后续运行条件变化后，该问题可能进一步放大，应结合现场工况持续跟踪并闭环整改。"


def v2_urgency(section_id: str, detail: str) -> str:
    if section_id in {"2.1.2", "2.1.3", "2.2.1.1", "2.2.2.1", "2.3.1", "2.3.2", "2.3.3"}:
        return "高"
    if any(token in detail for token in ["触电", "火灾", "爆炸", "超标", "未投入", "路径单一"]):
        return "高"
    return "中"


def v2_missing_fields_hint(section_id: str) -> str:
    hints = {
        "2.1.1": "变压器容量、运行电流、负荷率曲线、上下级回路容量匹配资料",
        "2.1.2": "关键负荷清单、供电路径图、备用电源配置、ATS配置及倒闸方案",
        "2.1.3": "ATS/双投开关配置、切换逻辑、切换时间和年度试验记录",
        "2.1.5": "电容柜型号、消谐电抗器配置、控制器报警信息、功率因数、THDu实测值和电容容量衰减记录",
        "2.2.1.1": "各母线THDu/THDi实测值、谐波源清单、有源滤波器输出电流和电容柜运行状态",
        "2.2.1.3": "大功率电机/冲击负荷清单、启动频次、启动电流、冲击负荷功率、电压波动记录及软启动/变频配置情况",
        "2.2.2.1": "红外测温记录、环境温度、温升计算值、复测照片和紧固记录",
        "2.3.2": "剩余电流实测值、漏保投运状态、零地混接排查记录和末端RCD测试记录",
    }
    return hints.get(section_id, "现场记录、检测数据、设备台账、照片证据和整改闭环记录")


def extract_load_rate(text: Any) -> float | None:
    clean = clean_text(text)
    patterns = [
        r"负荷率\s*[:：]?\s*([0-9]+(?:\.[0-9]+)?)\s*%",
        r"计算负荷率\s*[:：]?\s*([0-9]+(?:\.[0-9]+)?)\s*%",
    ]
    for pattern in patterns:
        match = re.search(pattern, clean)
        if match:
            try:
                return float(match.group(1))
            except ValueError:
                return None
    return None


def format_percent(value: float) -> str:
    return f"{value:.2f}".rstrip("0").rstrip(".") + "%"


def matched_risk_analysis(section_id: str, problem: str, risk_analysis: str, status_label: str) -> str:
    if status_label in {"OK", "数据缺失", "待核实"}:
        return ""
    detail = clean_text(problem)
    if section_id == "2.1.1":
        load_rate = extract_load_rate(problem)
        if load_rate is not None and load_rate <= 100:
            return (
                "【根因】当前运行负荷率已进入80%-100%关注区间，容量裕度收窄。\n"
                "【直接影响】变压器长期接近高负荷运行时，温升裕度下降、绝缘老化速度加快，后续新增负荷或高温季节可能推高运行风险。\n"
                "【恶化条件】若继续增产、负荷转移不及时或散热条件下降，负荷率可能进一步升高并进入过载状态，应提前监测并预留负荷调整方案。"
            )
        if load_rate is not None:
            return (
                "【根因】当前运行负荷率已超过额定容量，存在实际过载。\n"
                "【直接影响】变压器过载会导致温升升高、绝缘加速劣化，严重时可能引发跳闸、设备损坏或电气火灾。\n"
                "【恶化条件】若高负荷持续运行且缺少负荷转移、增容或温度监测，故障概率和停产影响会进一步放大。"
            )
    if section_id == "2.1.5":
        return (
            f"【根因】{detail}。\n"
            "【直接影响】谐波电流进入普通电容柜后，可能与系统电感形成谐振回路，导致电容电流放大、过载发热，严重时引发电容器爆裂或电气火灾。\n"
            "【恶化条件】当THDu持续高于3%且电容柜未配置消谐电抗器时，风险会明显升高；若后续负荷提升或变频设备增加，应将谐波治理与电容柜整改合并处理。"
        )
    if section_id == "2.2.1.1":
        return (
            f"【根因】{detail}。\n"
            "【直接影响】THDu超标会增加变压器、电缆和电机的附加损耗，加速绝缘老化；同时谐波与电容柜共振时可能引发电容过载、爆炸或火灾风险。\n"
            "【恶化条件】若产能提升、新增变频负荷或有源滤波器输出不足，THDu可能继续升高，保护误动作和设备过热概率会增加。"
        )
    if section_id == "2.2.2.1":
        return (
            f"【根因】{detail}。\n"
            "【直接影响】温升异常会加速接点氧化、绝缘老化和回路电阻升高，形成发热正反馈。\n"
            "【恶化条件】在高负荷、夏季高温或散热不良条件下，温升可能突破缺陷阈值并诱发停电或火灾风险。"
        )
    if section_id == "2.3.2":
        return (
            f"【根因】{detail}。\n"
            "【直接影响】漏电或零地混接未被及时发现时，故障点可能持续发热，并削弱人身触电和电气火灾防护能力。\n"
            "【恶化条件】潮湿环境、绝缘劣化或负荷增加会使剩余电流继续增大，应先启用保护再逐级定位根因。"
        )
    fallback = clean_text(risk_analysis) or "该问题会降低配电系统安全裕度和运行可靠性。"
    return (
        f"【根因】{detail}。\n"
        f"【直接影响】{fallback}\n"
        "【恶化条件】当负荷、环境或运行频次进一步增加时，问题可能从关注项演变为停电、设备损坏或人身安全风险。"
    )


def status_block_lines(
    situation: str,
    conclusion: str,
    risk: str,
    suggestions: list[str],
) -> list[str]:
    lines: list[str] = []
    situation_items = _v2_line_items(situation)
    if situation_items:
        lines.extend(["【现状描述】", *situation_items, ""])

    lines.append("【结论】")
    conclusion_items = [
        truncate_policy_text(item, "conclusion_policy", "max_words_per_item")
        for item in (_v2_line_items(conclusion) or ["该检测项本次资料不足，暂按待核实处理。"])
    ]
    lines.extend(conclusion_items)
    lines.append("")

    risk_items = [
        truncate_policy_text(item, "risk_policy", "max_words")
        for item in _v2_line_items(risk)
    ]
    if risk_items:
        lines.extend(["【风险分析】", *risk_items, ""])

    grouped = v2_group_suggestions(suggestions, conclusion)
    lines.append("【建议】")
    for group_title, items in grouped:
        if group_title:
            lines.append(group_title)
        for item in items:
            lines.append(f"• {v2_enrich_suggestion(item)}")
    lines.append("")
    return lines


def v2_group_suggestions(suggestions: list[str], conclusion: str) -> list[tuple[str, list[str]]]:
    cleaned = [clean_text(item).strip("。；;") for item in suggestions if clean_text(item)]
    if not cleaned:
        cleaned = ["补充现场检测数据并形成整改闭环记录"]
    if not policy_bool("suggestion_policy", "group_by_time", True):
        return [("", cleaned[:4])]
    high = "不合格" in conclusion or "治理紧迫性：高" in conclusion or "数据存在矛盾" in conclusion
    if high:
        return [
            ("【立即执行（2周内）】", cleaned[:2]),
            ("【短期整改（3个月内）】", cleaned[2:4] or cleaned[:1]),
            ("【中期规划（1年内）】", cleaned[4:5] or ["将该问题纳入年度停电检修和专项治理计划"]),
        ]
    return [
        ("【短期整改（3个月内）】", cleaned[:2]),
        ("【中期规划（1年内）】", cleaned[2:3] or ["建立周期性复测和趋势跟踪机制"]),
    ]


def v2_enrich_suggestion(suggestion: str) -> str:
    text = clean_bullet_item_text(suggestion)
    if not text:
        return "补充现场检测数据并形成整改闭环记录，预期提高后续判定准确性。"
    if not policy_bool("suggestion_policy", "include_expected_effect", True):
        return text.rstrip("。") + "。"
    if "依据" in text and "预期" in text:
        return text + "。"
    if len(text) < 15 or text in {"完善", "加强巡检", "维护时消缺"}:
        text = f"将“{text}”展开为具体整改任务，明确责任人、实施位置、验收标准和复核记录"
    return f"{text}，依据相关标准和现场评估结果执行，预期降低该项风险并形成可追溯闭环。"


def chapter3_summary_text(model: dict[str, Any], issues: list[dict[str, Any]], action_rows: list[dict[str, Any]]) -> str:
    counts = {"OK": 0, "一般": 0, "NG": 0, "数据缺失": 0, "数据冲突": 0, "待核实": 0}
    for issue in issues:
        counts[_v2_status(issue)] = counts.get(_v2_status(issue), 0) + 1
    total = sum(counts.values())
    ng_count = counts.get("NG", 0)
    top_risk_count = policy_int("chapter_3_policy", "top_risk_count", 3)
    urgent = sorted(
        [item for item in issues if _v2_status(item) in {"NG", "数据冲突"}],
        key=lambda item: v2_priority_rank(item),
    )[:top_risk_count]
    lines = [
        "#### 3.1.1 风险全景图",
        f"本次评估共覆盖{total}个检测项，其中：OK {counts.get('OK', 0)}项、一般 {counts.get('一般', 0)}项、NG {ng_count}项、数据缺失 {counts.get('数据缺失', 0) + counts.get('待核实', 0)}项、数据异常 {counts.get('数据冲突', 0)}项。",
        f"整体判断：{v2_overall_judgement(ng_count)}",
        "",
        f"最需紧急关注的{min(top_risk_count, len(urgent))}个问题：",
    ]
    for index, issue in enumerate(urgent, 1):
        lines.append(f"{index}. {v2_issue_sentence(issue)}")
    if not urgent:
        lines.append("1. 本次未识别到明确NG项，建议跟踪一般项和数据缺口。")
    lines.extend(["", "#### 3.1.2 各维度风险分析"])
    for prefix, title in [
        ("2.1", "系统架构"),
        ("2.2", "环境工况"),
        ("2.3", "故障保护"),
        ("2.4", "设备/元件状态"),
        ("2.5", "运维管理"),
    ]:
        items = [item for item in issues if str(item.get("section_id") or "").startswith(prefix)]
        status_counts = v2_dimension_counts(items)
        lines.extend([
            f"【{title}】",
            f"状态统计：OK {status_counts.get('OK', 0)}项 / 一般 {status_counts.get('一般', 0)}项 / NG {status_counts.get('NG', 0)}项 / 数据缺失 {status_counts.get('数据缺失', 0) + status_counts.get('待核实', 0)}项",
            f"评价：{dimension_evaluation_word(status_counts.get('NG', 0), status_counts.get('一般', 0))}",
            "关键发现：",
        ])
        for item in sorted(items, key=v2_priority_rank)[:3]:
            lines.append(f"• {v2_issue_sentence(item)}")
        if not items:
            lines.append("• 本维度未识别到明确异常，建议按常规周期复核。")
        lines.append("")

    if policy_bool("chapter_3_policy", "include_data_gap_analysis", True):
        lines.extend(["", "#### 3.1.3 数据缺口分析"])
        missing_items = [item for item in issues if _v2_status(item) in {"数据缺失", "待核实"}]
        if missing_items:
            lines.append("本次评估中以下数据项尚未收到，建议按优先级补充：")
            lines.append("【优先补充（影响NG判定精度）】")
            priority_missing = sorted(missing_items, key=v2_priority_rank)
            for item in priority_missing[:5]:
                lines.append(f"• {v2_missing_fields_hint(str(item.get('section_id') or ''))}（{clean_text(item.get('section_title')) or '相关检测项'}）：用于完善该项判定，补充方法为现场检测、台账核对或试验报告复核。")
            if len(priority_missing) > 5:
                lines.append("【建议补充（完善评估精度）】")
                for item in priority_missing[5:10]:
                    lines.append(f"• {v2_missing_fields_hint(str(item.get('section_id') or ''))}（{clean_text(item.get('section_title')) or '相关检测项'}）：用于提升评估精度，补充方法为资料归档或现场复核。")
        else:
            lines.append("本次必需资料基本覆盖，后续建议补充照片、复测记录和整改闭环证据，以提高报告可追溯性。")
    return "\n".join(lines).strip()


def v2_priority_rank(issue: dict[str, Any]) -> tuple[int, str]:
    text = " ".join(clean_text(issue.get(key)) for key in ("priority", "risk_level", "finding_priority", "problem"))
    if _v2_status(issue) == "NG" or "高" in text:
        return (0, clean_text(issue.get("section_id")))
    if _v2_status(issue) == "数据冲突":
        return (1, clean_text(issue.get("section_id")))
    if "中" in text or _v2_status(issue) == "一般":
        return (2, clean_text(issue.get("section_id")))
    return (3, clean_text(issue.get("section_id")))


def v2_issue_sentence(issue: dict[str, Any]) -> str:
    section = clean_text(issue.get("section_id"))
    title = clean_text(issue.get("section_title") or issue.get("category") or section)
    problem = clean_text(issue.get("problem") or issue.get("description"))
    recommendation = clean_text(issue.get("recommendation"))
    problem = shorten_table_text(problem, 90)
    recommendation = concise_recommendation_text(recommendation)
    return f"{section} {title}：{problem}。→ {recommendation or '建议制定专项整改措施'}"


def v2_overall_judgement(ng_count: int) -> str:
    if ng_count == 0:
        return "系统整体安全风险可控，重点应关注一般项的跟踪改善。"
    if ng_count <= 2:
        return "系统存在若干安全隐患，建议按优先级制定整改计划。"
    if ng_count <= 5:
        return "系统存在多项安全隐患，部分问题之间存在关联放大效应，建议尽快启动高优先级项整改。"
    if ng_count <= 8:
        return "系统安全风险较高，存在多项NG项亟需整改，建议立即启动高优先级项并制定全面整改方案。"
    return "系统安全风险严重，多项安全隐患可能相互叠加放大，建议立即采取应急措施并制定全面整改方案。"


def v2_dimension_counts(items: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = {"OK": 0, "一般": 0, "NG": 0, "数据缺失": 0, "数据冲突": 0, "待核实": 0}
    for item in items:
        status = _v2_status(item)
        counts[status] = counts.get(status, 0) + 1
    return counts


def v2_cross_link_records(issues: list[dict[str, Any]]) -> list[dict[str, str]]:
    status_by_section = {str(item.get("section_id") or ""): _v2_status(item) for item in issues}
    text_by_section: dict[str, str] = {}
    for issue in issues:
        section_id = str(issue.get("section_id") or "")
        text_by_section.setdefault(section_id, "")
        text_by_section[section_id] += " " + clean_text(issue.get("problem") or issue.get("description") or "")

    def status(section_id: str) -> str | None:
        return status_by_section.get(section_id)

    def in_status(section_id: str, allowed: set[str]) -> bool:
        return status(section_id) in allowed

    def not_ok(section_id: str) -> bool:
        return status(section_id) not in {None, "OK"}

    records: list[dict[str, str]] = []

    if in_status("2.2.1.1", {"NG", "一般", "数据冲突"}) and not_ok("2.1.5"):
        heat_tail = "，并会叠加设备温升风险" if in_status("2.2.2.1", {"NG", "一般", "数据冲突"}) else ""
        records.append({
            "id": "LINK-01",
            "main_section": "2.2.1.1",
            "related_name": "系统无功补偿与电容柜运行状态",
            "related_status": status("2.1.5") or "待核实",
            "title": "谐波×电容柜×温升连锁风险",
            "logic": f"THDu超标环境下普通电容柜可能形成谐波共振，造成电容过载、发热甚至爆炸{heat_tail}。",
            "effect": "谐波是根因，电容柜和温升是放大环节，分开治理容易遗漏根因。",
            "advice": "先停用高谐波区段普通电容柜并核实现有有源滤波容量，再复测THDu和温升。",
        })

    if in_status("2.1.1", {"NG", "一般", "数据冲突"}) and in_status("2.2.2.1", {"NG", "一般", "数据冲突"}):
        records.append({
            "id": "LINK-02",
            "main_section": "2.1.1",
            "related_name": "低压配电设备温升与发热",
            "related_status": status("2.2.2.1") or "待核实",
            "title": "负荷率×温升连锁风险",
            "logic": "高负荷率会降低温升裕量，同样的连接缺陷在高负荷设备上更容易突破温升限值。",
            "effect": "负荷接近满载时，连接点发热、散热不良和绝缘老化会相互促进。",
            "advice": "负荷转移、红外热点整改和复测应同步推进，避免只处理单一表现。",
        })

    if in_status("2.3.2", {"NG", "一般", "数据冲突"}) and in_status("2.4.3.1", {"NG", "一般", "数据冲突"}):
        records.append({
            "id": "LINK-03",
            "main_section": "2.3.2",
            "related_name": "低压回路剩余电流专项排查",
            "related_status": status("2.4.3.1") or "待核实",
            "title": "漏电保护×剩余电流体系风险",
            "logic": "漏电保护配置或投运不足叠加实测剩余电流偏大，会形成“有故障但保护不足”的双重风险。",
            "effect": "仅排查漏电点而不启用保护，或仅投入保护而不定位根因，都难以形成完整防护。",
            "advice": "先启用剩余电流监测和保护，再逐级排查零地混接、设备漏电和假漏电。",
        })

    if in_status("2.4.2.2", {"NG", "一般", "数据冲突"}) and in_status("2.4.3.1", {"NG", "一般", "数据冲突"}):
        records.append({
            "id": "LINK-04",
            "main_section": "2.4.2.2",
            "related_name": "低压回路剩余电流专项排查",
            "related_status": status("2.4.3.1") or "待核实",
            "title": "接地×剩余电流连锁风险",
            "logic": "接地不良或零地混接可能使工作电流经接地线返回，导致剩余电流异常升高。",
            "effect": "若未同步排查接地系统，剩余电流治理可能反复出现或误判为设备漏电。",
            "advice": "剩余电流专项排查时优先检查PE/N分离、接地连续性和桥架跨接。",
        })

    if status("2.1.2") == "NG" and status("2.1.3") == "NG":
        records.append({
            "id": "LINK-05",
            "main_section": "2.1.2",
            "related_name": "配网自动化与备用电源自动切换",
            "related_status": status("2.1.3") or "待核实",
            "title": "供电路径×ATS双重脆弱性",
            "logic": "关键负荷供电路径单一且无ATS时，故障恢复完全依赖人工倒闸。",
            "effect": "故障恢复时间可能从秒级或分钟级延长到小时级，生产连续性风险被放大。",
            "advice": "新增备用供电路径时同步配置ATS，并通过年度切换试验验证恢复时间。",
        })

    physical_text = text_by_section.get("2.2.2.3", "")
    physical_env_triggered = status("2.2.2.3") == "NG" and bool(re.search(r"潮湿|积水|冷凝水|防水|渗水|排水", physical_text))
    if physical_env_triggered and not_ok("2.2.2.2"):
        records.append({
            "id": "LINK-06",
            "main_section": "2.2.2.3",
            "related_name": "高压配电设备局部放电",
            "related_status": status("2.2.2.2") or "待核实",
            "title": "物理环境×局放绝缘劣化风险",
            "logic": "潮湿、积水或冷凝水环境会加速绝缘劣化，进而增加局部放电活动。",
            "effect": "环境问题不消除时，局放治理效果会下降，绝缘劣化可能持续累积。",
            "advice": "先处理防水、排水和除湿，再开展局放检测、在线监测和绝缘缺陷治理。",
        })

    return records


def v2_cross_link_map(issues: list[dict[str, Any]]) -> dict[str, list[dict[str, str]]]:
    link_map: dict[str, list[dict[str, str]]] = {}
    for record in v2_cross_link_records(issues):
        link_map.setdefault(record["main_section"], []).append(record)
    return link_map


def v2_section_cross_link_lines(records: list[dict[str, str]]) -> list[str]:
    if not records:
        return []
    lines = ["【关联风险】", "本项问题与以下检测项存在关联，建议统筹考虑整改："]
    for record in records:
        lines.extend([
            f"• {record['related_name']}（{record['related_status']}）：",
            f"  {record['logic']}",
            f"  建议将两项整改纳入同一计划，{record['advice']}",
        ])
    lines.append("")
    return lines


def v2_action_section(model: dict[str, Any]) -> str:
    rows = [row for row in (model.get("action_rows") or []) if isinstance(row, dict)]
    lines = [
        "### 3.2 改善行动速查表",
        "",
        "改善行动列表见下表。建议按紧急、重要、一般分层推进，并在实施后复核整改效果。",
        "",
    ]
    if not rows:
        lines.append("本次未形成明确改善行动项，建议补充资料后重新生成行动清单。")
        lines.extend(v2_special_analysis_lines(model))
        return "\n".join(lines)
    table_rows = []
    sorted_rows = sorted(rows, key=lambda row: (
        v2_priority_rank({"priority": row.get("优先级") or row.get("priority"), "section_id": row.get("section_id") or row.get("章节") or row.get("类别")}),
        clean_text(row.get("问题/主题") or row.get("problem") or ""),
    ))
    for row in sorted_rows[:40]:
        priority = clean_text(row.get("优先级") or row.get("priority") or "待核实")
        problem = clean_text(row.get("问题/主题") or row.get("problem") or "")
        recommendation = clean_text(row.get("建议措施") or row.get("recommendation") or "结合现场资料制定整改措施")
        table_rows.append([
            len(table_rows) + 1,
            shorten_table_text(
                clean_text(row.get("section_id") or row.get("章节") or row.get("类别") or row.get("category") or "待归类"),
                42,
            ),
            shorten_table_text(problem, 95),
            priority,
            concise_recommendation_text(recommendation, max_points=2, max_point_chars=28),
            v2_expected_benefit(problem),
        ])
    lines.extend(markdown_table(
        ["序号", "章节", "问题", "优先级", "建议", "预期收益"],
        table_rows,
    ))
    lines.extend(v2_special_analysis_lines(model))
    return "\n".join(lines)


def v2_special_analysis_lines(model: dict[str, Any]) -> list[str]:
    """Render the optional legacy special-analysis block only when supplied."""

    supplied = model.get("special_analysis") or {}
    if not isinstance(supplied, dict) or not any(
        clean_text(value) for value in supplied.values()
    ):
        return []
    titles = {
        "4.1": "新工厂建厂时规划建议",
        "4.2": "增容建议",
        "4.3": "日常用电管理建议",
        "4.4": "应急管理及合规性管理建议",
    }
    lines = ["", "## 4. 专项问题分析", ""]
    for number, title in titles.items():
        body = clean_text(supplied.get(number))
        if not body:
            continue
        lines.extend(
            [
                f"### {number} {title}",
                "",
                body,
                "",
            ]
        )
    return lines


def v2_action_category(priority: str) -> str:
    if "高" in priority:
        return "紧急"
    if "中" in priority:
        return "重要"
    return "一般"


def v2_time_requirement(priority: str) -> str:
    if "高" in priority:
        return "2周内/3个月内分阶段落实"
    if "中" in priority:
        return "1年内纳入计划"
    return "纳入年度计划"


def v2_standard_hint(problem: str) -> str:
    if "谐波" in problem or "THD" in problem:
        return "GB/T 14549-1993"
    if "剩余电流" in problem or "漏电" in problem:
        return "GB 13955-2017"
    if "供电" in problem or "负荷" in problem:
        return "GB 50052-2009"
    return "相关国家标准/行业规范"


def v2_expected_benefit(problem: str) -> str:
    if "谐波" in problem or "电容" in problem:
        return "降低谐波共振、设备过热和非计划停电风险。"
    if "路径" in problem or "ATS" in problem or "供电" in problem:
        return "提升供电连续性，将故障恢复时间从人工恢复缩短至自动切换级别。"
    if "剩余电流" in problem or "漏电" in problem or "接地" in problem:
        return "降低电气火灾和人身触电风险。"
    if "温升" in problem or "过热" in problem:
        return "降低连接点劣化和设备损毁风险。"
    return "提升运行可靠性和整改闭环可追溯性。"


def remove_v2_parent_missing_blocks(text: str) -> str:
    parent_ids = {"2.2.1", "2.2.2", "2.4.1", "2.4.2", "2.4.3", "2.5.3"}
    lines = text.splitlines()
    output: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        heading_match = re.match(r"^(#{1,6}\s*)?(\d+\.\d+\.\d+)\.?\s+(.+)$", line.strip())
        if heading_match and heading_match.group(2) in parent_ids:
            output.append(line)
            section_id = heading_match.group(2)
            lookahead = "\n".join(lines[index + 1:index + 6])
            if "因数据缺失" in lookahead or "未能给出明确结论" in lookahead:
                intro = v2_parent_intro(section_id)
                if intro:
                    output.extend(["", intro, ""])
                index += 1
                while index < len(lines):
                    next_line = lines[index].strip()
                    if re.match(r"^(#{1,6}\s*)?\d+\.\d+\.\d+(?:\.\d+)?\.?\s+", next_line):
                        break
                    index += 1
                continue
        output.append(line)
        index += 1
    return "\n".join(output)


def v2_parent_intro(section_id: str) -> str:
    intros = {
        "2.2.1": "电气环境工况主要从谐波、电压扰动和冲击负荷三个方面审视，以下按检测项逐项说明。",
        "2.2.2": "其他运行工况主要从低压发热、高压局放和物理环境风险三个方面审视，以下按检测项逐项说明。",
        "2.4.1": "配置与选型问题主要关注容量、分隔形式、联锁和位置显示等设备本质安全条件。",
        "2.4.2": "安装规范性问题主要关注裸露导体、接地、连接、标识、电缆桥架和封堵等现场缺陷。",
        "2.4.3": "带病运行问题主要关注剩余电流、照明、除湿等长期运行状态异常。",
        "2.5.3": "运维实施与组织主要关注人员配置、维护计划、巡检执行和维保覆盖情况。",
    }
    return intros.get(section_id, "")


_LEGACY_BUILD_REPORT_TEXT_FROM_MODEL = build_report_text_from_model


def build_report_text_from_model(model: dict[str, Any], fallback_text: str) -> str:
    text = _LEGACY_BUILD_REPORT_TEXT_FROM_MODEL(model, fallback_text)
    text = remove_v2_parent_missing_blocks(text)
    action_section = v2_action_section(model)
    return re.sub(
        r"### 3\.2\s+(?:改善行动列表与优先级|改善行动速查表)[\s\S]*$",
        action_section,
        text,
    )


# ---------------------------------------------------------------------------
# V3.0.2.1 output hygiene overrides.
#
# These final definitions are intentionally placed at the end of the module so
# they override the earlier compatibility definitions above.
# ---------------------------------------------------------------------------

_LEGACY_NORMALIZE_REPORT_TEXT = normalize_report_text
_LEGACY_NORMALIZE_TABLE_CELL_VALUE = normalize_table_cell_value


SOURCE_REFERENCE_PATTERNS = [
    r"[（(]\s*(?:来源|资料来源|文件来源|source|evidence)?\s*S\d+-\d+[^）)\n]{0,80}[）)]",
    r"(?:来源|资料来源|文件来源|source|evidence)\s*[:：]\s*[^。；;\n]{0,120}",
    r"\bS\d+-\d+[^:：。；;\n]{0,50}[:：]\s*",
    r"[\w\u4e00-\u9fff（）()#\- ]{0,80}\.(?:xlsx|xls|xlsm|docx|doc|zip|pdf)\b\s*(?:/|／|\\)?\s*",
    r"(?:^|[\s，,；;])(?:S\d+-\d+[^/。；;\n]{0,60}\s*/\s*)+",
]

POSITIVE_ASSESSMENT_TOKENS = [
    "OK",
    "满足要求",
    "符合要求",
    "符合标准",
    "未发现异常",
    "无异常",
    "正常",
    "合格",
    "通过",
]

NEGATIVE_ASSESSMENT_TOKENS = [
    "NG",
    "不合格",
    "异常",
    "超标",
    "风险",
    "隐患",
    "数据冲突",
    "数据异常",
    "数据缺失",
    "待核实",
    "未满足",
    "不满足",
    "缺失",
    "未配置",
    "未投入",
    "过载",
]


def strip_source_references(text: Any) -> str:
    cleaned = clean_text(text)
    if not cleaned:
        return ""
    for pattern in SOURCE_REFERENCE_PATTERNS:
        cleaned = re.sub(pattern, " ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*[/／\\]\s*(?:第?\d+行|行\d+|Sheet\d*|工作表\d*)", " ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned)
    cleaned = re.sub(r"\s+([，。；：、])", r"\1", cleaned)
    cleaned = re.sub(r"([：:])\s+", r"\1", cleaned)
    cleaned = re.sub(r"([：:])\s*([，。；;])", r"\2", cleaned)
    return cleaned.strip(" ，,；;：:")


def is_positive_assessment_block(*parts: Any) -> bool:
    text = strip_source_references(" ".join(clean_text(part) for part in parts if clean_text(part)))
    if not text:
        return False
    upper = text.upper()
    has_positive = ("OK" in upper) or any(token in text for token in POSITIVE_ASSESSMENT_TOKENS if token != "OK")
    has_negative = ("NG" in upper) or any(token in text for token in NEGATIVE_ASSESSMENT_TOKENS if token != "NG")
    return has_positive and not has_negative


def is_advice_label(line: str) -> bool:
    stripped = clean_inline_marks(line).strip()
    return bool(re.match(r"^【\s*(建议|整改建议|后续建议)\s*】\s*$", stripped))


def is_structural_boundary(line: str) -> bool:
    stripped = clean_inline_marks(line).strip()
    if not stripped:
        return False
    if re.match(r"^#{1,6}\s+", stripped):
        return True
    if re.match(r"^\d+(?:\.\d+){1,5}\.?\s+", stripped):
        return True
    return bool(re.match(r"^【\s*(现状描述|结论|数据|趋势|风险分析|判定|问题描述)\s*】", stripped))


def strip_positive_advice_sections(report_text: str) -> str:
    lines = report_text.splitlines()
    output: list[str] = []
    context_since_heading: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if is_advice_label(line) and is_positive_assessment_block(*context_since_heading):
            index += 1
            while index < len(lines):
                next_line = lines[index]
                if is_structural_boundary(next_line):
                    break
                index += 1
            continue
        output.append(line)
        stripped = clean_inline_marks(line).strip()
        if re.match(r"^#{1,6}\s+|^\d+(?:\.\d+){1,5}\.?\s+", stripped):
            context_since_heading = [line]
        else:
            context_since_heading.append(line)
            context_since_heading = context_since_heading[-20:]
        index += 1
    return "\n".join(output)


def normalize_report_text(report_text: str) -> str:
    """Normalize layout syntax without rewriting approved report prose.

    Manyselves only calls this vendored renderer with ``report_model=None`` and
    already-approved Markdown.  The former V3 hygiene override removed S2/S4
    provenance and whole positive-advice blocks at render time.  That made the
    DOCX disagree with the Chief Editor artifact and correctly tripped the
    protected-prose check.  Evidence cleanup belongs before approval, never in
    the deterministic renderer.
    """

    return _LEGACY_NORMALIZE_REPORT_TEXT(report_text)


def normalize_table_cell_value(value: Any, header: Any = "", is_header: bool = False) -> str:
    # Table cells are approved prose too; retain their evidence/source text.
    return _LEGACY_NORMALIZE_TABLE_CELL_VALUE(value, header, is_header)


def status_block_lines(
    situation: str,
    conclusion: str,
    risk: str,
    suggestions: list[str],
) -> list[str]:
    lines: list[str] = []
    situation_items = [
        strip_source_references(item)
        for item in _v2_line_items(situation)
        if strip_source_references(item)
    ]
    if situation_items:
        lines.extend(["【现状描述】", *situation_items, ""])

    lines.append("【结论】")
    conclusion_items = [
        truncate_policy_text(strip_source_references(item), "conclusion_policy", "max_words_per_item")
        for item in (_v2_line_items(conclusion) or ["该检测项本次资料不足，暂按待核实处理。"])
    ]
    conclusion_items = [item for item in conclusion_items if item]
    lines.extend(conclusion_items or ["该检测项本次资料不足，暂按待核实处理。"])
    lines.append("")

    risk_items = [
        truncate_policy_text(strip_source_references(item), "risk_policy", "max_words")
        for item in _v2_line_items(risk)
    ]
    risk_items = [item for item in risk_items if item]
    if risk_items:
        lines.extend(["【风险分析】", *risk_items, ""])

    if is_positive_assessment_block(situation, conclusion, risk) and not risk_items:
        return lines
    if not risk_items and not any(clean_text(item) for item in suggestions):
        return lines

    grouped = v2_group_suggestions(suggestions, conclusion)
    lines.append("【建议】")
    for group_title, items in grouped:
        if group_title:
            lines.append(group_title)
        for item in items:
            cleaned = strip_source_references(v2_enrich_suggestion(item))
            if cleaned:
                lines.append(f"• {cleaned}")
    lines.append("")
    return lines
