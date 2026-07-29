"""Map S4-4 diagnostic workbook observations into module 2.4 evidence."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .common import MappingGap, MappingResult, dispimg_refs, make_evidence


def _text(value: Any) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _is_placeholder(value: Any) -> bool:
    return _text(value) in {"", "/", "\\", "--"}


def _subject(room: Any, cabinet: Any = None) -> str:
    room_text = _text(room) or "未注明配电房"
    cabinet_text = _text(cabinet)
    return f"{room_text}/{cabinet_text}" if cabinet_text else room_text


def _numeric_value(sheet: Any, value_sheet: Any, row: int, column: int) -> Any:
    value = sheet.cell(row, column).value
    if isinstance(value, (int, float)):
        return value
    cached_value = value_sheet.cell(row, column).value
    return cached_value if isinstance(cached_value, (int, float)) else value


def _map_low_voltage(
    path: Path,
    file_id: str,
    sheet: Any,
    value_sheet: Any,
) -> tuple[list, list[MappingGap]]:
    evidence = []
    gaps: list[MappingGap] = []
    current_only_rows: list[int] = []
    for row in range(3, sheet.max_row + 1):
        room = sheet.cell(row, 1).value
        cabinet = sheet.cell(row, 2).value
        if not cabinet:
            continue
        subject = _subject(room, cabinet)
        current = _numeric_value(sheet, value_sheet, row, 3)
        capacity = _numeric_value(sheet, value_sheet, row, 4)
        if isinstance(current, (int, float)):
            value = float(current)
            unit = "A"
            fact = f"运行电流={current:g}A"
            if isinstance(capacity, (int, float)) and capacity:
                value = current * 0.4 * 1.732 / capacity * 100
                unit = "%"
                fact += f"；变压器容量={capacity:g}kVA；计算负荷率={value:.2f}%"
                routes = (("2.4", "2.4.1.1"), ("2.1", "2.1.1"))
                needs_confirmation = False
            else:
                current_only_rows.append(row)
                fact += "；未提供额定容量，不能据此计算负荷率或判断过载"
                routes = (("2.1", "2.1.1"), ("2.4", "2.4.4"))
                needs_confirmation = True
            for module_id, submodule_id in routes:
                evidence.append(
                    make_evidence(
                        file_id=file_id,
                        path=path,
                        sheet=sheet.title,
                        cell=f"C{row}:E{row}",
                        row=row,
                        column="C",
                        subject=subject,
                        fact=fact,
                        module_id=module_id,
                        submodule_id=submodule_id,
                        value=value,
                        unit=unit,
                        needs_confirmation=needs_confirmation,
                    )
                )

        paired_fields = (
            (6, 7, "安全连锁", "2.4.1.3"),
            (8, 9, "裸露导体绝缘防护", "2.4.2.1"),
            (13, 14, "等电位连接与接地", "2.4.2.2"),
            (15, 16, "色标标识", "2.4.2.4"),
            (21, None, "变压器噪声震动", "2.2.2.3"),
            (22, 23, "变压器温度", "2.2.2.1"),
        )
        for value_column, photo_column, label, submodule_id in paired_fields:
            value = sheet.cell(row, value_column).value
            if _is_placeholder(value):
                continue
            end_column = photo_column or value_column
            photo_value = sheet.cell(row, end_column).value if photo_column else None
            start_letter = get_column_letter(value_column)
            end_letter = get_column_letter(end_column)
            cell = f"{start_letter}{row}:{end_letter}{row}"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=cell,
                    row=row,
                    column=start_letter,
                    subject=subject,
                    fact=f"{label}={_text(value)}",
                    module_id=submodule_id.rsplit(".", 2)[0],
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(photo_value),
                    needs_confirmation=_text(value).upper().startswith("NG")
                    and not dispimg_refs(photo_value),
                )
            )
        for value_column, photo_column, label, module_id, submodule_id in (
            (17, 18, "谐波", "2.2", "2.2.1.1"),
            (19, 20, "电涌保护装置", "2.3", "2.3.3"),
        ):
            value = sheet.cell(row, value_column).value
            if _is_placeholder(value):
                continue
            start_letter = get_column_letter(value_column)
            end_letter = get_column_letter(photo_column)
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"{start_letter}{row}:{end_letter}{row}",
                    row=row,
                    column=start_letter,
                    subject=subject,
                    fact=f"{label}={_text(value)}",
                    module_id=module_id,
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(sheet.cell(row, photo_column).value),
                    needs_confirmation=_text(value).upper().startswith("NG"),
                )
            )

        residual_current = _numeric_value(sheet, value_sheet, row, 10)
        if isinstance(residual_current, (int, float)):
            ratio = (
                residual_current / current
                if isinstance(current, (int, float)) and current
                else None
            )
            fact = f"实测剩余电流={residual_current:g}A"
            if ratio is not None:
                fact += f"；占运行电流={ratio * 100:.2f}%"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"J{row}:L{row}",
                    row=row,
                    column="J",
                    subject=subject,
                    fact=fact,
                    module_id="2.4",
                    submodule_id="2.4.3.1",
                    photo_refs=dispimg_refs(sheet.cell(row, 11).value),
                    value=residual_current,
                    unit="A",
                    needs_confirmation=residual_current > 10
                    and not dispimg_refs(sheet.cell(row, 11).value),
                )
            )

        other = _text(sheet.cell(row, 24).value)
        if other:
            if "裸露" in other:
                module_id, submodule_id = "2.4", "2.4.2.1"
            elif any(keyword in other for keyword in ("过热", "温升", "发热")):
                module_id, submodule_id = "2.2", "2.2.2.1"
            else:
                module_id, submodule_id = "2.4", "2.4.4"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"X{row}:Z{row}",
                    row=row,
                    column="X",
                    subject=subject,
                    fact=f"其他检查发现={other}",
                    module_id=module_id,
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(
                        sheet.cell(row, 25).value,
                        sheet.cell(row, 26).value,
                    ),
                )
            )
    if current_only_rows:
        row_range = f"C{min(current_only_rows)}:E{max(current_only_rows)}"
        message = (
            f"低配表 {len(current_only_rows)} 条运行电流缺少额定容量；"
            "已保留原始电流，但不能计算负荷率或判断过载"
        )
        for module_id, submodule_id in (
            ("2.1", "2.1.1"),
            ("2.4", "2.4.1.1"),
        ):
            gaps.append(
                MappingGap(
                    code="missing_capacity_for_load_rate",
                    message=message,
                    sheet=sheet.title,
                    cell=row_range,
                    module_id=module_id,
                    submodule_id=submodule_id,
                )
            )
    return evidence, gaps


def _map_total_distribution(
    path: Path,
    file_id: str,
    sheet: Any,
    value_sheet: Any,
) -> tuple[list, list[MappingGap]]:
    evidence = []
    gaps: list[MappingGap] = []
    capacity_only_rows: list[int] = []
    for row in range(3, sheet.max_row + 1):
        room = sheet.cell(row, 1).value
        cabinet = sheet.cell(row, 2).value
        if not room and not cabinet:
            continue
        subject = _subject(room, cabinet)
        current = _numeric_value(sheet, value_sheet, row, 3)
        capacity = _numeric_value(sheet, value_sheet, row, 4)
        if isinstance(current, (int, float)):
            value = float(current)
            unit = "A"
            fact = f"运行电流={current:g}A"
            if isinstance(capacity, (int, float)) and capacity:
                value = current * 0.4 * 1.732 / capacity * 100
                unit = "%"
                fact += f"；变压器容量={capacity:g}kVA；计算负荷率={value:.2f}%"
            for module_id, submodule_id in (
                ("2.4", "2.4.1.1"),
                ("2.1", "2.1.1"),
            ):
                evidence.append(
                    make_evidence(
                        file_id=file_id,
                        path=path,
                        sheet=sheet.title,
                        cell=f"C{row}:E{row}",
                        row=row,
                        column="C",
                        subject=subject,
                        fact=fact,
                        module_id=module_id,
                        submodule_id=submodule_id,
                        value=value,
                        unit=unit,
                    )
                )
        elif isinstance(capacity, (int, float)) and capacity:
            capacity_only_rows.append(row)
            fact = (
                f"变压器容量={capacity:g}kVA；"
                "未提供运行电流，不能计算负荷率或判断过载"
            )
            for module_id, submodule_id in (
                ("2.1", "2.1.1"),
                ("2.4", "2.4.1.1"),
            ):
                evidence.append(
                    make_evidence(
                        file_id=file_id,
                        path=path,
                        sheet=sheet.title,
                        cell=f"C{row}:E{row}",
                        row=row,
                        column="D",
                        subject=subject,
                        fact=fact,
                        module_id=module_id,
                        submodule_id=submodule_id,
                        value=float(capacity),
                        unit="kVA",
                        needs_confirmation=True,
                    )
                )

        for value_column, photo_column, label, module_id, submodule_id in (
            (6, None, "安全连锁", "2.4", "2.4.1.3"),
            (7, 8, "裸露导体绝缘防护", "2.4", "2.4.2.1"),
            (9, 10, "实测剩余电流", "2.4", "2.4.3.1"),
            (12, 13, "等电位连接与接地", "2.4", "2.4.2.2"),
            (14, 15, "色标标识", "2.4", "2.4.2.4"),
            (16, 17, "谐波", "2.2", "2.2.1.1"),
            (18, 19, "电涌保护装置", "2.3", "2.3.3"),
            (20, None, "变压器噪声震动", "2.2", "2.2.2.3"),
        ):
            value = sheet.cell(row, value_column).value
            if _is_placeholder(value):
                continue
            photo_value = sheet.cell(row, photo_column).value if photo_column else None
            start_letter = get_column_letter(value_column)
            end_column = 11 if label == "实测剩余电流" else (photo_column or value_column)
            end_letter = get_column_letter(end_column)
            numeric_value = float(value) if isinstance(value, (int, float)) else None
            unit = None
            if numeric_value is not None and label == "实测剩余电流":
                unit = "A"
            elif numeric_value is not None and label == "谐波":
                unit = "%"
            fact_value = (
                f"{value:g}{unit}"
                if numeric_value is not None and unit
                else _text(value)
            )
            fact = f"{label}={fact_value}"
            if (
                label == "实测剩余电流"
                and numeric_value is not None
                and isinstance(current, (int, float))
                and current
            ):
                fact += f"；占运行电流={numeric_value / current * 100:.2f}%"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"{start_letter}{row}:{end_letter}{row}",
                    row=row,
                    column=start_letter,
                    subject=subject,
                    fact=fact,
                    module_id=module_id,
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(photo_value),
                    value=numeric_value,
                    unit=unit,
                    needs_confirmation=(
                        _text(value).upper().startswith("NG")
                        or (
                            label == "实测剩余电流"
                            and numeric_value is not None
                            and numeric_value > 10
                        )
                    )
                    and not dispimg_refs(photo_value),
                )
            )
    if capacity_only_rows:
        row_range = f"C{min(capacity_only_rows)}:E{max(capacity_only_rows)}"
        message = (
            f"总配表 {len(capacity_only_rows)} 条变压器容量缺少运行电流；"
            "已保留额定容量，但不能计算负荷率或判断过载"
        )
        for module_id, submodule_id in (
            ("2.1", "2.1.1"),
            ("2.4", "2.4.1.1"),
        ):
            gaps.append(
                MappingGap(
                    code="missing_current_for_load_rate",
                    message=message,
                    sheet=sheet.title,
                    cell=row_range,
                    module_id=module_id,
                    submodule_id=submodule_id,
                )
            )
    return evidence, gaps


def _map_general_inspection(path: Path, file_id: str, sheet: Any) -> list:
    evidence = []
    labels = {
        2: "门窗合规",
        4: "出口应急灯",
        6: "防鼠板",
        8: "消防设施",
        10: "安全用具",
        12: "操作工具",
        14: "整机备件",
        16: "通道及杂物",
        18: "模拟屏或站控单元",
        20: "其他",
    }
    routes = {
        2: ("2.2", "2.2.2.3"),
        4: ("2.2", "2.2.2.3"),
        6: ("2.2", "2.2.2.3"),
        8: ("2.2", "2.2.2.3"),
        10: ("2.5", "2.5.5"),
        12: ("2.5", "2.5.5"),
        14: ("2.5", "2.5.7"),
        16: ("2.2", "2.2.2.3"),
        18: ("2.5", "2.5.4"),
        20: ("2.2", "2.2.2.3"),
    }
    for row in range(4, sheet.max_row + 1):
        room = sheet.cell(row, 1).value
        populated = [sheet.cell(row, column).value for column in labels]
        if not room or not any(value not in (None, "") for value in populated):
            continue
        for column, label in labels.items():
            value = sheet.cell(row, column).value
            if _is_placeholder(value):
                continue
            photo_column = column + 1 if column < 20 else None
            photo_value = sheet.cell(row, photo_column).value if photo_column else None
            photo_refs = dispimg_refs(photo_value)
            companion_note = (
                _text(photo_value)
                if photo_value and not photo_refs and not _is_placeholder(photo_value)
                else ""
            )
            start_letter = get_column_letter(column)
            end_letter = get_column_letter(photo_column or column)
            route_module, route_submodule = routes[column]
            fact = f"{label}={_text(value)}"
            if companion_note:
                fact += f"；补充记录={companion_note}"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"{start_letter}{row}:{end_letter}{row}",
                    row=row,
                    column=start_letter,
                    subject=_subject(room),
                    fact=fact,
                    module_id=route_module,
                    submodule_id=route_submodule,
                    photo_refs=photo_refs,
                    needs_confirmation=_text(value).upper().startswith("NG")
                    and not photo_refs,
                )
            )
    return evidence


def _map_busbar(path: Path, file_id: str, sheet: Any) -> list:
    evidence = []
    fields = (
        (2, 3, "密闭性", "2.4.2.6"),
        (4, 5, "等电位连接与接地", "2.4.2.2"),
        (6, 7, "吊架及安装平直度", "2.4.2.5"),
        (8, 9, "母线紧固", "2.4.2.3"),
        (10, 11, "其他缺陷", "2.4.2.5"),
    )
    for row in range(4, sheet.max_row + 1):
        location = sheet.cell(row, 1).value
        if not location:
            continue
        for column, photo_column, label, submodule_id in fields:
            value = sheet.cell(row, column).value
            if value in (None, ""):
                continue
            start_letter = get_column_letter(column)
            end_letter = get_column_letter(photo_column)
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"{start_letter}{row}:{end_letter}{row}",
                    row=row,
                    column=start_letter,
                    subject=_subject(location),
                    fact=f"{label}={_text(value)}",
                    module_id="2.4",
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(sheet.cell(row, photo_column).value),
                )
            )
    return evidence


def _map_thermal(
    path: Path,
    file_id: str,
    sheet: Any,
    value_sheet: Any,
) -> tuple[list, list[MappingGap]]:
    evidence = []
    gaps: list[MappingGap] = []
    current_room = ""
    current_cabinet = ""
    current_environment: float | None = None
    current_humidity: float | None = None
    room_row: int | None = None
    cabinet_row: int | None = None
    environment_row: int | None = None
    humidity_row: int | None = None
    missing_position_rows: list[int] = []
    for row in range(3, sheet.max_row + 1):
        room = _text(sheet.cell(row, 1).value)
        if room:
            current_room = room
            room_row = row
        cabinet = _text(sheet.cell(row, 4).value)
        if cabinet:
            current_cabinet = cabinet
            cabinet_row = row
        environment = _numeric_value(sheet, value_sheet, row, 2)
        if isinstance(environment, (int, float)):
            current_environment = float(environment)
            environment_row = row
        humidity = _numeric_value(sheet, value_sheet, row, 3)
        if isinstance(humidity, (int, float)):
            current_humidity = float(humidity)
            humidity_row = row
        position = _text(sheet.cell(row, 6).value)
        measured = _numeric_value(sheet, value_sheet, row, 7)
        if not isinstance(measured, (int, float)):
            continue
        fact_parts = []
        if position:
            fact_parts.append(f"检测部位={position}")
        else:
            missing_position_rows.append(row)
            fact_parts.append("检测部位=未填写")
        fact_parts.append(f"测点温度={measured:g}℃")
        load_rate = _numeric_value(sheet, value_sheet, row, 5)
        if isinstance(load_rate, (int, float)):
            fact_parts.append(f"检测时负荷率={load_rate * 100:.2f}%")
        if current_environment is not None:
            fact_parts.append(f"环境温度={current_environment:g}℃")
            fact_parts.append(f"相对环境温升={measured - current_environment:.1f}K")
        if current_humidity is not None:
            fact_parts.append(f"环境湿度={current_humidity:g}%RH")
        fact = "；".join(fact_parts)
        context_rows = [
            context_row
            for context_row in (room_row, cabinet_row, environment_row, humidity_row, row)
            if context_row is not None
        ]
        source_start_row = min(context_rows)
        source_cell = f"A{source_start_row}:K{row}"
        for module_id, submodule_id in (
            ("2.4", "2.4.4"),
            ("2.2", "2.2.2.1"),
        ):
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=source_cell,
                    row=row,
                    column="G",
                    subject=_subject(current_room, current_cabinet),
                    fact=fact,
                    module_id=module_id,
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(sheet.cell(row, 11).value),
                    value=float(measured),
                    unit="℃",
                    needs_confirmation=not position,
                )
            )
        if not current_room or not current_cabinet:
            gaps.append(
                MappingGap(
                    code="incomplete_thermal_context",
                    message=f"红外实测温度缺少配电房或柜号上下文：第 {row} 行",
                    sheet=sheet.title,
                    cell=f"G{row}",
                    module_id="2.4",
                    submodule_id="2.4.4",
                )
            )
        if current_environment is None:
            gaps.append(
                MappingGap(
                    code="missing_thermal_reference",
                    message=f"红外实测温度缺少环境参考温度：第 {row} 行",
                    sheet=sheet.title,
                    cell=f"G{row}",
                    module_id="2.2",
                    submodule_id="2.2.2.1",
                )
            )
    if missing_position_rows:
        message = (
            f"红外表 {len(missing_position_rows)} 条实测温度未填写检测部位；"
            "温度与温升已保留，但不能直接归因到具体触头、端子或连接点"
        )
        cell = f"F{min(missing_position_rows)}:F{max(missing_position_rows)}"
        for module_id, submodule_id in (
            ("2.2", "2.2.2.1"),
            ("2.4", "2.4.4"),
        ):
            gaps.append(
                MappingGap(
                    code="missing_thermal_position",
                    message=message,
                    sheet=sheet.title,
                    cell=cell,
                    module_id=module_id,
                    submodule_id=submodule_id,
                )
            )
    return evidence, gaps


def _missing_sheet_gap(sheet: str, submodule_id: str) -> MappingGap:
    return MappingGap(
        code="missing_sheet",
        message=f"缺少{sheet}",
        sheet=sheet,
        module_id="2.4",
        submodule_id=submodule_id,
    )


def map_s4_4(path: Path, *, file_id: str) -> MappingResult:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    value_workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        evidence = []
        gaps: list[MappingGap] = []
        if "低配评估详情" in workbook.sheetnames:
            low_evidence, low_gaps = _map_low_voltage(
                path,
                file_id,
                workbook["低配评估详情"],
                value_workbook["低配评估详情"],
            )
            evidence.extend(low_evidence)
            gaps.extend(low_gaps)
        else:
            gaps.append(_missing_sheet_gap("低配评估详情", "2.4.1.1"))
        if "总配评估详情表" in workbook.sheetnames:
            total_evidence, total_gaps = _map_total_distribution(
                path,
                file_id,
                workbook["总配评估详情表"],
                value_workbook["总配评估详情表"],
            )
            evidence.extend(total_evidence)
            gaps.extend(total_gaps)
        else:
            gaps.append(_missing_sheet_gap("总配评估详情表", "2.4.1.1"))
        if "配电房合规性" in workbook.sheetnames:
            evidence.extend(_map_general_inspection(path, file_id, workbook["配电房合规性"]))
        if "母线与桥架" in workbook.sheetnames:
            evidence.extend(_map_busbar(path, file_id, workbook["母线与桥架"]))
        if "红外热成像检测记录表" in workbook.sheetnames:
            thermal_evidence, thermal_gaps = _map_thermal(
                path,
                file_id,
                workbook["红外热成像检测记录表"],
                value_workbook["红外热成像检测记录表"],
            )
            evidence.extend(thermal_evidence)
            gaps.extend(thermal_gaps)
        else:
            gaps.append(_missing_sheet_gap("红外热成像检测记录表", "2.4.4"))
        return MappingResult(evidence_items=evidence, gaps=gaps)
    finally:
        value_workbook.close()
        workbook.close()
