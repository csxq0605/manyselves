"""Map S4-6 assessment summaries while preserving their derived status."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..taxonomy import REPORT_TAXONOMY, resolve_submodule
from .common import MappingGap, MappingResult, dispimg_refs, make_evidence


_PATH_ALIASES = (
    ("配电系统负荷分配与过载风险", "2.1.1"),
    ("额定/分断能力", "2.4.1.1"),
    ("电压事件（过压）的防范", "2.3.3"),
    ("设备分合/储能/工作位置显示", "2.4.1.4"),
    ("带病运行问题", "2.4.4"),
)


def _text(value: Any) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _is_placeholder(value: Any) -> bool:
    return _text(value) in {"", "/", "\\", "--"}


def _known_submodule(value: Any) -> str | None:
    candidate = _text(value)
    if not candidate.startswith(("2.1.", "2.2.", "2.3.", "2.4.", "2.5.")):
        return None
    try:
        resolve_submodule(candidate)
    except ValueError:
        return None
    return candidate


def _submodule_from_path(value: Any) -> str | None:
    path_text = _text(value)
    if not path_text:
        return None
    for title, submodule_id in _PATH_ALIASES:
        if title in path_text:
            return submodule_id
    definitions = sorted(
        (
            definition
            for module in REPORT_TAXONOMY.values()
            for definition in module.submodules.values()
        ),
        key=lambda definition: len(definition.title),
        reverse=True,
    )
    for definition in definitions:
        if definition.title in path_text:
            return definition.id
    return None


def _submodule_from_conclusion(path_value: Any, observation: str) -> str | None:
    path_text = _text(path_value)
    if "配置与选型问题" in path_text and any(
        keyword in observation
        for keyword in ("开关", "断路器", "熔断器", "额定", "容量")
    ):
        return "2.4.1.1"
    if "带病运行问题" in path_text:
        if "无法实现分合闸" in observation or "无法分合" in observation:
            return "2.4.1.4"
        if any(
            keyword in observation
            for keyword in ("剩余电流", "接地线的电流过高", "漏电流")
        ):
            return "2.4.3.1"
        if any(keyword in observation for keyword in ("异响", "震动")):
            return "2.2.2.3"
        if any(
            keyword in observation
            for keyword in ("零线腐蚀", "固定在零排", "接线端螺帽")
        ):
            return "2.4.2.3"
        return "2.4.4"
    if "电缆、桥架、母线安装问题" in path_text and any(
        keyword in observation
        for keyword in ("导体裸露", "铜线裸露", "母排裸露", "母线接头裸露")
    ):
        return "2.4.2.1"
    if "设备外壳IP等级与封堵问题" in path_text and "锁闭管理" in observation:
        return "2.5.5"
    return _submodule_from_path(path_text)


def map_s4_6(path: Path, *, file_id: str) -> MappingResult:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        evidence = []
        gaps: list[MappingGap] = []
        if "评估信息汇总表" in workbook.sheetnames:
            sheet = workbook["评估信息汇总表"]
            for row in range(4, sheet.max_row + 1):
                submodule_id = _known_submodule(sheet.cell(row, 2).value)
                if not submodule_id:
                    continue
                finding = _text(sheet.cell(row, 6).value)
                recommendation = _text(sheet.cell(row, 7).value)
                if _is_placeholder(finding) and _is_placeholder(recommendation):
                    gaps.append(
                        MappingGap(
                            code="empty_summary_assessment",
                            message=f"{submodule_id} 汇总评估仅填写占位符，不能作为结论证据",
                            sheet=sheet.title,
                            cell=f"F{row}:G{row}",
                            module_id=resolve_submodule(submodule_id).module_id,
                            submodule_id=submodule_id,
                        )
                    )
                    continue
                parts = []
                if not _is_placeholder(finding):
                    parts.append(f"既有评估结论={finding}")
                if not _is_placeholder(recommendation):
                    parts.append(f"既有行动建议={recommendation}")
                fact = "；".join(parts)
                expected_submodule: str | None = None
                if submodule_id == "2.4.1.1" and "无法分合" in finding:
                    expected_submodule = "2.4.1.4"
                elif submodule_id == "2.4.1.3" and "被遮挡" in finding:
                    expected_submodule = "2.4.4"
                if expected_submodule:
                    gaps.append(
                        MappingGap(
                            code="summary_taxonomy_conflict",
                            message=(
                                f"汇总表将“{finding}”列入 {submodule_id}，"
                                f"但内容更接近 {expected_submodule}，需专业复核"
                            ),
                            sheet=sheet.title,
                            cell=f"B{row}:G{row}",
                            module_id=resolve_submodule(submodule_id).module_id,
                            submodule_id=submodule_id,
                        )
                    )
                evidence.append(
                    make_evidence(
                        file_id=file_id,
                        path=path,
                        sheet=sheet.title,
                        cell=f"F{row}:G{row}",
                        row=row,
                        column="F",
                        subject=f"评估总表 {submodule_id}",
                        fact=fact,
                        module_id=resolve_submodule(submodule_id).module_id,
                        submodule_id=submodule_id,
                        confidence=0.65,
                        needs_confirmation=True,
                    )
                )
        else:
            gaps.append(
                MappingGap(
                    code="missing_sheet",
                    message="缺少评估信息汇总表",
                    sheet="评估信息汇总表",
                )
            )

        if "结论建议汇总表" in workbook.sheetnames:
            sheet = workbook["结论建议汇总表"]
            for row in range(1, sheet.max_row + 1):
                actual_path = _text(sheet.cell(row, 2).value)
                legacy_path = _text(sheet.cell(row, 8).value)
                path_value = actual_path or legacy_path
                observation = _text(sheet.cell(row, 3).value)
                risk = _text(sheet.cell(row, 4).value)
                action = _text(sheet.cell(row, 5).value)
                if not any((observation, risk, action)):
                    continue
                if (
                    actual_path.upper().startswith(("=DISPIMG(", "=_XLFN.DISPIMG("))
                    or "CATEGRY" in actual_path.upper()
                    or observation in {"结论建议汇总表", "发现问题 Topics"}
                ):
                    continue
                submodule_id = _submodule_from_conclusion(path_value, observation)
                if not submodule_id:
                    gaps.append(
                        MappingGap(
                            code="unrouted_conclusion",
                            message=f"结论建议未匹配报告子模块：{path_value or '未填写分类路径'}",
                            sheet=sheet.title,
                            cell=f"B{row}:F{row}",
                        )
                    )
                    continue
                if (
                    "配置与选型问题" in path_value
                    and _submodule_from_path(path_value) is None
                    and submodule_id == "2.4.1.1"
                ):
                    gaps.append(
                        MappingGap(
                            code="inferred_conclusion_submodule",
                            message=(
                                "结论路径仅填写到 2.4.1 配置与选型问题；"
                                f"根据“{observation}”暂按 2.4.1.1 保留，需专业复核"
                            ),
                            sheet=sheet.title,
                            cell=f"B{row}:H{row}",
                            module_id="2.4",
                            submodule_id="2.4.1.1",
                        )
                    )
                fact = "；".join(
                    part
                    for part in (
                        f"既有问题={observation}" if observation else "",
                        f"既有风险={risk}" if risk else "",
                        f"既有建议={action}" if action else "",
                    )
                    if part
                )
                source_cell = f"B{row}:F{row}" if actual_path else f"C{row}:E{row}"
                evidence.append(
                    make_evidence(
                        file_id=file_id,
                        path=path,
                        sheet=sheet.title,
                        cell=source_cell,
                        row=row,
                        column="C",
                        subject=f"评估总表 {submodule_id}",
                        fact=fact,
                        module_id=resolve_submodule(submodule_id).module_id,
                        submodule_id=submodule_id,
                        photo_refs=dispimg_refs(sheet.cell(row, 6).value),
                        confidence=0.6,
                        needs_confirmation=True,
                    )
                )
        return MappingResult(evidence_items=evidence, gaps=gaps)
    finally:
        workbook.close()
