"""Map the S2-1 collection checklist into traceable availability evidence."""

from pathlib import Path

from openpyxl import load_workbook

from ..taxonomy import resolve_submodule
from .common import MappingGap, MappingResult, make_evidence

_DOCUMENT_EXACT_ROUTES = {
    "中低压电气系统图": "2.5.2",
    "竣工电气单线及控制原理图": "2.5.2",
    "主要开关（断路器）技术参数": "2.4.1.1",
    "母线统计表": "2.1.1",
    "双电源装置统计表": "2.1.3",
    "设备标签与挂牌实际照片": "2.4.2.4",
    "关键负荷清单": "2.1.2",
    "短路及时间/电流配合研究": "2.3.1",
    "保护定值表": "2.3.1",
    "保护动作/报警事件记录": "2.3.1",
    "设备巡检/红外成像历史记录": "2.4.4",
    "维护试验与巡检报告": "2.5.3.2",
    "电能质量分析报告": "2.2.1.1",
    "网压波动历史事件统计": "2.2.1.2",
    "近三年电气事故/未遂事件报告": "2.5.1",
    "承包商/临时用电管理记录": "2.5.5",
    "电气安全培训程序": "2.5.3.1",
    "培训记录": "2.5.3.1",
    "现有人员资质": "2.5.3.1",
    "工作票/操作票": "2.5.1",
    "上锁挂牌程序": "2.5.5",
    "EOP文件（应急操作程序）": "2.5.1",
    "SOP文件（标准操作程序）": "2.5.1",
    "现有维护计划": "2.5.3.2",
    "应急预案及演练记录": "2.5.1",
}

_DOCUMENT_ROUTES = (
    (("SOP", "EOP", "操作规程", "应急预案"), "2.5.1"),
    (("保护定值", "定值计算"), "2.3.1"),
    (("自动切换", "ATS"), "2.1.3"),
    (("无功补偿", "电容柜"), "2.1.5"),
    (("组织架构", "人员配备"), "2.5.3.1"),
    (("维护计划", "维护记录", "维护工作"), "2.5.3.2"),
    (("维保覆盖", "维保记录"), "2.5.3.3"),
    (("单线图", "系统图", "图纸"), "2.5.2"),
    (("智能化", "站控", "监控系统"), "2.5.4"),
    (("LOTO", "安全用具", "操作工具"), "2.5.5"),
    (("退市", "生命周期"), "2.5.6"),
    (("备件",), "2.5.7"),
)


def _route_document(document_name: str) -> str | None:
    if exact := _DOCUMENT_EXACT_ROUTES.get(document_name):
        return exact
    folded = document_name.upper()
    for keywords, submodule_id in _DOCUMENT_ROUTES:
        if any(keyword.upper() in folded for keyword in keywords):
            return submodule_id
    return None


def map_s2_1(path: Path, *, file_id: str) -> MappingResult:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if "收资表" not in workbook.sheetnames:
            return MappingResult(
                evidence_items=[],
                gaps=[MappingGap(code="missing_sheet", message="缺少收资表", sheet="收资表")],
            )
        sheet = workbook["收资表"]
        evidence = []
        gaps: list[MappingGap] = []
        category: str | None = None
        for row in range(2, sheet.max_row + 1):
            category = str(sheet.cell(row, 1).value or category or "未分类").strip()
            document_name = str(sheet.cell(row, 2).value or "").strip()
            if not document_name:
                continue
            availability = str(sheet.cell(row, 7).value or "未填写").strip()
            effectiveness = str(sheet.cell(row, 8).value or "未填写").strip()
            note = str(sheet.cell(row, 9).value or "").strip()
            fact = f"具备情况={availability}；有效性={effectiveness}"
            if note:
                fact += f"；备注={note}"
            submodule_id = _route_document(document_name)
            module_id = resolve_submodule(submodule_id).module_id if submodule_id else None
            if submodule_id is None:
                gaps.append(
                    MappingGap(
                        code="unrouted_document",
                        message=f"收资项未匹配报告子模块：{document_name}",
                        sheet=sheet.title,
                        cell=f"B{row}",
                    )
                )
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"G{row}:H{row}",
                    row=row,
                    column="G",
                    subject=f"{category}/{document_name}",
                    fact=fact,
                    module_id=module_id,
                    submodule_id=submodule_id,
                    needs_confirmation=availability.upper() != "OK"
                    or effectiveness.upper() != "OK",
                )
            )
        return MappingResult(evidence_items=evidence, gaps=gaps)
    finally:
        workbook.close()
