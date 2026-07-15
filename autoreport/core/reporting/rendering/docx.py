"""Render an audited ReportState into a template-backed DOCX file."""

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from ..models import ClaimKind, ReportingModel
from ..report_state import ReportState
from ..taxonomy import REPORT_TAXONOMY

_KIND_LABELS = {
    ClaimKind.FACT: "事实",
    ClaimKind.CONCLUSION: "判断",
    ClaimKind.RISK: "风险",
    ClaimKind.RECOMMENDATION: "建议",
}


def _evidence_note(evidence_ids: list[str]) -> str:
    if not evidence_ids:
        return "未核实"
    preview = ", ".join(f"{evidence_id[:13]}…" for evidence_id in evidence_ids[:3])
    remaining = len(evidence_ids) - 3
    return f"{preview}（另 {remaining} 条）" if remaining > 0 else preview


def _review_message(message: str) -> str:
    return re.sub(r"(ev-[0-9a-f]{10})[0-9a-f]+", r"\1…", message)


class RenderResult(ReportingModel):
    output_path: Path
    template_sha256: str
    output_sha256: str
    issue_count: int
    rendered_photo_count: int
    missing_photo_ids: list[str]


class DocxRenderer:
    def __init__(self, template_path: Path, *, asset_root: Path):
        self.template_path = Path(template_path)
        self.asset_root = Path(asset_root)

    def render(self, state: ReportState, output_path: Path) -> RenderResult:
        if not self.template_path.is_file():
            raise FileNotFoundError(f"DOCX template not found: {self.template_path}")
        document = Document(self.template_path)
        self._clear_template_body(document)
        self._set_document_fonts(document)
        self._add_title(document, state.title)
        self._add_overview(document, state)
        document.add_heading("2 评估内容描述", level=1)

        evidence_by_id = {item.id: item for item in state.evidence_items}
        assets_by_id = {asset.id: asset for asset in state.photo_assets}
        rendered_photo_ids: set[str] = set()
        missing_photo_ids: list[str] = []

        for draft in sorted(state.module_drafts, key=lambda item: item.module_id):
            module = REPORT_TAXONOMY[draft.module_id]
            document.add_heading(f"{module.id} {module.title}", level=1)
            current_submodule = None
            for claim in draft.claims:
                if claim.submodule_id != current_submodule:
                    submodule = module.submodules[claim.submodule_id]
                    document.add_heading(
                        f"{submodule.id} {submodule.title}",
                        level=2,
                    )
                    current_submodule = claim.submodule_id
                paragraph = document.add_paragraph()
                label = paragraph.add_run(f"{_KIND_LABELS[claim.kind]}：")
                label.bold = True
                paragraph.add_run(claim.text)
                citation = document.add_paragraph()
                self._set_optional_style(citation, document, "Caption", "脚注")
                citation.add_run(
                    "证据："
                    + (", ".join(claim.evidence_ids) or "未核实")
                    + "；Skill："
                    + ", ".join(claim.skill_ids)
                )

                for evidence_id in claim.evidence_ids:
                    evidence = evidence_by_id[evidence_id]
                    for photo_id in evidence.photo_refs:
                        if photo_id in rendered_photo_ids:
                            continue
                        rendered_photo_ids.add(photo_id)
                        asset = assets_by_id.get(photo_id)
                        asset_path = None
                        if asset is not None:
                            asset_path = (
                                asset.path
                                if asset.path.is_absolute()
                                else self.asset_root / asset.path
                            )
                        if asset_path is None or not asset_path.is_file():
                            document.add_paragraph(f"[图片缺失：{photo_id}]")
                            missing_photo_ids.append(photo_id)
                            continue
                        try:
                            document.add_picture(str(asset_path), width=Inches(5.8))
                            caption = document.add_paragraph(
                                f"图：{evidence.subject}（{photo_id}）"
                            )
                            self._set_optional_style(caption, document, "Caption", "脚注")
                            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        except Exception:
                            document.add_paragraph(f"[图片缺失：{photo_id}]")
                            missing_photo_ids.append(photo_id)

        self._add_conclusions(document, state)
        self._add_review_table(document, state)
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        self._validate_rendered_document(output_path, state)
        return RenderResult(
            output_path=output_path,
            template_sha256=hashlib.sha256(self.template_path.read_bytes()).hexdigest(),
            output_sha256=hashlib.sha256(output_path.read_bytes()).hexdigest(),
            issue_count=len(state.review_issues),
            rendered_photo_count=len(rendered_photo_ids) - len(missing_photo_ids),
            missing_photo_ids=missing_photo_ids,
        )

    @staticmethod
    def _clear_template_body(document: Document) -> None:
        body = document._element.body
        section_properties = body.sectPr
        for child in list(body):
            if child is not section_properties:
                body.remove(child)

    @staticmethod
    def _add_title(document: Document, title: str) -> None:
        paragraph = document.add_paragraph()
        DocxRenderer._set_optional_style(paragraph, document, "Title", "Heading 1")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(24)
        document.add_paragraph("基于已审计客户证据生成")

    @staticmethod
    def _set_document_fonts(document: Document) -> None:
        for style in document.styles:
            if style.type != WD_STYLE_TYPE.PARAGRAPH:
                continue
            properties = style.element.get_or_add_rPr()
            fonts = properties.get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
            fonts.set(qn("w:ascii"), "Arial")
            fonts.set(qn("w:hAnsi"), "Arial")

    @staticmethod
    def _set_optional_style(paragraph, document: Document, *names: str) -> None:
        available = {style.name for style in document.styles}
        for name in names:
            if name in available:
                paragraph.style = name
                return

    @staticmethod
    def _add_overview(document: Document, state: ReportState) -> None:
        document.add_heading("1 配电评估概述", level=1)
        document.add_heading("1.1 评估背景", level=2)
        document.add_paragraph(state.editorial.overview)
        document.add_heading("1.2 待提升问题与建议概览", level=2)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for index, header in enumerate(("序号", "类别", "问题/主题", "优先级", "备注")):
            table.cell(0, index).text = header
        for index, item in enumerate(state.editorial.risk_summary, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = item.module_id
            cells[2].text = item.text
            cells[3].text = item.priority
            cells[4].text = _evidence_note(item.evidence_ids)

    @staticmethod
    def _add_conclusions(document: Document, state: ReportState) -> None:
        document.add_heading("3 结论与建议", level=1)
        document.add_heading("3.1 风险/问题汇总与概览", level=2)
        document.add_paragraph(state.editorial.overview)
        for item in state.editorial.risk_summary:
            paragraph = document.add_paragraph(
                f"[{item.priority}] {item.module_id}/{item.submodule_id}：{item.text}"
            )
            DocxRenderer._set_optional_style(paragraph, document, "List Bullet", "正文")
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
        table = document.add_table(rows=2, cols=5)
        table.style = "Table Grid"
        title_cell = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
        title_paragraph = title_cell.paragraphs[0]
        title_paragraph.text = "3.2 改善行动列表与优先级"
        title_paragraph.paragraph_format.keep_with_next = True
        title_run = title_paragraph.runs[0]
        title_run.bold = True
        title_run.font.size = Pt(16)
        for index, header in enumerate(("类别", "问题/主题", "优先级", "建议措施", "备注")):
            table.cell(1, index).text = header
        for action in state.editorial.actions:
            cells = table.add_row().cells
            cells[0].text = action.module_id
            cells[1].text = action.submodule_id
            cells[2].text = action.priority
            cells[3].text = action.text
            cells[4].text = _evidence_note(action.evidence_ids)

    @staticmethod
    def _add_review_table(document: Document, state: ReportState) -> None:
        table = document.add_table(rows=2, cols=4)
        table.style = "Table Grid"
        title_cell = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
        title_paragraph = title_cell.paragraphs[0]
        title_paragraph.text = "审校与补证事项"
        title_paragraph.paragraph_format.keep_with_next = True
        title_run = title_paragraph.runs[0]
        title_run.bold = True
        title_run.font.size = Pt(18)
        headers = ("子模块", "类型", "级别", "说明")
        for index, header in enumerate(headers):
            table.cell(1, index).text = header
        for issue in state.review_issues:
            cells = table.add_row().cells
            cells[0].text = issue.submodule_id or issue.module_id
            cells[1].text = issue.kind
            cells[2].text = issue.severity
            cells[3].text = _review_message(issue.message)

    @staticmethod
    def _validate_rendered_document(output_path: Path, state: ReportState) -> None:
        if not output_path.is_file() or output_path.stat().st_size == 0:
            raise ValueError("DOCX renderer did not produce an output file")
        rendered = Document(output_path)
        texts = {paragraph.text.strip() for paragraph in rendered.paragraphs}
        if state.title not in texts:
            raise ValueError("rendered DOCX is missing the report title")
        for draft in state.module_drafts:
            expected = f"{draft.module_id} {REPORT_TAXONOMY[draft.module_id].title}"
            if expected not in texts:
                raise ValueError(f"rendered DOCX is missing module heading {expected}")
        for heading in (
            "1 配电评估概述",
            "2 评估内容描述",
            "3 结论与建议",
        ):
            if heading not in texts:
                raise ValueError(f"rendered DOCX is missing fixed heading {heading}")
