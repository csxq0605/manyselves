"""Map the S2-1 collection checklist into traceable availability evidence."""

from pathlib import Path

from openpyxl import load_workbook

from .common import MappingGap, MappingResult, make_evidence

_MODULE_24_DOCUMENTS = {
    "主要开关（断路器）技术参数": "2.4.1.1",
    "设备标签与挂牌实际照片": "2.4.2.4",
    "设备巡检/红外成像历史记录": "2.4.4",
    "维护试验与巡检报告": "2.4.4",
}


def map_s2_1(path: Path, *, file_id: str) -> MappingResult:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if "收资表" not in workbook.sheetnames:
            return MappingResult(
                evidence_items=[],
                gaps=[MappingGap(code="missing_sheet", message="缺少收资表", sheet="收资表")],
            )
        sheet = workbook["收资表"]
        evidence = []
        category: str | None = None
        for row in range(2, sheet.max_row + 1):
            category = str(sheet.cell(row, 1).value or category or "未分类").strip()
            document_name = str(sheet.cell(row, 2).value or "").strip()
            if not document_name:
                continue
            availability = str(sheet.cell(row, 7).value or "未填写").strip()
            effectiveness = str(sheet.cell(row, 8).value or "未填写").strip()
            note = str(sheet.cell(row, 9).value or "").strip()
            fact = f"具备情况={availability}；有效性={effectiveness}"
            if note:
                fact += f"；备注={note}"
            submodule_id = _MODULE_24_DOCUMENTS.get(document_name)
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"G{row}:H{row}",
                    row=row,
                    column="G",
                    subject=f"{category}/{document_name}",
                    fact=fact,
                    module_id="2.4" if submodule_id else None,
                    submodule_id=submodule_id,
                    needs_confirmation=availability.upper() != "OK"
                    or effectiveness.upper() != "OK",
                )
            )
        return MappingResult(evidence_items=evidence, gaps=[])
    finally:
        workbook.close()
