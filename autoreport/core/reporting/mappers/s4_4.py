"""Map S4-4 diagnostic workbook observations into module 2.4 evidence."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .common import MappingGap, MappingResult, dispimg_refs, make_evidence


def _text(value: Any) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _subject(room: Any, cabinet: Any = None) -> str:
    room_text = _text(room) or "未注明配电房"
    cabinet_text = _text(cabinet)
    return f"{room_text}/{cabinet_text}" if cabinet_text else room_text


def _map_low_voltage(path: Path, file_id: str, sheet: Any) -> list:
    evidence = []
    for row in range(3, sheet.max_row + 1):
        room = sheet.cell(row, 1).value
        cabinet = sheet.cell(row, 2).value
        if not cabinet:
            continue
        subject = _subject(room, cabinet)
        current = sheet.cell(row, 3).value
        capacity = sheet.cell(row, 4).value
        if isinstance(current, (int, float)) and isinstance(capacity, (int, float)) and capacity:
            load_rate = current * 0.4 * 1.732 / capacity * 100
            fact = f"运行电流={current:g}A；变压器容量={capacity:g}kVA；计算负荷率={load_rate:.2f}%"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"C{row}:E{row}",
                    row=row,
                    column="C",
                    subject=subject,
                    fact=fact,
                    module_id="2.4",
                    submodule_id="2.4.1.1",
                    value=load_rate,
                    unit="%",
                )
            )

        paired_fields = (
            (6, 7, "安全连锁", "2.4.1.3"),
            (8, 9, "电缆状态", "2.4.2.5"),
            (13, 14, "IP与接地", "2.4.2.2"),
            (15, 16, "色标标识", "2.4.2.4"),
            (21, None, "变压器噪声震动", "2.4.4"),
            (22, 23, "变压器温度", "2.4.4"),
        )
        for value_column, photo_column, label, submodule_id in paired_fields:
            value = sheet.cell(row, value_column).value
            if value in (None, ""):
                continue
            end_column = photo_column or value_column
            photo_value = sheet.cell(row, end_column).value if photo_column else None
            start_letter = get_column_letter(value_column)
            end_letter = get_column_letter(end_column)
            cell = f"{start_letter}{row}:{end_letter}{row}"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=cell,
                    row=row,
                    column=start_letter,
                    subject=subject,
                    fact=f"{label}={_text(value)}",
                    module_id="2.4",
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(photo_value),
                    needs_confirmation=_text(value).upper().startswith("NG")
                    and not dispimg_refs(photo_value),
                )
            )

        residual_current = sheet.cell(row, 10).value
        if isinstance(residual_current, (int, float)):
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"J{row}:K{row}",
                    row=row,
                    column="J",
                    subject=subject,
                    fact=f"实测剩余电流={residual_current:g}A",
                    module_id="2.4",
                    submodule_id="2.4.3.1",
                    photo_refs=dispimg_refs(sheet.cell(row, 11).value),
                    value=residual_current,
                    unit="A",
                    needs_confirmation=residual_current > 10
                    and not dispimg_refs(sheet.cell(row, 11).value),
                )
            )

        other = _text(sheet.cell(row, 24).value)
        if other:
            submodule_id = "2.4.2.1" if "裸露" in other else "2.4.4"
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"X{row}:Z{row}",
                    row=row,
                    column="X",
                    subject=subject,
                    fact=f"其他检查发现={other}",
                    module_id="2.4",
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(
                        sheet.cell(row, 25).value,
                        sheet.cell(row, 26).value,
                    ),
                )
            )
    return evidence


def _map_general_inspection(path: Path, file_id: str, sheet: Any) -> list:
    evidence = []
    labels = {
        2: "门窗合规",
        4: "出口应急灯",
        6: "防鼠板",
        8: "消防设施",
        10: "安全用具",
        12: "操作工具",
        14: "整机备件",
        16: "通道及杂物",
        18: "模拟屏或站控单元",
        20: "其他",
    }
    for row in range(4, sheet.max_row + 1):
        room = sheet.cell(row, 1).value
        populated = [sheet.cell(row, column).value for column in labels]
        if not room or not any(value not in (None, "") for value in populated):
            continue
        for column, label in labels.items():
            value = sheet.cell(row, column).value
            if value in (None, "", "\\"):
                continue
            photo_column = column + 1 if column < 20 else None
            photo_value = sheet.cell(row, photo_column).value if photo_column else None
            start_letter = get_column_letter(column)
            end_letter = get_column_letter(photo_column or column)
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"{start_letter}{row}:{end_letter}{row}",
                    row=row,
                    column=start_letter,
                    subject=_subject(room),
                    fact=f"{label}={_text(value)}",
                    module_id="2.4",
                    submodule_id="2.4.4",
                    photo_refs=dispimg_refs(photo_value),
                )
            )
    return evidence


def _map_busbar(path: Path, file_id: str, sheet: Any) -> list:
    evidence = []
    fields = (
        (2, 3, "密闭性", "2.4.2.6"),
        (4, 5, "等电位连接与接地", "2.4.2.2"),
        (6, 7, "吊架及安装平直度", "2.4.2.5"),
        (8, 9, "母线紧固", "2.4.2.3"),
        (10, 11, "其他缺陷", "2.4.2.5"),
    )
    for row in range(4, sheet.max_row + 1):
        location = sheet.cell(row, 1).value
        if not location:
            continue
        for column, photo_column, label, submodule_id in fields:
            value = sheet.cell(row, column).value
            if value in (None, ""):
                continue
            start_letter = get_column_letter(column)
            end_letter = get_column_letter(photo_column)
            evidence.append(
                make_evidence(
                    file_id=file_id,
                    path=path,
                    sheet=sheet.title,
                    cell=f"{start_letter}{row}:{end_letter}{row}",
                    row=row,
                    column=start_letter,
                    subject=_subject(location),
                    fact=f"{label}={_text(value)}",
                    module_id="2.4",
                    submodule_id=submodule_id,
                    photo_refs=dispimg_refs(sheet.cell(row, photo_column).value),
                )
            )
    return evidence


def _map_thermal(path: Path, file_id: str, sheet: Any) -> list:
    evidence = []
    current_room = ""
    current_cabinet = ""
    current_environment: float | None = None
    for row in range(3, sheet.max_row + 1):
        current_room = _text(sheet.cell(row, 1).value) or current_room
        current_cabinet = _text(sheet.cell(row, 4).value) or current_cabinet
        environment = sheet.cell(row, 2).value
        if isinstance(environment, (int, float)):
            current_environment = float(environment)
        position = _text(sheet.cell(row, 6).value)
        measured = sheet.cell(row, 7).value
        if not position or not isinstance(measured, (int, float)):
            continue
        fact = f"检测部位={position}；测点温度={measured:g}℃"
        if current_environment is not None:
            fact += f"；相对环境温升={measured - current_environment:.1f}K"
        evidence.append(
            make_evidence(
                file_id=file_id,
                path=path,
                sheet=sheet.title,
                cell=f"F{row}:K{row}",
                row=row,
                column="F",
                subject=_subject(current_room, current_cabinet),
                fact=fact,
                module_id="2.4",
                submodule_id="2.4.4",
                photo_refs=dispimg_refs(sheet.cell(row, 11).value),
                value=float(measured),
                unit="℃",
            )
        )
    return evidence


def map_s4_4(path: Path, *, file_id: str) -> MappingResult:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        evidence = []
        gaps: list[MappingGap] = []
        if "低配评估详情" in workbook.sheetnames:
            evidence.extend(_map_low_voltage(path, file_id, workbook["低配评估详情"]))
        else:
            gaps.append(
                MappingGap(code="missing_sheet", message="缺少低配评估详情", sheet="低配评估详情")
            )
        if "配电房合规性" in workbook.sheetnames:
            evidence.extend(_map_general_inspection(path, file_id, workbook["配电房合规性"]))
        if "母线与桥架" in workbook.sheetnames:
            evidence.extend(_map_busbar(path, file_id, workbook["母线与桥架"]))
        if "红外热成像检测记录表" in workbook.sheetnames:
            evidence.extend(_map_thermal(path, file_id, workbook["红外热成像检测记录表"]))
        return MappingResult(evidence_items=evidence, gaps=gaps)
    finally:
        workbook.close()
