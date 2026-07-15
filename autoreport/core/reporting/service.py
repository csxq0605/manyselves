"""Complete V2 reporting service executed inside the AutoReport runtime."""

import asyncio
import uuid
from collections.abc import Awaitable, Callable
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from ...config.schema import AgentDefaults
from ...interfaces.types import AgentType, SystemNotice
from ..loops.bus import MessageBus
from ..providers.base import LLMProvider
from ..tools.task_board import TaskBoard
from .config import AgentDefinition, load_packaged_workflow
from .coverage import evaluate_coverage
from .intake.manifest import build_manifest
from .intake.wps_images import extract_wps_images
from .mappers import map_s2_1, map_s4_4, map_s4_6
from .models import (
    EvidenceItem,
    ModuleDraft,
    ModuleTask,
    OutputArtifact,
    ParsedArtifact,
    PhotoAsset,
    ProjectManifest,
    ReportRequest,
    ReviewIssue,
    SourceLocation,
)
from .planner import CoveragePlanningError, plan_modules
from .rendering.docx import DocxRenderer
from .report_state import build_report_state
from .review.auditor import audit_draft
from .review.cross_module import cross_module_review
from .review.revisions import RevisionLimitError, RevisionRouter
from .skills.resolver import SkillResolver
from .store import ReportingStore
from .workers.generic import GenericModuleWorker
from .workers.module_24 import Module24Worker
from .workers.orchestrator import draft_modules_parallel
from .agent_runner import ReportingAgentRunner
from .workflow import ReportWorkflowRunner


class ReportingRunResult(BaseModel):
    model_config = ConfigDict(arbitrary_types_allowed=True)

    run_id: str
    status: str
    phases: list[str] = Field(default_factory=list)
    output_paths: list[Path] = Field(default_factory=list)
    missing_evidence: list[str] = Field(default_factory=list)
    error: str | None = None


class MissingEvidenceError(RuntimeError):
    def __init__(self, module_ids: list[str]):
        super().__init__(f"missing evidence for modules: {', '.join(module_ids)}")
        self.module_ids = module_ids


Handler = Callable[[dict], Awaitable[None]]


class ReportingService:
    """Run the packaged declarative workflow against the current project."""

    def __init__(
        self,
        workspace: Path,
        *,
        bus: MessageBus,
        task_board: TaskBoard,
        llm_provider: LLMProvider | None = None,
        agent_defaults: AgentDefaults | None = None,
    ):
        self.workspace = Path(workspace).resolve()
        self.bus = bus
        self.task_board = task_board
        self.llm_provider = llm_provider
        self.agent_defaults = agent_defaults or AgentDefaults()
        self.store = ReportingStore(self.workspace)
        self.agents, self.workflow = load_packaged_workflow()
        self.skills = SkillResolver.packaged()
        self.module_24_worker = Module24Worker(self.skills)
        self.module_workers = {
            module_id: GenericModuleWorker(module_id, self.skills)
            for module_id in ("2.1", "2.2", "2.3", "2.5")
        }
        self.module_workers["2.4"] = self.module_24_worker
        self.revision_router = RevisionRouter(max_rounds=2)
        template_root = Path(__file__).resolve().parents[2] / "templates" / "reporting"
        self.report_template_path = template_root / "report_template.docx"
        self.report_template_hash_path = template_root / "report_template.sha256"
        self._handlers: dict[str, Handler] = {
            "manifest-builder": self._build_manifest,
            "intake-parser": self._parse_artifacts,
            "evidence-normalizer": self._normalize_evidence,
            "coverage-evaluator": self._evaluate_coverage,
            "report-planner": self._plan_modules,
            "module-worker": self._draft_modules,
            "evidence-auditor": self._audit_evidence,
            "revision-router": self._route_revisions,
            "project-delivery": self._deliver,
        }

    async def run(self, request: ReportRequest) -> ReportingRunResult:
        self.store.ensure_layout()
        run_id = f"report-{uuid.uuid4().hex[:10]}"
        phases: list[str] = []
        state: dict = {"request": request, "run_id": run_id}
        await self._notice(f"配电报告流程 {run_id} 已启动。")

        try:
            if self.llm_provider is not None:
                phases.extend(phase.id for phase in self.workflow.phases)
                runner = ReportWorkflowRunner(
                    self,
                    ReportingAgentRunner(
                        self.workspace,
                        self.bus,
                        self.llm_provider,
                        self.agent_defaults,
                    ),
                )
                await runner.run(state)
            else:
                # Compatibility path for offline legacy fixtures only. The GUI always
                # injects its active provider and therefore never enters this writer.
                compatibility_task = self.task_board.create_task(
                    AgentType.MAIN, AgentType.MAIN, "offline-reporting-fixture"
                )
                self.task_board.start_task(
                    compatibility_task.task_id, target_agent=AgentType.MAIN
                )
                phases.extend(["intake", "coverage", "module", "quality"])
                try:
                    await self._build_manifest(state)
                    await self._parse_artifacts(state)
                    await self._normalize_evidence(state)
                    await self._evaluate_coverage(state)
                    await self._plan_modules(state)
                    await self._draft_modules(state)
                    await self._audit_evidence(state)
                    await self._route_revisions(state)
                    await self._deliver(state)
                except MissingEvidenceError:
                    self.task_board.block_task(
                        compatibility_task.task_id, target_agent=AgentType.MAIN
                    )
                    raise
                except Exception:
                    self.task_board.fail_task(
                        compatibility_task.task_id, target_agent=AgentType.MAIN
                    )
                    raise
                self.task_board.complete_task(
                    compatibility_task.task_id, target_agent=AgentType.MAIN
                )
        except MissingEvidenceError as exc:
            result = ReportingRunResult(
                run_id=run_id,
                status="blocked",
                phases=phases,
                missing_evidence=exc.module_ids,
            )
            self._save_run(result)
            await self._notice(f"配电报告流程因缺少模块 {', '.join(exc.module_ids)} 的证据而阻塞。")
            return result
        except Exception as exc:
            result = ReportingRunResult(
                run_id=run_id,
                status="failed",
                phases=phases,
                error=str(exc),
            )
            self._save_run(result)
            await self._notice(f"配电报告流程失败：{exc}")
            return result

        artifacts: list[OutputArtifact] = state.get("output_artifacts", [])
        result = ReportingRunResult(
            run_id=run_id,
            status="completed",
            phases=phases,
            output_paths=[self.workspace / artifact.path for artifact in artifacts],
        )
        self._save_run(result)
        await self._notice(
            "配电报告流程已完成："
            + ", ".join(str(path.relative_to(self.workspace)) for path in result.output_paths)
        )
        return result

    async def _run_agent(self, agent: AgentDefinition, state: dict) -> None:
        task = self.task_board.create_task(
            AgentType.MAIN,
            AgentType.MAIN,
            f"reporting:{agent.id}",
        )
        self.task_board.start_task(task.task_id, target_agent=AgentType.MAIN)
        try:
            await self._handlers[agent.id](state)
        except MissingEvidenceError:
            self.task_board.block_task(task.task_id, target_agent=AgentType.MAIN)
            raise
        except Exception:
            self.task_board.fail_task(task.task_id, target_agent=AgentType.MAIN)
            raise
        self.task_board.complete_task(task.task_id, target_agent=AgentType.MAIN)

    async def _notice(self, content: str) -> None:
        await self.bus.publish(SystemNotice(agent_type=AgentType.MAIN, content=content))

    def _save_run(self, result: ReportingRunResult) -> Path:
        return self.store.write_json(
            f"Work/runs/{result.run_id}.json",
            result.model_dump(mode="json"),
        )

    async def _build_manifest(self, state: dict) -> None:
        manifest = build_manifest(self.workspace)
        state["project_manifest"] = manifest
        self.store.write_json("Work/manifest.json", manifest.model_dump(mode="json"))

    async def _parse_artifacts(self, state: dict) -> None:
        manifest: ProjectManifest = state["project_manifest"]
        artifacts: list[ParsedArtifact] = []
        for manifest_file in manifest.files:
            if manifest_file.purpose in {"s2-1", "s4-4", "s4-6"}:
                continue
            try:
                workbook = load_workbook(
                    self.workspace / manifest_file.path,
                    read_only=True,
                    data_only=True,
                )
                for sheet in workbook.worksheets:
                    rows = list(sheet.iter_rows(values_only=True))
                    if not rows:
                        continue
                    headers = [
                        str(value or f"column_{index + 1}") for index, value in enumerate(rows[0])
                    ]
                    for row_index, row in enumerate(rows[1:], start=2):
                        if not any(value not in (None, "") for value in row):
                            continue
                        artifacts.append(
                            ParsedArtifact(
                                id=f"artifact-{len(artifacts) + 1:04d}",
                                kind="workbook_row",
                                source=SourceLocation(
                                    file_id=manifest_file.id,
                                    path=manifest_file.path,
                                    sheet=sheet.title,
                                    cell=f"A{row_index}",
                                ),
                                payload={"headers": headers, "values": list(row)},
                            )
                        )
                workbook.close()
                manifest_file.parse_status = "parsed"
            except Exception as exc:
                manifest_file.parse_status = "failed"
                manifest_file.error = str(exc)
        state["parsed_artifacts"] = artifacts
        self.store.write_json("Work/manifest.json", manifest.model_dump(mode="json"))

    async def _normalize_evidence(self, state: dict) -> None:
        evidence: list[EvidenceItem] = []
        manifest: ProjectManifest = state["project_manifest"]
        photo_assets: list[PhotoAsset] = []
        mapping_gaps: list[dict] = []
        mappers = {
            "s2-1": map_s2_1,
            "s4-4": map_s4_4,
            "s4-6": map_s4_6,
        }
        for manifest_file in manifest.files:
            mapper = mappers.get(manifest_file.purpose or "")
            if mapper is None:
                continue
            input_path = self.workspace / manifest_file.path
            try:
                if manifest_file.purpose == "s4-4":
                    extracted = extract_wps_images(
                        input_path,
                        output_dir=self.workspace / "Work" / "assets" / manifest_file.id,
                    )
                    photo_assets.extend(
                        asset.model_copy(update={"path": asset.path.relative_to(self.workspace)})
                        for asset in extracted.values()
                    )
                mapped = mapper(input_path, file_id=manifest_file.id)
                evidence.extend(
                    item.model_copy(
                        update={
                            "source": item.source.model_copy(update={"path": manifest_file.path})
                        }
                    )
                    for item in mapped.evidence_items
                )
                mapping_gaps.extend(
                    {
                        "file_id": manifest_file.id,
                        **gap.model_dump(mode="json"),
                    }
                    for gap in mapped.gaps
                )
                manifest_file.parse_status = "parsed"
            except Exception as exc:
                manifest_file.parse_status = "failed"
                manifest_file.error = str(exc)

        for artifact in state.get("parsed_artifacts", []):
            pairs = [
                f"{header}={value}"
                for header, value in zip(
                    artifact.payload["headers"], artifact.payload["values"], strict=False
                )
                if value not in (None, "")
            ]
            if not pairs:
                continue
            evidence.append(
                EvidenceItem(
                    id=f"ev-{len(evidence) + 1:04d}",
                    subject=str(artifact.payload["values"][0]),
                    fact="; ".join(pairs),
                    source=artifact.source,
                )
            )
        # Runtime source tools and ClaimLedger use stable E-* identifiers. Mapper
        # internals may emit legacy ev-* ids, so normalize once at the boundary.
        evidence = [
            item.model_copy(update={"id": f"E-{index:04d}"})
            for index, item in enumerate(evidence, start=1)
        ]
        state["evidence_items"] = evidence
        state["photo_assets"] = photo_assets
        state["mapping_gaps"] = mapping_gaps
        self.store.write_jsonl(
            "Work/evidence.jsonl",
            [item.model_dump(mode="json") for item in evidence],
        )
        self.store.write_json(
            "Work/photo-manifest.json",
            {"assets": [asset.model_dump(mode="json") for asset in photo_assets]},
        )
        self.store.write_json("Work/mapping-gaps.json", {"gaps": mapping_gaps})
        self.store.write_json("Work/manifest.json", manifest.model_dump(mode="json"))

    async def _evaluate_coverage(self, state: dict) -> None:
        request: ReportRequest = state["request"]
        evidence: list[EvidenceItem] = state.get("evidence_items", [])
        coverage = evaluate_coverage(request, evidence)
        state["coverage_matrix"] = coverage
        self.store.write_json("Work/coverage.json", coverage.model_dump(mode="json"))

    async def _plan_modules(self, state: dict) -> None:
        request: ReportRequest = state["request"]
        try:
            state["module_tasks"] = plan_modules(
                request,
                state["coverage_matrix"],
                state.get("evidence_items", []),
            )
        except CoveragePlanningError as exc:
            raise MissingEvidenceError(exc.missing_submodules) from exc

    async def _draft_modules(self, state: dict) -> None:
        evidence_items = list(state.get("evidence_items", []))
        drafts, executions = await draft_modules_parallel(
            state.get("module_tasks", []),
            evidence_items,
            self.module_workers,
        )
        state["module_drafts"] = drafts
        state["module_executions"] = executions
        self.store.write_json(
            "Work/module-execution.json",
            {"modules": [execution.model_dump(mode="json") for execution in executions]},
        )
        for draft in drafts:
            self.store.write_json(
                f"Work/drafts/{draft.module_id}.json",
                draft.model_dump(mode="json"),
            )

    async def _audit_evidence(self, state: dict) -> None:
        issues: list[ReviewIssue] = []
        for draft in state.get("module_drafts", []):
            issues.extend(
                audit_draft(
                    draft,
                    state.get("evidence_items", []),
                    self.skills,
                )
            )
        cross_issues = cross_module_review(
            state.get("module_drafts", []),
            state.get("evidence_items", []),
        )
        issues.extend(cross_issues)
        state["cross_module_issues"] = cross_issues
        state["review_issues"] = issues
        self.store.write_json(
            "Work/cross-module-review.json",
            {"issues": [issue.model_dump(mode="json") for issue in cross_issues]},
        )

    async def _route_revisions(self, state: dict) -> None:
        evidence = state.get("evidence_items", [])
        tasks = {task.module_id: task for task in state.get("module_tasks", [])}
        final_drafts: list[ModuleDraft] = []
        final_issues: list[ReviewIssue] = []
        for draft in state.get("module_drafts", []):
            module_issues = [
                issue
                for issue in state.get("review_issues", [])
                if issue.module_id == draft.module_id
            ]
            parent_task = tasks[draft.module_id]
            current = draft
            persistent_cross = [
                issue
                for issue in module_issues
                if issue.kind in {"metric_conflict", "action_conflict"}
                and issue.severity == "blocking"
            ]
            while any(issue.severity == "blocking" for issue in module_issues):

                def revise_submodule(
                    submodule_id: str,
                    _claims: list,
                    _issues: list[ReviewIssue],
                ) -> list:
                    evidence_ids = parent_task.submodule_evidence.get(submodule_id, [])
                    local_task = ModuleTask(
                        id=f"{parent_task.id}-{submodule_id}-r{current.revision + 1}",
                        module_id=draft.module_id,
                        evidence_ids=evidence_ids,
                        submodule_evidence={submodule_id: evidence_ids},
                        revision=current.revision + 1,
                    )
                    return (
                        self.module_workers[draft.module_id]
                        .run(
                            local_task,
                            evidence,
                        )
                        .claims
                    )

                try:
                    current = self.revision_router.route(
                        current,
                        module_issues,
                        revise_submodule,
                    )
                except RevisionLimitError as exc:
                    state["revision_escalation"] = exc.payload
                    self.store.write_json("Work/revision-escalation.json", exc.payload)
                    raise RuntimeError(
                        f"module {draft.module_id} exceeded local revision limit"
                    ) from exc
                module_issues = audit_draft(current, evidence, self.skills) + persistent_cross

            current = self.revision_router.route(
                current,
                module_issues,
                lambda _submodule_id, claims, _issues: claims,
            )
            final_drafts.append(current)
            final_issues.extend(module_issues)
            self.store.write_json(
                f"Work/drafts/{current.module_id}.json",
                current.model_dump(mode="json"),
            )
        state["module_drafts"] = final_drafts
        state["review_issues"] = final_issues
        for final_draft in final_drafts:
            self.store.write_json(
                f"Work/drafts/{final_draft.module_id}.json",
                final_draft.model_dump(mode="json"),
            )

    async def _deliver(self, state: dict) -> None:
        output_artifacts: list[OutputArtifact] = []
        for draft in state.get("module_drafts", []):
            relative = Path("Outputs") / "Modules" / f"{draft.module_id}.md"
            self.store.write_text(relative.as_posix(), draft.markdown)
            output_artifacts.append(
                OutputArtifact(kind="module", path=relative, module_id=draft.module_id)
            )
        review_path = self.store.write_json(
            "Outputs/Reviews/phase-a.json",
            {"issues": [issue.model_dump(mode="json") for issue in state.get("review_issues", [])]},
        )
        self.store.write_json(
            "Outputs/Reviews/full-review.json",
            {"issues": [issue.model_dump(mode="json") for issue in state.get("review_issues", [])]},
        )
        output_artifacts.append(
            OutputArtifact(kind="review", path=review_path.relative_to(self.workspace))
        )

        report_state = build_report_state(
            title="配电安全专家咨询报告",
            request=state["request"],
            manifest=state["project_manifest"],
            coverage=state["coverage_matrix"],
            evidence_items=state.get("evidence_items", []),
            photo_assets=state.get("photo_assets", []),
            module_drafts=state.get("module_drafts", []),
            review_issues=state.get("review_issues", []),
            module_executions=state.get("module_executions", []),
        )
        self.store.write_json(
            "Work/report-state.json",
            report_state.model_dump(mode="json"),
        )
        self.store.write_json(
            "Work/editorial.json",
            report_state.editorial.model_dump(mode="json"),
        )
        expected_template_hash = self.report_template_hash_path.read_text(encoding="utf-8").split()[
            0
        ]
        report_relative = Path("Outputs") / "Reports" / "配电安全专家咨询报告.docx"
        render_result = DocxRenderer(
            self.report_template_path,
            asset_root=self.workspace,
        ).render(report_state, self.workspace / report_relative)
        if render_result.template_sha256 != expected_template_hash:
            raise ValueError("packaged DOCX template hash does not match approved template")
        self.store.write_json(
            "Outputs/Reports/render-log.json",
            render_result.model_dump(mode="json"),
        )
        output_artifacts.append(OutputArtifact(kind="report", path=report_relative))
        state["output_artifacts"] = output_artifacts
