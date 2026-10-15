from pathlib import Path
import re,json,zipfile,hashlib
import openpyxl
R=Path(__file__).resolve().parents[1]
report=(R/'test/Outputs/Reports/配电安全专家咨询报告.md').read_text(encoding='utf-8')
w=openpyxl.load_workbook(R/'test/Inputs/S4-4诊断工作用表.xlsx',data_only=True)
loadlines=[x for x in report.splitlines() if re.search(r'\| E-(0073|0082|0091|0100|0109|0118|0127|0136|0145) \|',x)]
checks=[]
for row,line in zip(range(3,12),loadlines):
    cols=[x.strip() for x in line.strip('|').split('|')]
    s=w['总配评估详情表']
    checks.append({'row':row,'capacity':float(cols[2])==s.cell(row,4).value,'current':float(cols[3])==s.cell(row,3).value,'load_pct':abs(float(cols[4].strip('%'))-s.cell(row,5).value*100)<0.005})
photos=json.loads((R/'test/Work/photo-manifest.json').read_text(encoding='utf-8'))['assets']
with zipfile.ZipFile(R/'test/Outputs/Reports/配电安全专家咨询报告.docx') as z:
    hashes={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    photoresults=[{'id':a['id'],'exists':(R/'test'/a['path']).exists(),'embedded_unchanged':a['sha256'] in hashes} for a in photos]
metrics={'load_table_rows':checks,'photo_count':len(photos),'photo_files_found':sum(x['exists'] for x in photoresults),'photos_embedded_unchanged':sum(x['embedded_unchanged'] for x in photoresults),'missing_photos':[x for x in photoresults if not x['exists'] or not x['embedded_unchanged']]}
(Path(__file__).parent/'metrics.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(metrics,ensure_ascii=False))
