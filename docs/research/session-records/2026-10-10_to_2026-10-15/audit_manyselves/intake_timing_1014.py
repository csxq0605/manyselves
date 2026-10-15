"""Offline before/after intake timings on synthetic bytes, no model calls."""
import importlib
import json
from pathlib import Path
import statistics
import subprocess
import time
import types

from docx import Document
from openpyxl import Workbook

from manyselves.capabilities.distribution_reporting.runtime.models.preparation import ManifestFile

ROOT = Path(__file__).parent
FIXTURES = ROOT / "intake_timing_fixtures"
FIXTURES.mkdir(exist_ok=True)
PREFIX = "manyselves.capabilities.distribution_reporting.runtime."


def baseline(name, file):
    module = types.ModuleType(PREFIX + name)
    module.__package__ = module.__name__.rsplit(".", 1)[0]
    source = subprocess.check_output(["git", "show", "cf29d06:" + file]).decode("utf-8")
    exec(compile(source, file, "exec"), module.__dict__)
    return module


old_adapter = baseline("intake.adapters", "manyselves/capabilities/distribution_reporting/runtime/intake/adapters.py")
old_mapper = baseline("mappers.s4_4", "manyselves/capabilities/distribution_reporting/runtime/mappers/s4_4.py")
new_adapter = importlib.import_module(PREFIX + "intake.adapters")
new_mapper = importlib.import_module(PREFIX + "mappers.s4_4")

doc = Document()
table = doc.add_table(rows=101, cols=2)
for index, row in enumerate(table.rows):
    row.cells[0].text, row.cells[1].text = f"柜{index}", f"{index} A"
doc_path = FIXTURES / "measurements.docx"
doc.save(doc_path)
book = Workbook()
for index in range(301):
    book.active.append([f"柜{index}", index, "=1+2"])
xlsx_path = FIXTURES / "measurements.xlsx"
book.save(xlsx_path)
thermal = Workbook()
sheet = thermal.active
sheet.title = "红外热成像检测记录表"
for row in range(3, 63):
    for column, value in {1:"A房", 2:20, 3:50, 4:f"柜{row}", 5:0.5, 6:"端子", 7:40}.items():
        sheet.cell(row, column, value)
thermal_path = FIXTURES / "S4-4.xlsx"
thermal.save(thermal_path)

result = {}
for label, path in (("docx_101_table_rows", doc_path), ("xlsx_301_formula_rows", xlsx_path), ("s4_4_60_measurements", thermal_path)):
    manifest = ManifestFile(id="fixture", path=path, sha256="a" * 64, media_type="test")
    arms = {"before": [], "after": []}
    counts = {}
    for repeat in range(3):
        for arm in (("before", "after") if repeat % 2 == 0 else ("after", "before")):
            start = time.perf_counter()
            if label.startswith("s4"):
                output = (old_mapper if arm == "before" else new_mapper).map_s4_4(path, file_id="fixture")
                counts[arm] = {"facts": len(output.evidence_items), "gaps": len(output.gaps)}
            else:
                output = (old_adapter if arm == "before" else new_adapter).IntakeAdapterRegistry().parse(path, manifest)
                counts[arm] = {"artifacts": len(output), "gaps": sum(a.kind == "manual_required" for a in output)}
            arms[arm].append(time.perf_counter() - start)
    result[label] = {"seconds": arms, "median_seconds": {arm: statistics.median(values) for arm, values in arms.items()}, "counts": counts}
(ROOT / "intake_timing_metrics_1014.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
