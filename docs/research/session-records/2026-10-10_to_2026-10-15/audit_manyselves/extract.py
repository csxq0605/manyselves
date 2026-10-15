from pathlib import Path
import json,re,zipfile
import openpyxl
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
OUT=Path(__file__).parent
for p in (ROOT/'test/Inputs').glob('*.xlsx'):
    w=openpyxl.load_workbook(p,data_only=True)
    lines=[]
    for s in w:
        lines.append(f'\n## {s.title} ({s.max_row}x{s.max_column})')
        for row in s:
            cells=[f'{c.coordinate}={str(c.value).replace(chr(10)," / ")}' for c in row if c.value is not None]
            if cells: lines.append(' | '.join(cells))
    (OUT/(p.stem+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
    print(p.name, [(s.title,s.max_row,s.max_column) for s in w])
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for p in [* (ROOT/'test/Templates').glob('*.docx'),*(ROOT/'test/Outputs/Reports').glob('*.docx')]:
    with zipfile.ZipFile(p) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        paras=[''.join(x.itertext()) for x in []]
        paras=[''.join(x.xpath('.//w:t/text()',namespaces=ns)) for x in xml.xpath('//w:p',namespaces=ns)]
        (OUT/(p.stem+'.docx.txt')).write_text('\n'.join(paras),encoding='utf-8')
        print(p.name,'paras',len(paras),'tables',len(xml.xpath('//w:tbl',namespaces=ns)),'drawings',len(xml.xpath('//w:drawing',namespaces=ns)),'footnotes','word/footnotes.xml' in z.namelist())
report=(ROOT/'test/Outputs/Reports/配电安全专家咨询报告.md').read_text(encoding='utf-8')
index=(ROOT/'test/Outputs/Reports/证据与来源索引.md').read_text(encoding='utf-8')
ev=[json.loads(x) for x in (ROOT/'test/Work/evidence.jsonl').read_text(encoding='utf-8').splitlines()]
ids={x['id'] for x in ev}
used=set(re.findall(r'E-\d{4}',report))
print('report lines',len(report.splitlines()),'chars',len(report),'evidence total',len(ev),'used',len(used),'undefined',sorted(used-ids))
(OUT/'evidence_compact.txt').write_text('\n'.join(f"{x['id']} | {x['subject']} | {x['fact']} | {json.dumps(x['source'],ensure_ascii=False)}" for x in ev),encoding='utf-8')
print('footnote refs',len(re.findall(r'\[\d+\]',report)),'photos',len(re.findall(r'!\[',report)))
