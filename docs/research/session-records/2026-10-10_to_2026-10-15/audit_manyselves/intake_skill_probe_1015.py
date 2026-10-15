"""Offline intake characterizations and local source/skill statistics."""
import hashlib
import json
import sys
import uuid
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent / 'manyselves'))
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from manyselves.capabilities.distribution_reporting.runtime.preparation import prepare_manifest_file
from manyselves.capabilities.distribution_reporting.runtime.preparation_tools import PreparationTools
from manyselves.capabilities.distribution_reporting.runtime.models.preparation import ManifestFile, PreparationContext, ProjectManifest
from manyselves.capabilities.distribution_reporting.runtime.models.reporting import ReportRequest
from manyselves.capabilities.distribution_reporting.runtime.storage import ReportingStore
from manyselves.capabilities.distribution_reporting.runtime.mappers import map_s4_4
from manyselves.capabilities.distribution_reporting.runtime.intake.manifest import _purpose_for
from manyselves.capabilities.distribution_reporting.runtime.research.project_evidence import ProjectEvidenceIndex
from manyselves.capabilities.distribution_reporting.runtime.research.fact_context import render_fact_packet
from manyselves.runtime.artifacts import parse_artifact
from manyselves.capabilities.distribution_reporting.runtime.module_authoring_preparation import load_template_skill
from manyselves.capabilities.distribution_reporting.runtime.models.agentic import TemplateSkillSubmission

workspace = ROOT / ('intake-probe-' + uuid.uuid4().hex[:8])
(workspace / 'Inputs').mkdir(parents=True)
doc = Document()
doc.add_paragraph('PARAGRAPH_PRESENT')
table = doc.add_table(rows=2, cols=2)
table.cell(0,0).text = 'Reading'
table.cell(1,0).text = 'TABLE_ONLY_READING_17'
doc.save(workspace/'Inputs/table.docx')
wb = Workbook()
ws = wb.active
ws.append(['Device', 'Measured'])
ws.append(['FICTIONAL', '=10+7'])
wb.save(workspace/'Inputs/formula.xlsx')
img = Image.new('RGB', (500, 100), 'white')
ImageDraw.Draw(img).text((10,10), 'IMAGE_ONLY_READING_17', fill='black')
img.save(workspace/'Inputs/observation.png')
img.save(workspace/'Inputs/scan.pdf', 'PDF')
(workspace/'Inputs/corrupt.docx').write_bytes(b'not-a-docx')

rows = []
for name in ['table.docx','formula.xlsx','observation.png','scan.pdf','corrupt.docx']:
    path = workspace/'Inputs'/name
    mf = ManifestFile(id='fixture', path=Path('Inputs')/name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), media_type='test')
    result = prepare_manifest_file(workspace, 'fixture', mf, 0)
    ctx = PreparationContext(run_id='fixture', request=ReportRequest(instruction='synthetic'),
                             project_manifest=ProjectManifest(files=[mf]),
                             preparation_worker_results=[result], parsed_artifacts=result.parsed_artifacts)
    tool = PreparationTools(workspace=workspace, input_snapshot=None, snapshot_content=lambda s,t: None,
                            runtime_photo_ids=lambda e,p: None, store=ReportingStore(workspace))
    tool.normalize_evidence(ctx)
    rows.append({'case':name, 'status':result.status, 'artifact_kinds':[a.kind for a in result.parsed_artifacts],
                 'facts':[e.fact for e in ctx.evidence_items], 'mapping_gaps':ctx.mapping_gaps,
                 'worker_has_error':result.error is not None})

wb = Workbook()
ws = wb.active
ws.title = '红外热成像检测记录表'
for row, vals in {3:['ROOM-A',20,50,'CAB-A',0.5,'POINT-A',40],
                  4:['ROOM-B',None,None,'CAB-B',0.5,'POINT-B',55]}.items():
    for col, val in enumerate(vals,1): ws.cell(row,col,val)
thermal_path = workspace/'Inputs/S4-4-synthetic.xlsx'
wb.save(thermal_path)
mapped = map_s4_4(thermal_path,file_id='thermal')
thermal_b = [e.model_dump(mode='json') for e in mapped.evidence_items if e.subject=='ROOM-B/CAB-B']

original = ROOT.parent/'test/Inputs/S4-4诊断工作用表.xlsx'
orig = map_s4_4(original, file_id='local-original')
raw = load_workbook(original, read_only=True, data_only=True)
low = raw['低配评估详情']
expected = {row:low.cell(row,10).value for row in range(3,low.max_row+1)
            if low.cell(row,2).value and isinstance(low.cell(row,10).value,(int,float))}
observed = {e.source.row:e.value for e in orig.evidence_items
            if e.source.sheet=='低配评估详情' and e.source.column=='J' and e.submodule_id=='2.4.3.1'}
raw.close()
groups=Counter((e.source.sheet,e.source.cell,e.subject,e.fact) for e in orig.evidence_items)
skills=ROOT.parent/'test/Inputs/report-template-role-skills'
skill_sizes={p.parent.name:len(p.read_text(encoding='utf-8')) for p in skills.glob('*/SKILL.md')}
unassigned = ctx.evidence_items
alternate = parse_artifact(workspace/'Inputs/table.docx')
skill_state = {}
load_template_skill(ROOT.parent/'test', skill_state)
validated_skills = TemplateSkillSubmission(skills=skill_state['template_skill_text'], boundary_manifest=skill_state['template_skill_boundary'])
report={'model_calls':0, 'production_code_modified':True, 'synthetic_intake':rows,
        'alternate_document_parser':{'table_fact_recovered':any('TABLE_ONLY_READING_17' in b.text for b in alternate.blocks),'block_count':len(alternate.blocks)},
        'actual_skill_structural_validation_passed':len(validated_skills.skills)==14,
        'synthetic_cross_room_thermal':thermal_b,
        'actual_s4_4_local_replay':{'total_mapped':len(orig.evidence_items),
            'unique_source_subject_fact':len(groups),'multi_route_extra_records':sum(v-1 for v in groups.values()),
            'residual_raw_rows':len(expected),'residual_mapped_rows':len(observed),'residual_values_match':expected==observed},
        'template_role_skill_chars':skill_sizes,
        'skill_total_chars':sum(skill_sizes.values()),
        'filename_routing':{n:_purpose_for(Path(n)) for n in ['S4-4-synthetic.xlsx','renamed.xlsx','S4-4-notes.docx']}}
(ROOT/'intake_skill_metrics_1015.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:report[k] for k in ['model_calls','actual_s4_4_local_replay','actual_skill_structural_validation_passed','template_role_skill_chars']},ensure_ascii=False,indent=2))
